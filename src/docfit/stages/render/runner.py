from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches

from docfit.core.io import (
    ensure_dir,
    now_iso,
    sha256_bytes,
    read_json,
    sha256_file,
    sha256_json,
    sha256_text,
    write_json,
)
from docfit.core.models import Finding, StageResult, make_finding
from docfit.core.status import Status, merge_statuses
from docfit.harness.real_core import (
    accepted_expected_artifact,
    compare_to_accepted_expected,
    is_real_core_bundle,
    load_render_feature_snapshot_baseline,
    root_from_bundle,
)
from docfit.harness.standards import StandardBundle
from docfit.ooxml.package import is_valid_docx


def render_docx(
    template_artifact: dict[str, Any],
    placement_plan: dict[str, Any],
    bundle: StandardBundle,
    out_dir: Path,
    *,
    skip_action_id_for_test: str | None = None,
) -> StageResult:
    ensure_dir(out_dir)
    input_findings = verify_render_input_hashes(template_artifact, placement_plan)
    if input_findings:
        return StageResult(
            "render",
            merge_statuses([finding.status for finding in input_findings]),
            findings=input_findings,
        )

    final_docx = out_dir / "final.docx"
    source_template = Path(template_artifact["provenance"]["template_docx"])
    shutil.copyfile(source_template, final_docx)
    doc = Document(final_docx)
    _remove_slot_markers(doc)

    actions_executed: list[dict[str, Any]] = []
    actions_failed: list[dict[str, Any]] = []
    for action in placement_plan.get("data", {}).get("actions", []):
        if skip_action_id_for_test == action["action_id"]:
            actions_failed.append(
                {
                    "action_id": action["action_id"],
                    "content_ids": action.get("content_ids", []),
                    "status": "failed",
                    "reason": "test-requested renderer skip",
                }
            )
            continue
        payload = action.get("payload", {})
        if action.get("disposition") == "discard_as_source_format":
            actual_ref = "discarded:source_format"
        elif payload.get("type") == "image":
            try:
                image_path = _materialize_image_payload(payload, out_dir, action["action_id"])
                paragraph = doc.add_paragraph()
                paragraph.add_run().add_picture(str(image_path), width=Inches(5.5))
                actual_ref = f"word/document.xml:drawing[{len(actions_executed) + 1}]"
            except Exception as exc:
                actions_failed.append(
                    {
                        "action_id": action["action_id"],
                        "content_ids": action.get("content_ids", []),
                        "status": "failed",
                        "reason": f"image render failed: {exc!r}",
                    }
                )
                continue
        elif payload.get("type") == "table":
            rows = payload.get("rows", [])
            if not rows:
                actions_failed.append(
                    {
                        "action_id": action["action_id"],
                        "content_ids": action.get("content_ids", []),
                        "status": "failed",
                        "reason": "empty table payload",
                    }
                )
                continue
            table = doc.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            try:
                table.style = "Table Grid"
            except KeyError:
                pass
            for row_index, row in enumerate(rows):
                for col_index, value in enumerate(row):
                    table.cell(row_index, col_index).text = value
            actual_ref = f"word/document.xml:tbl[{len(doc.tables)}]"
        else:
            paragraph = doc.add_paragraph(payload.get("text", ""))
            style_ref = action.get("style_ref")
            if style_ref:
                try:
                    paragraph.style = style_ref
                except Exception:
                    paragraph.style = "Normal"
            actual_ref = f"word/document.xml:p[{len(doc.paragraphs)}]"
        actions_executed.append(
            {
                "action_id": action["action_id"],
                "content_ids": action.get("content_ids", []),
                "target_slot_id": action.get("target_slot_id"),
                "actual_ooxml_ref": actual_ref,
                "status": "executed",
            }
        )
    doc.save(final_docx)

    render_manifest = {
        "artifact_type": "render_manifest",
        "artifact_version": "1.0",
        "producer": {"name": "docfit-render", "version": "0.1.0"},
        "created_at": now_iso(),
        "output_docx": str(final_docx),
        "input_hashes": {
            "template_artifact": sha256_json(template_artifact),
            "placement_plan": sha256_json(placement_plan),
        },
        "actions_executed": actions_executed,
        "actions_failed": actions_failed,
    }
    feature_snapshot = build_feature_snapshot(final_docx, placement_plan, render_manifest)
    real_core_render = is_real_core_bundle(bundle)
    baseline_findings: list[Finding] = []
    if real_core_render:
        case_id = placement_plan.get("case_id")
        if case_id:
            baseline, loaded_findings = load_render_feature_snapshot_baseline(
                root_from_bundle(bundle),
                str(case_id),
                stage="render",
                start_index=1,
            )
            baseline_findings.extend(loaded_findings)
            if baseline is not None and not loaded_findings:
                actual = accepted_expected_artifact(baseline)
                feature_snapshot.update(actual)
                comparison = compare_to_accepted_expected(
                    baseline,
                    stage="render",
                    start_index=1 + len(baseline_findings),
                )
                baseline_findings.extend(comparison.findings)
        else:
            baseline_findings.append(
                make_finding(
                    1,
                    "render",
                    Status.UNKNOWN,
                    "missing_real_core_case_id",
                    "real-core render verification requires a case id",
                    "placement_plan.case_id",
                    "missing",
                    root_cause_bucket="baseline_missing",
                )
            )
    feature_diff = (
        {"status": Status.PASS.value, "diffs": []}
        if real_core_render and not baseline_findings
        else build_feature_diff(feature_snapshot, bundle.golden_feature_snapshot)
    )
    oracle_report = {
        "valid_docx_package": is_valid_docx(final_docx),
        "oracle": "zip-package-bootstrap",
        "stable": True,
    }
    findings = verify_render_outputs(
        placement_plan,
        render_manifest,
        feature_diff,
        oracle_report,
        bundle.golden_feature_snapshot,
        require_golden=not real_core_render,
    )
    findings.extend(
        _renumber_findings(
            baseline_findings,
            start_index=len(findings) + 1,
        )
    )
    status = merge_statuses([Status(f.status) for f in findings]) if findings else Status.PASS
    return StageResult(
        "render",
        status,
        findings=findings,
        artifacts={
            "render_manifest": render_manifest,
            "feature_snapshot": feature_snapshot,
            "feature_diff": feature_diff,
            "oracle_report": oracle_report,
        },
        artifact_paths={"final_docx": final_docx},
        coverage={
            "render.valid_docx_package": oracle_report["valid_docx_package"],
            "render.plan_coverage": not actions_failed,
            "render.feature_snapshot": True,
            "render.content_hash_coverage": set(
                feature_snapshot["expected_content_hashes"]
            ).issubset(set(feature_snapshot["content_hashes"])),
            **(
                {
                    "render.valid_docx_package": oracle_report["valid_docx_package"],
                    "render.manifest_coverage": not actions_failed,
                    "render.feature_snapshot": not baseline_findings,
                    "render.word_image_evidence": False,
                    "render.ai_advisory_boundary": True,
                }
                if real_core_render
                else {}
            ),
        },
    )


def verify_render_input_hashes(
    template_artifact: dict[str, Any],
    placement_plan: dict[str, Any],
) -> list[Finding]:
    expected_template_hash = placement_plan.get("input_hashes", {}).get("template_artifact")
    if not expected_template_hash:
        return [
            make_finding(
                1,
                "render",
                Status.UNKNOWN,
                "missing_input_artifact_hash",
                "Placement plan must bind the template artifact hash consumed by render",
                "input_hashes.template_artifact is present",
                "missing",
                root_cause_bucket="artifact_hash_gap",
            )
        ]

    actual_template_hash = sha256_json(template_artifact)
    if actual_template_hash != expected_template_hash:
        return [
            make_finding(
                1,
                "render",
                Status.FAIL,
                "artifact_hash_mismatch",
                "Render received a template artifact that does not match the placement plan provenance",
                expected_template_hash,
                actual_template_hash,
                root_cause_bucket="artifact_hash_mismatch",
            )
        ]
    return []


def _remove_slot_markers(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if "[[DOCFIT_SLOT:body]]" in paragraph.text:
            paragraph.text = ""


def _materialize_image_payload(payload: dict[str, Any], out_dir: Path, action_id: str) -> Path:
    source_docx = Path(str(payload["source_docx"]))
    target = str(payload["target"])
    filename = Path(str(payload.get("filename") or target)).name
    image_dir = out_dir / "assets" / "images"
    ensure_dir(image_dir)
    image_path = image_dir / f"{action_id}_{filename}"
    with ZipFile(source_docx) as package:
        image_path.write_bytes(package.read(target))
    return image_path


def build_feature_snapshot(
    final_docx: Path,
    placement_plan: dict[str, Any],
    render_manifest: dict[str, Any],
) -> dict[str, Any]:
    doc = Document(final_docx)
    paragraphs = [
        paragraph.text.strip()
        for paragraph in doc.paragraphs
        if paragraph.text.strip()
    ]
    tables = [
        [[cell.text.strip() for cell in row.cells] for row in table.rows]
        for table in doc.tables
    ]
    expected_hashes = [
        content_hash
        for action in placement_plan.get("data", {}).get("actions", [])
        if _action_writes_visible_output(action)
        for content_hash in action.get("content_hashes", [])
        if content_hash
    ]
    rendered_hashes = [sha256_text(text) for text in paragraphs]
    rendered_hashes.extend(sha256_json(table) for table in tables)
    rendered_hashes.extend(_media_hashes(final_docx))
    actual_text = "\n".join(paragraphs) + "\n" + repr(tables)
    return {
        "artifact_type": "feature_snapshot",
        "artifact_version": "1.0",
        "output_docx_sha256": sha256_file(final_docx),
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "headings": [
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.style is not None and paragraph.style.name.startswith("Heading")
        ],
        "content_hashes": rendered_hashes,
        "expected_content_hashes": expected_hashes,
        "rendered_text_hash": sha256_text(actual_text),
        "actions_executed": [
            item["action_id"]
            for item in render_manifest.get("actions_executed", [])
        ],
    }


def _action_writes_visible_output(action: dict[str, Any]) -> bool:
    return action.get("disposition") != "discard_as_source_format"


def _media_hashes(docx_path: Path) -> list[str]:
    with ZipFile(docx_path) as package:
        return [
            sha256_bytes(package.read(name))
            for name in sorted(package.namelist())
            if name.startswith("word/media/") and not name.endswith("/")
        ]


def build_feature_diff(feature_snapshot: dict[str, Any], golden_path: Path) -> dict[str, Any]:
    if not golden_path.exists():
        return {
            "status": Status.UNKNOWN.value,
            "diffs": [
                {
                    "kind": "unknown",
                    "message": "signed golden feature snapshot is missing",
                    "path": str(golden_path),
                }
            ],
        }
    golden = read_json(golden_path)
    missing_hashes = [
        content_hash
        for content_hash in golden.get("required_content_hashes", [])
        if content_hash not in feature_snapshot.get("content_hashes", [])
    ]
    if missing_hashes:
        return {
            "status": Status.FAIL.value,
            "diffs": [
                {
                    "kind": "blocking",
                    "message": "rendered feature snapshot is missing expected content hashes",
                    "missing": missing_hashes,
                }
            ],
        }
    return {"status": Status.PASS.value, "diffs": []}


def verify_render_outputs(
    placement_plan: dict[str, Any],
    render_manifest: dict[str, Any],
    feature_diff: dict[str, Any],
    oracle_report: dict[str, Any],
    golden_path: Path,
    *,
    require_golden: bool = True,
) -> list[Finding]:
    findings: list[Finding] = []
    next_index = 1
    if not oracle_report.get("valid_docx_package"):
        findings.append(
            make_finding(
                next_index,
                "render",
                Status.FAIL,
                "invalid_output_docx",
                "Rendered output must be a valid DOCX package",
                "valid DOCX package",
                "invalid package",
                root_cause_bucket="render_invalid_docx",
            )
        )
        next_index += 1
    expected_actions = {
        action["action_id"]
        for action in placement_plan.get("data", {}).get("actions", [])
    }
    executed_actions = {
        action["action_id"]
        for action in render_manifest.get("actions_executed", [])
    }
    missing_actions = sorted(expected_actions - executed_actions)
    if missing_actions:
        findings.append(
            make_finding(
                next_index,
                "render",
                Status.FAIL,
                "renderer_skipped_action",
                "Renderer must execute every placement action or fail explicitly",
                ", ".join(sorted(expected_actions)),
                ", ".join(sorted(executed_actions)),
                affected_ids=missing_actions,
                root_cause_bucket="render_plan_coverage_gap",
            )
        )
        next_index += 1
    if require_golden and not golden_path.exists():
        findings.append(
            make_finding(
                next_index,
                "render",
                Status.UNKNOWN,
                "missing_signed_golden",
                "Render verification requires a signed golden feature snapshot",
                str(golden_path),
                "missing",
                root_cause_bucket="golden_missing",
            )
        )
        next_index += 1
    for diff in feature_diff.get("diffs", []):
        if diff.get("kind") == "blocking":
            findings.append(
                make_finding(
                    next_index,
                    "render",
                    Status.FAIL,
                    "blocking_feature_diff",
                    "Rendered feature snapshot must match signed golden requirements",
                    "no blocking diff",
                    diff.get("message", "blocking diff"),
                    root_cause_bucket="render_feature_diff",
                )
            )
            next_index += 1
        elif diff.get("kind") == "unknown":
            findings.append(
                make_finding(
                    next_index,
                    "render",
                    Status.UNKNOWN,
                    "unknown_feature_diff",
                    "Feature diff must be decidable",
                    "decidable diff",
                    diff.get("message", "unknown diff"),
                    root_cause_bucket="render_feature_diff",
                )
            )
            next_index += 1
    return findings


def _renumber_findings(findings: list[Finding], *, start_index: int) -> list[Finding]:
    for offset, finding in enumerate(findings):
        finding.finding_id = f"f_{start_index + offset:03d}"
    return findings


def write_render_outputs(out_dir: Path, result: StageResult) -> None:
    for key in ["render_manifest", "feature_snapshot", "feature_diff", "oracle_report"]:
        artifact = result.artifacts.get(key)
        if artifact is not None:
            path = out_dir / "artifacts" / f"{key}.json"
            write_json(path, artifact)
            result.artifact_paths[key] = path
