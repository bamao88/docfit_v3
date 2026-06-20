from __future__ import annotations

from copy import deepcopy
import re
import shutil
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from docfit.core.io import now_iso, sha256_file, sha256_json, write_json
from docfit.core.models import StageResult, make_finding
from docfit.core.status import Status
from docfit.harness import template_units
from docfit.harness.generated_template_inspector import (
    inspect_generated_template_docx,
    iter_visible_text_entries,
)
from docfit.ooxml.package import is_valid_docx


DEFAULT_TEMPLATE_GENERATION_STRATEGY = "source_copy_scaffold"
BODY_SLOT_MARKER = "[[DOCFIT_SLOT:body]]"

UNIT_DEFINITIONS = (
    ("cover", "封面", ("封面", "题名", "论文题目", "学校", "学号", "指导教师")),
    ("integrity_statement", "诚信声明", ("诚信声明", "原创性声明", "授权书")),
    ("toc", "目录", ("目录", "目 录")),
    ("abstract_cn", "中文摘要", ("摘要", "摘 要", "关键词")),
    ("abstract_en", "英文摘要", ("abstract", "key words", "keywords")),
    ("body_main", "正文", ("正文", "绪论", "第一章", "1 ")),
    ("references", "参考文献", ("参考文献", "references")),
    ("acknowledgement", "致谢", ("致谢", "acknowledgement")),
    ("appendix", "附录", ("附录", "appendix")),
    ("post_forms", "后置固定表单", ("任务书", "开题", "评审", "答辩", "成绩评定")),
)

FILLABLE_MARKERS = ("××", "□□", "____", "——", "：", ":")
FILLABLE_LABELS = (
    "题名",
    "题目",
    "姓名",
    "学号",
    "学院",
    "专业",
    "班级",
    "教师",
    "日期",
    "摘要正文",
    "关键词",
)
MANUAL_ONLY_MARKERS = ("签名", "年月日", "年  月  日", "意见", "成绩", "评定")
GENERATED_MARKERS = ("目录", "页码", "编号", "图目录", "表目录", "公式")
INSTRUCTION_MARKERS = (
    "格式",
    "要求",
    "说明",
    "模板",
    "几号",
    "号字",
    "空一行",
    "倍行距",
    "页边距",
    "附件",
)


def generate_template(
    source_template_docx: Path,
    out_dir: Path,
    *,
    strategy: str = DEFAULT_TEMPLATE_GENERATION_STRATEGY,
    target_units: list[dict[str, Any]] | None = None,
) -> StageResult:
    if not source_template_docx.exists():
        return StageResult(
            "template_generate",
            Status.UNKNOWN,
            findings=[
                make_finding(
                    1,
                    "template_generate",
                    Status.UNKNOWN,
                    "template_generation_source_missing",
                    "缺少学校原始模板 Word，无法生成可填写模板",
                    "existing source template DOCX",
                    str(source_template_docx),
                    root_cause_bucket="input_missing",
                )
            ],
            coverage=_coverage(input_exists=False),
            blocked_at="template_generate",
        )
    if not is_valid_docx(source_template_docx):
        return StageResult(
            "template_generate",
            Status.FAIL,
            findings=[
                make_finding(
                    1,
                    "template_generate",
                    Status.FAIL,
                    "template_generation_source_invalid",
                    "学校原始模板不是有效 DOCX 包，无法生成可填写模板",
                    "valid source template DOCX",
                    str(source_template_docx),
                    evidence_refs=[str(source_template_docx)],
                    root_cause_bucket="input_invalid",
                )
            ],
            coverage=_coverage(input_exists=True, input_valid_docx=False),
            blocked_at="template_generate",
        )

    request = build_template_generation_request(
        source_template_docx,
        out_dir,
        strategy=strategy,
    )
    source_tree = inspect_source_template_docx(source_template_docx)
    discovered_rules = infer_template_rules(source_tree)
    if target_units:
        discovered_rules = align_target_units_to_source_tree(
            source_tree,
            target_units,
            discovered_rules=discovered_rules,
        )
    template_artifact = build_template_artifact(request, source_tree, discovered_rules)
    decisions = build_template_unit_decisions(template_artifact)
    plan = build_template_generation_plan(
        request,
        template_artifact=template_artifact,
        decisions=decisions,
    )
    generated_template_docx = out_dir / "generated_template.docx"
    execution = execute_template_generation_plan(
        source_template_docx,
        generated_template_docx,
        plan,
    )
    manifest = build_template_generation_manifest(
        request=request,
        source_tree=source_tree,
        discovered_rules=discovered_rules,
        template_artifact=template_artifact,
        decisions=decisions,
        plan=plan,
        generated_template_docx=generated_template_docx,
        execution=execution,
    )

    return StageResult(
        "template_generate",
        Status.PASS,
        artifacts={
            "template_generation_request": request,
            "source_template_tree": source_tree,
            "discovered_template_rules": discovered_rules,
            "template_artifact": template_artifact,
            "template_unit_decisions": decisions,
            "template_generation_plan": plan,
            "template_generation_manifest": manifest,
        },
        artifact_paths={"generated_template_docx": generated_template_docx},
        coverage=_coverage(
            input_exists=True,
            input_valid_docx=True,
            source_tree=bool(source_tree.get("layers", {}).get("body_flow")),
            discovered_rules=bool(discovered_rules.get("units")),
            template_artifact=bool(template_artifact.get("data", {}).get("units")),
            decisions=bool(decisions.get("units")),
            generation_plan=bool(plan.get("actions")),
            output_docx=generated_template_docx.exists(),
            manifest=True,
            body_slot=bool(manifest.get("slots")),
        ),
        user_message=(
            "template_generate produced a complete stage artifact chain; "
            "template quality still belongs to template-gap and real-core gates."
        ),
    )


def build_template_generation_request(
    source_template_docx: Path,
    out_dir: Path,
    *,
    strategy: str = DEFAULT_TEMPLATE_GENERATION_STRATEGY,
) -> dict[str, Any]:
    return {
        "artifact_type": "template_generation_request",
        "artifact_version": "1.0",
        "created_at": now_iso(),
        "source_template_docx": str(source_template_docx),
        "source_template_hash": sha256_file(source_template_docx),
        "out_dir": str(out_dir),
        "strategy": strategy,
        "optional_labels": {},
    }


def inspect_source_template_docx(source_template_docx: Path) -> dict[str, Any]:
    inspected = inspect_generated_template_docx(source_template_docx)
    body_flow = _body_flow_from_inspection(inspected)
    return {
        "artifact_type": "source_template_tree",
        "artifact_version": "1.0",
        "producer": {"name": "docfit-template-generate", "version": "0.2.0"},
        "created_at": now_iso(),
        "metadata": {
            "source_template_docx": str(source_template_docx),
            "source_template_hash": sha256_file(source_template_docx),
            "input_exists": inspected.get("input_exists"),
            "input_valid_docx": inspected.get("input_valid_docx"),
        },
        "layers": {
            "package_global": {
                "numbering_definitions": inspected.get("data", {}).get(
                    "numbering_definitions", []
                ),
            },
            "section_rules": inspected.get("data", {}).get("sections", []),
            "header_footer": inspected.get("data", {}).get("headers_footers", []),
            "body_flow": body_flow,
            "embedded_resources": [],
            "unknown_objects": inspected.get("data", {}).get(
                "unknown_visible_objects", []
            ),
        },
        "indexes": {
            "by_source_ref": {
                item.get("source_ref"): item.get("node_id")
                for item in body_flow
                if item.get("source_ref")
            },
            "body_order": [item.get("node_id") for item in body_flow],
        },
        "warnings": _source_tree_warnings(inspected),
        "data": inspected.get("data", {}),
    }


def infer_template_rules(source_tree: dict[str, Any]) -> dict[str, Any]:
    entries = _body_entries(source_tree)
    units = _infer_units(entries)
    return {
        "artifact_type": "discovered_template_rules",
        "artifact_version": "1.0",
        "producer": {"name": "docfit-template-generate", "version": "0.2.0"},
        "created_at": now_iso(),
        "source_template_hash": source_tree.get("metadata", {}).get(
            "source_template_hash"
        ),
        "discovery_method": "deterministic_keyword_and_structure_heuristics",
        "units": units,
        "unknowns": _rule_unknowns(source_tree, units),
    }


def align_target_units_to_source_tree(
    source_tree: dict[str, Any],
    target_units: list[dict[str, Any]],
    *,
    discovered_rules: dict[str, Any] | None = None,
) -> dict[str, Any]:
    entries = _body_entries(source_tree)
    aligned_units: list[dict[str, Any]] = []
    for index, target_unit in enumerate(target_units):
        aligned_units.append(
            _aligned_target_unit(
                target_unit,
                entries,
                fallback_order=(index + 1) * 10,
            )
        )
    return {
        "artifact_type": "discovered_template_rules",
        "artifact_version": "1.0",
        "producer": {"name": "docfit-template-generate", "version": "0.2.0"},
        "created_at": now_iso(),
        "source_template_hash": source_tree.get("metadata", {}).get(
            "source_template_hash"
        ),
        "discovery_method": "signed_target_units_aligned_to_source_tree",
        "units": aligned_units,
        "unknowns": [
            {
                "source_ref": item.get("source_ref"),
                "reason": item.get("reason", "unknown visible object"),
                "recommended_disposition": "preserve_or_review",
            }
            for item in source_tree.get("layers", {}).get("unknown_objects", [])
        ],
        "source_discovery": discovered_rules,
    }


def build_template_artifact(
    request: dict[str, Any],
    source_tree: dict[str, Any],
    discovered_rules: dict[str, Any],
) -> dict[str, Any]:
    units = discovered_rules.get("units", [])
    paragraphs = source_tree.get("data", {}).get("paragraphs", [])
    instruction_paragraphs = _dedupe_by_key(
        [
            *_instruction_paragraphs_from_units(units),
            *_instruction_paragraphs_from_source_tree(source_tree),
        ],
        "source_ref",
    )
    template_units.apply_instruction_policy(paragraphs, instruction_paragraphs)
    slots: list[dict[str, Any]] = []
    regions: list[dict[str, Any]] = []
    template_units.extend_slots_and_regions_from_units(slots, regions, units)
    if not any(slot.get("slot_id") == "slot_body_start" for slot in slots):
        slots.append(
            {
                "slot_id": "slot_body_start",
                "unit_id": "body_main",
                "element_id": "slot_body_start",
                "kind": "body_content",
                "writable": True,
                "required": True,
                "accepted_content_kinds": ["heading", "paragraph", "table", "image"],
                "source_ref": "template-generate:body-slot",
                "policy": "fill",
            }
        )
    if not any(region.get("region_id") == "body_main" for region in regions):
        regions.append(
            {
                "region_id": "body_main",
                "kind": "body",
                "required": True,
                "anchors": ["slot_body_start"],
                "source_ref": "template-generate:body-slot",
                "policy": "fill",
            }
        )
    return {
        "artifact_type": "template_artifact",
        "artifact_version": "1.0",
        "producer": {"name": "docfit-template-generate", "version": "0.2.0"},
        "created_at": now_iso(),
        "input_hashes": {
            "template_docx": request.get("source_template_hash"),
            "source_template_tree": sha256_json(source_tree),
            "discovered_template_rules": sha256_json(discovered_rules),
        },
        "provenance": {"template_docx": request.get("source_template_docx")},
        "status_notes": [
            "units are inferred from source Word structure and deterministic keywords",
            "formal quality still requires template-gap against accepted standards",
        ],
        "data": {
            "source_template_tree": "source_template_tree.json",
            "discovered_template_rules": "discovered_template_rules.json",
            "page_setup": {
                "sections": source_tree.get("layers", {}).get("section_rules", [])
            },
            "styles": _style_inventory(source_tree),
            "paragraphs": paragraphs,
            "units": units,
            "instruction_paragraphs": instruction_paragraphs,
            "regions": regions,
            "slots": slots,
            "protected_zones": template_units.protected_zones_from_units(units),
            "numbering": source_tree.get("data", {}).get("numbering_definitions", []),
            "headers_footers": source_tree.get("layers", {}).get("header_footer", []),
            "required_fields": template_units.required_fields_from_units(units),
            "unsupported": source_tree.get("layers", {}).get("unknown_objects", []),
        },
    }


def build_template_unit_decisions(template_artifact: dict[str, Any]) -> dict[str, Any]:
    units: list[dict[str, Any]] = []
    for unit in template_artifact.get("data", {}).get("units", []):
        unit_id = str(unit.get("unit_id"))
        unit_anchor_ref = _first_source_ref(unit)
        decisions: list[dict[str, Any]] = []
        if any(
            element.get("policy") in {"fixed", "manual_only"}
            for element in unit.get("elements", [])
        ):
            decisions.append(
                {
                    "decision_id": f"{unit_id}.copy_fixed_block",
                    "decision_type": "copy_fixed_block",
                    "unit_id": unit_id,
                    "element_id": None,
                    "source_ref": unit_anchor_ref,
                    "reason": "preserve fixed or manual-only source template content",
                }
            )
        for element in unit.get("elements", []):
            policy = element.get("policy")
            element_id = element.get("element_id")
            source_ref = _first_source_ref(element)
            if policy == "remove_instruction":
                decision_type = "remove_instruction_text"
            elif policy == "fill":
                decision_type = "create_fillable_slot"
            elif policy == "generated":
                decision_type = "create_generated_field_placeholder"
            elif policy in {"fixed", "manual_only"} and (
                source_ref is None
                and _should_synthesize_visible_text(element)
            ):
                decision_type = "insert_fixed_text"
                source_ref = unit_anchor_ref
            elif policy == "manual_only":
                decision_type = "create_manual_placeholder"
            else:
                continue
            decisions.append(
                {
                    "decision_id": f"{unit_id}.{element_id}.{decision_type}",
                    "decision_type": decision_type,
                    "unit_id": unit_id,
                    "element_id": element_id,
                    "element_name": element.get("name"),
                    "content": element.get("content") or element.get("name") or "",
                    "source_ref": source_ref,
                    "reason": _decision_reason(decision_type),
                }
            )
        units.append(
            {
                "unit_id": unit_id,
                "unit_name": unit.get("name"),
                "source_policy": unit.get("policy"),
                "generation_policy": "unit_actions",
                "decisions": decisions,
                "unresolved_questions": [],
            }
        )
    return {
        "artifact_type": "template_unit_decisions",
        "artifact_version": "1.0",
        "producer": {"name": "docfit-template-generate", "version": "0.2.0"},
        "created_at": now_iso(),
        "input_hashes": {"template_artifact": sha256_json(template_artifact)},
        "units": units,
    }


def build_template_generation_plan(
    request: dict[str, Any],
    *,
    template_artifact: dict[str, Any],
    decisions: dict[str, Any],
) -> dict[str, Any]:
    actions: list[dict[str, Any]] = [
        {
            "action_id": "a_001",
            "action_type": "copy_source_docx",
            "unit_id": None,
            "element_id": None,
            "source_ref": request.get("source_template_docx"),
            "target_ref": "generated_template.docx",
            "status": "planned",
            "reason": "create the generated Word from the source template package",
        }
    ]
    next_id = 2
    page_boundary_refs: set[str] = set()
    section_boundary_refs: set[str] = set()
    for index, unit in enumerate(template_artifact.get("data", {}).get("units", [])):
        if index == 0:
            continue
        page = unit.get("page") or {}
        page_break_rule = str(page.get("page_break") or "")
        source_ref = _first_source_ref(unit)
        if not source_ref:
            continue
        if _page_break_rule_requires_break(page_break_rule) and (
            source_ref not in page_boundary_refs
        ):
            actions.append(
                {
                    "action_id": f"a_{next_id:03d}",
                    "action_type": "insert_page_break_before_unit",
                    "unit_id": unit.get("unit_id"),
                    "element_id": None,
                    "source_ref": source_ref,
                    "target_ref": source_ref,
                    "status": "planned",
                    "reason": "unit page rule requires a deterministic page break before this unit",
                }
            )
            page_boundary_refs.add(source_ref)
            next_id += 1
        section_isolation_rule = str(page.get("section_isolation") or "")
        if _page_break_rule_requires_break(section_isolation_rule) and (
            source_ref not in section_boundary_refs
        ):
            actions.append(
                {
                    "action_id": f"a_{next_id:03d}",
                    "action_type": "insert_section_break_before_unit",
                    "unit_id": unit.get("unit_id"),
                    "element_id": None,
                    "source_ref": source_ref,
                    "target_ref": source_ref,
                    "status": "planned",
                    "reason": "unit page rule requires a deterministic section boundary before this unit",
                }
            )
            section_boundary_refs.add(source_ref)
            next_id += 1
    for action in _synthetic_unit_title_actions(template_artifact):
        action["action_id"] = f"a_{next_id:03d}"
        actions.append(action)
        next_id += 1
    for unit in decisions.get("units", []):
        for decision in unit.get("decisions", []):
            action_type = _action_type_for_decision(decision["decision_type"])
            actions.append(
                {
                    "action_id": f"a_{next_id:03d}",
                    "action_type": action_type,
                    "unit_id": decision.get("unit_id"),
                    "element_id": decision.get("element_id"),
                    "source_ref": decision.get("source_ref"),
                    "target_ref": _target_ref_for_decision(decision),
                    "status": "planned",
                    "reason": decision.get("reason"),
                }
            )
            next_id += 1
    planned_instruction_refs = {
        action.get("source_ref")
        for action in actions
        if action.get("action_type") == "remove_instruction_text"
    }
    for instruction in template_artifact.get("data", {}).get(
        "instruction_paragraphs",
        [],
    ):
        source_ref = instruction.get("source_ref")
        if not source_ref or source_ref in planned_instruction_refs:
            continue
        actions.append(
            {
                "action_id": f"a_{next_id:03d}",
                "action_type": "remove_instruction_text",
                "unit_id": "template_instructions",
                "element_id": None,
                "source_ref": source_ref,
                "target_ref": source_ref,
                "status": "planned",
                "reason": instruction.get(
                    "reason",
                    "detected template instruction text should not appear in the generated fillable template",
                ),
            }
        )
        planned_instruction_refs.add(source_ref)
        next_id += 1
    if not any(action.get("action_type") == "ensure_body_slot" for action in actions):
        actions.append(
            {
                "action_id": f"a_{next_id:03d}",
                "action_type": "ensure_body_slot",
                "unit_id": "body_main",
                "element_id": "slot_body_start",
                "source_ref": None,
                "target_ref": BODY_SLOT_MARKER,
                "status": "planned",
                "reason": "guarantee a stable write position for later content placement",
            }
        )
    return {
        "artifact_type": "template_generation_plan",
        "artifact_version": "1.0",
        "producer": {"name": "docfit-template-generate", "version": "0.2.0"},
        "created_at": now_iso(),
        "strategy": request.get("strategy"),
        "source_template_docx": request.get("source_template_docx"),
        "input_hashes": {
            "source_template_docx": request.get("source_template_hash"),
            "template_artifact": sha256_json(template_artifact),
            "template_unit_decisions": sha256_json(decisions),
        },
        "actions": actions,
    }


def _synthetic_unit_title_actions(template_artifact: dict[str, Any]) -> list[dict[str, Any]]:
    units = template_artifact.get("data", {}).get("units", [])
    actions: list[dict[str, Any]] = []
    for index, unit in enumerate(units):
        if str(unit.get("unit_id") or "") != "toc":
            continue
        title = _generated_unit_title_text(unit)
        if not title:
            continue
        source_ref = _first_source_ref(unit)
        next_ref = _next_unit_source_ref(units, index)
        if not next_ref:
            continue
        source_order = _paragraph_index(source_ref)
        previous_order = _paragraph_index(_previous_unit_source_ref(units, index))
        next_order = _paragraph_index(next_ref)
        if next_order is None:
            continue
        source_is_in_expected_window = (
            source_order is not None
            and (previous_order is None or source_order > previous_order)
            and source_order < next_order
        )
        if source_is_in_expected_window:
            continue
        actions.append(
            {
                "action_id": "",
                "action_type": "insert_synthetic_unit_title_before",
                "unit_id": unit.get("unit_id"),
                "element_id": "e_001",
                "source_ref": next_ref,
                "target_ref": title,
                "status": "planned",
                "reason": (
                    "toc has no visible title in the expected unit position; "
                    "insert a generated title before the next unit"
                ),
            }
        )
    return actions


def _generated_unit_title_text(unit: dict[str, Any]) -> str | None:
    for element in unit.get("elements", []):
        if str(element.get("element_id") or "") != "e_001":
            continue
        if str(element.get("policy") or "") != "generated":
            continue
        content = _normalize_text(str(element.get("content") or element.get("name") or ""))
        if content and len(content) <= 12:
            return content
    return None


def _previous_unit_source_ref(units: list[dict[str, Any]], index: int) -> str | None:
    for unit in reversed(units[:index]):
        ref = _first_source_ref(unit)
        if ref:
            return ref
    return None


def _next_unit_source_ref(units: list[dict[str, Any]], index: int) -> str | None:
    for unit in units[index + 1 :]:
        ref = _first_source_ref(unit)
        if ref:
            return ref
    return None


def execute_template_generation_plan(
    source_template_docx: Path,
    generated_template_docx: Path,
    plan: dict[str, Any],
) -> dict[str, Any]:
    generated_template_docx.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source_template_docx, generated_template_docx)
    doc = Document(generated_template_docx)
    paragraph_map = {
        index: paragraph for index, paragraph in enumerate(doc.paragraphs, start=1)
    }
    executed: list[dict[str, Any]] = []
    review: list[dict[str, Any]] = []
    slots: list[dict[str, Any]] = []
    generated_fields: list[dict[str, Any]] = []
    paragraphs_to_remove: list[Paragraph] = []

    for action in plan.get("actions", []):
        action_type = action.get("action_type")
        if action_type == "copy_source_docx":
            executed.append(_executed(action, output_ref=str(generated_template_docx)))
        elif action_type == "insert_page_break_before_unit":
            output_ref = _insert_page_break_before(
                doc,
                paragraph_map,
                action.get("source_ref"),
            )
            if output_ref is None:
                review.append(_needs_review(action, "source node not found"))
                continue
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "insert_section_break_before_unit":
            output_ref = _insert_section_break_before(
                doc,
                paragraph_map,
                action.get("source_ref"),
            )
            if output_ref is None:
                review.append(_needs_review(action, "source node not found"))
                continue
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "remove_instruction_text":
            target = _paragraph_for_ref(paragraph_map, action.get("source_ref"))
            target_cell = _cell_for_ref(doc, action.get("source_ref"))
            if target is None and target_cell is None:
                review.append(_needs_review(action, "source node not found"))
                continue
            if target is not None:
                paragraphs_to_remove.append(target)
            if target_cell is not None:
                _clear_cell(target_cell)
            executed.append(_executed(action, output_ref=action.get("source_ref")))
        elif action_type == "create_fillable_slot":
            marker = _slot_marker(action)
            output_ref = _insert_marker(doc, paragraph_map, action.get("source_ref"), marker)
            slot = _slot_from_action(action, marker=marker, output_ref=output_ref)
            slots.append(slot)
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "create_generated_field_placeholder":
            marker = _generated_marker(action)
            output_ref = _insert_marker(doc, paragraph_map, action.get("source_ref"), marker)
            generated_fields.append(
                {
                    "field_id": marker.strip("[]"),
                    "unit_id": action.get("unit_id"),
                    "element_id": action.get("element_id"),
                    "marker": marker,
                    "output_ref": output_ref,
                }
            )
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "create_manual_placeholder":
            executed.append(_executed(action, output_ref=action.get("source_ref")))
        elif action_type == "insert_fixed_text":
            text = str(action.get("target_ref") or "")
            output_ref = _insert_marker(
                doc,
                paragraph_map,
                action.get("source_ref"),
                text,
            )
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "insert_synthetic_unit_title_before":
            text = str(action.get("target_ref") or "")
            output_ref = _insert_styled_paragraph_before(
                paragraph_map,
                action.get("source_ref"),
                text,
                page_break_before=True,
            )
            if output_ref is None:
                review.append(_needs_review(action, "source node not found"))
                continue
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "protect_block":
            executed.append(_executed(action, output_ref=action.get("source_ref")))
        elif action_type == "ensure_body_slot":
            marker = BODY_SLOT_MARKER
            existing_ref = _find_marker_ref(doc, marker)
            output_ref = existing_ref or _append_marker(doc, marker)
            slots.append(
                {
                    "slot_id": "slot_body_start",
                    "unit_id": "body_main",
                    "element_id": "slot_body_start",
                    "kind": "body_content",
                    "marker": marker,
                    "output_ref": output_ref,
                    "required": True,
                }
            )
            executed.append(_executed(action, output_ref=output_ref))
        else:
            review.append(_needs_review(action, f"unsupported action type: {action_type}"))

    for paragraph in dict.fromkeys(paragraphs_to_remove):
        _remove_paragraph(paragraph)
    doc.save(generated_template_docx)
    return {
        "actions_executed": executed,
        "actions_requiring_review": review,
        "slots": _dedupe_by_key(slots, "slot_id"),
        "generated_fields": generated_fields,
        "page_breaks": [
            {
                "unit_id": action.get("unit_id"),
                "source_ref": action.get("source_ref"),
                "output_ref": action.get("output_ref"),
            }
            for action in executed
            if action.get("action_type") == "insert_page_break_before_unit"
        ],
        "section_breaks": [
            {
                "unit_id": action.get("unit_id"),
                "source_ref": action.get("source_ref"),
                "output_ref": action.get("output_ref"),
            }
            for action in executed
            if action.get("action_type") == "insert_section_break_before_unit"
        ],
        "synthesized_texts": [
            {
                "unit_id": action.get("unit_id"),
                "element_id": action.get("element_id"),
                "text": action.get("target_ref"),
                "output_ref": action.get("output_ref"),
            }
            for action in executed
            if action.get("action_type")
            in {"insert_fixed_text", "insert_synthetic_unit_title_before"}
        ],
    }


def build_template_generation_manifest(
    *,
    request: dict[str, Any],
    source_tree: dict[str, Any],
    discovered_rules: dict[str, Any],
    template_artifact: dict[str, Any],
    decisions: dict[str, Any],
    plan: dict[str, Any],
    generated_template_docx: Path,
    execution: dict[str, Any],
) -> dict[str, Any]:
    return {
        "artifact_type": "template_generation_manifest",
        "artifact_version": "1.0",
        "producer": {"name": "docfit-template-generate", "version": "0.2.0"},
        "created_at": now_iso(),
        "strategy": request.get("strategy"),
        "input_hashes": {
            "source_template_docx": request.get("source_template_hash"),
            "source_template_tree": sha256_json(source_tree),
            "discovered_template_rules": sha256_json(discovered_rules),
            "template_artifact": sha256_json(template_artifact),
            "template_unit_decisions": sha256_json(decisions),
            "template_generation_plan": sha256_json(plan),
        },
        "output": {
            "generated_template_docx": str(generated_template_docx),
            "generated_template_docx_hash": sha256_file(generated_template_docx),
        },
        "slots": execution.get("slots", []),
        "generated_fields": execution.get("generated_fields", []),
        "page_breaks": execution.get("page_breaks", []),
        "section_breaks": execution.get("section_breaks", []),
        "synthesized_texts": execution.get("synthesized_texts", []),
        "actions_executed": execution.get("actions_executed", []),
        "actions_requiring_review": execution.get("actions_requiring_review", []),
    }


def write_template_generation_outputs(out_dir: Path, result: StageResult) -> None:
    for key in [
        "template_generation_request",
        "source_template_tree",
        "discovered_template_rules",
        "template_artifact",
        "template_unit_decisions",
        "template_generation_plan",
        "template_generation_manifest",
    ]:
        artifact = result.artifacts.get(key)
        if artifact is None:
            continue
        path = out_dir / "artifacts" / f"{key}.json"
        write_json(path, artifact)
        result.artifact_paths[key] = path


def _body_flow_from_inspection(tree: dict[str, Any]) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for entry in iter_visible_text_entries(tree):
        order = entry.get("order") or 0
        node_id = f"body_{len(items) + 1:04d}"
        items.append(
            {
                "node_id": node_id,
                "structure_layer": "header_footer"
                if entry.get("kind") in {"header", "footer"}
                else "body_flow",
                "flow_item_type": entry.get("kind"),
                "kind": entry.get("kind"),
                "source_ref": entry.get("source_ref"),
                "part_name": _part_name(entry.get("source_ref")),
                "order": order,
                "parent_ref": None,
                "container_ref": entry.get("table_source_ref"),
                "visible": bool(entry.get("text")),
                "text": entry.get("text", ""),
                "style": entry.get("style", ""),
                "style_details": entry.get("style_details", {}),
                "structural_signals": _structural_signals(entry),
            }
        )
    return sorted(items, key=lambda item: (float(item.get("order") or 0), item["node_id"]))


def _source_tree_warnings(inspected: dict[str, Any]) -> list[dict[str, Any]]:
    warnings: list[dict[str, Any]] = []
    for item in inspected.get("data", {}).get("unknown_visible_objects", []):
        warnings.append(
            {
                "code": "unknown_visible_object",
                "message": item.get("reason", "unknown visible object"),
                "source_ref": item.get("source_ref"),
                "severity": "review",
            }
        )
    return warnings


def _body_entries(source_tree: dict[str, Any]) -> list[dict[str, Any]]:
    entries = []
    for item in source_tree.get("layers", {}).get("body_flow", []):
        if item.get("structure_layer") != "body_flow":
            continue
        if item.get("text"):
            entries.append(item)
    return entries


def _infer_units(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not entries:
        return [
            {
                "unit_id": "body_main",
                "name": "正文",
                "order": 10,
                "status": "required",
                "policy": "fill",
                "source_refs": [],
                "elements": [],
            }
        ]
    anchors = _unit_anchors(entries)
    units: list[dict[str, Any]] = []
    for anchor_index, anchor in enumerate(anchors):
        next_start = (
            anchors[anchor_index + 1]["entry_index"]
            if anchor_index + 1 < len(anchors)
            else len(entries)
        )
        region_entries = entries[anchor["entry_index"] : next_start]
        units.append(
            {
                "unit_id": anchor["unit_id"],
                "name": anchor["name"],
                "order": (anchor_index + 1) * 10,
                "status": "required",
                "policy": _unit_policy(anchor["unit_id"]),
                "source_refs": [anchor["source_ref"]],
                "page": {},
                "elements": _infer_elements(anchor, region_entries),
            }
        )
    return units


def _unit_anchors(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    anchors: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, entry in enumerate(entries):
        unit_id, name = _unit_for_text(str(entry.get("text", "")), index)
        if unit_id is None or unit_id in seen:
            continue
        anchors.append(
            {
                "entry_index": index,
                "unit_id": unit_id,
                "name": name,
                "source_ref": entry.get("source_ref"),
                "text": entry.get("text", ""),
            }
        )
        seen.add(unit_id)
    if not anchors or anchors[0]["entry_index"] != 0:
        anchors.insert(
            0,
            {
                "entry_index": 0,
                "unit_id": "cover",
                "name": "封面",
                "source_ref": entries[0].get("source_ref"),
                "text": entries[0].get("text", ""),
            },
        )
    if "body_main" not in {anchor["unit_id"] for anchor in anchors}:
        body_index = _first_body_like_index(entries)
        anchors.append(
            {
                "entry_index": body_index,
                "unit_id": "body_main",
                "name": "正文",
                "source_ref": entries[body_index].get("source_ref"),
                "text": entries[body_index].get("text", ""),
            }
        )
    return sorted(
        _dedupe_anchors(anchors),
        key=lambda item: (int(item["entry_index"]), item["unit_id"]),
    )


def _infer_elements(anchor: dict[str, Any], entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    elements: list[dict[str, Any]] = []
    for entry in entries:
        text = str(entry.get("text", "")).strip()
        if not text:
            continue
        policy = _element_policy(anchor["unit_id"], text, entry)
        elements.append(
            {
                "element_id": f"e_{len(elements) + 1:03d}",
                "name": _element_name(anchor["unit_id"], text, policy),
                "order": len(elements) + 1,
                "policy": policy,
                "type": _element_type(policy),
                "fill": "yes" if policy == "fill" else "no",
                "content": text if policy != "remove_instruction" else "",
                "style": _style_summary(entry),
                "position": entry.get("source_ref", ""),
                "relationship": "",
                "source_refs": [entry.get("source_ref", "")],
            }
        )
    if not elements:
        elements.append(
            {
                "element_id": "e_001",
                "name": anchor["name"],
                "order": 1,
                "policy": "fill" if anchor["unit_id"] == "body_main" else "fixed",
                "content": anchor.get("text", ""),
                "style": "",
                "source_refs": [anchor.get("source_ref", "")],
            }
        )
    return elements


def _aligned_target_unit(
    target_unit: dict[str, Any],
    entries: list[dict[str, Any]],
    *,
    fallback_order: int,
) -> dict[str, Any]:
    aligned_elements: list[dict[str, Any]] = []
    source_refs: list[str] = []
    unit_anchor = _find_source_entry_for_target_unit(entries, target_unit)
    if unit_anchor and unit_anchor.get("source_ref"):
        source_refs.append(str(unit_anchor["source_ref"]))
    for fallback_element_order, element in enumerate(
        target_unit.get("elements", []),
        start=1,
    ):
        match = _find_source_entry_for_target_element(entries, element)
        policy = _normalize_target_policy(str(element.get("policy") or ""))
        refs = [match["source_ref"]] if match and match.get("source_ref") else []
        source_refs.extend(refs)
        aligned = {
            **element,
            "element_id": str(
                element.get("element_id") or f"e_{fallback_element_order:03d}"
            ),
            "order": element.get("order")
            or element.get("element_order")
            or fallback_element_order,
            "policy": policy,
            "source_refs": refs,
        }
        if match is not None:
            aligned["source_text"] = match.get("text", "")
            aligned["position"] = match.get("source_ref", "")
        aligned_elements.append(aligned)
    unit_source_refs = _dedupe(source_refs) or list(target_unit.get("source_refs", []))
    return {
        **target_unit,
        "unit_id": str(target_unit.get("unit_id") or "unknown_unit"),
        "name": target_unit.get("name") or target_unit.get("unit_id") or "未知单元",
        "order": target_unit.get("order") or fallback_order,
        "status": target_unit.get("status") or "required",
        "policy": target_unit.get("policy")
        or _unit_policy(str(target_unit.get("unit_id") or "")),
        "source_refs": unit_source_refs,
        "elements": aligned_elements,
    }


def _find_source_entry_for_target_unit(
    entries: list[dict[str, Any]],
    target_unit: dict[str, Any],
) -> dict[str, Any] | None:
    if str(target_unit.get("unit_id") or "") == "body_main":
        body_anchor = _find_body_main_source_entry(entries)
        if body_anchor is not None:
            return body_anchor
    query = _target_unit_query(target_unit)
    if not _query_has_needles(query):
        return None
    return _find_entry_by_query(entries, query)


def _target_unit_query(target_unit: dict[str, Any]) -> dict[str, Any]:
    text = _normalize_text(str(target_unit.get("name") or target_unit.get("unit_id") or ""))
    if not text or _looks_like_descriptor(text):
        return {"full": [], "tokens": [], "min_tokens": 0}
    normalized = _normalize_for_match(text)
    tokens = _dedupe(
        [
            _normalize_for_match(token)
            for token in _split_match_tokens(text)
            if _normalize_for_match(token)
        ]
    )
    return {
        "full": [normalized] if normalized else [],
        "tokens": tokens,
        "min_tokens": 1 if tokens else 0,
    }


def _find_source_entry_for_target_element(
    entries: list[dict[str, Any]],
    element: dict[str, Any],
) -> dict[str, Any] | None:
    query = _target_element_query(element)
    if not _query_has_needles(query):
        return None
    return _find_first_entry_by_query(entries, query)


def _target_element_query(element: dict[str, Any]) -> dict[str, Any]:
    full: list[str] = []
    tokens: list[str] = []
    for value in (
        element.get("content"),
        element.get("name"),
        element.get("raw"),
    ):
        text = _normalize_text(str(value or ""))
        if not text or _looks_like_descriptor(text):
            continue
        full.append(text)
        tokens.extend(_split_match_tokens(text))
    normalized_tokens = _dedupe(
        [_normalize_for_match(token) for token in tokens if _normalize_for_match(token)]
    )
    return {
        "full": _dedupe(
            [
                _normalize_for_match(candidate)
                for candidate in full
                if _normalize_for_match(candidate)
            ]
        ),
        "tokens": normalized_tokens,
        "min_tokens": min(2, len(normalized_tokens)) if normalized_tokens else 0,
    }


def _find_entry_by_query(
    entries: list[dict[str, Any]],
    query: dict[str, Any],
) -> dict[str, Any] | None:
    candidates: list[tuple[int, int, dict[str, Any]]] = []
    for entry in entries:
        text = _normalize_for_match(entry.get("text", ""))
        if not text:
            continue
        score = _query_entry_score(entry, query, text)
        if score <= 0:
            continue
        candidates.append((score, -int(entry.get("order") or 0), entry))
    if not candidates:
        return None
    return max(candidates, key=lambda item: (item[0], item[1]))[2]


def _find_first_entry_by_query(
    entries: list[dict[str, Any]],
    query: dict[str, Any],
) -> dict[str, Any] | None:
    for entry in entries:
        text = _normalize_for_match(entry.get("text", ""))
        if not text:
            continue
        if any(needle and needle in text for needle in query.get("full", [])):
            return entry
        tokens = [token for token in query.get("tokens", []) if token and token in text]
        if tokens and len(tokens) >= int(query.get("min_tokens") or 1):
            return entry
    return None


def _query_entry_score(
    entry: dict[str, Any],
    query: dict[str, Any],
    normalized_text: str,
) -> int:
    score = 0
    for needle in query.get("full", []):
        if not needle:
            continue
        if normalized_text == needle:
            score = max(score, 120)
        elif _allow_partial_query_match(needle, normalized_text):
            score = max(score, 70)
    tokens = [
        token
        for token in query.get("tokens", [])
        if token and _allow_partial_query_match(token, normalized_text)
    ]
    if tokens and len(tokens) >= int(query.get("min_tokens") or 1):
        score = max(score, 35 + 8 * len(tokens))
    if score <= 0:
        return 0
    text = str(entry.get("text") or "").strip()
    style = str(entry.get("style") or "").lower()
    signals = entry.get("structural_signals") or {}
    if signals.get("short_text"):
        score += 30
    if signals.get("centered"):
        score += 20
    if signals.get("large_font"):
        score += 15
    if "heading" in style or "标题" in style:
        score += 20
    if "正文前标题" in style or "正文尾标题" in style:
        score += 25
    if signals.get("looks_like_instruction_text"):
        score -= 90
    if len(text) > 80:
        score -= 35
    if len(text) > 140:
        score -= 45
    if "\t" in text:
        score -= 20
    return score


def _allow_partial_query_match(needle: str, normalized_text: str) -> bool:
    if needle == normalized_text:
        return True
    if len(needle) <= 2 and len(normalized_text) <= 6:
        return False
    return needle in normalized_text


def _find_body_main_source_entry(
    entries: list[dict[str, Any]],
) -> dict[str, Any] | None:
    candidates: list[tuple[int, int, dict[str, Any]]] = []
    for entry in entries:
        text = str(entry.get("text") or "").strip()
        normalized = _normalize_for_match(text)
        style = str(entry.get("style") or "").lower()
        signals = entry.get("structural_signals") or {}
        chapter_heading = _looks_like_body_chapter_heading(normalized)
        if _body_main_anchor_excluded(text) and not chapter_heading:
            continue
        score = 0
        if style in {"heading 1", "标题 1"} or "heading 1" in style:
            score += 100
        if chapter_heading:
            score += 80
        if score <= 0:
            continue
        if signals.get("short_text"):
            score += 20
        if _has_placeholder_chapter_number(text):
            score -= 70
        if signals.get("looks_like_instruction_text") and not chapter_heading:
            score -= 80
        elif signals.get("looks_like_instruction_text"):
            score -= 15
        if len(text) > 60:
            score -= 40
        if score <= 0:
            continue
        candidates.append((score, -int(entry.get("order") or 0), entry))
    if not candidates:
        return None
    return max(candidates, key=lambda item: (item[0], item[1]))[2]


def _looks_like_body_chapter_heading(normalized_text: str) -> bool:
    return bool(re.match(r"^第[一二三四五六七八九十0-9]+章", normalized_text))


def _has_placeholder_chapter_number(text: str) -> bool:
    return bool(re.search(r"第\s*[Xx]\s*章", text))


def _body_main_anchor_excluded(text: str) -> bool:
    normalized = _normalize_for_match(text)
    if not normalized:
        return True
    excluded = (
        "目录",
        "摘要",
        "abstract",
        "参考文献",
        "致谢",
        "附录",
        "声明",
        "封面",
        "图目录",
        "表目录",
        "正文基本格式",
        "正文标题",
        "格式",
        "说明",
        "黑体",
        "三号",
        "第x章",
    )
    return any(marker in normalized for marker in excluded)


def _query_has_needles(query: dict[str, Any]) -> bool:
    return bool(query.get("full") or query.get("tokens"))


def _normalize_target_policy(policy: str) -> str:
    return {
        "fillable": "fill",
        "fill": "fill",
        "fixed": "fixed",
        "manual_only": "manual_only",
        "generated": "generated",
        "template_default_optional": "template_default_optional",
    }.get(policy, policy or "fixed")


def _rule_unknowns(source_tree: dict[str, Any], units: list[dict[str, Any]]) -> list[dict[str, Any]]:
    unknowns = [
        {
            "source_ref": item.get("source_ref"),
            "reason": item.get("reason", "unknown visible object"),
            "recommended_disposition": "preserve_or_review",
        }
        for item in source_tree.get("layers", {}).get("unknown_objects", [])
    ]
    if len(units) <= 1:
        unknowns.append(
            {
                "source_ref": None,
                "reason": "template unit discovery found one or fewer units",
                "recommended_disposition": "review_template_rule_discovery",
            }
        )
    return unknowns


def _instruction_paragraphs_from_units(units: list[dict[str, Any]]) -> list[dict[str, Any]]:
    paragraphs: list[dict[str, Any]] = []
    for unit in units:
        for element in unit.get("elements", []):
            if element.get("policy") != "remove_instruction":
                continue
            source_ref = _first_source_ref(element)
            paragraph_index = _paragraph_index(source_ref)
            if paragraph_index is None:
                continue
            paragraphs.append(
                {
                    "source_ref": source_ref,
                    "paragraph_index": paragraph_index,
                    "text": element.get("content") or element.get("name", ""),
                    "policy": "strip",
                    "final_disposition": "omit_from_final",
                    "reason": "detected template instruction text should not appear in the generated fillable template",
                }
            )
    return paragraphs


def _instruction_paragraphs_from_source_tree(source_tree: dict[str, Any]) -> list[dict[str, Any]]:
    paragraphs: list[dict[str, Any]] = []
    for entry in _body_entries(source_tree):
        text = str(entry.get("text", ""))
        if not _looks_like_instruction(text):
            continue
        source_ref = entry.get("source_ref")
        if not source_ref:
            continue
        paragraphs.append(
            {
                "source_ref": source_ref,
                "paragraph_index": _paragraph_index(source_ref),
                "text": text,
                "policy": "strip",
                "final_disposition": "omit_from_final",
                "reason": "detected template instruction text should not appear in the generated fillable template",
            }
        )
    return paragraphs


def _style_inventory(source_tree: dict[str, Any]) -> list[dict[str, Any]]:
    styles: dict[str, dict[str, Any]] = {}
    for paragraph in source_tree.get("data", {}).get("paragraphs", []):
        style = paragraph.get("style")
        if not style:
            continue
        styles.setdefault(str(style), {"name": style, "count": 0})
        styles[str(style)]["count"] += 1
    return list(styles.values())


def _decision_reason(decision_type: str) -> str:
    return {
        "remove_instruction_text": "instruction/example text should not enter the fillable template",
        "create_fillable_slot": "fillable source element needs a stable marker for later placement",
        "create_generated_field_placeholder": "generated element needs a marker for later field generation",
        "create_manual_placeholder": "manual-only content is preserved but not automatically filled",
        "insert_fixed_text": "visible standard text is missing from the aligned source region and should be present in the generated template",
    }.get(decision_type, "template generation decision")


def _action_type_for_decision(decision_type: str) -> str:
    return {
        "copy_fixed_block": "protect_block",
        "remove_instruction_text": "remove_instruction_text",
        "create_fillable_slot": "create_fillable_slot",
        "create_generated_field_placeholder": "create_generated_field_placeholder",
        "create_manual_placeholder": "create_manual_placeholder",
        "insert_fixed_text": "insert_fixed_text",
    }[decision_type]


def _target_ref_for_decision(decision: dict[str, Any]) -> str:
    decision_type = decision.get("decision_type")
    if decision_type == "create_fillable_slot":
        return f"[[DOCFIT_SLOT:{decision.get('unit_id')}.{decision.get('element_id')}]]"
    if decision_type == "create_generated_field_placeholder":
        return f"[[DOCFIT_GENERATED:{decision.get('unit_id')}.{decision.get('element_id')}]]"
    if decision_type == "insert_fixed_text":
        return str(decision.get("content") or "")
    return str(decision.get("source_ref") or "")


def _unit_for_text(text: str, index: int) -> tuple[str | None, str | None]:
    normalized = _normalize_text(text)
    for unit_id, name, needles in UNIT_DEFINITIONS:
        if unit_id == "cover" and index > 12:
            continue
        for needle in needles:
            if _normalize_text(needle) in normalized:
                return unit_id, name
    return None, None


def _unit_policy(unit_id: str) -> str:
    if unit_id in {"body_main", "abstract_cn", "abstract_en"}:
        return "fill"
    if unit_id in {"toc"}:
        return "generated"
    return "fixed"


def _element_policy(unit_id: str, text: str, entry: dict[str, Any]) -> str:
    lowered = text.lower()
    if _looks_like_instruction(text):
        return "remove_instruction"
    if unit_id == "toc" or any(marker.lower() in lowered for marker in GENERATED_MARKERS):
        return "generated"
    if any(marker in text for marker in MANUAL_ONLY_MARKERS):
        return "manual_only"
    if any(marker in text for marker in FILLABLE_MARKERS) and any(
        label in text for label in FILLABLE_LABELS
    ):
        return "fill"
    if unit_id in {"abstract_cn", "abstract_en", "body_main"} and not _looks_like_heading(entry):
        return "fill"
    return "fixed"


def _looks_like_instruction(text: str) -> bool:
    if _has_substantive_template_text(text) and re.search(
        r"[（(].*(宋体|黑体|楷体|居中|行距|字号|号字|pt).*[）)]",
        text,
    ):
        return False
    if template_units.contains_instruction_marker(text):
        return True
    if any(marker in text for marker in INSTRUCTION_MARKERS):
        return True
    return bool(re.search(r"[（(].*(宋体|黑体|楷体|居中|行距|字号|号字|pt).*[）)]", text))


def _page_break_rule_requires_break(rule: str) -> bool:
    normalized = _normalize_text(rule)
    return normalized == "是" or normalized.startswith("是；")


def _should_synthesize_visible_text(element: dict[str, Any]) -> bool:
    order = element.get("order") or element.get("element_order")
    if order not in {1, "1"}:
        return False
    text = str(element.get("content") or element.get("name") or "").strip()
    if not text:
        return False
    if len(text) > 120:
        return False
    normalized = _normalize_for_match(text)
    if not normalized or _looks_like_nonvisible_requirement(normalized):
        return False
    return True


def _looks_like_nonvisible_requirement(normalized: str) -> bool:
    markers = (
        "页眉",
        "页脚",
        "页码",
        "页边距",
        "装订线",
        "纸张",
        "section",
        "schoolyaml",
        "ooxml",
        "审查口径",
        "全局规则",
        "源模板",
        "源文件",
        "当前阶段",
        "目标输出",
        "生成机制",
        "标题编号体系",
        "样式",
        "字体",
        "字号",
        "行距",
        "大纲级别",
        "保留学校封面本体",
        "不属于模板的说明文字",
        "markdown",
        "自动化测试",
        "渲染测试",
        "验收",
    )
    return any(marker in normalized for marker in markers)


def _has_substantive_template_text(text: str) -> bool:
    cleaned = _normalize_for_match(text)
    if not cleaned:
        return False
    if cleaned in {
        "目录",
        "摘要",
        "abstract",
        "keywords",
        "keyword",
        "论文题目",
        "毕业论文设计中文题目",
        "titleofgraduationpaper",
    }:
        return True
    if any(marker in cleaned for marker in ("目录", "摘要", "论文", "题目")):
        return len(cleaned) <= 24
    return False


def _element_name(unit_id: str, text: str, policy: str) -> str:
    if policy == "remove_instruction":
        return "模板说明文字"
    if policy == "generated":
        return "系统生成占位"
    if policy == "fill":
        for label in FILLABLE_LABELS:
            if label in text:
                return label
        return "可填写内容"
    if policy == "manual_only":
        return "人工填写位置"
    if len(text) <= 24:
        return text
    return f"{UNIT_DEFINITION_NAMES.get(unit_id, unit_id)}固定内容"


UNIT_DEFINITION_NAMES = {unit_id: name for unit_id, name, _ in UNIT_DEFINITIONS}


def _element_type(policy: str) -> str:
    return {
        "fill": "fillable",
        "generated": "generated",
        "manual_only": "manual_only",
        "remove_instruction": "instruction_text",
    }.get(policy, "fixed_text")


def _style_summary(entry: dict[str, Any]) -> str:
    details = entry.get("style_details") or {}
    dominant = details.get("dominant_run") or {}
    paragraph = details.get("paragraph") or {}
    parts = []
    if dominant.get("font_names"):
        parts.append("/".join(str(name) for name in dominant["font_names"]))
    if dominant.get("font_size_pt"):
        parts.append(f"{dominant['font_size_pt']}pt")
    if dominant.get("bold"):
        parts.append("加粗")
    if paragraph.get("alignment"):
        parts.append(str(paragraph["alignment"]))
    return "；".join(parts)


def _structural_signals(entry: dict[str, Any]) -> dict[str, Any]:
    text = str(entry.get("text", ""))
    details = entry.get("style_details") or {}
    paragraph = details.get("paragraph") or {}
    dominant = details.get("dominant_run") or {}
    return {
        "centered": paragraph.get("alignment") == "center",
        "short_text": len(text.strip()) <= 20,
        "large_font": (dominant.get("font_size_pt") or 0) >= 16,
        "bold": bool(dominant.get("bold")),
        "looks_like_instruction_text": _looks_like_instruction(text),
        "likely_unit_heading": _looks_like_heading(entry),
    }


def _looks_like_heading(entry: dict[str, Any]) -> bool:
    text = str(entry.get("text", ""))
    unit_id, _ = _unit_for_text(text, int(float(entry.get("order") or 0)))
    details = entry.get("style_details") or {}
    paragraph = details.get("paragraph") or {}
    dominant = details.get("dominant_run") or {}
    short_text = len(text.strip()) <= 20
    centered = paragraph.get("alignment") == "center"
    large_font = (dominant.get("font_size_pt") or 0) >= 16
    return bool(unit_id) or (
        short_text
        and (centered or large_font)
    )


def _first_body_like_index(entries: list[dict[str, Any]]) -> int:
    for index, entry in enumerate(entries):
        unit_id, _ = _unit_for_text(str(entry.get("text", "")), index)
        if unit_id == "body_main":
            return index
    return max(len(entries) - 1, 0)


def _dedupe_anchors(anchors: list[dict[str, Any]]) -> list[dict[str, Any]]:
    deduped: list[dict[str, Any]] = []
    seen: set[str] = set()
    for anchor in anchors:
        unit_id = str(anchor.get("unit_id"))
        if unit_id in seen:
            continue
        deduped.append(anchor)
        seen.add(unit_id)
    return deduped


def _first_source_ref(item: dict[str, Any]) -> str | None:
    refs = item.get("source_refs")
    if isinstance(refs, list) and refs:
        return str(refs[0])
    ref = item.get("source_ref")
    return str(ref) if ref else None


def _paragraph_index(source_ref: str | None) -> int | None:
    if not source_ref:
        return None
    match = re.search(r"word/document\.xml:p\[(\d+)\]", source_ref)
    return int(match.group(1)) if match else None


def _paragraph_for_ref(
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
) -> Paragraph | None:
    index = _paragraph_index(source_ref)
    if index is None:
        return None
    return paragraph_map.get(index)


def _cell_for_ref(doc: Document, source_ref: str | None) -> _Cell | None:
    if not source_ref:
        return None
    match = re.search(
        r"word/document\.xml:tbl\[(\d+)\]/tr\[(\d+)\]/tc\[(\d+)\]",
        source_ref,
    )
    if not match:
        return None
    table_index, row_index, cell_index = (int(value) - 1 for value in match.groups())
    try:
        return doc.tables[table_index].rows[row_index].cells[cell_index]
    except IndexError:
        return None


def _clear_cell(cell: _Cell) -> None:
    for paragraph in list(cell.paragraphs):
        _remove_paragraph(paragraph)
    if not cell.paragraphs:
        cell.add_paragraph("")


def _insert_marker(
    doc: Document,
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
    marker: str,
) -> str:
    target = _paragraph_for_ref(paragraph_map, source_ref)
    if target is not None:
        _insert_paragraph_after(target, marker)
        return f"{source_ref}/after:{marker}"
    target_cell = _cell_for_ref(doc, source_ref)
    if target_cell is not None:
        target_cell.add_paragraph(marker)
        return f"{source_ref}/p[last]:{marker}"
    return _append_marker(doc, marker)


def _insert_page_break_before(
    doc: Document,
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
) -> str | None:
    target = _paragraph_for_ref(paragraph_map, source_ref)
    if target is not None:
        target.paragraph_format.page_break_before = True
        return f"{source_ref}/pageBreakBefore"
    target_cell = _cell_for_ref(doc, source_ref)
    if target_cell is not None and target_cell.paragraphs:
        target_cell.paragraphs[0].paragraph_format.page_break_before = True
        return f"{source_ref}/p[1]/pageBreakBefore"
    return None


def _insert_section_break_before(
    doc: Document,
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
) -> str | None:
    target = _paragraph_for_ref(paragraph_map, source_ref)
    if target is None:
        return None
    boundary = _insert_paragraph_before(target, "")
    paragraph_properties = boundary._p.get_or_add_pPr()
    paragraph_properties.append(_next_page_section_properties(doc))
    return f"{source_ref}/before:sectPr"


def _insert_styled_paragraph_before(
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
    text: str,
    *,
    page_break_before: bool = False,
) -> str | None:
    target = _paragraph_for_ref(paragraph_map, source_ref)
    if target is None:
        return None
    inserted = _insert_paragraph_before(target, text)
    inserted.style = target.style
    inserted.alignment = target.alignment
    inserted.paragraph_format.page_break_before = page_break_before
    return f"{source_ref}/before:{text}"


def _next_page_section_properties(doc: Document) -> Any:
    section_properties = deepcopy(doc.sections[0]._sectPr)
    for child in list(section_properties):
        if child.tag == qn("w:type"):
            section_properties.remove(child)
    section_type = OxmlElement("w:type")
    section_type.set(qn("w:val"), "nextPage")
    section_properties.insert(0, section_type)
    return section_properties


def _insert_paragraph_before(paragraph: Paragraph, text: str) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._p.addprevious(new_element)
    new_paragraph = Paragraph(new_element, paragraph._parent)
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def _insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = Paragraph(new_element, paragraph._parent)
    new_paragraph.add_run(text)
    return new_paragraph


def _append_marker(doc: Document, marker: str) -> str:
    doc.add_paragraph(marker)
    return f"word/document.xml:p[{len(doc.paragraphs)}]"


def _find_marker_ref(doc: Document, marker: str) -> str | None:
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        if marker in paragraph.text:
            return f"word/document.xml:p[{index}]"
    return None


def _remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _slot_marker(action: dict[str, Any]) -> str:
    return f"[[DOCFIT_SLOT:{action.get('unit_id')}.{action.get('element_id')}]]"


def _generated_marker(action: dict[str, Any]) -> str:
    return f"[[DOCFIT_GENERATED:{action.get('unit_id')}.{action.get('element_id')}]]"


def _slot_from_action(
    action: dict[str, Any],
    *,
    marker: str,
    output_ref: str,
) -> dict[str, Any]:
    return {
        "slot_id": f"{action.get('unit_id')}.{action.get('element_id')}",
        "unit_id": action.get("unit_id"),
        "element_id": action.get("element_id"),
        "kind": "body_content",
        "marker": marker,
        "output_ref": output_ref,
        "required": True,
    }


def _executed(action: dict[str, Any], *, output_ref: str | None) -> dict[str, Any]:
    return {
        **action,
        "output_ref": output_ref,
        "status": "executed",
    }


def _needs_review(action: dict[str, Any], reason: str) -> dict[str, Any]:
    return {
        **action,
        "status": "needs_review",
        "reason": reason,
    }


def _dedupe_by_key(items: list[dict[str, Any]], key: str) -> list[dict[str, Any]]:
    deduped: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in items:
        value = str(item.get(key))
        if value in seen:
            continue
        deduped.append(item)
        seen.add(value)
    return deduped


def _part_name(source_ref: str | None) -> str:
    if not source_ref:
        return ""
    return source_ref.split(":", 1)[0]


def _normalize_text(text: str) -> str:
    return re.sub(r"\s+", "", text).lower()


def _normalize_for_match(value: Any) -> str:
    text = _strip_format_annotations(str(value or ""))
    text = re.sub(r"[□×Xx_＿]+", "", text)
    text = re.sub(r"[…·•.。．]{2,}", "", text)
    text = re.sub(r"[\s:：;；,，.。!！?？、（）()《》<>“”\"'‘’\[\]【】]", "", text)
    return text.strip().lower()


def _strip_format_annotations(text: str) -> str:
    format_markers = (
        "号",
        "黑体",
        "宋体",
        "楷体",
        "仿宋",
        "Times",
        "居中",
        "加粗",
        "行距",
        "字号",
        "字体",
        "页边距",
        "厘米",
        "空格",
        "格式",
        "pt",
        "表示",
    )

    def replace_annotation(match: re.Match[str]) -> str:
        content = match.group(1)
        if any(marker in content for marker in format_markers):
            return ""
        return match.group(0)

    text = re.sub(r"（([^（）]*)）", replace_annotation, text)
    text = re.sub(r"\(([^()]*)\)", replace_annotation, text)
    return text


def _split_match_tokens(text: str) -> list[str]:
    stripped = _strip_format_annotations(str(text or ""))
    raw_tokens = re.split(r"[:：;；,，.。、\s/]+", stripped)
    tokens: list[str] = []
    for token in raw_tokens:
        normalized = _normalize_for_match(token)
        if len(normalized) >= 2 and not _looks_like_descriptor(normalized):
            tokens.append(normalized)
    return _dedupe(tokens)


def _looks_like_descriptor(value: str) -> bool:
    descriptor_markers = (
        "当前阶段",
        "目标输出",
        "系统生成",
        "××",
        "模板",
        "正文",
        "标题",
        "内容流",
    )
    if value in {"固定", "填充", "生成", "手工", "学生"}:
        return True
    return any(marker in value for marker in descriptor_markers) and "：" not in value


def _dedupe(items: list[str]) -> list[str]:
    deduped: list[str] = []
    seen: set[str] = set()
    for item in items:
        if item in seen:
            continue
        deduped.append(item)
        seen.add(item)
    return deduped


def _coverage(
    *,
    input_exists: bool,
    input_valid_docx: bool | None = None,
    source_tree: bool = False,
    discovered_rules: bool = False,
    template_artifact: bool = False,
    decisions: bool = False,
    generation_plan: bool = False,
    output_docx: bool = False,
    manifest: bool = False,
    body_slot: bool = False,
) -> dict[str, bool]:
    return {
        "template_generation.input_exists": input_exists,
        "template_generation.input_valid_docx": bool(input_valid_docx),
        "template_generation.source_tree": source_tree,
        "template_generation.discovered_rules": discovered_rules,
        "template_generation.template_artifact": template_artifact,
        "template_generation.decisions": decisions,
        "template_generation.plan": generation_plan,
        "template_generation.output_docx": output_docx,
        "template_generation.manifest": manifest,
        "template_generation.body_slot": body_slot,
    }
