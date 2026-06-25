from __future__ import annotations

from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from docx import Document

from docfit.core.io import now_iso, sha256_bytes, sha256_file
from docfit.ooxml.package import detect_unsupported_visible_objects, is_valid_docx


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
RUN_CONTAINER_TAGS = {"hyperlink", "fldSimple", "sdt", "smartTag", "customXml", "ins"}
INVISIBLE_CONTENT_TAGS = {"del", "moveFrom"}


def inspect_generated_template_docx(generated_template: Path) -> dict[str, Any]:
    tree: dict[str, Any] = {
        "artifact_type": "generated_template_tree",
        "artifact_version": "1.0",
        "producer": {"name": "docfit-generated-template-inspector", "version": "0.1.0"},
        "created_at": now_iso(),
        "input_docx": str(generated_template),
        "input_exists": generated_template.exists(),
        "input_valid_docx": is_valid_docx(generated_template),
        "input_hashes": {},
        "source_kind": "generated_template_docx",
        "data": {
            "paragraphs": [],
            "tables": [],
            "headers_footers": [],
            "fields": [],
            "content_controls": [],
            "footnotes": [],
            "text_boxes": [],
            "images": [],
            "breaks": [],
            "sections": [],
            "numbering_refs": [],
            "numbering_definitions": [],
            "unknown_visible_objects": [],
        },
    }
    if not tree["input_exists"] or not tree["input_valid_docx"]:
        return tree

    tree["input_hashes"]["generated_template_docx"] = sha256_file(generated_template)
    doc = Document(generated_template)
    paragraph_styles = _paragraph_style_details_by_index(generated_template)
    table_details = _table_details_by_index(generated_template)
    tree["data"]["paragraphs"] = _paragraphs(doc, paragraph_styles)
    tree["data"]["tables"] = _tables(doc, table_details, paragraph_styles)
    tree["data"].update(_inspect_ooxml_parts(generated_template))
    unsupported = detect_unsupported_visible_objects(generated_template)
    if tree["data"].get("footnotes"):
        unsupported = [
            item for item in unsupported if item.get("object_type") != "footnote"
        ]
    if tree["data"].get("text_boxes"):
        unsupported = [
            item for item in unsupported if item.get("object_type") != "text_box"
        ]
    if tree["data"].get("images"):
        unsupported = [
            item for item in unsupported if item.get("object_type") != "image"
        ]
    tree["data"]["unknown_visible_objects"] = unsupported
    return tree


def iter_visible_text_entries(tree: dict[str, Any]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    data = tree.get("data", {})
    for paragraph in data.get("paragraphs", []):
        if paragraph.get("text"):
            entries.append(
                {
                    "kind": "paragraph",
                    "text": paragraph.get("text", ""),
                    "style": paragraph.get("style", ""),
                    "style_details": paragraph.get("style_details", {}),
                    "runs": paragraph.get("runs", []),
                    "source_ref": paragraph.get("source_ref", ""),
                    "order": paragraph.get("xml_index", paragraph.get("index", 0)),
                    "paragraph_index": paragraph.get("xml_index"),
                    "python_docx_index": paragraph.get("python_docx_index")
                    or paragraph.get("index"),
                }
            )
    for table in data.get("tables", []):
        for cell in table.get("cells", []):
            if cell.get("text"):
                entries.append(
                    {
                        "kind": "table_cell",
                        "text": cell.get("text", ""),
                        "style": table.get("style", ""),
                        "style_details": cell.get("style_details", {}),
                        "source_ref": cell.get("source_ref", ""),
                        "order": (
                            cell.get("first_paragraph_index")
                            or cell.get("row_first_paragraph_index")
                            or table.get("first_paragraph_index")
                            or 10_000 + cell.get("global_index", 0)
                        ),
                        "paragraph_index": cell.get("first_paragraph_index"),
                        "end_paragraph_index": cell.get("last_paragraph_index"),
                        "paragraph_indices": cell.get("paragraph_indices", []),
                        "cell_paragraph_refs": cell.get("cell_paragraph_refs", []),
                        "table_index": table.get("index"),
                        "table_source_ref": table.get("source_ref", ""),
                    }
                )
    for part in data.get("headers_footers", []):
        if part.get("text"):
            entries.append(
                {
                    "kind": part.get("kind", "header_footer"),
                    "text": part.get("text", ""),
                    "style": "",
                    "source_ref": part.get("source_ref", ""),
                    "part_paragraph_refs": part.get("paragraph_refs", []),
                    "order": 20_000 + part.get("index", 0),
                }
            )
    for control in data.get("content_controls", []):
        tag = control.get("tag")
        if tag:
            entries.append(
                {
                    "kind": "content_control",
                    "text": control.get("text") or f"sdt:{tag}",
                    "sdt_tag": tag,
                    "style": "",
                    "source_ref": control.get("source_ref", ""),
                    "order": 15_000 + control.get("index", 0),
                }
            )
    return entries


def _paragraphs(
    doc: Document,
    paragraph_styles: dict[int, dict[str, Any]],
) -> list[dict[str, Any]]:
    paragraphs: list[dict[str, Any]] = []
    xml_indices = _xml_paragraph_indices(doc)
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        xml_index = xml_indices.get(_element_path(paragraph._p), index)
        style_details = paragraph_styles.get(xml_index, {})
        text = str(style_details.get("text") or paragraph.text).strip()
        if not text:
            continue
        runs = style_details.get("runs") or [
            {
                "text": run.text,
                "bold": run.bold,
                "italic": run.italic,
                "underline": run.underline,
                "color": str(run.font.color.rgb)
                if run.font.color.rgb is not None
                else None,
                "font_names": [run.font.name] if run.font.name else [],
                "font_size_pt": (
                    run.font.size.pt if run.font.size is not None else None
                ),
            }
            for run in paragraph.runs
            if run.text
        ]
        paragraphs.append(
            {
                "index": index,
                "python_docx_index": index,
                "xml_index": xml_index,
                "text": text,
                "style": paragraph.style.name if paragraph.style is not None else "",
                "style_details": style_details,
                "runs": runs,
                "source_ref": f"word/document.xml:p[{xml_index}]",
            }
        )
    return paragraphs


def _xml_paragraph_indices(doc: Document) -> dict[str, int]:
    return {
        _element_path(paragraph): index
        for index, paragraph in enumerate(doc.element.body.iter(f"{W_NS}p"), start=1)
    }


def _tables(
    doc: Document,
    table_details: dict[int, dict[str, Any]],
    paragraph_styles: dict[int, dict[str, Any]],
) -> list[dict[str, Any]]:
    tables: list[dict[str, Any]] = []
    global_cell_index = 0
    for table_index, table in enumerate(doc.tables, start=1):
        table_detail = table_details.get(table_index, {})
        cell_details = table_detail.get("cells", {})
        row_details = table_detail.get("rows", {})
        cells: list[dict[str, Any]] = []
        cell_trace_by_tc_path: dict[str, dict[str, Any]] = {}
        for row_index, row in enumerate(table.rows, start=1):
            row_detail = row_details.get(row_index, {})
            for physical_cell_index, tc in enumerate(row._tr.tc_lst, start=1):
                physical_detail = cell_details.get((row_index, physical_cell_index), {})
                if physical_detail:
                    cell_trace_by_tc_path[_element_path(tc)] = physical_detail
            for cell_index, cell in enumerate(row.cells, start=1):
                global_cell_index += 1
                layout_cell_detail = cell_details.get((row_index, cell_index), {})
                cell_detail = cell_trace_by_tc_path.get(
                    _element_path(cell._tc),
                    {},
                ) or layout_cell_detail
                first_paragraph_index = cell_detail.get("first_paragraph_index")
                row_first_paragraph_index = cell_detail.get(
                    "row_first_paragraph_index"
                ) or row_detail.get("row_first_paragraph_index")
                row_last_paragraph_index = cell_detail.get(
                    "row_last_paragraph_index"
                ) or row_detail.get("row_last_paragraph_index")
                text = "\n".join(
                    _cell_paragraph_text(
                        paragraph_styles,
                        cell_detail.get("paragraph_indices", []),
                        local_index,
                        paragraph,
                    )
                    for local_index, paragraph in enumerate(cell.paragraphs)
                    if _cell_paragraph_text(
                        paragraph_styles,
                        cell_detail.get("paragraph_indices", []),
                        local_index,
                        paragraph,
                    )
                )
                cell_paragraphs = _cell_paragraph_facts(
                    paragraph_styles,
                    cell_detail.get("paragraph_indices", []),
                    cell.paragraphs,
                )
                cells.append(
                    {
                        "row": row_index,
                        "column": cell_index,
                        "global_index": global_cell_index,
                        "text": text,
                        "style_details": paragraph_styles.get(
                            first_paragraph_index or row_first_paragraph_index,
                            {},
                        ),
                        "paragraph_indices": cell_detail.get("paragraph_indices", []),
                        "paragraphs": cell_paragraphs,
                        "cell_paragraph_refs": [
                            paragraph["source_ref"] for paragraph in cell_paragraphs
                        ],
                        "first_paragraph_index": first_paragraph_index,
                        "last_paragraph_index": cell_detail.get("last_paragraph_index"),
                        "row_first_paragraph_index": row_first_paragraph_index,
                        "row_last_paragraph_index": row_last_paragraph_index,
                        "keep_refs": cell_detail.get("keep_refs", []),
                        "row_cant_split": bool(
                            cell_detail.get("row_cant_split")
                            or row_detail.get("row_cant_split")
                        ),
                        "row_source_ref": cell_detail.get("row_source_ref")
                        or row_detail.get("row_source_ref", ""),
                        "source_ref": layout_cell_detail.get("source_ref")
                        or (
                            f"word/document.xml:tbl[{table_index}]"
                            f"/tr[{row_index}]/tc[{cell_index}]"
                        ),
                    }
                )
        tables.append(
            {
                "index": table_index,
                "style": table.style.name if table.style is not None else "",
                "row_count": len(table.rows),
                "column_count": max((len(row.cells) for row in table.rows), default=0),
                "first_paragraph_index": table_detail.get("first_paragraph_index"),
                "last_paragraph_index": table_detail.get("last_paragraph_index"),
                "cant_split_row_refs": table_detail.get("cant_split_row_refs", []),
                "keep_refs": table_detail.get("keep_refs", []),
                "cells": cells,
                "source_ref": f"word/document.xml:tbl[{table_index}]",
            }
        )
    return tables


def _cell_paragraph_text(
    paragraph_styles: dict[int, dict[str, Any]],
    paragraph_indices: list[int],
    local_index: int,
    paragraph: Any,
) -> str:
    paragraph_index = (
        int(paragraph_indices[local_index])
        if local_index < len(paragraph_indices)
        else None
    )
    if paragraph_index is not None:
        text = str(paragraph_styles.get(paragraph_index, {}).get("text") or "").strip()
        if text:
            return text
    return str(getattr(paragraph, "text", "") or "").strip()


def _cell_paragraph_facts(
    paragraph_styles: dict[int, dict[str, Any]],
    paragraph_indices: list[int],
    paragraphs: list[Any],
) -> list[dict[str, Any]]:
    facts: list[dict[str, Any]] = []
    for local_index, paragraph in enumerate(paragraphs):
        if local_index >= len(paragraph_indices):
            continue
        paragraph_index = int(paragraph_indices[local_index])
        style_details = paragraph_styles.get(paragraph_index, {})
        text = _cell_paragraph_text(
            paragraph_styles,
            paragraph_indices,
            local_index,
            paragraph,
        )
        if not text:
            continue
        facts.append(
            {
                "index": paragraph_index,
                "xml_index": paragraph_index,
                "text": text,
                "style": (
                    paragraph.style.name
                    if getattr(paragraph, "style", None) is not None
                    else ""
                ),
                "style_details": style_details,
                "runs": style_details.get("runs", []),
                "source_ref": f"word/document.xml:p[{paragraph_index}]",
            }
        )
    return facts


def _table_details_by_index(path: Path) -> dict[int, dict[str, Any]]:
    try:
        with ZipFile(path) as package:
            root = ET.fromstring(package.read("word/document.xml"))
    except (KeyError, ET.ParseError):
        return {}

    paragraph_indices = {
        id(paragraph): index
        for index, paragraph in enumerate(root.iter(f"{W_NS}p"), start=1)
    }
    details: dict[int, dict[str, Any]] = {}
    body = root.find(f"{W_NS}body")
    if body is None:
        return details
    for table_index, table in enumerate(body.findall(f"{W_NS}tbl"), start=1):
        table_paragraphs = [
            paragraph_indices[id(paragraph)]
            for paragraph in table.iter(f"{W_NS}p")
            if id(paragraph) in paragraph_indices
        ]
        table_keep_refs: list[str] = []
        cant_split_refs: list[str] = []
        cell_details: dict[tuple[int, int], dict[str, Any]] = {}
        row_details: dict[int, dict[str, Any]] = {}
        for row_index, row in enumerate(table.findall(f"{W_NS}tr"), start=1):
            row_source_ref = f"word/document.xml:tbl[{table_index}]/tr[{row_index}]"
            row_paragraphs = [
                paragraph_indices[id(paragraph)]
                for paragraph in row.iter(f"{W_NS}p")
                if id(paragraph) in paragraph_indices
            ]
            row_first_paragraph_index = min(row_paragraphs) if row_paragraphs else None
            row_last_paragraph_index = max(row_paragraphs) if row_paragraphs else None
            row_cant_split = (
                row.find(f"{W_NS}trPr/{W_NS}cantSplit") is not None
            )
            if row_cant_split:
                cant_split_refs.append(f"{row_source_ref}/cantSplit")
            row_details[row_index] = {
                "row_first_paragraph_index": row_first_paragraph_index,
                "row_last_paragraph_index": row_last_paragraph_index,
                "row_cant_split": row_cant_split,
                "row_source_ref": row_source_ref,
            }
            for cell_index, cell in enumerate(row.findall(f"{W_NS}tc"), start=1):
                cell_paragraphs = [
                    paragraph_indices[id(paragraph)]
                    for paragraph in cell.iter(f"{W_NS}p")
                    if id(paragraph) in paragraph_indices
                ]
                cell_keep_refs: list[str] = []
                for paragraph in cell.iter(f"{W_NS}p"):
                    paragraph_index = paragraph_indices.get(id(paragraph))
                    if paragraph_index is None:
                        continue
                    properties = paragraph.find(f"{W_NS}pPr")
                    if properties is None:
                        continue
                    if properties.find(f"{W_NS}keepNext") is not None:
                        cell_keep_refs.append(
                            f"word/document.xml:p[{paragraph_index}]/keepNext"
                        )
                    if properties.find(f"{W_NS}keepLines") is not None:
                        cell_keep_refs.append(
                            f"word/document.xml:p[{paragraph_index}]/keepLines"
                        )
                table_keep_refs.extend(cell_keep_refs)
                cell_details[(row_index, cell_index)] = {
                    "paragraph_indices": cell_paragraphs,
                    "first_paragraph_index": min(cell_paragraphs)
                    if cell_paragraphs
                    else None,
                    "last_paragraph_index": max(cell_paragraphs)
                    if cell_paragraphs
                    else None,
                    "row_first_paragraph_index": row_first_paragraph_index,
                    "row_last_paragraph_index": row_last_paragraph_index,
                    "keep_refs": cell_keep_refs,
                    "row_cant_split": row_cant_split,
                    "row_source_ref": row_source_ref,
                    "source_ref": (
                        f"{row_source_ref}/tc[{cell_index}]"
                    ),
                }
        details[table_index] = {
            "first_paragraph_index": min(table_paragraphs)
            if table_paragraphs
            else None,
            "last_paragraph_index": max(table_paragraphs)
            if table_paragraphs
            else None,
            "cant_split_row_refs": cant_split_refs,
            "keep_refs": _dedupe(table_keep_refs),
            "rows": row_details,
            "cells": cell_details,
        }
    return details


def _paragraph_style_details_by_index(path: Path) -> dict[int, dict[str, Any]]:
    try:
        with ZipFile(path) as package:
            root = ET.fromstring(package.read("word/document.xml"))
            style_catalog = _style_catalog(package)
    except (KeyError, ET.ParseError):
        return {}

    details: dict[int, dict[str, Any]] = {}
    for index, paragraph in enumerate(root.iter(f"{W_NS}p"), start=1):
        text = _visible_text(paragraph)
        if not text:
            continue
        paragraph_properties = paragraph.find(f"{W_NS}pPr")
        direct_paragraph_style = _paragraph_style(paragraph_properties)
        style_id = direct_paragraph_style.get("style_id")
        inherited = style_catalog.get(style_id, {})
        inherited_paragraph = inherited.get("paragraph", {})
        inherited_run = inherited.get("run", {})
        direct_paragraph_run = _run_properties(
            paragraph_properties.find(f"{W_NS}rPr")
            if paragraph_properties is not None
            else None
        )
        paragraph_run_properties = _merge_run_styles(
            inherited_run,
            direct_paragraph_run,
        )
        runs = _paragraph_runs(
            paragraph,
            paragraph_run_properties,
            part_name="word/document.xml",
            paragraph_index=index,
        )
        details[index] = {
            "text": text,
            "paragraph": _merge_paragraph_styles(
                inherited_paragraph,
                direct_paragraph_style,
            ),
            "paragraph_run_properties": paragraph_run_properties,
            "runs": runs,
            "dominant_run": _dominant_run_style(runs),
            "style_inheritance": {
                **inherited.get("inheritance", {}),
                "run": inherited_run,
            },
        }
    return details


def _style_catalog(package: ZipFile) -> dict[str | None, dict[str, Any]]:
    try:
        root = ET.fromstring(package.read("word/styles.xml"))
    except (KeyError, ET.ParseError):
        return {}

    raw_styles: dict[str, dict[str, Any]] = {}
    default_style_id: str | None = None
    for style in root.findall(f"{W_NS}style"):
        if _attr(style, "type") != "paragraph":
            continue
        style_id = _attr(style, "styleId")
        if not style_id:
            continue
        if _attr(style, "default") in {"1", "true", "on"}:
            default_style_id = style_id
        style_name = style.find(f"{W_NS}name")
        based_on = style.find(f"{W_NS}basedOn")
        paragraph_properties = style.find(f"{W_NS}pPr")
        run_properties = style.find(f"{W_NS}rPr")
        raw_styles[style_id] = {
            "style_id": style_id,
            "style_name": _attr(style_name, "val"),
            "based_on": _attr(based_on, "val"),
            "paragraph": _paragraph_style(paragraph_properties),
            "run": _run_properties(run_properties),
            "source_ref": f"word/styles.xml:style[{style_id}]",
        }

    resolved: dict[str | None, dict[str, Any]] = {
        style_id: _resolve_style(style_id, raw_styles, [])
        for style_id in raw_styles
    }
    if default_style_id and None not in resolved:
        resolved[None] = resolved.get(default_style_id, {})
    return resolved


def _resolve_style(
    style_id: str,
    raw_styles: dict[str, dict[str, Any]],
    seen: list[str],
) -> dict[str, Any]:
    if style_id in seen:
        return {}
    style = raw_styles.get(style_id)
    if not style:
        return {}
    base = _resolve_style(
        str(style.get("based_on")),
        raw_styles,
        [*seen, style_id],
    ) if style.get("based_on") else {}
    source_refs = [
        *base.get("inheritance", {}).get("source_refs", []),
        style.get("source_ref", ""),
    ]
    style_chain = [
        *base.get("inheritance", {}).get("style_chain", []),
        style_id,
    ]
    return {
        "paragraph": _merge_paragraph_styles(
            base.get("paragraph", {}),
            style.get("paragraph", {}),
        ),
        "run": _merge_run_styles(base.get("run", {}), style.get("run", {})),
        "inheritance": {
            "style_id": style_id,
            "style_name": style.get("style_name"),
            "based_on": style.get("based_on"),
            "style_chain": style_chain,
            "source_refs": [ref for ref in source_refs if ref],
        },
    }


def _merge_paragraph_styles(
    inherited: dict[str, Any],
    override: dict[str, Any],
) -> dict[str, Any]:
    return {
        "style_id": override.get("style_id") or inherited.get("style_id"),
        "alignment": override.get("alignment") or inherited.get("alignment"),
        "spacing": _merge_optional_dict(
            inherited.get("spacing", {}),
            override.get("spacing", {}),
        ),
        "indent": _merge_optional_dict(
            inherited.get("indent", {}),
            override.get("indent", {}),
        ),
        "page_break_before": bool(
            override.get("page_break_before")
            or inherited.get("page_break_before")
        ),
        "keep_next": bool(override.get("keep_next") or inherited.get("keep_next")),
        "keep_lines": bool(override.get("keep_lines") or inherited.get("keep_lines")),
    }


def _merge_run_styles(
    inherited: dict[str, Any],
    override: dict[str, Any],
) -> dict[str, Any]:
    return {
        "font_names": override.get("font_names") or inherited.get("font_names", []),
        "font_size_pt": _first_not_none(
            override.get("font_size_pt"),
            inherited.get("font_size_pt"),
        ),
        "bold": _first_not_none(override.get("bold"), inherited.get("bold")),
        "italic": _first_not_none(override.get("italic"), inherited.get("italic")),
        "underline": _first_not_none(
            override.get("underline"),
            inherited.get("underline"),
        ),
        "color": _first_not_none(override.get("color"), inherited.get("color")),
    }


def _merge_optional_dict(
    inherited: dict[str, Any],
    override: dict[str, Any],
) -> dict[str, Any]:
    keys = set(inherited) | set(override)
    merged: dict[str, Any] = {}
    for key in keys:
        value = _first_not_none(override.get(key), inherited.get(key))
        if value is not None:
            merged[key] = value
    return merged


def _first_not_none(*values: Any) -> Any:
    for value in values:
        if value is not None:
            return value
    return None


def _paragraph_style(properties: ET.Element | None) -> dict[str, Any]:
    if properties is None:
        return {}
    style = properties.find(f"{W_NS}pStyle")
    alignment = properties.find(f"{W_NS}jc")
    spacing = properties.find(f"{W_NS}spacing")
    indent = properties.find(f"{W_NS}ind")
    return {
        "style_id": _attr(style, "val"),
        "alignment": _attr(alignment, "val"),
        "spacing": _spacing_style(spacing),
        "indent": _indent_style(indent),
        "page_break_before": properties.find(f"{W_NS}pageBreakBefore") is not None,
        "keep_next": properties.find(f"{W_NS}keepNext") is not None,
        "keep_lines": properties.find(f"{W_NS}keepLines") is not None,
    }


def _run_style(
    run: ET.Element,
    inherited_run_properties: dict[str, Any] | None = None,
    *,
    source_ref: str | None = None,
    container_refs: list[str] | None = None,
) -> dict[str, Any]:
    properties = run.find(f"{W_NS}rPr")
    style = {
        "text": _visible_text(run, strip=False),
        **_merge_run_styles(
            inherited_run_properties or {},
            _run_properties(properties),
        ),
    }
    if source_ref:
        style["source_ref"] = source_ref
    if container_refs:
        style["container_refs"] = list(container_refs)
    return style


def _paragraph_runs(
    paragraph: ET.Element,
    paragraph_run_properties: dict[str, Any],
    *,
    part_name: str,
    paragraph_index: int,
) -> list[dict[str, Any]]:
    runs: list[dict[str, Any]] = []
    container_counts: dict[str, int] = {}

    def walk(
        node: ET.Element,
        *,
        path_parts: list[str],
        container_refs: list[str],
    ) -> None:
        for child in list(node):
            local_name = _local_name(child.tag)
            if local_name in INVISIBLE_CONTENT_TAGS:
                continue
            next_path_parts = path_parts
            next_container_refs = container_refs
            if local_name in RUN_CONTAINER_TAGS:
                container_counts[local_name] = container_counts.get(local_name, 0) + 1
                container_ref = (
                    f"{part_name}:p[{paragraph_index}]"
                    f"/{local_name}[{container_counts[local_name]}]"
                )
                next_path_parts = [*path_parts, f"{local_name}[{container_counts[local_name]}]"]
                next_container_refs = [*container_refs, container_ref]
            if child.tag == f"{W_NS}r":
                if _visible_text(child, strip=False):
                    run_index = len(runs) + 1
                    source_path = "/".join([*path_parts, f"r[{run_index}]"])
                    runs.append(
                        _run_style(
                            child,
                            paragraph_run_properties,
                            source_ref=f"{part_name}:p[{paragraph_index}]/{source_path}",
                            container_refs=container_refs,
                        )
                    )
                continue
            walk(
                child,
                path_parts=next_path_parts,
                container_refs=next_container_refs,
            )

    walk(paragraph, path_parts=[], container_refs=[])
    return runs


def _run_properties(properties: ET.Element | None) -> dict[str, Any]:
    if properties is None:
        return {
            "font_names": [],
            "font_size_pt": None,
            "bold": None,
            "italic": None,
            "underline": None,
            "color": None,
        }
    fonts = properties.find(f"{W_NS}rFonts")
    size = properties.find(f"{W_NS}sz")
    underline = properties.find(f"{W_NS}u")
    color = properties.find(f"{W_NS}color")
    return {
        "font_names": _font_names(fonts),
        "font_size_pt": _half_points_to_points(_attr(size, "val")),
        "bold": _toggle_value(properties.find(f"{W_NS}b")),
        "italic": _toggle_value(properties.find(f"{W_NS}i")),
        "underline": _underline_value(underline),
        "color": _color_value(color),
    }


def _dominant_run_style(runs: list[dict[str, Any]]) -> dict[str, Any]:
    candidates = [
        run
        for run in runs
        if _run_visible_for_style(str(run.get("text", "")))
    ]
    if not candidates:
        candidates = runs
    if not candidates:
        return {}
    return max(candidates, key=lambda run: len(str(run.get("text", "")).strip()))


def _run_visible_for_style(text: str) -> bool:
    stripped = text.strip()
    return bool(stripped) and not (
        stripped.startswith("（") and stripped.endswith("）")
    )


def _spacing_style(spacing: ET.Element | None) -> dict[str, Any]:
    if spacing is None:
        return {}
    line = _attr(spacing, "line")
    line_rule = _attr(spacing, "lineRule")
    return {
        "before_pt": _twips_to_points(_attr(spacing, "before")),
        "after_pt": _twips_to_points(_attr(spacing, "after")),
        "line": line,
        "line_rule": line_rule,
        "line_spacing": _line_spacing(line, line_rule),
    }


def _indent_style(indent: ET.Element | None) -> dict[str, Any]:
    if indent is None:
        return {}
    return {
        "first_line_twips": _int_or_none(_attr(indent, "firstLine")),
        "left_twips": _int_or_none(_attr(indent, "left")),
        "right_twips": _int_or_none(_attr(indent, "right")),
    }


def _font_names(fonts: ET.Element | None) -> list[str]:
    if fonts is None:
        return []
    names = [
        _attr(fonts, "eastAsia"),
        _attr(fonts, "ascii"),
        _attr(fonts, "hAnsi"),
        _attr(fonts, "cs"),
    ]
    return [name for name in dict.fromkeys(names) if name]


def _toggle_value(element: ET.Element | None) -> bool | None:
    if element is None:
        return None
    value = _attr(element, "val")
    if value in {"0", "false", "False", "off"}:
        return False
    return True


def _underline_value(element: ET.Element | None) -> str | bool | None:
    if element is None:
        return None
    value = _attr(element, "val")
    if value in {"0", "false", "False", "none"}:
        return False
    return value or "single"


def _color_value(element: ET.Element | None) -> str | None:
    value = _attr(element, "val")
    if value is None:
        return None
    return value.lower()


def _attr(element: ET.Element | None, name: str) -> str | None:
    if element is None:
        return None
    return element.attrib.get(f"{W_NS}{name}")


def _half_points_to_points(value: str | None) -> float | None:
    number = _int_or_none(value)
    return number / 2 if number is not None else None


def _twips_to_points(value: str | None) -> float | None:
    number = _int_or_none(value)
    return number / 20 if number is not None else None


def _line_spacing(line: str | None, line_rule: str | None) -> str | None:
    number = _int_or_none(line)
    if number is None:
        return None
    if line_rule in {None, "auto"}:
        if number == 240:
            return "single"
        if number == 360:
            return "1.5"
        if number == 480:
            return "double"
        return f"multiple:{number / 240:g}"
    if line_rule == "exact":
        return f"exact:{number / 20:g}pt"
    return str(number)


def _int_or_none(value: str | None) -> int | None:
    if value is None:
        return None
    try:
        return int(value)
    except ValueError:
        return None


def _inspect_ooxml_parts(generated_template: Path) -> dict[str, list[dict[str, Any]]]:
    headers_footers: list[dict[str, Any]] = []
    fields: list[dict[str, Any]] = []
    content_controls: list[dict[str, Any]] = []
    footnotes: list[dict[str, Any]] = []
    text_boxes: list[dict[str, Any]] = []
    images: list[dict[str, Any]] = []
    breaks: list[dict[str, Any]] = []
    sections: list[dict[str, Any]] = []
    numbering_refs: list[dict[str, Any]] = []
    numbering_definitions: list[dict[str, Any]] = []

    with ZipFile(generated_template) as package:
        relationships = _document_relationships(package)
        style_catalog = _style_catalog(package)
        numbering_catalog = _numbering_catalog(package)
        numbering_definitions = numbering_catalog["definitions"]
        footnotes = _footnotes(package)
        part_names = [
            name
            for name in package.namelist()
            if name == "word/document.xml"
            or name.startswith("word/header")
            or name.startswith("word/footer")
        ]
        for part_index, part_name in enumerate(sorted(part_names), start=1):
            try:
                root = ET.fromstring(package.read(part_name))
            except ET.ParseError:
                continue
            if part_name.startswith("word/header") or part_name.startswith("word/footer"):
                paragraphs = _part_paragraphs(
                    root,
                    part_name=part_name,
                    style_catalog=style_catalog,
                )
                headers_footers.append(
                    {
                        "index": part_index,
                        "kind": (
                            "header" if part_name.startswith("word/header") else "footer"
                        ),
                        "part_name": part_name,
                        "text": _visible_text(root),
                        "paragraphs": paragraphs,
                        "paragraph_refs": [
                            paragraph["source_ref"] for paragraph in paragraphs
                        ],
                        "source_ref": part_name,
                    }
                )
            fields.extend(_fields(root, part_name))
            breaks.extend(_breaks(root, part_name))
            if part_name == "word/document.xml":
                content_controls.extend(_content_controls(root, part_name))
                text_boxes.extend(_text_boxes(root, part_name))
                images.extend(_images(root, part_name, relationships, package))
                sections.extend(_sections(root, relationships))
                numbering_refs.extend(_numbering_refs(root, numbering_catalog))

    return {
        "headers_footers": headers_footers,
        "fields": fields,
        "content_controls": content_controls,
        "footnotes": footnotes,
        "text_boxes": text_boxes,
        "images": images,
        "breaks": breaks,
        "sections": sections,
        "numbering_refs": numbering_refs,
        "numbering_definitions": numbering_definitions,
    }


def _part_paragraphs(
    root: ET.Element,
    *,
    part_name: str,
    style_catalog: dict[str | None, dict[str, Any]],
) -> list[dict[str, Any]]:
    paragraphs: list[dict[str, Any]] = []
    for index, paragraph in enumerate(root.iter(f"{W_NS}p"), start=1):
        text = _visible_text(paragraph)
        if not text:
            continue
        paragraph_properties = paragraph.find(f"{W_NS}pPr")
        direct_paragraph_style = _paragraph_style(paragraph_properties)
        style_id = direct_paragraph_style.get("style_id")
        inherited = style_catalog.get(style_id, {})
        inherited_paragraph = inherited.get("paragraph", {})
        inherited_run = inherited.get("run", {})
        direct_paragraph_run = _run_properties(
            paragraph_properties.find(f"{W_NS}rPr")
            if paragraph_properties is not None
            else None
        )
        paragraph_run_properties = _merge_run_styles(
            inherited_run,
            direct_paragraph_run,
        )
        runs = _paragraph_runs(
            paragraph,
            paragraph_run_properties,
            part_name=part_name,
            paragraph_index=index,
        )
        paragraphs.append(
            {
                "index": index,
                "xml_index": index,
                "part_name": part_name,
                "text": text,
                "style": "",
                "style_details": {
                    "text": text,
                    "paragraph": _merge_paragraph_styles(
                        inherited_paragraph,
                        direct_paragraph_style,
                    ),
                    "paragraph_run_properties": paragraph_run_properties,
                    "runs": runs,
                    "dominant_run": _dominant_run_style(runs),
                    "style_inheritance": {
                        **inherited.get("inheritance", {}),
                        "run": inherited_run,
                    },
                },
                "runs": runs,
                "source_ref": f"{part_name}:p[{index}]",
            }
        )
    return paragraphs


def _visible_text(root: ET.Element, *, strip: bool = True) -> str:
    chunks: list[str] = []

    def walk(node: ET.Element) -> None:
        local_name = _local_name(node.tag)
        if local_name in INVISIBLE_CONTENT_TAGS:
            return
        if node.tag == f"{W_NS}t":
            chunks.append(node.text or "")
        elif node.tag == f"{W_NS}tab":
            chunks.append("\t")
        elif node.tag in {f"{W_NS}br", f"{W_NS}cr"}:
            chunks.append("\n")
        for child in list(node):
            walk(child)

    walk(root)
    text = "".join(chunks)
    return text.strip() if strip else text


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _element_path(element: Any) -> str:
    return element.getroottree().getpath(element)


def _document_relationships(package: ZipFile) -> dict[str, str]:
    try:
        root = ET.fromstring(package.read("word/_rels/document.xml.rels"))
    except (KeyError, ET.ParseError):
        return {}
    relationships: dict[str, str] = {}
    for node in root.findall(f"{REL_NS}Relationship"):
        relationship_id = node.attrib.get("Id")
        target = node.attrib.get("Target")
        if relationship_id and target:
            relationships[relationship_id] = _normalize_word_target(target)
    return relationships


def _normalize_word_target(target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    if target.startswith("word/"):
        return target
    return f"word/{target}"


def _sections(
    root: ET.Element,
    relationships: dict[str, str],
) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    effective_refs: dict[tuple[str, str], dict[str, Any]] = {}

    def append_section(
        section_properties: ET.Element,
        paragraph_index: int | None,
        source_ref: str,
    ) -> None:
        nonlocal effective_refs
        direct_refs = _section_references(
            section_properties,
            relationships,
            source_ref,
        )
        for ref in direct_refs:
            effective_refs[(ref["kind"], ref["type"])] = ref
        section_index = len(sections) + 1
        sections.append(
            {
                "index": section_index,
                "paragraph_index": paragraph_index,
                "source_ref": source_ref,
                "references": direct_refs,
                "effective_references": sorted(
                    effective_refs.values(),
                    key=lambda item: (item["kind"], item["type"], item["part_name"]),
                ),
                "page_numbering": _page_numbering(section_properties),
                "page_size": _page_size(section_properties),
                "page_margins": _page_margins(section_properties),
            }
        )

    for paragraph_index, paragraph in enumerate(root.iter(f"{W_NS}p"), start=1):
        paragraph_properties = paragraph.find(f"{W_NS}pPr")
        section_properties = (
            paragraph_properties.find(f"{W_NS}sectPr")
            if paragraph_properties is not None
            else None
        )
        if section_properties is not None:
            append_section(
                section_properties,
                paragraph_index,
                f"word/document.xml:p[{paragraph_index}]/sectPr",
            )
    for section_properties in root.findall(f"./{W_NS}body/{W_NS}sectPr"):
        append_section(section_properties, None, "word/document.xml:body/sectPr")
    return sections


def _section_references(
    section_properties: ET.Element,
    relationships: dict[str, str],
    section_source_ref: str,
) -> list[dict[str, Any]]:
    refs: list[dict[str, Any]] = []
    for index, node in enumerate(section_properties, start=1):
        local_name = node.tag.rsplit("}", 1)[-1]
        if local_name not in {"headerReference", "footerReference"}:
            continue
        relationship_id = node.attrib.get(f"{R_NS}id", "")
        part_name = relationships.get(relationship_id, "")
        refs.append(
            {
                "kind": "header" if local_name == "headerReference" else "footer",
                "type": node.attrib.get(f"{W_NS}type", "default"),
                "relationship_id": relationship_id,
                "part_name": part_name,
                "source_ref": f"{section_source_ref}/{local_name}[{index}]",
            }
        )
    return refs


def _page_numbering(section_properties: ET.Element) -> dict[str, Any]:
    node = section_properties.find(f"{W_NS}pgNumType")
    if node is None:
        return {}
    return {
        "format": _attr(node, "fmt"),
        "start": _int_or_none(_attr(node, "start")),
    }


def _page_margins(section_properties: ET.Element) -> dict[str, Any]:
    node = section_properties.find(f"{W_NS}pgMar")
    if node is None:
        return {}
    return {
        "gutter_pt": _twips_to_points(_attr(node, "gutter")),
        "header_pt": _twips_to_points(_attr(node, "header")),
        "footer_pt": _twips_to_points(_attr(node, "footer")),
        "top_pt": _twips_to_points(_attr(node, "top")),
        "bottom_pt": _twips_to_points(_attr(node, "bottom")),
        "left_pt": _twips_to_points(_attr(node, "left")),
        "right_pt": _twips_to_points(_attr(node, "right")),
    }


def _page_size(section_properties: ET.Element) -> dict[str, Any]:
    node = section_properties.find(f"{W_NS}pgSz")
    if node is None:
        return {}
    return {
        "width_pt": _twips_to_points(_attr(node, "w")),
        "height_pt": _twips_to_points(_attr(node, "h")),
        "orientation": _attr(node, "orient") or "portrait",
    }


def _fields(root: ET.Element, part_name: str) -> list[dict[str, Any]]:
    fields: list[dict[str, Any]] = []
    field_stack: list[dict[str, Any]] = []
    next_index = 1

    def append_field(
        *,
        kind: str,
        instruction: str,
        paragraph_index: int | None,
        end_paragraph_index: int | None = None,
    ) -> None:
        nonlocal next_index
        normalized_instruction = _normalize_instruction(instruction)
        if not normalized_instruction:
            return
        source_ref = f"{part_name}:field[{next_index}]"
        if paragraph_index is not None:
            source_ref = f"{part_name}:p[{paragraph_index}]/field[{next_index}]"
        item = {
            "index": next_index,
            "kind": kind,
            "field_type": _field_type(normalized_instruction),
            "instruction": normalized_instruction,
            "part_name": part_name,
            "paragraph_index": paragraph_index,
            "end_paragraph_index": end_paragraph_index,
            "source_ref": source_ref,
        }
        if end_paragraph_index is not None and end_paragraph_index != paragraph_index:
            item["end_source_ref"] = f"{part_name}:p[{end_paragraph_index}]"
        fields.append(item)
        next_index += 1

    for paragraph_index, paragraph in enumerate(root.iter(f"{W_NS}p"), start=1):
        for node in paragraph.iter():
            if node.tag == f"{W_NS}fldSimple":
                append_field(
                    kind="fldSimple",
                    instruction=node.attrib.get(f"{W_NS}instr", ""),
                    paragraph_index=paragraph_index,
                    end_paragraph_index=paragraph_index,
                )
                continue
            if node.tag == f"{W_NS}fldChar":
                field_char_type = node.attrib.get(f"{W_NS}fldCharType", "")
                if field_char_type == "begin":
                    field_stack.append(
                        {
                            "paragraph_index": paragraph_index,
                            "instruction_parts": [],
                        }
                    )
                elif field_char_type == "end" and field_stack:
                    field = field_stack.pop()
                    append_field(
                        kind="complexField",
                        instruction="".join(field["instruction_parts"]),
                        paragraph_index=field["paragraph_index"],
                        end_paragraph_index=paragraph_index,
                    )
                continue
            if node.tag == f"{W_NS}instrText":
                text = node.text or ""
                if field_stack:
                    field_stack[-1]["instruction_parts"].append(text)
                else:
                    append_field(
                        kind="instrText",
                        instruction=text,
                        paragraph_index=paragraph_index,
                        end_paragraph_index=paragraph_index,
                    )

    for field in reversed(field_stack):
        append_field(
            kind="complexField_unclosed",
            instruction="".join(field["instruction_parts"]),
            paragraph_index=field["paragraph_index"],
        )
    return fields


def _content_controls(root: ET.Element, part_name: str) -> list[dict[str, Any]]:
    controls: list[dict[str, Any]] = []
    for index, sdt in enumerate(root.iter(f"{W_NS}sdt"), start=1):
        tag = sdt.find(f"./{W_NS}sdtPr/{W_NS}tag")
        alias = sdt.find(f"./{W_NS}sdtPr/{W_NS}alias")
        controls.append(
            {
                "index": index,
                "tag": _attr(tag, "val"),
                "alias": _attr(alias, "val"),
                "text": _visible_text(sdt),
                "part_name": part_name,
                "source_ref": f"{part_name}:sdt[{index}]",
            }
        )
    return controls


def _footnotes(package: ZipFile) -> list[dict[str, Any]]:
    try:
        root = ET.fromstring(package.read("word/footnotes.xml"))
    except (KeyError, ET.ParseError):
        return []
    footnotes: list[dict[str, Any]] = []
    for index, node in enumerate(root.findall(f"{W_NS}footnote"), start=1):
        footnote_id = _attr(node, "id")
        if footnote_id in {"-1", "0"}:
            continue
        text = _visible_text(node)
        if not text:
            continue
        footnotes.append(
            {
                "index": len(footnotes) + 1,
                "footnote_id": footnote_id,
                "text": text,
                "source_ref": f"word/footnotes.xml:footnote[{footnote_id}]",
            }
        )
    return footnotes


def _text_boxes(root: ET.Element, part_name: str) -> list[dict[str, Any]]:
    text_boxes: list[dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(root.iter(f"{W_NS}p"), start=1):
        for node in paragraph.iter():
            if node.tag.rsplit("}", 1)[-1] != "txbxContent":
                continue
            text = _visible_text(node)
            if not text:
                continue
            text_boxes.append(
                {
                    "index": len(text_boxes) + 1,
                    "paragraph_index": paragraph_index,
                    "text": text,
                    "source_ref": (
                        f"{part_name}:p[{paragraph_index}]"
                        f"/textbox[{len(text_boxes) + 1}]"
                    ),
                }
            )
    return text_boxes


def _images(
    root: ET.Element,
    part_name: str,
    relationships: dict[str, str],
    package: ZipFile,
) -> list[dict[str, Any]]:
    images: list[dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(root.iter(f"{W_NS}p"), start=1):
        for node in paragraph.iter():
            if node.tag.rsplit("}", 1)[-1] != "blip":
                continue
            relationship_id = (
                node.attrib.get(f"{R_NS}embed") or node.attrib.get(f"{R_NS}link") or ""
            )
            target = relationships.get(relationship_id, "")
            item: dict[str, Any] = {
                "index": len(images) + 1,
                "paragraph_index": paragraph_index,
                "relationship_id": relationship_id,
                "target": target,
                "source_ref": (
                    f"{part_name}:p[{paragraph_index}]"
                    f"/image[{len(images) + 1}]"
                ),
            }
            if target:
                try:
                    payload = package.read(target)
                except KeyError:
                    payload = b""
                if payload:
                    item["sha256"] = sha256_bytes(payload)
                    item["byte_count"] = len(payload)
            images.append(item)
    return images


def _normalize_instruction(instruction: str) -> str:
    return " ".join(instruction.strip().split())


def _field_type(instruction: str) -> str:
    normalized = instruction.strip().upper()
    if normalized.startswith("TOC"):
        return "TOC"
    if normalized.startswith("PAGEREF"):
        return "PAGEREF"
    if normalized.startswith("PAGE") or normalized.startswith("NUMPAGES"):
        return "PAGE"
    if normalized.startswith("HYPERLINK"):
        return "HYPERLINK"
    return normalized.split(" ", 1)[0] if normalized else "UNKNOWN"


def _breaks(root: ET.Element, part_name: str) -> list[dict[str, Any]]:
    breaks: list[dict[str, Any]] = []
    next_index = 1
    for paragraph_index, paragraph in enumerate(root.iter(f"{W_NS}p"), start=1):
        for node in paragraph.iter(f"{W_NS}br"):
            breaks.append(
                {
                    "index": next_index,
                    "kind": "break",
                    "paragraph_index": paragraph_index,
                    "type": node.attrib.get(f"{W_NS}type", "line"),
                    "source_ref": f"{part_name}:p[{paragraph_index}]/br[{next_index}]",
                }
            )
            next_index += 1
        paragraph_properties = paragraph.find(f"{W_NS}pPr")
        if (
            paragraph_properties is not None
            and paragraph_properties.find(f"{W_NS}sectPr") is not None
        ):
            breaks.append(
                {
                    "index": next_index,
                    "kind": "section",
                    "paragraph_index": paragraph_index,
                    "type": "section_properties",
                    "source_ref": f"{part_name}:p[{paragraph_index}]/sectPr",
                }
            )
            next_index += 1
    for _node in root.findall(f"./{W_NS}body/{W_NS}sectPr"):
        breaks.append(
            {
                "index": next_index,
                "kind": "section",
                "paragraph_index": None,
                "type": "section_properties",
                "source_ref": f"{part_name}:body/sectPr",
            }
        )
        next_index += 1
    return breaks


def _numbering_catalog(package: ZipFile) -> dict[str, Any]:
    catalog: dict[str, Any] = {
        "definitions": [],
        "definitions_by_num": {},
        "definitions_by_style": {},
        "style_numbering": {},
        "style_names": {},
    }
    abstract_levels: dict[str, dict[str, dict[str, Any]]] = {}
    nums: dict[str, str] = {}
    try:
        root = ET.fromstring(package.read("word/numbering.xml"))
    except (KeyError, ET.ParseError):
        return catalog

    for abstract in root.findall(f"{W_NS}abstractNum"):
        abstract_num_id = _attr(abstract, "abstractNumId")
        if not abstract_num_id:
            continue
        abstract_levels[abstract_num_id] = {}
        for level in abstract.findall(f"{W_NS}lvl"):
            ilvl = _attr(level, "ilvl") or "0"
            level_source_ref = (
                f"word/numbering.xml:abstractNum[{abstract_num_id}]/lvl[{ilvl}]"
            )
            abstract_levels[abstract_num_id][ilvl] = {
                "abstract_num_id": abstract_num_id,
                "ilvl": ilvl,
                "num_fmt": _attr(level.find(f"{W_NS}numFmt"), "val"),
                "lvl_text": _attr(level.find(f"{W_NS}lvlText"), "val"),
                "start": _int_or_none(_attr(level.find(f"{W_NS}start"), "val")),
                "suffix": _attr(level.find(f"{W_NS}suff"), "val"),
                "paragraph_style_id": _attr(level.find(f"{W_NS}pStyle"), "val"),
                "definition_source_ref": level_source_ref,
            }

    for num in root.findall(f"{W_NS}num"):
        num_id = _attr(num, "numId")
        abstract_num_id = _attr(num.find(f"{W_NS}abstractNumId"), "val")
        if num_id and abstract_num_id:
            nums[num_id] = abstract_num_id

    for num_id, abstract_num_id in nums.items():
        for ilvl, level in abstract_levels.get(abstract_num_id, {}).items():
            definition = {
                **level,
                "num_id": num_id,
                "num_source_ref": f"word/numbering.xml:num[{num_id}]",
                "source_ref": level["definition_source_ref"],
            }
            catalog["definitions"].append(definition)
            catalog["definitions_by_num"][(num_id, ilvl)] = definition
            style_id = definition.get("paragraph_style_id")
            if style_id:
                catalog["definitions_by_style"].setdefault(style_id, []).append(
                    definition
                )

    style_numbering, style_names = _style_numbering_catalog(package)
    catalog["style_numbering"] = style_numbering
    catalog["style_names"] = style_names
    return catalog


def _style_numbering_catalog(
    package: ZipFile,
) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    try:
        root = ET.fromstring(package.read("word/styles.xml"))
    except (KeyError, ET.ParseError):
        return {}, {}

    raw_styles: dict[str, dict[str, Any]] = {}
    style_names: dict[str, str] = {}
    for style in root.findall(f"{W_NS}style"):
        if _attr(style, "type") != "paragraph":
            continue
        style_id = _attr(style, "styleId")
        if not style_id:
            continue
        name = _attr(style.find(f"{W_NS}name"), "val")
        if name:
            style_names[style_id] = name
        based_on = _attr(style.find(f"{W_NS}basedOn"), "val")
        paragraph_properties = style.find(f"{W_NS}pPr")
        num_pr = (
            paragraph_properties.find(f"{W_NS}numPr")
            if paragraph_properties is not None
            else None
        )
        raw_styles[style_id] = {
            "style_id": style_id,
            "style_name": name,
            "based_on": based_on,
            "num_id": _attr(
                num_pr.find(f"{W_NS}numId") if num_pr is not None else None,
                "val",
            ),
            "ilvl": _attr(
                num_pr.find(f"{W_NS}ilvl") if num_pr is not None else None,
                "val",
            ),
            "source_ref": (
                f"word/styles.xml:style[{style_id}]/numPr"
                if num_pr is not None
                else ""
            ),
        }

    return (
        {
            style_id: _resolve_style_numbering(style_id, raw_styles, [])
            for style_id in raw_styles
        },
        style_names,
    )


def _resolve_style_numbering(
    style_id: str,
    raw_styles: dict[str, dict[str, Any]],
    seen: list[str],
) -> dict[str, Any]:
    if style_id in seen:
        return {}
    style = raw_styles.get(style_id)
    if not style:
        return {}
    base = (
        _resolve_style_numbering(
            str(style.get("based_on")),
            raw_styles,
            [*seen, style_id],
        )
        if style.get("based_on")
        else {}
    )
    num_id = _first_not_none(style.get("num_id"), base.get("num_id"))
    ilvl = _first_not_none(style.get("ilvl"), base.get("ilvl"))
    source_refs = [
        *base.get("source_refs", []),
        style.get("source_ref", ""),
    ]
    return {
        "style_id": style_id,
        "style_name": style.get("style_name"),
        "num_id": num_id,
        "ilvl": ilvl,
        "source_refs": [ref for ref in source_refs if ref],
    }


def _numbering_refs(
    root: ET.Element,
    numbering_catalog: dict[str, Any],
) -> list[dict[str, Any]]:
    refs: list[dict[str, Any]] = []
    paragraph_index = 0
    for paragraph in root.iter(f"{W_NS}p"):
        paragraph_index += 1
        paragraph_properties = paragraph.find(f"{W_NS}pPr")
        if paragraph_properties is None:
            continue
        paragraph_style_id = _attr(paragraph_properties.find(f"{W_NS}pStyle"), "val")
        num_pr = paragraph_properties.find(f"{W_NS}numPr")
        direct_num_id = _attr(
            num_pr.find(f"{W_NS}numId") if num_pr is not None else None,
            "val",
        )
        direct_ilvl = _attr(
            num_pr.find(f"{W_NS}ilvl") if num_pr is not None else None,
            "val",
        )
        style_numbering = numbering_catalog["style_numbering"].get(
            paragraph_style_id,
            {},
        )
        definition = _numbering_definition_for_paragraph(
            numbering_catalog,
            paragraph_style_id=paragraph_style_id,
            num_id=direct_num_id or style_numbering.get("num_id"),
            ilvl=direct_ilvl or style_numbering.get("ilvl"),
        )
        if num_pr is None and not definition:
            continue
        source_refs = []
        if num_pr is not None:
            source_refs.append(f"word/document.xml:p[{paragraph_index}]/numPr")
        source_refs.extend(style_numbering.get("source_refs", []))
        if definition and definition.get("source_ref"):
            source_refs.append(definition["source_ref"])
        source_kind = "direct"
        if num_pr is None:
            source_kind = "paragraph_style"
        elif style_numbering:
            source_kind = "direct_with_style"
        ref = {
            "paragraph_index": paragraph_index,
            "paragraph_style_id": paragraph_style_id,
            "paragraph_style_name": numbering_catalog["style_names"].get(
                paragraph_style_id,
                "",
            ),
            "source_kind": source_kind,
            "source_ref": f"word/document.xml:p[{paragraph_index}]/numPr"
            if num_pr is not None
            else f"word/document.xml:p[{paragraph_index}]/pStyle",
            "source_refs": _dedupe(source_refs),
            "text": _visible_text(paragraph),
        }
        if definition:
            ref.update(
                {
                    "num_id": definition.get("num_id"),
                    "abstract_num_id": definition.get("abstract_num_id"),
                    "ilvl": definition.get("ilvl"),
                    "num_fmt": definition.get("num_fmt"),
                    "lvl_text": definition.get("lvl_text"),
                    "paragraph_numbering_style_id": definition.get(
                        "paragraph_style_id"
                    ),
                    "definition_source_ref": definition.get("source_ref"),
                }
            )
        else:
            ref.update(
                {
                    "num_id": direct_num_id or style_numbering.get("num_id"),
                    "ilvl": direct_ilvl or style_numbering.get("ilvl"),
                    "num_fmt": None,
                    "lvl_text": None,
                }
            )
        refs.append(ref)
    return refs


def _numbering_definition_for_paragraph(
    numbering_catalog: dict[str, Any],
    *,
    paragraph_style_id: str | None,
    num_id: str | None,
    ilvl: str | None,
) -> dict[str, Any] | None:
    by_num = numbering_catalog["definitions_by_num"]
    if num_id:
        if ilvl and (num_id, ilvl) in by_num:
            return by_num[(num_id, ilvl)]
        style_candidates = [
            definition
            for definition in numbering_catalog["definitions"]
            if definition.get("num_id") == num_id
            and definition.get("paragraph_style_id") == paragraph_style_id
        ]
        if style_candidates:
            return _first_numbering_definition(style_candidates)
        num_candidates = [
            definition
            for definition in numbering_catalog["definitions"]
            if definition.get("num_id") == num_id
        ]
        if num_candidates:
            return _first_numbering_definition(num_candidates)

    if paragraph_style_id:
        style_candidates = numbering_catalog["definitions_by_style"].get(
            paragraph_style_id,
            [],
        )
        if ilvl:
            style_candidates = [
                definition
                for definition in style_candidates
                if definition.get("ilvl") == ilvl
            ]
        if style_candidates:
            return _first_numbering_definition(style_candidates)
    return None


def _first_numbering_definition(
    definitions: list[dict[str, Any]],
) -> dict[str, Any] | None:
    if not definitions:
        return None
    return sorted(
        definitions,
        key=lambda item: (
            _int_or_none(str(item.get("num_id"))) or 0,
            _int_or_none(str(item.get("ilvl"))) or 0,
        ),
    )[0]


def _dedupe(values: list[str]) -> list[str]:
    return [value for value in dict.fromkeys(values) if value]
