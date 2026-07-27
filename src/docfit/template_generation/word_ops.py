from __future__ import annotations

from copy import deepcopy
import re
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .refs import _cell_for_ref, _paragraph_for_ref


def _clear_cell(cell: _Cell) -> None:
    for paragraph in list(cell.paragraphs):
        _remove_paragraph(paragraph)
    if not cell.paragraphs:
        cell.add_paragraph("")


def _clear_runs_by_raw_run_ids(
    paragraph: Paragraph,
    source_ref: str | None,
    raw_run_ids: list[str],
) -> str | None:
    source_paragraph_index = _source_ref_paragraph_index(source_ref)
    run_indices: list[int] = []
    for raw_run_id in raw_run_ids:
        parsed = _parse_raw_run_id(raw_run_id)
        if parsed is None:
            return None
        paragraph_index, run_index = parsed
        if (
            source_paragraph_index is not None
            and paragraph_index != source_paragraph_index
        ):
            return None
        run_indices.append(run_index)
    runs = list(paragraph.runs)
    if not run_indices or any(index < 1 or index > len(runs) for index in run_indices):
        return None
    for run_index in sorted(set(run_indices)):
        runs[run_index - 1].text = ""
    joined = ",".join(raw_run_ids)
    return f"{source_ref}/runs[{joined}]:cleared"


def _replace_run_text_ranges(
    paragraph: Paragraph,
    source_ref: str | None,
    ranges: list[dict[str, Any]],
) -> str | None:
    source_paragraph_index = _source_ref_paragraph_index(source_ref)
    runs = list(paragraph.runs)
    grouped: dict[int, list[dict[str, Any]]] = {}
    for item in ranges:
        raw_run_id = str(item.get("raw_run_id") or "")
        parsed = _parse_raw_run_id(raw_run_id)
        if parsed is None:
            return None
        paragraph_index, run_index = parsed
        if (
            source_paragraph_index is not None
            and paragraph_index != source_paragraph_index
        ):
            return None
        if run_index < 1 or run_index > len(runs):
            return None
        grouped.setdefault(run_index, []).append(item)
    if not grouped:
        return None
    for run_index, replacements in grouped.items():
        run = runs[run_index - 1]
        text = run.text
        for item in sorted(replacements, key=lambda value: int(value.get("start") or 0), reverse=True):
            start = int(item.get("start") or 0)
            end = int(item.get("end") or start)
            if start < 0 or end < start or end > len(text):
                return None
            replacement = str(item.get("replacement") or "")
            text = f"{text[:start]}{replacement}{text[end:]}"
        run.text = text
    joined = ",".join(
        f"{item.get('raw_run_id')}:{item.get('start')}-{item.get('end')}"
        for item in ranges
    )
    return f"{source_ref}/char_ranges[{joined}]:replaced"


def _source_ref_paragraph_index(source_ref: str | None) -> int | None:
    if not source_ref:
        return None
    match = re.search(r"word/document\.xml:p\[(\d+)\]", source_ref)
    return int(match.group(1)) if match else None


def _parse_raw_run_id(raw_run_id: str) -> tuple[int, int] | None:
    match = re.search(r"(?:^|\.)p_(\d{4})\.r_(\d{3})$", raw_run_id)
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def _insert_marker(
    doc: Document,
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
    marker: str,
) -> str:
    target = _paragraph_for_ref(paragraph_map, source_ref)
    if target is not None:
        _insert_paragraph_after(target, marker)
        return f"{source_ref}/after:{marker}"
    target_cell = _cell_for_ref(doc, source_ref)
    if target_cell is not None:
        target_cell.add_paragraph(marker)
        return f"{source_ref}/p[last]:{marker}"
    return _append_marker(doc, marker)


def _insert_sdt(
    doc: Document,
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
    tag: str,
    *,
    alias: str | None = None,
    placeholder: str = "",
) -> str:
    target = _paragraph_for_ref(paragraph_map, source_ref)
    if target is not None:
        _insert_inline_sdt_at_paragraph_end(
            target,
            tag,
            alias=alias,
            placeholder=placeholder,
        )
        return f"{source_ref}/inline:sdt[{tag}]"
    target_cell = _cell_for_ref(doc, source_ref)
    if target_cell is not None:
        paragraph = target_cell.paragraphs[-1] if target_cell.paragraphs else target_cell.add_paragraph("")
        _insert_inline_sdt_at_paragraph_end(
            paragraph,
            tag,
            alias=alias,
            placeholder=placeholder,
        )
        return f"{source_ref}/p[last]/inline:sdt[{tag}]"
    return _append_sdt(doc, tag, alias=alias, placeholder=placeholder)


def _replace_run_text_ranges_with_inline_sdt(
    paragraph: Paragraph,
    source_ref: str | None,
    ranges: list[dict[str, Any]],
    tag: str,
    *,
    alias: str | None = None,
    placeholder: str = "",
) -> str | None:
    if not ranges:
        return None
    anchors: list[tuple[int, int]] = []
    for item in ranges:
        parsed = _parse_raw_run_id(str(item.get("raw_run_id") or ""))
        if parsed is None:
            return None
        _, run_index = parsed
        anchors.append((run_index, int(item.get("start") or 0)))
    anchor_run_index, anchor_offset = min(anchors)
    output_ref = _replace_run_text_ranges(paragraph, source_ref, ranges)
    if output_ref is None:
        return None
    runs = list(paragraph.runs)
    if anchor_run_index < 1 or anchor_run_index > len(runs):
        return None
    anchor_run = runs[anchor_run_index - 1]
    remaining_text = anchor_run.text
    if anchor_offset < 0 or anchor_offset > len(remaining_text):
        return None
    prefix = remaining_text[:anchor_offset]
    suffix = remaining_text[anchor_offset:]
    inline_sdt = _sdt_inline(
        tag,
        alias=alias,
        placeholder=placeholder,
        template_run=anchor_run,
    )
    if anchor_offset == 0:
        anchor_run._r.addprevious(inline_sdt)
    elif anchor_offset == len(remaining_text):
        anchor_run._r.addnext(inline_sdt)
    else:
        anchor_run.text = prefix
        suffix_element = deepcopy(anchor_run._r)
        suffix_run = Run(suffix_element, paragraph)
        suffix_run.text = suffix
        anchor_run._r.addnext(suffix_element)
        anchor_run._r.addnext(inline_sdt)
    return f"{output_ref}/inline:sdt[{tag}]"


def _insert_page_break_before(
    doc: Document,
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
) -> str | None:
    target = _paragraph_for_ref(paragraph_map, source_ref)
    if target is not None:
        existing_boundary = _existing_page_boundary_before(
            target,
            paragraph_map=paragraph_map,
            source_ref=source_ref,
        )
        if existing_boundary is not None:
            return existing_boundary
        target.paragraph_format.page_break_before = True
        return f"{source_ref}/pageBreakBefore"
    target_cell = _cell_for_ref(doc, source_ref)
    if target_cell is not None and target_cell.paragraphs:
        target_cell.paragraphs[0].paragraph_format.page_break_before = True
        return f"{source_ref}/p[1]/pageBreakBefore"
    return None


def _existing_page_boundary_before(
    target: Paragraph,
    *,
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
) -> str | None:
    """Return an existing effective page boundary without adding another one."""

    ppr = target._p.pPr
    if (
        ppr is not None
        and ppr.find(qn("w:pageBreakBefore")) is not None
    ):
        return f"{source_ref}/pageBreakBefore"

    index_by_element = {
        id(paragraph._p): index for index, paragraph in paragraph_map.items()
    }
    sibling = target._p.getprevious()
    while sibling is not None:
        if sibling.tag == qn("w:tbl"):
            break
        if sibling.tag != qn("w:p"):
            sibling = sibling.getprevious()
            continue
        sibling_index = index_by_element.get(id(sibling))
        sibling_ref = (
            f"word/document.xml:p[{sibling_index}]"
            if sibling_index is not None
            else None
        )
        sibling_ppr = sibling.find(qn("w:pPr"))
        sect_pr = (
            sibling_ppr.find(qn("w:sectPr"))
            if sibling_ppr is not None
            else None
        )
        if sect_pr is not None:
            section_type = sect_pr.find(qn("w:type"))
            section_value = (
                str(section_type.get(qn("w:val")) or "")
                if section_type is not None
                else "nextPage"
            )
            if section_value != "continuous":
                return (
                    f"{sibling_ref}/sectPr"
                    if sibling_ref is not None
                    else "word/document.xml:sectPr"
                )
        page_break = sibling.find(".//" + qn("w:br"))
        if (
            page_break is not None
            and str(page_break.get(qn("w:type")) or "") == "page"
        ):
            return (
                f"{sibling_ref}/br"
                if sibling_ref is not None
                else "word/document.xml:br"
            )
        visible_text = "".join(sibling.itertext()).strip()
        has_visible_object = bool(
            sibling.findall(".//" + qn("w:drawing"))
            or sibling.findall(".//" + qn("w:object"))
        )
        if visible_text or has_visible_object:
            break
        sibling = sibling.getprevious()
    return None


def _insert_section_break_before(
    doc: Document,
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
) -> str | None:
    target = _paragraph_for_ref(paragraph_map, source_ref)
    if target is None:
        return None
    target_index = _source_ref_paragraph_index(source_ref)
    if target_index is not None and target_index > 1:
        previous = paragraph_map.get(target_index - 1)
        if previous is not None:
            paragraph_properties = previous._p.get_or_add_pPr()
            if paragraph_properties.find(qn("w:sectPr")) is None:
                paragraph_properties.append(_next_page_section_properties(doc))
            return f"{source_ref}/before:sectPr"
    boundary = _insert_paragraph_before(target, "")
    paragraph_properties = boundary._p.get_or_add_pPr()
    paragraph_properties.append(_next_page_section_properties(doc))
    return f"{source_ref}/before:sectPr"


def _set_keep_together_for_refs(
    doc: Document,
    paragraph_map: dict[int, Paragraph],
    source_refs: list[str],
    *,
    output_ref_by_paragraph: dict[int, str] | None = None,
) -> list[str]:
    output_refs: list[str] = []
    paragraph_targets: list[tuple[str, Paragraph]] = []
    for source_ref in source_refs:
        target = _paragraph_for_ref(paragraph_map, source_ref)
        if target is not None:
            paragraph_targets.append((source_ref, target))
            continue
        target_cell = _cell_for_ref(doc, source_ref)
        if target_cell is not None:
            cell_ref = _cell_ref_from_source_ref(source_ref) or source_ref
            for index, paragraph in enumerate(target_cell.paragraphs, start=1):
                paragraph_targets.append((f"{cell_ref}/p[{index}]", paragraph))
        for row, row_ref in _table_rows_for_ref(doc, source_ref):
            _ensure_row_cant_split(row)
            output_refs.append(f"{row_ref}/cantSplit")

    for index, (target_ref, paragraph) in enumerate(paragraph_targets):
        output_target_ref = (
            output_ref_by_paragraph.get(id(paragraph._p))
            if output_ref_by_paragraph is not None
            else target_ref
        )
        if output_target_ref is None and "word/document.xml:tbl[" in target_ref:
            output_target_ref = target_ref
        if not output_target_ref:
            continue
        paragraph.paragraph_format.keep_together = True
        is_last_unit_paragraph = index == len(paragraph_targets) - 1
        if is_last_unit_paragraph:
            output_refs.append(f"{output_target_ref}/keepLines")
        else:
            paragraph.paragraph_format.keep_with_next = True
            output_refs.append(f"{output_target_ref}/keepWithNext+keepLines")
    return _dedupe_str(output_refs)


def _table_rows_for_ref(doc: Document, source_ref: str | None) -> list[tuple[Any, str]]:
    if not source_ref:
        return []
    text = str(source_ref)
    row_match = re.search(r"word/document\.xml:tbl\[(\d+)\]/tr\[(\d+)\]", text)
    if row_match:
        table_index = int(row_match.group(1)) - 1
        row_index = int(row_match.group(2)) - 1
        try:
            row = doc.tables[table_index].rows[row_index]
        except IndexError:
            return []
        return [
            (
                row,
                f"word/document.xml:tbl[{table_index + 1}]/tr[{row_index + 1}]",
            )
        ]
    table_match = re.search(r"word/document\.xml:tbl\[(\d+)\](?!/tr)", text)
    if not table_match:
        return []
    table_index = int(table_match.group(1)) - 1
    try:
        table = doc.tables[table_index]
    except IndexError:
        return []
    return [
        (row, f"word/document.xml:tbl[{table_index + 1}]/tr[{row_index}]")
        for row_index, row in enumerate(table.rows, start=1)
    ]


def _cell_ref_from_source_ref(source_ref: str | None) -> str | None:
    if not source_ref:
        return None
    match = re.search(
        r"(word/document\.xml:tbl\[\d+\]/tr\[\d+\]/tc\[\d+\])",
        str(source_ref),
    )
    return match.group(1) if match else None


def _ensure_row_cant_split(row: Any) -> None:
    row_properties = row._tr.get_or_add_trPr()
    cant_split = row_properties.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        row_properties.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def _dedupe_str(values: list[Any]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = str(value or "").strip()
        if not text or text in seen:
            continue
        result.append(text)
        seen.add(text)
    return result


def _insert_styled_paragraph_before(
    paragraph_map: dict[int, Paragraph],
    source_ref: str | None,
    text: str,
    *,
    page_break_before: bool = False,
) -> str | None:
    target = _paragraph_for_ref(paragraph_map, source_ref)
    if target is None:
        return None
    inserted = _insert_paragraph_before(target, text)
    inserted.style = target.style
    inserted.alignment = target.alignment
    inserted.paragraph_format.page_break_before = page_break_before
    return f"{source_ref}/before:{text}"


def _next_page_section_properties(doc: Document) -> Any:
    section_properties = deepcopy(doc.sections[0]._sectPr)
    for child in list(section_properties):
        if child.tag == qn("w:type"):
            section_properties.remove(child)
    section_type = OxmlElement("w:type")
    section_type.set(qn("w:val"), "nextPage")
    section_properties.insert(0, section_type)
    return section_properties


def _insert_paragraph_before(paragraph: Paragraph, text: str) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._p.addprevious(new_element)
    new_paragraph = Paragraph(new_element, paragraph._parent)
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def _insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = Paragraph(new_element, paragraph._parent)
    new_paragraph.add_run(text)
    return new_paragraph


def _insert_sdt_after(
    paragraph: Paragraph,
    tag: str,
    *,
    alias: str | None = None,
    placeholder: str = "",
) -> None:
    paragraph._p.addnext(_sdt_block(tag, alias=alias, placeholder=placeholder))


def _insert_inline_sdt_at_paragraph_end(
    paragraph: Paragraph,
    tag: str,
    *,
    alias: str | None = None,
    placeholder: str = "",
) -> None:
    template_run = paragraph.runs[-1] if paragraph.runs else paragraph.add_run("")
    template_run._r.addnext(
        _sdt_inline(
            tag,
            alias=alias,
            placeholder=placeholder,
            template_run=template_run,
        )
    )


def _sdt_inline(
    tag: str,
    *,
    alias: str | None,
    placeholder: str,
    template_run: Run,
) -> Any:
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")
    alias_node = OxmlElement("w:alias")
    alias_node.set(qn("w:val"), alias or tag)
    tag_node = OxmlElement("w:tag")
    tag_node.set(qn("w:val"), tag)
    sdt_pr.append(alias_node)
    sdt_pr.append(tag_node)
    sdt_content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    if template_run._r.rPr is not None:
        run.append(deepcopy(template_run._r.rPr))
    text = OxmlElement("w:t")
    if placeholder:
        text.text = placeholder
    run.append(text)
    sdt_content.append(run)
    sdt.append(sdt_pr)
    sdt.append(sdt_content)
    return sdt


def _append_sdt(doc: Document, tag: str, *, alias: str | None = None, placeholder: str = "") -> str:
    block = _sdt_block(tag, alias=alias, placeholder=placeholder)
    section_properties = doc.element.body.sectPr
    if section_properties is not None:
        section_properties.addprevious(block)
    else:
        doc.element.body.append(block)
    return f"word/document.xml:sdt[{tag}]"


def _append_sdt_to_parent(
    parent: Any,
    tag: str,
    *,
    alias: str | None = None,
    placeholder: str = "",
) -> None:
    parent.append(_sdt_block(tag, alias=alias, placeholder=placeholder))


def _sdt_block(tag: str, *, alias: str | None = None, placeholder: str = "") -> Any:
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")
    alias_node = OxmlElement("w:alias")
    alias_node.set(qn("w:val"), alias or tag)
    tag_node = OxmlElement("w:tag")
    tag_node.set(qn("w:val"), tag)
    sdt_pr.append(alias_node)
    sdt_pr.append(tag_node)
    sdt_content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    if placeholder:
        text.text = placeholder
    run.append(text)
    paragraph.append(run)
    sdt_content.append(paragraph)
    sdt.append(sdt_pr)
    sdt.append(sdt_content)
    return sdt


def _append_marker(doc: Document, marker: str) -> str:
    doc.add_paragraph(marker)
    return f"word/document.xml:p[{_body_paragraph_count(doc)}]"


def _body_paragraph_count(doc: Document) -> int:
    return sum(1 for _ in doc.element.body.iter(qn("w:p")))


def _find_sdt_tag_ref(doc: Document, tag: str) -> str | None:
    for index, sdt in enumerate(doc.element.body.iter(qn("w:sdt")), start=1):
        tag_node = next(iter(sdt.iter(qn("w:tag"))), None)
        if tag_node is not None and tag_node.get(qn("w:val")) == tag:
            return f"word/document.xml:sdt[{index}:{tag}]"
    return None


def _remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)
