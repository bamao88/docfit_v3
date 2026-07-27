from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .constants import BODY_SLOT_MARKER
from .identity_resolver import L1IdentityResolver
from .refs import _cell_for_ref, _paragraph_for_ref, _paragraph_map_by_ooxml_index
from .text_utils import _dedupe_by_key
from .word_ops import (
    _append_sdt,
    _clear_runs_by_raw_run_ids,
    _clear_cell,
    _find_sdt_tag_ref,
    _insert_page_break_before,
    _insert_section_break_before,
    _insert_sdt,
    _insert_styled_paragraph_before,
    _remove_paragraph,
    _replace_run_text_ranges,
    _replace_run_text_ranges_with_inline_sdt,
    _set_keep_together_for_refs,
)


_LAYOUT_ACTION_TYPES = {
    "insert_page_break_before_unit",
    "insert_section_break_before_unit",
    "set_keep_together_unit",
}


def execute_template_generation_plan(
    source_template_docx: Path,
    generated_template_docx: Path,
    plan: dict[str, Any],
    *,
    l1_input_contract: dict[str, Any],
) -> dict[str, Any]:
    resolver = L1IdentityResolver(l1_input_contract)
    source_precondition = resolver.validate_source_package(source_template_docx)
    if source_precondition["status"] != "PASS":
        return {
            "actions_executed": [],
            "actions_requiring_review": [
                _needs_review(action, str(source_precondition["reason"]))
                for action in plan.get("actions", [])
            ],
            "slots": [],
            "generated_fields": [],
            "page_breaks": [],
            "section_breaks": [],
            "keep_together": [],
            "synthesized_texts": [],
            "identity_resolution": {
                "source_precondition": source_precondition,
                "action_preconditions": [],
            },
        }
    generated_template_docx.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source_template_docx, generated_template_docx)
    doc = Document(generated_template_docx)
    paragraph_map = _paragraph_map_by_ooxml_index(doc)
    executed: list[dict[str, Any]] = []
    review: list[dict[str, Any]] = []
    slots: list[dict[str, Any]] = []
    generated_fields: list[dict[str, Any]] = []
    paragraphs_to_remove: list[Paragraph] = []
    action_preconditions: list[dict[str, Any]] = []
    deferred_layout_actions: list[dict[str, Any]] = []

    for action in plan.get("actions", []):
        action_type = action.get("action_type")
        if action_type in _LAYOUT_ACTION_TYPES:
            deferred_layout_actions.append(action)
            continue
        action_precondition = resolver.validate_action(action)
        action_preconditions.append(action_precondition)
        if action_precondition["status"] != "PASS":
            review.append(
                _needs_review(
                    action,
                    "; ".join(action_precondition["errors"]),
                )
            )
            continue
        if action_type == "copy_source_docx":
            executed.append(_executed(action, output_ref=str(generated_template_docx)))
        elif action_type == "remove_instruction_text":
            target = _paragraph_for_ref(paragraph_map, action.get("source_ref"))
            target_cell = _cell_for_ref(doc, action.get("source_ref"))
            char_ranges = [
                item
                for item in action.get("affected_char_ranges", [])
                if isinstance(item, dict)
            ]
            if char_ranges and target is not None and _single_source_seq_action(action):
                output_ref = _replace_run_text_ranges(
                    target,
                    action.get("source_ref"),
                    char_ranges,
                )
                if output_ref is None:
                    review.append(_needs_review(action, "source char ranges not found for span removal"))
                    continue
                executed.append(_executed(action, output_ref=output_ref))
                continue
            raw_run_ids = [
                str(raw_run_id)
                for raw_run_id in action.get("affected_raw_run_ids", [])
                if raw_run_id
            ]
            if raw_run_ids and target is not None and _single_source_seq_action(action):
                output_ref = _clear_runs_by_raw_run_ids(
                    target,
                    action.get("source_ref"),
                    raw_run_ids,
                )
                if output_ref is None:
                    review.append(_needs_review(action, "source runs not found for run-level removal"))
                    continue
                executed.append(_executed(action, output_ref=output_ref))
                continue
            if target is None and target_cell is None:
                review.append(_needs_review(action, "source node not found"))
                continue
            if target is not None:
                paragraphs_to_remove.append(target)
            if target_cell is not None:
                _clear_cell(target_cell)
            executed.append(_executed(action, output_ref=action.get("source_ref")))
        elif action_type == "create_fillable_slot":
            tag = _sdt_tag(action)
            output_ref = _insert_sdt(
                doc,
                paragraph_map,
                action.get("source_ref"),
                tag,
                alias=_sdt_alias(action),
            )
            slot = _slot_from_action(action, sdt_tag=tag, output_ref=output_ref)
            slots.append(slot)
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "replace_span_with_slot":
            target = _paragraph_for_ref(paragraph_map, action.get("source_ref"))
            char_ranges = [
                item
                for item in action.get("affected_char_ranges", [])
                if isinstance(item, dict)
            ]
            if target is None:
                review.append(_needs_review(action, "source node not found"))
                continue
            tag = _span_sdt_tag(action)
            if char_ranges:
                output_ref = _replace_run_text_ranges_with_inline_sdt(
                    target,
                    action.get("source_ref"),
                    char_ranges,
                    tag,
                    alias=_sdt_alias(action),
                )
            else:
                output_ref = _clear_runs_by_raw_run_ids(
                    target,
                    action.get("source_ref"),
                    [
                        str(raw_run_id)
                        for raw_run_id in action.get("affected_raw_run_ids", [])
                        if raw_run_id
                    ],
                )
            if output_ref is None:
                review.append(_needs_review(action, "source span not found for replacement"))
                continue
            slot_ref = output_ref
            if not char_ranges:
                slot_ref = _insert_sdt(
                    doc,
                    paragraph_map,
                    action.get("source_ref"),
                    tag,
                    alias=_sdt_alias(action),
                )
            slots.append(_slot_from_action(action, sdt_tag=tag, output_ref=slot_ref))
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "create_generated_field_placeholder":
            tag = _generated_tag(action)
            output_ref = _insert_sdt(
                doc,
                paragraph_map,
                action.get("source_ref"),
                tag,
                alias=_sdt_alias(action),
            )
            generated_fields.append(
                {
                    "field_id": tag,
                    "unit_id": action.get("unit_id"),
                    "element_id": action.get("element_id"),
                    "sdt_tag": tag,
                    "output_ref": output_ref,
                    "source_seq_refs": action.get("affected_source_seq_refs", []),
                }
            )
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "insert_fixed_text":
            text = str(action.get("target_ref") or "")
            output_ref = _insert_marker(
                doc,
                paragraph_map,
                action.get("source_ref"),
                text,
            )
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "insert_synthetic_unit_title_before":
            text = str(action.get("target_ref") or "")
            output_ref = _insert_styled_paragraph_before(
                paragraph_map,
                action.get("source_ref"),
                text,
                page_break_before=True,
            )
            if output_ref is None:
                review.append(_needs_review(action, "source node not found"))
                continue
            executed.append(_executed(action, output_ref=output_ref))
        elif action_type == "ensure_body_slot":
            tag = "slot_body_start"
            existing_ref = _find_sdt_tag_ref(doc, tag)
            output_ref = existing_ref or _append_sdt(
                doc,
                tag,
                alias="正文内容",
            )
            slots.append(
                {
                    "slot_id": "slot_body_start",
                    "owner_scope": "document_body",
                    "element_id": "slot_body_start",
                    "kind": "body_content",
                    "sdt_tag": tag,
                    "output_ref": output_ref,
                    "required": True,
                }
            )
            executed.append(_executed(action, output_ref=output_ref))
        else:
            review.append(_needs_review(action, f"unsupported action type: {action_type}"))

    for paragraph in dict.fromkeys(paragraphs_to_remove):
        _remove_paragraph(paragraph)
    _strip_internal_markers(doc)
    _execute_deferred_layout_actions(
        doc=doc,
        paragraph_map=paragraph_map,
        layout_actions=deferred_layout_actions,
        resolver=resolver,
        action_preconditions=action_preconditions,
        executed=executed,
        review=review,
    )
    doc.save(generated_template_docx)
    return {
        "actions_executed": executed,
        "actions_requiring_review": review,
        "identity_resolution": {
            "source_precondition": source_precondition,
            "action_preconditions": action_preconditions,
        },
        "slots": _dedupe_by_key(slots, "slot_id"),
        "generated_fields": generated_fields,
        "page_breaks": [
            {
                **_page_action_summary(action),
            }
            for action in executed
            if action.get("action_type") == "insert_page_break_before_unit"
        ],
        "section_breaks": [
            {
                **_page_action_summary(action),
            }
            for action in executed
            if action.get("action_type") == "insert_section_break_before_unit"
        ],
        "keep_together": [
            {
                **_page_action_summary(action),
                "keep_scope": action.get("keep_scope"),
            }
            for action in executed
            if action.get("action_type") == "set_keep_together_unit"
        ],
        "synthesized_texts": [
            {
                "unit_id": action.get("unit_id"),
                "element_id": action.get("element_id"),
                "text": action.get("target_ref"),
                "output_ref": action.get("output_ref"),
            }
            for action in executed
            if action.get("action_type")
            in {"insert_fixed_text", "insert_synthetic_unit_title_before"}
        ],
    }


def _sdt_tag(action: dict[str, Any]) -> str:
    return f"{action.get('unit_id')}.{action.get('element_id')}"


def _span_sdt_tag(action: dict[str, Any]) -> str:
    return f"{action.get('unit_id')}.{action.get('element_id')}"


def _single_source_seq_action(action: dict[str, Any]) -> bool:
    return len(action.get("affected_source_seq_refs", []) or []) == 1


def _generated_tag(action: dict[str, Any]) -> str:
    return f"generated.{action.get('unit_id')}.{action.get('element_id')}"


def _sdt_alias(action: dict[str, Any]) -> str:
    return str(action.get("element_name") or action.get("target_ref") or _sdt_tag(action))


def _slot_from_action(
    action: dict[str, Any],
    *,
    sdt_tag: str,
    output_ref: str,
) -> dict[str, Any]:
    return {
        "slot_id": f"{action.get('unit_id')}.{action.get('element_id')}",
        "unit_id": action.get("unit_id"),
        "element_id": action.get("element_id"),
        "span_id": action.get("span_id"),
        "kind": "body_content",
        "sdt_tag": sdt_tag,
        "output_ref": output_ref,
        "source_seq_refs": action.get("affected_source_seq_refs", []),
        "required": True,
    }


def _executed(action: dict[str, Any], *, output_ref: str | None) -> dict[str, Any]:
    return {
        **action,
        "output_ref": output_ref,
        "status": "executed",
    }


def _needs_review(action: dict[str, Any], reason: str) -> dict[str, Any]:
    return {
        **action,
        "status": "needs_review",
        "reason": reason,
    }


def _execute_deferred_layout_actions(
    *,
    doc: Document,
    paragraph_map: dict[int, Paragraph],
    layout_actions: list[dict[str, Any]],
    resolver: L1IdentityResolver,
    action_preconditions: list[dict[str, Any]],
    executed: list[dict[str, Any]],
    review: list[dict[str, Any]],
) -> None:
    boundary_actions = [
        action
        for action in layout_actions
        if action.get("action_type")
        in {"insert_page_break_before_unit", "insert_section_break_before_unit"}
    ]
    keep_actions = [
        action
        for action in layout_actions
        if action.get("action_type") == "set_keep_together_unit"
    ]
    boundary_executed: list[dict[str, Any]] = []
    for action in boundary_actions:
        action_precondition = resolver.validate_action(action)
        action_preconditions.append(action_precondition)
        if action_precondition["status"] != "PASS":
            review.append(
                _needs_review(
                    action,
                    "; ".join(action_precondition["errors"]),
                )
            )
            continue
        action_type = action.get("action_type")
        if action_type == "insert_page_break_before_unit":
            output_ref = _insert_page_break_before(
                doc,
                paragraph_map,
                action.get("source_ref"),
            )
        else:
            output_ref = _insert_section_break_before(
                doc,
                paragraph_map,
                action.get("source_ref"),
            )
        if output_ref is None:
            review.append(_needs_review(action, "source node not found"))
            continue
        executed_action = action
        if (
            action_type == "insert_page_break_before_unit"
            and not str(output_ref).endswith("/pageBreakBefore")
        ):
            executed_action = {
                **action,
                "action_type": "preserve_existing_page_boundary",
                "requested_action_type": action_type,
            }
        boundary_executed.append(
            _executed(executed_action, output_ref=output_ref)
        )

    executed.extend(
        _with_final_layout_output_ref(doc, paragraph_map, action)
        for action in boundary_executed
    )
    output_ref_by_paragraph = _paragraph_ref_by_element_id(doc)
    for action in keep_actions:
        action_precondition = resolver.validate_action(action)
        action_preconditions.append(action_precondition)
        if action_precondition["status"] != "PASS":
            review.append(
                _needs_review(
                    action,
                    "; ".join(action_precondition["errors"]),
                )
            )
            continue
        source_refs = [
            str(ref)
            for ref in action.get("affected_source_refs", []) or []
            if ref
        ]
        if not source_refs and action.get("source_ref"):
            source_refs = [str(action.get("source_ref"))]
        output_refs = _set_keep_together_for_refs(
            doc,
            paragraph_map,
            source_refs,
            output_ref_by_paragraph=output_ref_by_paragraph,
        )
        if not output_refs:
            review.append(_needs_review(action, "no paragraph refs found for keep_together"))
            continue
        executed.append(_executed(action, output_ref=";".join(output_refs)))


def _with_final_layout_output_ref(
    doc: Document,
    paragraph_map: dict[int, Paragraph],
    action: dict[str, Any],
) -> dict[str, Any]:
    action_type = action.get("action_type")
    suffix = {
        "insert_page_break_before_unit": "pageBreakBefore",
        "insert_section_break_before_unit": "before:sectPr",
    }.get(str(action_type))
    if suffix is None:
        return action
    target = _paragraph_for_ref(paragraph_map, action.get("source_ref"))
    if target is None:
        return action
    output_ref_by_paragraph = _paragraph_ref_by_element_id(doc)
    final_ref = output_ref_by_paragraph.get(id(target._p))
    if not final_ref:
        return action
    return {**action, "output_ref": f"{final_ref}/{suffix}"}


def _paragraph_ref_by_element_id(doc: Document) -> dict[int, str]:
    return {
        id(paragraph_element): f"word/document.xml:p[{index}]"
        for index, paragraph_element in enumerate(
            doc.element.body.iter(qn("w:p")),
            start=1,
        )
    }


def _page_action_summary(action: dict[str, Any]) -> dict[str, Any]:
    return {
        "unit_id": action.get("unit_id"),
        "page_policy_unit_id": action.get("page_policy_unit_id"),
        "satisfies_page_policy_unit_ids": action.get(
            "satisfies_page_policy_unit_ids",
            [],
        ),
        "supersedes_action_ids": action.get("supersedes_action_ids", []),
        "source_ref": action.get("source_ref"),
        "output_ref": action.get("output_ref"),
        "action_id": action.get("action_id"),
        "page_policy": action.get("page_policy"),
        "page_policy_origin": action.get("page_policy_origin"),
        "page_policy_confidence": action.get("page_policy_confidence"),
        "page_policy_evidence_refs": action.get("page_policy_evidence_refs", []),
        "page_policy_proposal_ids": action.get("page_policy_proposal_ids", []),
        "agent_proposal_id": action.get("agent_proposal_id"),
        "status": action.get("status"),
    }


def _strip_internal_markers(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if "[[DOCFIT_" in paragraph.text:
            paragraph.text = paragraph.text.replace(BODY_SLOT_MARKER, "").strip()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "[[DOCFIT_" in paragraph.text:
                        paragraph.text = paragraph.text.replace(BODY_SLOT_MARKER, "").strip()
