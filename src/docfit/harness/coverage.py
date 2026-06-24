from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document

from docfit.core.io import read_json, sha256_file
from docfit.core.models import Finding, make_finding
from docfit.core.status import Status, merge_statuses
from docfit.ooxml.package import is_valid_docx, read_document_xml
from docfit.harness.baselines import load_baseline_file
from docfit.harness.profiles import (
    BOOTSTRAP_PROFILE,
    BOOTSTRAP_TEMPLATE_DOCX,
    REAL_CORE_PROFILE,
    REAL_CORE_SCHOOLS,
    REAL_CORE_STUDENTS,
    get_eval_cases_for_profile,
    get_eval_profile,
)
from docfit.harness.product_quality import (
    audit_e2e_case,
    business_acceptance_coverage,
    missing_e2e_audit_inputs,
)
from docfit.harness.word_evidence import verify_word_image_evidence


def bootstrap_required_capabilities() -> list[str]:
    return BOOTSTRAP_PROFILE.all_required_capabilities()


def validate_standard_coverage_requirements(
    profile: str | None,
    required_capabilities: list[str],
    *,
    stage: str = "standards",
    start_index: int = 1,
) -> list[Finding]:
    profile_def = get_eval_profile(profile or "")
    if profile_def is None:
        return [
            make_finding(
                start_index,
                stage,
                Status.UNKNOWN,
                "unknown_coverage_profile",
                "Signed standard references an unknown coverage profile",
                ", ".join(
                    sorted([BOOTSTRAP_PROFILE.profile_id, REAL_CORE_PROFILE.profile_id])
                ),
                profile or "missing",
                root_cause_bucket="coverage_gap",
            )
        ]

    expected = set(profile_def.all_required_capabilities())
    actual = set(required_capabilities)
    missing = sorted(expected - actual)
    unknown = sorted(actual - expected)
    if not missing and not unknown:
        return []

    details = []
    if missing:
        details.append("missing: " + ", ".join(missing))
    if unknown:
        details.append("unknown: " + ", ".join(unknown))
    return [
        make_finding(
            start_index,
            stage,
            Status.UNKNOWN,
            "coverage_requirements_drift",
            "Signed standard coverage requirements do not match the declared capability profile",
            ", ".join(sorted(expected)),
            "; ".join(details),
            root_cause_bucket="coverage_gap",
        )
    ]


def coverage_gate_findings(
    stage: str,
    coverage: dict[str, Any],
    required_capabilities: list[str],
    *,
    start_index: int = 1,
) -> list[Finding]:
    missing = [
        capability
        for capability in required_capabilities
        if coverage.get(capability) is not True
    ]
    if not missing:
        return []
    actual = [
        f"{capability}={coverage.get(capability, 'missing')}"
        for capability in missing
    ]
    return [
        make_finding(
            start_index,
            stage,
            Status.UNKNOWN,
            "coverage_insufficient",
            f"{stage} coverage is insufficient for the active contract",
            "all required capabilities covered",
            ", ".join(actual),
            affected_ids=missing,
            root_cause_bucket="coverage_gap",
        )
    ]


def _safe_document(path: Path):
    if not path.exists() or not is_valid_docx(path):
        return None
    try:
        return Document(path)
    except Exception:
        return None


def _has_slot_marker(path: Path) -> bool:
    if not path.exists() or not is_valid_docx(path):
        return False
    try:
        return "[[DOCFIT_SLOT:body]]" in read_document_xml(path)
    except Exception:
        return False


def _has_required_content_hashes(path: Path) -> bool:
    if not path.exists():
        return False
    try:
        return bool(read_json(path).get("required_content_hashes"))
    except Exception:
        return False


def evaluate_bootstrap_coverage(root: Path) -> tuple[dict[str, Any], list[Finding]]:
    template_docx = root / BOOTSTRAP_TEMPLATE_DOCX
    student_docx = root / "inputs/students/bootstrap-demo-pass/raw/source_document.docx"
    expected_paths = BOOTSTRAP_PROFILE.expected_paths(root)
    expected_placement = expected_paths["placement_plan"]
    expected_snapshot = expected_paths["feature_snapshot"]
    template_doc = _safe_document(template_docx)
    student_doc = _safe_document(student_docx)
    has_visible_content = bool(
        student_doc
        and (
            any(paragraph.text.strip() for paragraph in student_doc.paragraphs)
            or student_doc.tables
        )
    )
    capability_checks = {
        "template.docx_openable": is_valid_docx(template_docx),
        "template.required_regions": _has_slot_marker(template_docx),
        "template.required_slots": _has_slot_marker(template_docx),
        "template.styles_inventory": bool(template_doc and list(template_doc.styles)),
        "content.visible_paragraphs": bool(
            student_doc
            and any(paragraph.text.strip() for paragraph in student_doc.paragraphs)
        ),
        "content.visible_tables": bool(student_doc and student_doc.tables),
        "content.reading_order": has_visible_content,
        "content.stable_ids": has_visible_content,
        "placement.no_silent_drop": expected_placement.exists(),
        "placement.slot_compatibility": expected_placement.exists(),
        "placement.required_slots": _has_slot_marker(template_docx),
        "render.valid_docx_package": is_valid_docx(template_docx) and is_valid_docx(student_docx),
        "render.plan_coverage": expected_placement.exists(),
        "render.feature_snapshot": expected_snapshot.exists(),
        "render.content_hash_coverage": _has_required_content_hashes(expected_snapshot),
    }
    missing = [
        capability
        for capability, is_covered in capability_checks.items()
        if not is_covered
    ]
    required_flat = BOOTSTRAP_PROFILE.all_required_capabilities()
    covered = sorted(
        capability
        for capability in required_flat
        if capability_checks.get(capability) is True
    )
    status = Status.UNKNOWN if missing else Status.PASS
    report = {
        "profile": BOOTSTRAP_PROFILE.profile_id,
        "status": status.value,
        "required": len(required_flat),
        "covered": len(covered),
        "missing": missing,
    }
    findings: list[Finding] = []
    if missing:
        findings.append(
            make_finding(
                1,
                "coverage",
                Status.UNKNOWN,
                "coverage_insufficient",
                "bootstrap-core input and expected artifact coverage is incomplete",
                "all required bootstrap capabilities have input assets",
                ", ".join(missing),
                root_cause_bucket="coverage_gap",
            )
        )
    return report, findings


def evaluate_profile_coverage(
    root: Path,
    profile_id: str,
) -> tuple[dict[str, Any], list[Finding]]:
    if profile_id == BOOTSTRAP_PROFILE.profile_id:
        return evaluate_bootstrap_coverage(root)
    if profile_id == REAL_CORE_PROFILE.profile_id:
        return evaluate_real_core_coverage(root)
    return (
        {
            "profile": profile_id,
            "status": Status.UNKNOWN.value,
            "required": 0,
            "covered": 0,
            "missing": [f"unknown profile {profile_id}"],
        },
        [
            make_finding(
                1,
                "coverage",
                Status.UNKNOWN,
                "unknown_coverage_profile",
                "Coverage profile is not registered",
                ", ".join(
                    sorted([BOOTSTRAP_PROFILE.profile_id, REAL_CORE_PROFILE.profile_id])
                ),
                profile_id,
                root_cause_bucket="coverage_gap",
            )
        ],
    )


def evaluate_real_core_coverage(root: Path) -> tuple[dict[str, Any], list[Finding]]:
    findings: list[Finding] = []
    next_index = 1
    cases = get_eval_cases_for_profile(REAL_CORE_PROFILE.profile_id)
    product_quality_findings: list[Finding] = []
    product_quality_report: dict[str, Any] = {
        "checked_cases": [],
        "failing_cases": [],
        "missing_cases": [],
        "coverage": business_acceptance_coverage([]),
    }
    case_counts = {
        "template": sum(1 for case in cases if case.stage == "template"),
        "content": sum(1 for case in cases if case.stage == "content"),
        "e2e": sum(1 for case in cases if case.stage == "e2e"),
    }

    case_registry_path = root / "eval_profiles/real-core-v0/profile.yaml"
    if not case_registry_path.exists():
        findings.append(
            make_finding(
                next_index,
                "coverage",
                Status.UNKNOWN,
                "missing_profile_case_registry",
                "real-core-v0 requires a checked-in case registry",
                str(case_registry_path),
                "missing",
                root_cause_bucket="coverage_gap",
            )
        )
        next_index += 1

    required_source_files = _real_core_source_files()
    missing_source_files = [
        str(path)
        for path in required_source_files
        if not (root / path).exists()
    ]
    if missing_source_files:
        findings.append(
            make_finding(
                next_index,
                "coverage",
                Status.UNKNOWN,
                "missing_real_core_source_evidence",
                "real-core-v0 requires the fixed source evidence set",
                "all fixed templates, students, and review evidence exist",
                ", ".join(missing_source_files),
                affected_ids=missing_source_files,
                root_cause_bucket="source_evidence_missing",
            )
        )
        next_index += 1

    for school in REAL_CORE_SCHOOLS:
        school_id = str(school["school_id"])
        standard_path = (
            root / "standards/targets" / school_id / "v1/target.standard.yaml"
        )
        if not standard_path.exists():
            findings.append(
                make_finding(
                    next_index,
                    "coverage",
                    Status.UNKNOWN,
                    "missing_signed_standard",
                    f"Signed standard is missing for {school_id}/v1",
                    "reviewed target.standard.yaml exists",
                    str(standard_path),
                    affected_ids=[school_id],
                    root_cause_bucket="standard_missing",
                )
            )
            next_index += 1
        template_contract = (
            root
            / "standards"
            / "targets"
            / school_id
            / "v1/template_quality/final_template.expected.yaml"
        )
        _, baseline_findings = load_baseline_file(
            template_contract,
            stage="coverage",
            start_index=next_index,
        )
        findings.extend(baseline_findings)
        next_index += len(baseline_findings)

        gap_findings = _real_core_template_gap_findings(
            root,
            school_id,
            start_index=next_index,
        )
        findings.extend(gap_findings)
        next_index += len(gap_findings)

    for student in REAL_CORE_STUDENTS:
        student_id = str(student["student_id"])
        _, baseline_findings = load_baseline_file(
            root
            / "standards"
            / "students"
            / student_id
            / "v1/content_extract/student_content_artifact.expected.yaml",
            stage="coverage",
            start_index=next_index,
        )
        findings.extend(baseline_findings)
        next_index += len(baseline_findings)

    for case in [case for case in cases if case.stage == "e2e"]:
        _, plan_findings = load_baseline_file(
            root
            / "standards"
            / "cases"
            / case.case_id
            / "v1/placement/placement_plan.expected.yaml",
            stage="coverage",
            start_index=next_index,
        )
        findings.extend(plan_findings)
        next_index += len(plan_findings)

        _, snapshot_findings = load_baseline_file(
            root
            / "standards"
            / "cases"
            / case.case_id
            / "v1/render/feature_snapshot.expected.json",
            stage="coverage",
            start_index=next_index,
        )
        findings.extend(snapshot_findings)
        next_index += len(snapshot_findings)

        case_dir = _real_core_eval_runs_root(root) / case.case_id
        evidence_path = case_dir / "evidence/word_image_evidence.json"
        if not evidence_path.exists():
            findings.append(
                make_finding(
                    next_index,
                    "coverage",
                    Status.UNKNOWN,
                    "missing_word_image_evidence",
                    "real-core-v0 render cases require Word image evidence packages",
                    str(evidence_path),
                    "missing",
                    affected_ids=[case.case_id],
                    root_cause_bucket="oracle_gap",
                )
            )
            next_index += 1
        else:
            try:
                evidence_manifest = read_json(evidence_path)
            except Exception as exc:  # pragma: no cover - defensive branch
                findings.append(
                    make_finding(
                        next_index,
                        "coverage",
                        Status.UNKNOWN,
                        "invalid_word_image_evidence",
                        "Word image evidence manifest must be valid JSON",
                        "valid JSON manifest",
                        repr(exc),
                        affected_ids=[case.case_id],
                        root_cause_bucket="oracle_gap",
                    )
                )
                next_index += 1
            else:
                evidence_findings = verify_word_image_evidence(
                    evidence_manifest,
                    expected_final_docx=case_dir / "final.docx",
                    stage="coverage",
                    start_index=next_index,
                )
                for finding in evidence_findings:
                    if case.case_id not in finding.affected_ids:
                        finding.affected_ids.append(case.case_id)
                findings.extend(evidence_findings)
                next_index += len(evidence_findings)

            missing_audit_inputs = missing_e2e_audit_inputs(case_dir)
            if missing_audit_inputs:
                findings.append(
                    make_finding(
                        next_index,
                        "coverage",
                        Status.UNKNOWN,
                        "missing_product_quality_evidence",
                        "real-core-v0 coverage requires auditable e2e business artifacts",
                        "final.docx and template/content/placement/render artifacts",
                        ", ".join(str(path) for path in missing_audit_inputs),
                        affected_ids=[case.case_id],
                        root_cause_bucket="business_acceptance_gap",
                    )
                )
                product_quality_report["missing_cases"].append(case.case_id)
                next_index += 1
            else:
                product_quality_report["checked_cases"].append(case.case_id)
                case_findings = audit_e2e_case(case_dir)
                if case_findings:
                    product_quality_report["failing_cases"].append(case.case_id)
                for offset, finding in enumerate(case_findings):
                    finding.finding_id = f"f_{next_index + offset:03d}"
                    if case.case_id not in finding.affected_ids:
                        finding.affected_ids.append(case.case_id)
                findings.extend(case_findings)
                product_quality_findings.extend(case_findings)
                next_index += len(case_findings)

    required_flat = REAL_CORE_PROFILE.all_required_capabilities()
    product_quality_report["coverage"] = business_acceptance_coverage(
        product_quality_findings
    )
    if product_quality_report["missing_cases"]:
        product_quality_report["coverage"] = {
            capability: False
            for capability in product_quality_report["coverage"]
        }
    status = (
        merge_statuses([finding.status for finding in findings])
        if findings
        else Status.PASS
    )
    missing_types = sorted({finding.type for finding in findings})
    if not findings:
        baseline_status = "signed_business_accepted"
    elif any(
        finding_type in missing_types
        for finding_type in {
            "missing_generated_template_gap_evidence",
            "generated_template_gap_blocking",
        }
    ):
        baseline_status = "generated_template_gap_pending"
    elif set(missing_types) == {"missing_word_image_evidence"}:
        baseline_status = "source_facts_signed_word_evidence_pending"
    elif product_quality_report["failing_cases"] or product_quality_report["missing_cases"]:
        baseline_status = "business_acceptance_blocked"
    else:
        baseline_status = "pending_review"

    report = {
        "profile": REAL_CORE_PROFILE.profile_id,
        "status": status.value,
        "required": len(required_flat),
        "covered": 0 if findings else len(required_flat),
        "missing": missing_types,
        "case_counts": case_counts,
        "source_files": {
            "required": [str(path) for path in required_source_files],
            "missing": missing_source_files,
        },
        "baseline_status": baseline_status,
        "product_quality": product_quality_report,
    }
    return report, findings


def _real_core_source_files() -> list[Path]:
    paths = [Path("inputs/targets/shared/raw/template_recognition_alignment_review.md")]
    for school in REAL_CORE_SCHOOLS:
        paths.append(school["template_docx"])
        generated_template_docx = school.get("generated_template_docx")
        if generated_template_docx is not None:
            paths.append(generated_template_docx)
        paths.append(school["review_source"])
    for student in REAL_CORE_STUDENTS:
        paths.append(student["student_docx"])
        paths.append(student["review_source"])
    return paths


def _real_core_eval_runs_root(root: Path) -> Path:
    return root / "runs/eval/real-core-v0"


def _real_core_template_gap_findings(
    root: Path,
    school_id: str,
    *,
    start_index: int,
) -> list[Finding]:
    case_id = f"real_core_v0_template_{school_id}"
    artifacts = _real_core_eval_runs_root(root) / case_id / "artifacts"
    required = [
        artifacts / "generated_template.docx",
        artifacts / "generated_template_tree.json",
        artifacts / "template_gap_report.json",
        artifacts / "template_gap_report.md",
        artifacts / "template_gap_report.docx",
    ]
    missing = [path for path in required if not path.exists()]
    if missing:
        return [
            make_finding(
                start_index,
                "coverage",
                Status.UNKNOWN,
                "missing_generated_template_gap_evidence",
                "real-core-v0 coverage requires generated template gap evidence",
                "generated_template.docx, generated_template_tree.json, and template_gap_report.*",
                ", ".join(str(path) for path in missing),
                affected_ids=[school_id, case_id],
                root_cause_bucket="template_generation_gap",
            )
        ]

    report_path = artifacts / "template_gap_report.json"
    try:
        report = read_json(report_path)
    except Exception as exc:  # pragma: no cover - defensive branch
        return [
            make_finding(
                start_index,
                "coverage",
                Status.UNKNOWN,
                "invalid_generated_template_gap_report",
                "template_gap_report.json must be valid JSON",
                "valid JSON report",
                repr(exc),
                affected_ids=[school_id, case_id],
                root_cause_bucket="template_generation_gap",
            )
        ]

    findings: list[Finding] = []
    next_index = start_index
    bound_hash = report.get("generated_template", {}).get("sha256")
    actual_hash = sha256_file(artifacts / "generated_template.docx")
    if bound_hash != actual_hash:
        findings.append(
            make_finding(
                next_index,
                "coverage",
                Status.FAIL,
                "generated_template_gap_hash_mismatch",
                "template_gap_report.json must bind the generated_template.docx hash",
                actual_hash,
                str(bound_hash),
                affected_ids=[school_id, case_id],
                root_cause_bucket="template_generation_gap",
            )
        )
        next_index += 1

    summary = report.get("summary", {})
    required_summary_fields = {
        "known_status",
        "display_status",
        "passed_count",
        "failed_count",
        "unknown_count",
        "blocking_status",
        "per_unit",
    }
    missing_summary = sorted(required_summary_fields - set(summary))
    if missing_summary:
        findings.append(
            make_finding(
                next_index,
                "coverage",
                Status.UNKNOWN,
                "generated_template_gap_summary_incomplete",
                "template_gap_report.json must include the status summary fields",
                ", ".join(sorted(required_summary_fields)),
                ", ".join(missing_summary),
                affected_ids=[school_id, case_id],
                root_cause_bucket="template_generation_gap",
            )
        )
        next_index += 1

    blocking_status = summary.get("blocking_status")
    if blocking_status in {Status.FAIL.value, Status.UNKNOWN.value}:
        findings.append(
            make_finding(
                next_index,
                "coverage",
                Status(blocking_status),
                "generated_template_gap_blocking",
                "生成模板 Word 差距报告仍然阻断 real-core-v0 coverage",
                "template gap blocking_status = PASS",
                str(summary.get("display_status") or blocking_status),
                affected_ids=[school_id, case_id],
                root_cause_bucket="template_generation_gap",
            )
        )
    return findings
