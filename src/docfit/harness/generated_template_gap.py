from __future__ import annotations

from pathlib import Path
from typing import Any
import re
import shutil

from docx import Document

from docfit.core.io import ensure_dir, now_iso, sha256_file, write_json, write_text
from docfit.core.models import Finding, StageResult, make_finding
from docfit.core.status import Status, merge_statuses
from docfit.harness.generated_template_inspector import (
    inspect_generated_template_docx,
    iter_visible_text_entries,
)
from docfit.harness.real_core import accepted_expected_artifact, load_template_unit_baseline
from docfit.harness.standards import StandardBundle


TEMPLATE_GENERATION_CAPABILITIES = (
    "template_generation.output_docx",
    "template_generation.actual_tree",
    "template_generation.unit_match",
    "template_generation.element_match",
    "template_generation.style_match",
    "template_generation.header_footer_match",
    "template_generation.page_rule_match",
    "template_generation.field_match",
    "template_generation.numbering_match",
    "template_generation.report",
)


def evaluate_generated_template_gap(
    bundle: StandardBundle,
    generated_template_docx: Path,
    out_dir: Path,
    *,
    stage: str = "template",
) -> StageResult:
    artifacts_dir = ensure_dir(out_dir / "artifacts")
    generated_copy = artifacts_dir / "generated_template.docx"
    tree_path = artifacts_dir / "generated_template_tree.json"
    report_json_path = artifacts_dir / "template_gap_report.json"
    report_md_path = artifacts_dir / "template_gap_report.md"
    report_docx_path = artifacts_dir / "template_gap_report.docx"

    if generated_template_docx.exists():
        if generated_template_docx.resolve() != generated_copy.resolve():
            shutil.copyfile(generated_template_docx, generated_copy)
        inspected_path = generated_copy
    else:
        inspected_path = generated_template_docx

    tree = inspect_generated_template_docx(inspected_path)
    write_json(tree_path, tree)

    baseline, baseline_findings = load_template_unit_baseline(
        bundle,
        stage=stage,
        start_index=1,
    )
    expected = accepted_expected_artifact(baseline or {})
    expected_units = expected.get("units", []) if isinstance(expected, dict) else []
    standard_path = _template_unit_contract_path(bundle)
    report = build_template_gap_report(
        bundle=bundle,
        generated_template=inspected_path,
        generated_template_source=generated_template_docx,
        standard_path=standard_path,
        expected_units=expected_units,
        tree=tree,
    )
    write_json(report_json_path, report)
    write_text(report_md_path, render_template_gap_markdown(report))
    write_template_gap_docx(report_docx_path, report)

    findings = list(baseline_findings)
    findings.extend(
        findings_from_template_gap_report(
            report,
            stage=stage,
            start_index=len(findings) + 1,
        )
    )
    status = Status(report["summary"]["blocking_status"])
    if baseline_findings:
        status = merge_statuses([status] + [finding.status for finding in baseline_findings])
    result = StageResult(
        stage,
        status,
        findings=findings,
        artifacts={
            "generated_template_tree": tree,
            "template_gap_report": report,
        },
        artifact_paths={
            "generated_template_docx": generated_copy
            if generated_copy.exists()
            else inspected_path,
            "generated_template_tree": tree_path,
            "template_gap_report_json": report_json_path,
            "template_gap_report_md": report_md_path,
            "template_gap_report_docx": report_docx_path,
        },
        coverage=template_generation_coverage(report, tree),
        blocked_at=stage if status != Status.PASS else None,
    )
    return result


def build_template_gap_report(
    *,
    bundle: StandardBundle,
    generated_template: Path,
    generated_template_source: Path,
    standard_path: Path,
    expected_units: list[dict[str, Any]],
    tree: dict[str, Any],
) -> dict[str, Any]:
    check_items: list[dict[str, Any]] = []

    if not tree.get("input_exists"):
        check_items.append(
            _check(
                "template_generation.output_docx",
                Status.UNKNOWN,
                "template_generation_output_missing",
                "缺少 generated_template.docx，被测生成模板 Word 还没有进入评测",
                str(generated_template_source),
                "missing",
                affected_ids=[bundle.school_id],
            )
        )
    elif not tree.get("input_valid_docx"):
        check_items.append(
            _check(
                "template_generation.output_docx",
                Status.FAIL,
                "template_generation_output_invalid",
                "generated_template.docx 不是有效 DOCX 包",
                "valid DOCX package",
                str(generated_template),
                evidence_refs=[str(generated_template)],
                affected_ids=[bundle.school_id],
            )
        )
    else:
        check_items.append(
            _check(
                "template_generation.output_docx",
                Status.PASS,
                "template_generation_output_docx_present",
                "generated_template.docx 已作为被测输入记录路径和 hash",
                "valid DOCX with sha256",
                tree.get("input_hashes", {}).get("generated_template_docx", ""),
                evidence_refs=[str(generated_template)],
                affected_ids=[bundle.school_id],
            )
        )

    if not expected_units:
        check_items.append(
            _check(
                "template_generation.unit_match",
                Status.UNKNOWN,
                "template_generation_expected_units_missing",
                "模板差距检查缺少 template_unit_contract.yaml 中的 expected.units",
                "expected.units",
                "missing",
                evidence_refs=[str(standard_path)],
                affected_ids=[bundle.school_id],
            )
        )
    elif tree.get("input_valid_docx"):
        check_items.extend(_compare_units(expected_units, tree))
        check_items.extend(_unknown_visible_object_checks(tree))

    summary = summarize_check_items(check_items)
    return {
        "artifact_type": "template_gap_report",
        "artifact_version": "1.0",
        "producer": {"name": "docfit-generated-template-gap", "version": "0.1.0"},
        "created_at": now_iso(),
        "school_id": bundle.school_id,
        "template_version": bundle.template_version,
        "generated_template": {
            "path": str(generated_template),
            "source_path": str(generated_template_source),
            "sha256": tree.get("input_hashes", {}).get("generated_template_docx"),
        },
        "standard": {
            "path": str(standard_path),
            "sha256": sha256_file(standard_path) if standard_path.exists() else None,
        },
        "generated_template_tree": "generated_template_tree.json",
        "summary": summary,
        "coverage": template_generation_coverage_from_items(check_items, tree),
        "check_items": check_items,
    }


def summarize_check_items(check_items: list[dict[str, Any]]) -> dict[str, Any]:
    passed_count = sum(1 for item in check_items if item.get("status") == Status.PASS.value)
    failed_count = sum(1 for item in check_items if item.get("status") == Status.FAIL.value)
    unknown_count = sum(
        1 for item in check_items if item.get("status") == Status.UNKNOWN.value
    )
    if failed_count:
        known_status = Status.FAIL.value
    elif passed_count:
        known_status = Status.PASS.value
    else:
        known_status = Status.UNKNOWN.value

    if known_status == Status.UNKNOWN.value:
        display_status = Status.UNKNOWN.value
    elif unknown_count:
        display_status = f"{known_status} + UNKNOWN"
    else:
        display_status = known_status

    if failed_count:
        blocking_status = Status.FAIL.value
    elif unknown_count:
        blocking_status = Status.UNKNOWN.value
    else:
        blocking_status = Status.PASS.value

    return {
        "known_status": known_status,
        "display_status": display_status,
        "passed_count": passed_count,
        "failed_count": failed_count,
        "unknown_count": unknown_count,
        "blocking_status": blocking_status,
    }


def findings_from_template_gap_report(
    report: dict[str, Any],
    *,
    stage: str = "template",
    start_index: int = 1,
) -> list[Finding]:
    findings: list[Finding] = []
    next_index = start_index
    for item in report.get("check_items", []):
        status = Status(item["status"])
        if status == Status.PASS:
            continue
        findings.append(
            make_finding(
                next_index,
                stage,
                status,
                str(item.get("type", "template_generation_gap")),
                str(item.get("message", "")),
                str(item.get("expected", "")),
                str(item.get("actual", "")),
                evidence_refs=list(item.get("evidence_refs", [])),
                affected_ids=list(item.get("affected_ids", [])),
                root_cause_bucket=str(
                    item.get("root_cause_bucket", "template_generation_gap")
                ),
            )
        )
        next_index += 1
    return findings


def template_generation_coverage(
    report: dict[str, Any],
    tree: dict[str, Any],
) -> dict[str, bool]:
    return {
        f"template_generation.{key}": value
        for key, value in template_generation_coverage_from_items(
            report.get("check_items", []),
            tree,
        ).items()
    }


def template_generation_coverage_from_items(
    check_items: list[dict[str, Any]],
    tree: dict[str, Any],
) -> dict[str, bool]:
    categories = {str(item.get("category", "")) for item in check_items}
    valid_tree = bool(tree.get("input_valid_docx"))
    has_source_refs = any(
        entry.get("source_ref") for entry in iter_visible_text_entries(tree)
    )
    return {
        "output_docx": bool(tree.get("input_exists") and tree.get("input_valid_docx")),
        "actual_tree": valid_tree and has_source_refs,
        "unit_match": "unit" in categories,
        "element_match": "element" in categories,
        "style_match": "style" in categories,
        "header_footer_match": "header_footer" in categories,
        "page_rule_match": "page_rule" in categories,
        "field_match": "field" in categories,
        "numbering_match": "numbering" in categories,
        "report": bool(check_items),
    }


def render_template_gap_markdown(report: dict[str, Any]) -> str:
    summary = report["summary"]
    lines = [
        f"# 生成模板 Word 差距报告：{report['school_id']}",
        "",
        f"- 已知状态：{summary['known_status']}",
        f"- 展示状态：{summary['display_status']}",
        (
            "- 检查计数："
            f"PASS {summary['passed_count']} / "
            f"FAIL {summary['failed_count']} / "
            f"UNKNOWN {summary['unknown_count']}"
        ),
        f"- Gate 状态：{summary['blocking_status']}",
        f"- 生成模板 Word：{report['generated_template']['path']}",
        f"- 生成模板 hash：{report['generated_template'].get('sha256') or 'missing'}",
        f"- 标准文件：{report['standard']['path']}",
        "",
        "## 差距清单",
    ]
    for item in report.get("check_items", []):
        lines.extend(
            [
                "",
                f"### [{item['status']}] {item['check_id']}",
                "",
                item["message"],
                "",
                f"- 标准期望：{item['expected']}",
                f"- 实际结果：{item['actual']}",
                f"- 来源位置：{', '.join(item.get('evidence_refs', [])) or 'missing'}",
                f"- 下一步：{item.get('next_step', '')}",
            ]
        )
    lines.append("")
    return "\n".join(lines)


def write_template_gap_docx(path: Path, report: dict[str, Any]) -> None:
    ensure_dir(path.parent)
    doc = Document()
    doc.add_heading(f"生成模板 Word 差距报告：{report['school_id']}", level=1)
    summary = report["summary"]
    for label, value in (
        ("已知状态", summary["known_status"]),
        ("展示状态", summary["display_status"]),
        ("Gate 状态", summary["blocking_status"]),
        (
            "检查计数",
            (
                f"PASS {summary['passed_count']} / "
                f"FAIL {summary['failed_count']} / "
                f"UNKNOWN {summary['unknown_count']}"
            ),
        ),
        ("生成模板 Word", report["generated_template"]["path"]),
        ("生成模板 hash", report["generated_template"].get("sha256") or "missing"),
        ("标准文件", report["standard"]["path"]),
    ):
        doc.add_paragraph(f"{label}：{value}")
    doc.add_heading("差距清单", level=2)
    for item in report.get("check_items", []):
        doc.add_heading(f"[{item['status']}] {item['check_id']}", level=3)
        doc.add_paragraph(item["message"])
        doc.add_paragraph(f"标准期望：{item['expected']}")
        doc.add_paragraph(f"实际结果：{item['actual']}")
        doc.add_paragraph(
            "来源位置：" + (", ".join(item.get("evidence_refs", [])) or "missing")
        )
        doc.add_paragraph(f"下一步：{item.get('next_step', '')}")
    doc.save(path)


def _compare_units(
    expected_units: list[dict[str, Any]],
    tree: dict[str, Any],
) -> list[dict[str, Any]]:
    entries = iter_visible_text_entries(tree)
    checks: list[dict[str, Any]] = []
    unit_first_orders: list[tuple[str, int]] = []
    for unit in expected_units:
        unit_id = str(unit.get("unit_id") or "unknown_unit")
        element_checks, matched_orders = _compare_elements(unit, entries, tree)
        checks.extend(element_checks)
        if matched_orders:
            first_order = min(matched_orders)
            unit_first_orders.append((unit_id, first_order))
            checks.append(
                _check(
                    "template_generation.unit_match",
                    Status.PASS,
                    "template_generation_unit_found",
                    f"生成模板 Word 中找到单元 {unit_id} 的可见来源证据",
                    unit_id,
                    f"first_source_order={first_order}",
                    category="unit",
                    evidence_refs=_matched_evidence(element_checks),
                    affected_ids=[unit_id],
                )
            )
        else:
            status = (
                Status.UNKNOWN
                if unit.get("status") == "template_default_optional"
                else Status.FAIL
            )
            checks.append(
                _check(
                    "template_generation.unit_match",
                    status,
                    "template_generation_unit_missing",
                    f"生成模板 Word 中没有找到单元 {unit_id} 的可见来源证据",
                    unit_id,
                    "missing",
                    category="unit",
                    affected_ids=[unit_id],
                    next_step="如果生成逻辑应输出该单元，修模板生成；如果解析器无法识别，补生成模板解析器。",
                )
            )
        checks.extend(_header_footer_checks(unit, tree))
        checks.extend(_page_rule_checks(unit))

    checks.extend(_field_checks(expected_units, tree))
    checks.extend(_numbering_checks(expected_units, tree))
    if len(unit_first_orders) >= 2:
        ordered = sorted(unit_first_orders, key=lambda item: item[1])
        checks.append(
            _check(
                "template_generation.unit_match",
                Status.PASS if ordered == unit_first_orders else Status.FAIL,
                (
                    "template_generation_unit_order_ok"
                    if ordered == unit_first_orders
                    else "template_generation_unit_order_mismatch"
                ),
                "生成模板 Word 的已识别单元顺序与标准顺序对齐",
                repr([unit_id for unit_id, _order in unit_first_orders]),
                repr([unit_id for unit_id, _order in ordered]),
                category="unit",
                affected_ids=[unit_id for unit_id, _order in unit_first_orders],
                next_step="若顺序不一致，修模板生成的单元输出顺序。",
            )
        )
    return checks


def _compare_elements(
    unit: dict[str, Any],
    entries: list[dict[str, Any]],
    tree: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[int]]:
    unit_id = str(unit.get("unit_id") or "unknown_unit")
    checks: list[dict[str, Any]] = []
    matched_orders: list[int] = []
    for element in unit.get("elements", []):
        element_id = str(element.get("element_id") or "unknown_element")
        affected_id = f"{unit_id}.{element_id}"
        needles = _candidate_needles(element)
        matches = _find_matches(entries, needles)
        policy = str(element.get("policy") or "")
        if matches:
            matched_orders.extend(int(match.get("order", 0)) for match in matches)
            checks.append(
                _check(
                    "template_generation.element_match",
                    Status.PASS,
                    "template_generation_element_found",
                    f"生成模板 Word 中找到元素 {affected_id} 的可见文本来源",
                    " / ".join(needles),
                    _preview(matches[0].get("text", "")),
                    category="element",
                    evidence_refs=[match.get("source_ref", "") for match in matches[:3]],
                    affected_ids=[affected_id],
                )
            )
            checks.append(_style_check(unit_id, element, matches[0]))
            continue

        if policy in {"fixed", "manual_only"}:
            checks.append(
                _check(
                    "template_generation.element_match",
                    Status.FAIL,
                    "template_generation_element_missing",
                    f"生成模板 Word 缺少标准要求保留的元素 {affected_id}",
                    element.get("content") or element.get("name") or affected_id,
                    "missing",
                    category="element",
                    affected_ids=[affected_id],
                    next_step="修模板生成逻辑，保留该固定内容或手工填写位置。",
                )
            )
        elif policy == "generated":
            field_status = (
                "field parsed"
                if _has_generated_field(tree, element)
                else "no field parsed"
            )
            checks.append(
                _check(
                    "template_generation.element_match",
                    Status.UNKNOWN,
                    "template_generation_generated_element_source_unknown",
                    f"生成元素 {affected_id} 需要可检查的 Word 字段或等价占位",
                    element.get("content") or element.get("name") or affected_id,
                    field_status,
                    category="element",
                    affected_ids=[affected_id],
                    next_step="补字段到单元/元素的绑定检查，或修模板生成逻辑输出字段。",
                )
            )
        else:
            checks.append(
                _check(
                    "template_generation.element_match",
                    Status.UNKNOWN,
                    "template_generation_element_source_unknown",
                    f"元素 {affected_id} 需要可写位置或生成规则，但当前解析树无法证明",
                    element.get("content") or element.get("name") or affected_id,
                    "no matched Word source",
                    category="element",
                    affected_ids=[affected_id],
                    next_step="补解析器对可写位置、content control、字段或占位符的识别。",
                )
            )
    return checks, matched_orders


def _field_checks(
    expected_units: list[dict[str, Any]],
    tree: dict[str, Any],
) -> list[dict[str, Any]]:
    checks: list[dict[str, Any]] = []
    generated_elements = [
        (unit, element)
        for unit in expected_units
        for element in unit.get("elements", [])
        if element.get("policy") == "generated"
    ]
    if not generated_elements:
        return [
            _check(
                "template_generation.field_match",
                Status.PASS,
                "template_generation_no_required_fields",
                "模板标准没有声明必须检查的 Word 生成字段",
                "no generated field requirements",
                "no generated field requirements",
                category="field",
            )
        ]

    fields = tree.get("data", {}).get("fields", [])
    for unit, element in generated_elements:
        unit_id = str(unit.get("unit_id") or "unknown_unit")
        element_id = str(element.get("element_id") or "unknown_element")
        status = Status.PASS if _has_generated_field(tree, element) else Status.FAIL
        checks.append(
            _check(
                "template_generation.field_match",
                status,
                (
                    "template_generation_field_match"
                    if status == Status.PASS
                    else "template_generation_field_missing"
                ),
                f"生成元素 {unit_id}.{element_id} 需要 Word 字段或可检查的生成占位",
                element.get("content") or element.get("name") or f"{unit_id}.{element_id}",
                (
                    ", ".join(str(field.get("instruction", "")) for field in fields)
                    or "no Word field parsed"
                ),
                category="field",
                evidence_refs=[field.get("source_ref", "") for field in fields],
                affected_ids=[f"{unit_id}.{element_id}.field"],
                next_step="修模板生成逻辑输出字段，或补等价生成占位的确定性解析。",
            )
        )
    return checks


def _numbering_checks(
    expected_units: list[dict[str, Any]],
    tree: dict[str, Any],
) -> list[dict[str, Any]]:
    checks: list[dict[str, Any]] = []
    requirements = [
        unit
        for unit in expected_units
        if _unit_mentions_numbering(unit)
    ]
    if not requirements:
        return [
            _check(
                "template_generation.numbering_match",
                Status.PASS,
                "template_generation_no_numbering_requirements",
                "模板标准没有声明必须检查的编号规则",
                "no numbering requirements",
                "no numbering requirements",
                category="numbering",
            )
        ]
    numbering_refs = tree.get("data", {}).get("numbering_refs", [])
    for unit in requirements:
        unit_id = str(unit.get("unit_id") or "unknown_unit")
        checks.append(
            _check(
                "template_generation.numbering_match",
                Status.UNKNOWN,
                "template_generation_numbering_unverified",
                f"单元 {unit_id} 声明了编号规则，但当前解析器还不能把编号绑定到单元/元素",
                _preview(_numbering_requirement_text(unit)),
                (
                    ", ".join(ref.get("source_ref", "") for ref in numbering_refs[:5])
                    or "no OOXML numbering refs parsed"
                ),
                category="numbering",
                evidence_refs=[ref.get("source_ref", "") for ref in numbering_refs[:5]],
                affected_ids=[f"{unit_id}.numbering"],
                next_step="补编号定义、段落 numPr 和标题层级到模板单元的绑定检查。",
            )
        )
    return checks


def _style_check(
    unit_id: str,
    element: dict[str, Any],
    match: dict[str, Any],
) -> dict[str, Any]:
    element_id = str(element.get("element_id") or "unknown_element")
    expected_style = _normalize_text(element.get("style"))
    if not expected_style:
        return _check(
            "template_generation.style_match",
            Status.PASS,
            "template_generation_style_not_required",
            f"元素 {unit_id}.{element_id} 没有声明必须检查的样式字段",
            "no required style",
            "no required style",
            category="style",
            affected_ids=[f"{unit_id}.{element_id}.style"],
        )
    actual_style = _style_actual_summary(match)
    style_result = _compare_style_details(expected_style, match)
    if _style_matches(expected_style, _normalize_text(match.get("style"))):
        status = Status.PASS
        type_ = "template_generation_style_match"
        message = f"元素 {unit_id}.{element_id} 的段落样式名称与标准可比对"
        next_step = "none"
    elif style_result["mismatches"]:
        status = Status.FAIL
        type_ = "template_generation_style_mismatch"
        message = f"元素 {unit_id}.{element_id} 的 OOXML 样式属性不符合标准"
        next_step = "修模板生成样式，或修解析器的样式继承规则。"
    elif style_result["matched"] and not style_result["unknown"]:
        status = Status.PASS
        type_ = "template_generation_style_match"
        message = f"元素 {unit_id}.{element_id} 的 OOXML 样式属性符合已声明标准"
        next_step = "none"
    else:
        status = Status.UNKNOWN
        type_ = "template_generation_style_unverified"
        message = f"元素 {unit_id}.{element_id} 的样式细节还不能由当前解析器完整证明"
        next_step = "补 OOXML 样式解析，比较字体、字号、加粗、对齐、缩进、段距和行距。"
    return _check(
        "template_generation.style_match",
        status,
        type_,
        message,
        expected_style,
        actual_style or "style details unavailable",
        category="style",
        evidence_refs=[match.get("source_ref", "")],
        affected_ids=[f"{unit_id}.{element_id}.style"],
        next_step=next_step,
    )


def _header_footer_checks(
    unit: dict[str, Any],
    tree: dict[str, Any],
) -> list[dict[str, Any]]:
    unit_id = str(unit.get("unit_id") or "unknown_unit")
    header_footer = unit.get("header_footer") or {}
    if not isinstance(header_footer, dict):
        return []
    checks: list[dict[str, Any]] = []
    header_rule = _normalize_text(header_footer.get("header"))
    page_rule = _normalize_text(header_footer.get("page_number"))
    parts = tree.get("data", {}).get("headers_footers", [])
    fields = tree.get("data", {}).get("fields", [])
    header_text = _normalize_text(" ".join(part.get("text", "") for part in parts))
    page_fields = [
        field
        for field in fields
        if "PAGE" in str(field.get("instruction", "")).upper()
        or "NUMPAGES" in str(field.get("instruction", "")).upper()
    ]

    if header_rule:
        no_header_expected = _means_none(header_rule)
        status = (
            Status.PASS
            if (no_header_expected and not header_text)
            or (not no_header_expected and header_rule in header_text)
            else Status.UNKNOWN
        )
        checks.append(
            _check(
                "template_generation.header_footer_match",
                status,
                (
                    "template_generation_header_footer_match"
                    if status == Status.PASS
                    else "template_generation_header_footer_unverified"
                ),
                f"单元 {unit_id} 的页眉规则需要和生成模板 Word 对齐",
                header_rule,
                header_text or "no header/footer text parsed",
                category="header_footer",
                evidence_refs=[part.get("source_ref", "") for part in parts],
                affected_ids=[f"{unit_id}.header_footer.header"],
                next_step="补 section 级页眉页脚继承和单元边界绑定检查。",
            )
        )
    if page_rule:
        no_page_expected = "无页码字段" in page_rule or page_rule == "无"
        status = (
            Status.PASS
            if no_page_expected and not page_fields
            else Status.UNKNOWN
        )
        checks.append(
            _check(
                "template_generation.header_footer_match",
                status,
                (
                    "template_generation_page_number_rule_match"
                    if status == Status.PASS
                    else "template_generation_page_number_rule_unverified"
                ),
                f"单元 {unit_id} 的页码规则需要和生成模板 Word 对齐",
                page_rule,
                (
                    ", ".join(field.get("instruction", "") for field in page_fields)
                    or "no page field parsed"
                ),
                category="header_footer",
                evidence_refs=[field.get("source_ref", "") for field in page_fields],
                affected_ids=[f"{unit_id}.header_footer.page_number"],
                next_step="补页脚页码字段、页码格式和 section 起始页检查。",
            )
        )
    return checks


def _page_rule_checks(unit: dict[str, Any]) -> list[dict[str, Any]]:
    unit_id = str(unit.get("unit_id") or "unknown_unit")
    page = unit.get("page") or {}
    if not isinstance(page, dict):
        return []
    checks: list[dict[str, Any]] = []
    for field in ("page_break", "section_isolation", "keep_together"):
        rule = _normalize_text(page.get(field))
        if not rule:
            continue
        checks.append(
            _check(
                "template_generation.page_rule_match",
                Status.UNKNOWN,
                "template_generation_page_rule_unverified",
                f"单元 {unit_id} 的分页/同页规则当前无法完整证明：{field}",
                rule,
                "unit-to-section/page binding not implemented",
                category="page_rule",
                affected_ids=[f"{unit_id}.page.{field}"],
                next_step="补单元边界到 OOXML 分页符、分节符、keep-with-next/keep-lines 的映射检查。",
            )
        )
    return checks


def _unknown_visible_object_checks(tree: dict[str, Any]) -> list[dict[str, Any]]:
    checks: list[dict[str, Any]] = []
    for index, item in enumerate(
        tree.get("data", {}).get("unknown_visible_objects", []),
        start=1,
    ):
        checks.append(
            _check(
                "template_generation.actual_tree",
                Status.UNKNOWN,
                "template_generation_visible_object_unmodeled",
                f"生成模板 Word 中存在当前解析器未建模的可见对象：{item.get('object_type')}",
                "visible object modeled or explicitly exempted",
                item.get("reason", "unmodeled"),
                category="actual_tree",
                evidence_refs=[item.get("source_ref", "")],
                affected_ids=[item.get("content_id", f"unknown_visible_object_{index}")],
                next_step="补解析器能力或登记明确例外，不能静默忽略可见对象。",
            )
        )
    return checks


def _candidate_needles(element: dict[str, Any]) -> list[str]:
    candidates: list[str] = []
    content = _normalize_text(element.get("content"))
    if content and not _looks_like_descriptor(content):
        candidates.append(content)
        if len(content) > 80:
            candidates.append(content[:80])
    name = _normalize_text(element.get("name"))
    label = re.sub(r"(标签|内容|正文|结果|机制)$", "", name).strip()
    if label and not _looks_like_descriptor(label) and len(label) >= 2:
        candidates.append(label)
    return _dedupe(candidates)


def _find_matches(
    entries: list[dict[str, Any]],
    needles: list[str],
) -> list[dict[str, Any]]:
    if not needles:
        return []
    matches: list[dict[str, Any]] = []
    for entry in entries:
        text = _normalize_text(entry.get("text"))
        if any(_normalize_text(needle) in text for needle in needles):
            matches.append(entry)
    return matches


def _has_generated_field(tree: dict[str, Any], element: dict[str, Any]) -> bool:
    fields = tree.get("data", {}).get("fields", [])
    if not fields:
        return False
    name = _normalize_text(element.get("name")).lower()
    content = _normalize_text(element.get("content")).lower()
    instructions = [
        str(field.get("instruction", "")).upper()
        for field in fields
    ]
    if "toc" in name or "目录" in name or "toc" in content or "目录" in content:
        return any("TOC" in instruction for instruction in instructions)
    if "页码" in name or "page" in name or "页码" in content:
        return any("PAGE" in instruction for instruction in instructions)
    return bool(fields)


def _compare_style_details(
    expected_style: str,
    match: dict[str, Any],
) -> dict[str, list[str]]:
    expected = _expected_style_requirements(expected_style)
    actual = _actual_style_properties(match)
    matched: list[str] = []
    mismatches: list[str] = []
    unknown: list[str] = []

    for font in expected["fonts"]:
        actual_fonts = actual["font_names"]
        if font in actual_fonts:
            matched.append(f"font={font}")
        elif actual_fonts:
            mismatches.append(f"font expected {font}, actual {'/'.join(actual_fonts)}")
        else:
            unknown.append(f"font expected {font}")

    expected_size = expected["font_size_pt"]
    if expected_size is not None:
        actual_size = actual["font_size_pt"]
        if actual_size is None:
            unknown.append(f"font_size expected {expected_size:g}pt")
        elif abs(actual_size - expected_size) <= 0.25:
            matched.append(f"font_size={actual_size:g}pt")
        else:
            mismatches.append(
                f"font_size expected {expected_size:g}pt, actual {actual_size:g}pt"
            )

    expected_bold = expected["bold"]
    if expected_bold is not None:
        actual_bold = actual["bold"]
        if actual_bold is None:
            unknown.append(f"bold expected {expected_bold}")
        elif actual_bold == expected_bold:
            matched.append(f"bold={actual_bold}")
        else:
            mismatches.append(f"bold expected {expected_bold}, actual {actual_bold}")

    expected_alignment = expected["alignment"]
    if expected_alignment:
        actual_alignment = actual["alignment"]
        if not actual_alignment:
            unknown.append(f"alignment expected {expected_alignment}")
        elif actual_alignment == expected_alignment:
            matched.append(f"alignment={actual_alignment}")
        else:
            mismatches.append(
                f"alignment expected {expected_alignment}, actual {actual_alignment}"
            )

    expected_line_spacing = expected["line_spacing"]
    if expected_line_spacing:
        actual_line_spacing = actual["line_spacing"]
        if not actual_line_spacing:
            unknown.append(f"line_spacing expected {expected_line_spacing}")
        elif actual_line_spacing == expected_line_spacing:
            matched.append(f"line_spacing={actual_line_spacing}")
        else:
            mismatches.append(
                "line_spacing expected "
                f"{expected_line_spacing}, actual {actual_line_spacing}"
            )

    return {
        "matched": matched,
        "mismatches": mismatches,
        "unknown": unknown,
    }


def _expected_style_requirements(expected_style: str) -> dict[str, Any]:
    normalized = _normalize_text(expected_style)
    fonts = [
        font
        for font in ("华文行楷", "黑体", "宋体", "Times New Roman")
        if font in normalized
    ]
    size_match = re.search(r"(?P<size>\d+(?:\.\d+)?)\s*pt", normalized)
    bold: bool | None = None
    if "不加粗" in normalized:
        bold = False
    elif "加粗" in normalized:
        bold = True
    alignment = None
    if "居中" in normalized:
        alignment = "center"
    elif "两端对齐" in normalized:
        alignment = "both"
    elif "左对齐" in normalized:
        alignment = "left"
    line_spacing = None
    if "单倍行距" in normalized:
        line_spacing = "single"
    elif "1.5 倍行距" in normalized or "1.5倍行距" in normalized:
        line_spacing = "1.5"
    exact_match = re.search(r"固定值\s*(?P<size>\d+(?:\.\d+)?)\s*pt", normalized)
    if exact_match:
        line_spacing = f"exact:{float(exact_match.group('size')):g}pt"
    return {
        "fonts": fonts,
        "font_size_pt": float(size_match.group("size")) if size_match else None,
        "bold": bold,
        "alignment": alignment,
        "line_spacing": line_spacing,
    }


def _actual_style_properties(match: dict[str, Any]) -> dict[str, Any]:
    details = match.get("style_details") or {}
    dominant = details.get("dominant_run") or {}
    paragraph_run = details.get("paragraph_run_properties") or {}
    paragraph = details.get("paragraph") or {}
    font_names = list(
        dict.fromkeys(
            [
                *dominant.get("font_names", []),
                *paragraph_run.get("font_names", []),
            ]
        )
    )
    return {
        "font_names": font_names,
        "font_size_pt": dominant.get("font_size_pt")
        or paragraph_run.get("font_size_pt"),
        "bold": _first_known(dominant.get("bold"), paragraph_run.get("bold")),
        "alignment": paragraph.get("alignment"),
        "line_spacing": (paragraph.get("spacing") or {}).get("line_spacing"),
    }


def _style_actual_summary(match: dict[str, Any]) -> str:
    details = match.get("style_details") or {}
    props = _actual_style_properties(match)
    paragraph = details.get("paragraph") or {}
    parts = [
        f"style={match.get('style') or paragraph.get('style_id') or 'missing'}",
        f"font={'/'.join(props['font_names']) or 'missing'}",
        (
            "font_size="
            + (
                f"{props['font_size_pt']:g}pt"
                if props["font_size_pt"] is not None
                else "missing"
            )
        ),
        f"bold={props['bold']}",
        f"alignment={props['alignment'] or 'missing'}",
        f"line_spacing={props['line_spacing'] or 'missing'}",
    ]
    return "; ".join(parts)


def _first_known(*values: Any) -> Any:
    for value in values:
        if value is not None:
            return value
    return None


def _unit_mentions_numbering(unit: dict[str, Any]) -> bool:
    requirement_text = _numbering_requirement_text(unit)
    return "编号" in requirement_text or "目录层级" in requirement_text


def _numbering_requirement_text(unit: dict[str, Any]) -> str:
    values: list[str] = [
        str(unit.get("name", "")),
        str(unit.get("element_order", "")),
        str(unit.get("layout_relation", "")),
        str(unit.get("missing_policy", "")),
    ]
    for element in unit.get("elements", []):
        values.extend(
            str(element.get(field, ""))
            for field in (
                "name",
                "content",
                "position",
                "relationship",
                "raw",
            )
        )
    return " ".join(values)


def _style_matches(expected_style: str, actual_style: str) -> bool:
    if not expected_style or not actual_style:
        return False
    expected = expected_style.lower()
    actual = actual_style.lower()
    return expected == actual or actual in expected or expected in actual


def _means_none(value: str) -> bool:
    return value in {"无", "默认无"} or "默认无" in value or "源模板无" in value


def _looks_like_descriptor(value: str) -> bool:
    descriptor_markers = (
        "学生",
        "当前阶段",
        "目标输出",
        "系统生成",
        "××",
        "模板",
        "正文",
        "标题",
        "内容流",
    )
    if value in {"固定", "填充", "生成", "手工"}:
        return True
    return any(marker in value for marker in descriptor_markers) and "：" not in value


def _matched_evidence(checks: list[dict[str, Any]]) -> list[str]:
    refs: list[str] = []
    for check in checks:
        if check.get("status") != Status.PASS.value:
            continue
        refs.extend(check.get("evidence_refs", []))
    return _dedupe(refs)[:5]


def _template_unit_contract_path(bundle: StandardBundle) -> Path:
    rel_path = bundle.signed_standard.get("evidence_baselines", {}).get(
        "template_unit_contract",
        "template_unit_contract.yaml",
    )
    return bundle.school_dir / rel_path


def _check(
    check_id: str,
    status: Status,
    type_: str,
    message: str,
    expected: Any,
    actual: Any,
    *,
    category: str | None = None,
    evidence_refs: list[str] | None = None,
    affected_ids: list[str] | None = None,
    next_step: str = "none",
    root_cause_bucket: str = "template_generation_gap",
) -> dict[str, Any]:
    return {
        "check_id": check_id,
        "category": category or check_id.split(".")[-1],
        "status": status.value,
        "type": type_,
        "message": message,
        "expected": _preview(expected),
        "actual": _preview(actual),
        "evidence_refs": [ref for ref in (evidence_refs or []) if ref],
        "affected_ids": [identifier for identifier in (affected_ids or []) if identifier],
        "next_step": next_step,
        "root_cause_bucket": root_cause_bucket,
    }


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _preview(value: Any, limit: int = 320) -> str:
    normalized = _normalize_text(value)
    if len(normalized) <= limit:
        return normalized
    return normalized[: limit - 3] + "..."


def _dedupe(values: list[str]) -> list[str]:
    seen: set[str] = set()
    deduped: list[str] = []
    for value in values:
        if not value or value in seen:
            continue
        deduped.append(value)
        seen.add(value)
    return deduped
