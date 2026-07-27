from __future__ import annotations

import ast
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any

from docx import Document

from docfit.core.io import read_json, read_yaml, sha256_file, sha256_json, write_json, write_text
from docfit.core.models import Finding, StageResult
from docfit.core.status import Status, merge_statuses
from docfit.harness.reports import write_report_bundle
from docfit.harness.template_generation_run_bundle import (
    BoundArtifact,
    RUN_ARTIFACT_SPECS,
    RunArtifactSpec,
    TemplateGenerationRunBundle,
)
from docfit.harness.template_generation_stage_verifiers import (
    StageCheck,
    judge_template_generation_stage,
)
from docfit.harness.template_generation_standard_quality import (
    StageStandardSpec,
    TemplateGenerationStandardQualityReport,
    TemplateGenerationStandardSet,
)
from docfit.template_generation.t3_action_projection import (
    project_t3_gold_item,
)

@dataclass
class TemplateGenerationJudgeReport:
    status: Status
    first_bad_stage: str | None
    standard_set: TemplateGenerationStandardSet
    standard_quality: TemplateGenerationStandardQualityReport
    run_bundle: TemplateGenerationRunBundle
    stage_checks: list[StageCheck]
    findings: list[Finding]

    def to_dict(self) -> dict[str, Any]:
        stage_standard_diffs = build_stage_standard_diffs(self)
        mismatches = build_mismatches(self)
        root_causes = build_diagnosis_root_causes(self, mismatches)
        owner_assignments = build_owner_assignments(root_causes)
        fix_plan = build_fix_plans(root_causes)
        _attach_diagnosis_layers(mismatches, root_causes, owner_assignments, fix_plan)
        acceptance = build_standard_acceptance_summary(
            self,
            stage_standard_diffs=stage_standard_diffs,
            root_causes=root_causes,
        )
        return {
            "artifact_type": "template_generation_judge_report",
            "artifact_version": "1.0",
            "status": self.status.value,
            "standard_acceptance_status": acceptance["standard_acceptance_status"],
            "signoff_status": acceptance["signoff_status"],
            "signoff_blockers": acceptance["signoff_blockers"],
            "first_bad_stage": self.first_bad_stage,
            "source_run_id": self.run_bundle.source_run_id,
            "source_run_dir": str(self.run_bundle.source_run_dir),
            "source_run_dir_name": self.run_bundle.source_run_dir_name,
            "standard_set": self.standard_set.to_dict(),
            "standard_quality": self.standard_quality.to_dict(),
            "run_bundle": self.run_bundle.to_dict(),
            "stage_checks": [check.to_dict() for check in self.stage_checks],
            "stage_standard_quality_reports": _stage_standard_quality_report_refs(self),
            "stage_standard_diff_reports": _stage_standard_diff_report_refs(self),
            "mismatches": mismatches,
            "stage_standard_diffs": stage_standard_diffs,
            "root_causes": root_causes,
            "owner_assignments": owner_assignments,
            "fix_plan": fix_plan,
            "owner_summary": acceptance["owner_summary"],
            "top_blockers": acceptance["top_blockers"],
            "findings": [finding.to_dict() for finding in self.findings],
        }

    def to_stage_result(self) -> StageResult:
        report_dict = self.to_dict()
        return StageResult(
            "template_generation_judge",
            self.status,
            run_status=Status.PASS,
            quality_status=self.status,
            findings=self.findings,
            artifacts={
                "template_generation_judge_report": report_dict,
                "template_generation_run_bundle": self.run_bundle.to_dict(),
                "template_generation_stage_checks": [
                    check.to_dict() for check in self.stage_checks
                ],
                "template_generation_stage_standard_quality_report": (
                    self.standard_quality.to_dict()
                ),
                "stage_standard_quality_reports": _stage_standard_quality_report_refs(self),
                "stage_standard_diff_reports": _stage_standard_diff_report_refs(self),
                "mismatches": report_dict["mismatches"],
                "stage_standard_diffs": report_dict["stage_standard_diffs"],
                "root_causes": report_dict["root_causes"],
                "owner_assignments": report_dict["owner_assignments"],
                "fix_plan": report_dict["fix_plan"],
                "owner_summary": report_dict["owner_summary"],
            },
            coverage={
                "template_generation_judge.standard_quality": (
                    self.standard_quality.status == Status.PASS
                ),
                "template_generation_judge.run_bundle": self.run_bundle.status == Status.PASS,
                "template_generation_judge.stage_checks": bool(self.stage_checks),
                "template_generation_judge.standard_acceptance_status": report_dict[
                    "standard_acceptance_status"
                ],
                "template_generation_judge.signoff_status": report_dict["signoff_status"],
            },
            blocked_at=self.first_bad_stage,
        )


def aggregate_template_generation_judgement(
    *,
    standard_set: TemplateGenerationStandardSet,
    standard_quality: TemplateGenerationStandardQualityReport,
    run_bundle: TemplateGenerationRunBundle,
    stage_checks: list[StageCheck],
) -> TemplateGenerationJudgeReport:
    findings = [
        *standard_quality.findings,
        *run_bundle.findings,
        *run_bundle.observed_t6_findings,
        *[
            finding
            for check in stage_checks
            for finding in check.findings
        ],
    ]
    t6_status = (
        merge_statuses(
            [finding.status for finding in run_bundle.observed_t6_findings]
        )
        if run_bundle.observed_t6_findings
        else Status.PASS
    )
    status = merge_statuses(
        [standard_quality.status, run_bundle.status, t6_status]
        + [check.status for check in stage_checks]
    )
    first_bad_stage = _first_bad_stage(stage_checks)
    if first_bad_stage is None and standard_quality.status != Status.PASS:
        first_bad_stage = "standard_quality"
    if first_bad_stage is None and run_bundle.status != Status.PASS:
        first_bad_stage = "run_bundle"
    if first_bad_stage is None and t6_status != Status.PASS:
        first_bad_stage = "T6"
    return TemplateGenerationJudgeReport(
        status=status,
        first_bad_stage=first_bad_stage,
        standard_set=standard_set,
        standard_quality=standard_quality,
        run_bundle=run_bundle,
        stage_checks=stage_checks,
        findings=findings,
    )


def write_template_generation_standard_quality_outputs(
    out_dir: Path,
    report: TemplateGenerationStandardQualityReport,
) -> StageResult:
    json_path = out_dir / "template_generation_stage_standard_quality_report.json"
    md_path = out_dir / "template_generation_stage_standard_quality_report.md"
    write_json(json_path, report.to_dict())
    write_text(md_path, build_standard_quality_markdown(report))
    write_report_bundle(
        out_dir,
        stage="template_generation_standard_quality",
        status=report.status,
        run_status=Status.PASS,
        quality_status=report.status,
        findings=[finding.to_dict() for finding in report.findings],
        artifacts={
            "template_generation_stage_standard_quality_report": json_path.name,
            "template_generation_stage_standard_quality_report_md": md_path.name,
        },
        stage_statuses=report.stage_statuses,
        blocked_at=_blocked_at_from_findings(report.findings),
    )
    return StageResult(
        "template_generation_standard_quality",
        report.status,
        run_status=Status.PASS,
        quality_status=report.status,
        findings=report.findings,
        artifacts={"report": report.to_dict()},
        artifact_paths={
            "template_generation_stage_standard_quality_report": json_path,
            "template_generation_stage_standard_quality_report_md": md_path,
        },
        blocked_at=_blocked_at_from_findings(report.findings),
    )


def write_template_generation_judge_outputs(
    out_dir: Path,
    report: TemplateGenerationJudgeReport,
) -> StageResult:
    run_bundle_path = out_dir / "template_generation_run_bundle.json"
    stage_checks_path = out_dir / "template_generation_stage_checks.json"
    quality_path = out_dir / "template_generation_stage_standard_quality_report.json"
    quality_md_path = out_dir / "template_generation_stage_standard_quality_report.md"
    judge_path = out_dir / "template_generation_judge_report.json"
    judge_md_path = out_dir / "template_generation_judge_report.md"
    stage_quality_paths = write_stage_standard_quality_reports(out_dir, report)
    stage_diff_paths = write_stage_standard_diff_reports(out_dir, report)
    root_cause_path = out_dir / "template_generation_root_cause_report.json"
    root_cause_md_path = out_dir / "template_generation_root_cause_report.md"
    route_eval_path = out_dir / "template_generation_route_eval_report.json"
    route_eval_md_path = out_dir / "template_generation_route_eval_report.md"

    write_json(run_bundle_path, report.run_bundle.to_dict())
    write_json(stage_checks_path, [check.to_dict() for check in report.stage_checks])
    write_json(quality_path, report.standard_quality.to_dict())
    write_text(quality_md_path, build_standard_quality_markdown(report.standard_quality))
    report_dict = report.to_dict()
    root_cause_report = build_template_generation_root_cause_report(report, report_dict)
    route_eval = build_template_generation_route_eval_report(
        report,
        report_dict,
    )
    write_json(root_cause_path, root_cause_report)
    write_text(root_cause_md_path, build_template_generation_root_cause_markdown(root_cause_report))
    write_json(route_eval_path, route_eval)
    write_text(route_eval_md_path, build_template_generation_route_eval_markdown(route_eval))
    write_json(judge_path, report_dict)
    write_text(judge_md_path, build_judge_markdown(report))

    stage_statuses = {
        check.stage_key: check.status.value for check in report.stage_checks
    }
    stage_quality_artifacts = {
        report_id: paths["json"].name
        for report_id, paths in stage_quality_paths.items()
    }
    stage_quality_artifacts.update(
        {
            f"{report_id}_md": paths["md"].name
            for report_id, paths in stage_quality_paths.items()
        }
    )
    stage_diff_artifacts = {
        report_id: paths["json"].name
        for report_id, paths in stage_diff_paths.items()
    }
    stage_diff_artifacts.update(
        {
            f"{report_id}_md": paths["md"].name
            for report_id, paths in stage_diff_paths.items()
        }
    )
    summary = write_report_bundle(
        out_dir,
        stage="template_generation_judge",
        status=report.status,
        run_status=Status.PASS,
        quality_status=report.status,
        findings=[finding.to_dict() for finding in report.findings],
        artifacts={
            "template_generation_run_bundle": run_bundle_path.name,
            "template_generation_stage_checks": stage_checks_path.name,
            "template_generation_stage_standard_quality_report": quality_path.name,
            "template_generation_stage_standard_quality_report_md": quality_md_path.name,
            "template_generation_judge_report": judge_path.name,
            "template_generation_judge_report_md": judge_md_path.name,
            "template_generation_root_cause_report": root_cause_path.name,
            "template_generation_root_cause_report_md": root_cause_md_path.name,
            "template_generation_route_eval_report": route_eval_path.name,
            "template_generation_route_eval_report_md": route_eval_md_path.name,
            **stage_quality_artifacts,
            **stage_diff_artifacts,
        },
        coverage={
            "standard_quality_status": report.standard_quality.status.value,
            "run_bundle_status": report.run_bundle.status.value,
            "source_run_id": report.run_bundle.source_run_id,
            "first_bad_stage": report.first_bad_stage,
            "standard_acceptance_status": report_dict["standard_acceptance_status"],
            "signoff_status": report_dict["signoff_status"],
            "mismatch_count": len(report_dict["mismatches"]),
            "root_cause_count": len(report_dict["root_causes"]),
            "owner_summary": report_dict["owner_summary"],
            "route_eval_mismatch_count": len(route_eval["mismatches"]),
        },
        stage_statuses=stage_statuses,
        blocked_at=report.first_bad_stage,
    )
    result = report.to_stage_result()
    result.artifact_paths.update(
        {
            "summary": out_dir / "summary.json",
            "findings": out_dir / "findings.json",
            "template_generation_run_bundle": run_bundle_path,
            "template_generation_stage_checks": stage_checks_path,
            "template_generation_stage_standard_quality_report": quality_path,
            "template_generation_stage_standard_quality_report_md": quality_md_path,
            "template_generation_judge_report": judge_path,
            "template_generation_judge_report_md": judge_md_path,
            "template_generation_root_cause_report": root_cause_path,
            "template_generation_root_cause_report_md": root_cause_md_path,
            "template_generation_route_eval_report": route_eval_path,
            "template_generation_route_eval_report_md": route_eval_md_path,
        }
    )
    for report_id, paths in stage_quality_paths.items():
        result.artifact_paths[report_id] = paths["json"]
        result.artifact_paths[f"{report_id}_md"] = paths["md"]
    for report_id, paths in stage_diff_paths.items():
        result.artifact_paths[report_id] = paths["json"]
        result.artifact_paths[f"{report_id}_md"] = paths["md"]
    result.artifacts["summary"] = summary
    result.artifacts["template_generation_route_eval_report"] = route_eval
    return result


def write_stage_standard_quality_reports(
    out_dir: Path,
    report: TemplateGenerationJudgeReport,
) -> dict[str, dict[str, Path]]:
    paths: dict[str, dict[str, Path]] = {}
    for spec in RUN_ARTIFACT_SPECS:
        stage_report = build_stage_standard_quality_report(report, spec)
        report_id = stage_report["report_id"]
        json_path = out_dir / f"{report_id}.json"
        md_path = out_dir / f"{report_id}.md"
        write_json(json_path, stage_report)
        write_text(md_path, build_stage_standard_quality_markdown(stage_report))
        paths[report_id] = {"json": json_path, "md": md_path}
    return paths


def write_stage_standard_diff_reports(
    out_dir: Path,
    report: TemplateGenerationJudgeReport,
) -> dict[str, dict[str, Path]]:
    paths: dict[str, dict[str, Path]] = {}
    for spec in RUN_ARTIFACT_SPECS:
        stage_report = build_stage_standard_diagnosis_report(report, spec)
        report_id = stage_report["report_id"]
        json_path = out_dir / f"{report_id}.json"
        md_path = out_dir / f"{report_id}.md"
        write_json(json_path, stage_report)
        write_text(md_path, build_stage_standard_diagnosis_markdown(stage_report))
        paths[report_id] = {"json": json_path, "md": md_path}
    return paths


def build_stage_standard_diffs(
    report: TemplateGenerationJudgeReport,
) -> list[dict[str, Any]]:
    diffs: list[dict[str, Any]] = []
    for finding in report.standard_quality.findings:
        stage_id, stage_key, report_id = _standard_quality_finding_stage_context(
            report,
            finding,
        )
        owner, owner_detail, next_action = _owner_for_finding(finding)
        diffs.append(
            {
                "diff_id": f"diff_{len(diffs) + 1:03d}",
                "stage_id": stage_id,
                "stage_key": stage_key,
                "report_ref": f"{report_id}.json" if report_id else None,
                "status": finding.status.value,
                "diff_kind": _diff_kind(finding),
                "finding_type": finding.type,
                "message": finding.message,
                "expected": finding.expected,
                "observed": finding.actual,
                "affected_ids": finding.affected_ids,
                "evidence_refs": finding.evidence_refs,
                "root_cause_bucket": finding.root_cause_bucket,
                "owner": owner,
                "owner_detail": owner_detail,
                "next_action": next_action,
            }
        )
    for check in report.stage_checks:
        report_id = _stage_quality_report_id_for_stage(check.stage_key)
        for finding in check.findings:
            owner, owner_detail, next_action = _owner_for_finding(finding)
            diffs.append(
                {
                    "diff_id": f"diff_{len(diffs) + 1:03d}",
                    "stage_id": check.stage_id,
                    "stage_key": check.stage_key,
                    "report_ref": f"{report_id}.json" if report_id else None,
                    "status": finding.status.value,
                    "diff_kind": _diff_kind(finding),
                    "finding_type": finding.type,
                    "message": finding.message,
                    "expected": finding.expected,
                    "observed": finding.actual,
                    "affected_ids": finding.affected_ids,
                    "evidence_refs": finding.evidence_refs,
                    "root_cause_bucket": finding.root_cause_bucket,
                    "owner": owner,
                    "owner_detail": owner_detail,
                    "next_action": next_action,
                }
            )
    return diffs


def build_root_causes(
    stage_standard_diffs: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    root_causes: list[dict[str, Any]] = []
    for diff in stage_standard_diffs:
        root_causes.append(
            {
                "root_cause_id": f"rc_{len(root_causes) + 1:03d}",
                "diff_id": diff["diff_id"],
                "stage_id": diff["stage_id"],
                "stage_key": diff["stage_key"],
                "status": diff["status"],
                "finding_type": diff["finding_type"],
                "root_cause_bucket": diff["root_cause_bucket"],
                "owner": diff["owner"],
                "owner_detail": diff["owner_detail"],
                "reason": _root_cause_reason(diff),
                "next_action": diff["next_action"],
            }
        )
    return root_causes


def build_mismatches(
    report: TemplateGenerationJudgeReport,
    *,
    stage_key: str | None = None,
) -> list[dict[str, Any]]:
    mismatches: list[dict[str, Any]] = []
    counters: dict[str, int] = {}
    for finding in report.standard_quality.findings:
        stage_id, quality_stage_key, _report_id = _standard_quality_finding_stage_context(
            report,
            finding,
        )
        if stage_key is not None and quality_stage_key != stage_key:
            continue
        standard = report.standard_set.stages.get(quality_stage_key)
        counters[stage_id] = counters.get(stage_id, 0) + 1
        mismatch_id = f"{stage_id}-MISMATCH-{counters[stage_id]:03d}"
        finding_type = _normalized_mismatch_type(finding.type)
        mismatches.append(
            {
                "id": mismatch_id,
                "mismatch_id": mismatch_id,
                "stage_id": stage_id,
                "stage_key": quality_stage_key,
                "status": finding.status.value,
                "finding_type": finding.type,
                "type": finding_type,
                "field": _mismatch_field(finding_type, finding),
                "expected": _parse_finding_value(finding.expected),
                "observed": _parse_finding_value(finding.actual),
                "problem": _mismatch_problem(finding_type, finding),
                "affected_ids": finding.affected_ids,
                "evidence": {
                    "artifact_path": None,
                    "standard_path": str(standard.path) if standard is not None else None,
                    "source_seq_refs": _source_seq_refs_from_finding(finding),
                    "evidence_refs": finding.evidence_refs,
                },
                "root_cause_bucket": finding.root_cause_bucket,
                "message": finding.message,
            }
        )
    for check in report.stage_checks:
        if stage_key is not None and check.stage_key != stage_key:
            continue
        standard = report.standard_set.stages.get(check.stage_key)
        artifact = report.run_bundle.artifact_for_stage(check.stage_key)
        for finding in check.findings:
            stage_id = check.stage_id
            counters[stage_id] = counters.get(stage_id, 0) + 1
            mismatch_id = f"{stage_id}-MISMATCH-{counters[stage_id]:03d}"
            finding_type = _normalized_mismatch_type(finding.type)
            field = _mismatch_field(finding_type, finding)
            expected = _parse_finding_value(finding.expected)
            observed = _parse_finding_value(finding.actual)
            mismatches.append(
                {
                    "id": mismatch_id,
                    "mismatch_id": mismatch_id,
                    "stage_id": stage_id,
                    "stage_key": check.stage_key,
                    "status": finding.status.value,
                    "finding_type": finding.type,
                    "type": finding_type,
                    "field": field,
                    "expected": expected,
                    "observed": observed,
                    "problem": _mismatch_problem(finding_type, finding),
                    "affected_ids": finding.affected_ids,
                    "evidence": {
                        "artifact_path": str(artifact.path)
                        if artifact is not None and artifact.path is not None
                        else None,
                        "standard_path": str(standard.path) if standard is not None else None,
                        "source_seq_refs": _source_seq_refs_from_finding(finding),
                        "evidence_refs": finding.evidence_refs,
                    },
                    "root_cause_bucket": finding.root_cause_bucket,
                    "message": finding.message,
                }
            )
    return mismatches


def build_diagnosis_root_causes(
    report: TemplateGenerationJudgeReport,
    mismatches: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    root_causes: list[dict[str, Any]] = []
    for mismatch in mismatches:
        category = _root_cause_category(report, mismatch)
        legacy_owner, owner_detail, next_action = _legacy_owner_for_mismatch(mismatch)
        root_causes.append(
            {
                "id": _derived_id(mismatch["id"], "ROOT-CAUSE"),
                "root_cause_id": f"rc_{len(root_causes) + 1:03d}",
                "mismatch_id": mismatch["id"],
                "stage_id": mismatch["stage_id"],
                "stage_key": mismatch["stage_key"],
                "status": mismatch["status"],
                "finding_type": mismatch["finding_type"],
                "category": category,
                "first_bad_stage": _diagnosis_first_bad_stage(report, mismatch),
                "reason": _diagnosis_reason(category, mismatch),
                "root_cause_bucket": mismatch["root_cause_bucket"],
                "owner": legacy_owner,
                "owner_detail": owner_detail,
                "next_action": next_action,
            }
        )
    return root_causes


def build_owner_assignments(
    root_causes: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    assignments: list[dict[str, Any]] = []
    for root_cause in root_causes:
        primary, secondary = _owner_assignment_for_category(root_cause["category"])
        assignments.append(
            {
                "id": _derived_id(root_cause["mismatch_id"], "OWNER"),
                "mismatch_id": root_cause["mismatch_id"],
                "root_cause_id": root_cause["id"],
                "primary": primary,
                "secondary": secondary,
                "rationale": _owner_assignment_rationale(root_cause, primary, secondary),
            }
        )
    return assignments


def build_fix_plans(
    root_causes: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    plans: list[dict[str, Any]] = []
    for root_cause in root_causes:
        plans.append(
            {
                "id": _derived_id(root_cause["mismatch_id"], "FIX"),
                "mismatch_id": root_cause["mismatch_id"],
                "root_cause_id": root_cause["id"],
                **_fix_plan_for_root_cause(root_cause),
            }
        )
    return plans


def build_stage_standard_diagnosis_report(
    report: TemplateGenerationJudgeReport,
    spec: RunArtifactSpec,
) -> dict[str, Any]:
    artifact = report.run_bundle.artifacts.get(spec.artifact_key)
    check = _stage_check_for_spec(report, spec)
    standard = (
        report.standard_set.stages.get(spec.stage_key)
        if spec.stage_key is not None
        else None
    )
    stage_key = spec.stage_key or _implicit_stage_key(spec)
    stage_id = spec.stage_id or "T0"
    status = check.status if check is not None else (
        artifact.status if artifact is not None else Status.UNKNOWN
    )
    report_id = _stage_diff_report_id(spec)
    mismatches = build_mismatches(report, stage_key=spec.stage_key) if spec.stage_key else []
    root_causes = build_diagnosis_root_causes(report, mismatches)
    owner_assignments = build_owner_assignments(root_causes)
    fix_plan = build_fix_plans(root_causes)
    _attach_diagnosis_layers(mismatches, root_causes, owner_assignments, fix_plan)
    stage_diffs = [
        diff
        for diff in build_stage_standard_diffs(report)
        if diff["stage_key"] == stage_key
    ]
    return {
        "artifact_type": "template_generation_stage_standard_diff_report",
        "artifact_version": "1.0",
        "report_kind": "standard_diff_report",
        "report_id": report_id,
        "status": status.value,
        "source_run_id": report.run_bundle.source_run_id,
        "source_run_dir": str(report.run_bundle.source_run_dir),
        "stage_id": stage_id,
        "stage_key": stage_key,
        "artifact_key": spec.artifact_key,
        "artifact_name": spec.top_level_name,
        "artifact_under_test": (
            standard.artifact_under_test if standard is not None else spec.artifact_key
        ),
        "artifact_path": str(artifact.path) if artifact and artifact.path else None,
        "artifact_sha256": artifact.sha256 if artifact is not None else None,
        "standard_path": str(standard.path) if standard is not None else None,
        "standard_sha256": standard.sha256 if standard is not None else None,
        "verifier_state": (
            check.verifier_state
            if check is not None
            else (standard.verifier_state if standard is not None else "not_applicable")
        ),
        "gate_enabled": (
            check.gate_enabled
            if check is not None
            else (standard.gate_enabled if standard is not None else None)
        ),
        "audit_status": check.audit_status if check is not None else None,
        "mismatches": mismatches,
        "root_causes": root_causes,
        "owner_assignments": owner_assignments,
        "fix_plan": fix_plan,
        "stage_standard_diffs": stage_diffs,
    }


def build_template_generation_root_cause_report(
    report: TemplateGenerationJudgeReport,
    report_dict: dict[str, Any] | None = None,
) -> dict[str, Any]:
    report_dict = report_dict or report.to_dict()
    return {
        "artifact_type": "template_generation_root_cause_report",
        "artifact_version": "1.0",
        "report_kind": "root_cause_report",
        "status": report_dict["status"],
        "standard_acceptance_status": report_dict["standard_acceptance_status"],
        "signoff_status": report_dict["signoff_status"],
        "first_bad_stage": report_dict["first_bad_stage"],
        "source_run_id": report_dict["source_run_id"],
        "mismatches": report_dict["mismatches"],
        "root_causes": report_dict["root_causes"],
        "owner_assignments": report_dict["owner_assignments"],
        "fix_plan": report_dict["fix_plan"],
        "owner_summary": report_dict["owner_summary"],
        "top_blockers": report_dict["top_blockers"],
    }


def _stage_accuracy_metric(
    report: TemplateGenerationJudgeReport,
    check: StageCheck,
) -> dict[str, Any]:
    dispatch = {
        "t2_unit_pagination": _t2_accuracy_metric,
        "t3_element_policy": _t3_accuracy_metric,
        "t4_global_layout": _t4_accuracy_metric,
        "t5_template_spec": _t5_accuracy_metric,
    }
    builder = dispatch.get(check.stage_key)
    if builder is None:
        return _base_accuracy_metric(check, primary_accuracy=None)
    return builder(report, check)


def _base_accuracy_metric(
    check: StageCheck,
    *,
    primary_accuracy: float | None,
    **extra: Any,
) -> dict[str, Any]:
    return {
        "stage_id": check.stage_id,
        "stage_key": check.stage_key,
        "status": check.status.value,
        "audit_status": check.audit_status,
        "primary_accuracy": primary_accuracy,
        **extra,
    }


def _t2_accuracy_metric(
    report: TemplateGenerationJudgeReport,
    check: StageCheck,
) -> dict[str, Any]:
    del report
    audit = check.audit
    expected = [str(item) for item in audit.get("expected_unit_ids", []) or []]
    actual = [str(item) for item in audit.get("actual_unit_ids", []) or []]
    precision, recall, f1 = _set_prf(actual, expected)
    return _base_accuracy_metric(
        check,
        primary_accuracy=f1,
        precision=precision,
        recall=recall,
        f1=f1,
        order_exact_match=bool(audit.get("unit_order_matches")),
        missing_units=audit.get("missing_unit_ids", []),
        unexpected_units=audit.get("unexpected_unit_ids", []),
    )


def _t3_accuracy_metric(
    report: TemplateGenerationJudgeReport,
    check: StageCheck,
) -> dict[str, Any]:
    core_action_accuracy = check.audit.get("core_action_accuracy")
    if isinstance(core_action_accuracy, (int, float)):
        gold_count = int(check.audit.get("core_action_gold_count") or 0)
        unknown_gold_count = int(
            check.audit.get("core_action_unknown_gold_count") or 0
        )
        match_count = int(check.audit.get("core_action_match_count") or 0)
        return _base_accuracy_metric(
            check,
            primary_accuracy=float(core_action_accuracy),
            primary_metric="exact_action_accuracy",
            accuracy_level="core_action",
            allowed_actions=["keep", "fill", "delete"],
            gold_run_count=gold_count,
            excluded_unknown_gold_run_count=unknown_gold_count,
            unknown_scoring="excluded_from_primary",
            unknown_execution_fallback="keep",
            matched_run_count=match_count,
            mismatch_run_count=max(gold_count - match_count, 0),
            grouping_invariant=True,
        )
    standard = report.standard_set.stages.get(check.stage_key)
    expected_policy_units = _expected_policy_units(standard.expected if standard else {})
    conflicts = [
        item
        for item in check.audit.get("policy_group_conflicts", []) or []
        if isinstance(item, dict)
    ]
    conflict_units = {
        str(item.get("unit_id") or "")
        for item in conflicts
        if item.get("unit_id") not in (None, "")
    }
    policy_conformance = (
        round((len(expected_policy_units) - len(conflict_units)) / len(expected_policy_units), 4)
        if expected_policy_units
        else None
    )
    artifact = report.run_bundle.artifact_for_stage(check.stage_key)
    elements = (
        artifact.payload.get("elements", [])
        if artifact is not None and isinstance(artifact.payload, dict)
        else []
    )
    element_count = len(elements) if isinstance(elements, list) else 0
    required_gaps = check.audit.get("required_policy_field_gaps", []) or []
    required_field_compliance = (
        round(max(element_count - len(required_gaps), 0) / element_count, 4)
        if element_count
        else (1.0 if not required_gaps else 0.0)
    )
    primary = _avg_numbers([policy_conformance, required_field_compliance])
    return _base_accuracy_metric(
        check,
        primary_accuracy=primary,
        policy_group_conformance=policy_conformance,
        required_field_compliance=required_field_compliance,
        policy_conflict_count=len(conflicts),
        policy_conflict_units=sorted(conflict_units),
        required_policy_field_gap_count=len(required_gaps),
    )


def _t4_accuracy_metric(
    report: TemplateGenerationJudgeReport,
    check: StageCheck,
) -> dict[str, Any]:
    del report
    audit = check.audit
    missing_fields = audit.get("missing_evidence_fields", []) or []
    components = [
        1.0 if audit.get("has_global_layout_contract") else 0.0,
        1.0 if not missing_fields else 0.0,
        1.0 if int(audit.get("section_profile_count") or 0) > 0 else 0.0,
        1.0 if audit.get("page_numbering_status") not in (None, "", "missing") else 0.0,
    ]
    return _base_accuracy_metric(
        check,
        primary_accuracy=round(sum(components) / len(components), 4),
        layout_contract_completeness=round(sum(components) / len(components), 4),
        missing_evidence_fields=missing_fields,
        section_profile_count=audit.get("section_profile_count"),
        page_numbering_status=audit.get("page_numbering_status"),
        metric_note="T4 standard currently exposes contract completeness, not visual layout precision.",
    )


def _t5_accuracy_metric(
    report: TemplateGenerationJudgeReport,
    check: StageCheck,
) -> dict[str, Any]:
    del report
    audit = check.audit
    expected = [str(item) for item in audit.get("expected_unit_order", []) or []]
    actual = [str(item) for item in audit.get("actual_unit_order", []) or []]
    precision, recall, f1 = _set_prf(actual, expected)
    missing_section_refs = audit.get("missing_section_profile_refs", []) or []
    section_ref_compliance = (
        round(max(len(actual) - len(missing_section_refs), 0) / len(actual), 4)
        if actual
        else None
    )
    primary = _avg_numbers([f1, section_ref_compliance])
    return _base_accuracy_metric(
        check,
        primary_accuracy=primary,
        precision=precision,
        recall=recall,
        f1=f1,
        section_ref_compliance=section_ref_compliance,
        missing_section_profile_refs=missing_section_refs,
        missing_input_hashes=audit.get("missing_input_hashes", []),
        missing_review_flags=audit.get("missing_review_flags", []),
    )


def _set_prf(actual: list[str], expected: list[str]) -> tuple[float, float, float]:
    actual_set = set(actual)
    expected_set = set(expected)
    hit = actual_set & expected_set
    precision = len(hit) / len(actual_set) if actual_set else 0.0
    recall = len(hit) / len(expected_set) if expected_set else 0.0
    f1 = (2 * precision * recall / (precision + recall)) if precision + recall else 0.0
    return round(precision, 4), round(recall, 4), round(f1, 4)


def _expected_policy_units(expected: dict[str, Any]) -> set[str]:
    groups = expected.get("policy_groups", {}) if isinstance(expected, dict) else {}
    if not isinstance(groups, dict):
        return set()
    result: set[str] = set()
    for values in groups.values():
        if isinstance(values, list):
            result.update(str(value) for value in values if value not in (None, ""))
    return result


def _avg_numbers(values: list[float | None]) -> float | None:
    numeric = [value for value in values if isinstance(value, (int, float))]
    if not numeric:
        return None
    return round(sum(numeric) / len(numeric), 4)


def _load_run_json(run_dir: Path, *relative_paths: str) -> dict[str, Any] | None:
    for relative_path in relative_paths:
        path = run_dir / relative_path
        if not path.exists():
            continue
        try:
            value = read_json(path)
        except Exception:
            continue
        if isinstance(value, dict):
            return value
    return None


def _load_run_yaml(run_dir: Path, *relative_paths: str) -> dict[str, Any] | None:
    for relative_path in relative_paths:
        path = run_dir / relative_path
        if not path.exists():
            continue
        try:
            value = read_yaml(path)
        except Exception:
            continue
        if isinstance(value, dict):
            return value
    return None


def build_standard_acceptance_summary(
    report: TemplateGenerationJudgeReport,
    *,
    stage_standard_diffs: list[dict[str, Any]] | None = None,
    root_causes: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    stage_standard_diffs = (
        stage_standard_diffs
        if stage_standard_diffs is not None
        else build_stage_standard_diffs(report)
    )
    root_causes = (
        root_causes
        if root_causes is not None
        else build_root_causes(stage_standard_diffs)
    )
    owner_summary = _owner_summary(root_causes)
    has_fail = (
        any(diff["status"] == Status.FAIL.value for diff in stage_standard_diffs)
        or any(
            finding.status == Status.FAIL
            for finding in report.run_bundle.observed_t6_findings
        )
    )
    has_unknown = (
        report.standard_quality.status != Status.PASS
        or report.run_bundle.status != Status.PASS
        or any(
            finding.status == Status.UNKNOWN
            for finding in report.run_bundle.observed_t6_findings
        )
        or any(diff["status"] == Status.UNKNOWN.value for diff in stage_standard_diffs)
        or any(check.status == Status.UNKNOWN for check in report.stage_checks)
    )
    if has_fail:
        standard_acceptance_status = Status.FAIL.value
    elif has_unknown:
        standard_acceptance_status = Status.UNKNOWN.value
    else:
        standard_acceptance_status = Status.PASS.value

    signoff_blockers = _signoff_blockers(
        report,
        standard_acceptance_status=standard_acceptance_status,
        stage_standard_diffs=stage_standard_diffs,
    )
    signoff_status = "SIGNABLE" if not signoff_blockers else "NOT_SIGNABLE"
    top_blockers = _top_blockers(root_causes)
    return {
        "standard_acceptance_status": standard_acceptance_status,
        "signoff_status": signoff_status,
        "signoff_blockers": signoff_blockers,
        "owner_summary": owner_summary,
        "top_blockers": top_blockers,
    }


def build_stage_standard_quality_report(
    report: TemplateGenerationJudgeReport,
    spec: RunArtifactSpec,
) -> dict[str, Any]:
    artifact = report.run_bundle.artifacts.get(spec.artifact_key)
    check = _stage_check_for_spec(report, spec)
    standard = (
        report.standard_set.stages.get(spec.stage_key)
        if spec.stage_key is not None
        else None
    )
    findings = check.findings if check is not None else _run_bundle_findings_for_spec(report, spec)
    status = check.status if check is not None else merge_statuses(
        [artifact.status if artifact is not None else Status.UNKNOWN]
        + [finding.status for finding in findings]
    )
    stage_key = spec.stage_key or _implicit_stage_key(spec)
    stage_id = spec.stage_id or "T0"
    report_id = _stage_quality_report_id(spec)
    stage_diffs = [
        diff
        for diff in build_stage_standard_diffs(report)
        if diff["stage_key"] == stage_key
    ]
    mismatches = build_mismatches(report, stage_key=spec.stage_key) if spec.stage_key else []
    root_causes = build_diagnosis_root_causes(report, mismatches)
    owner_assignments = build_owner_assignments(root_causes)
    fix_plan = build_fix_plans(root_causes)
    _attach_diagnosis_layers(mismatches, root_causes, owner_assignments, fix_plan)
    acceptance = _artifact_acceptance_summary(
        status=status,
        check=check,
        stage_diffs=stage_diffs,
        artifact_status=artifact.status if artifact is not None else Status.UNKNOWN,
    )
    return {
        "artifact_type": "template_generation_artifact_standard_quality_report",
        "artifact_version": "1.0",
        "report_kind": "standard_quality_report",
        "report_id": report_id,
        "status": status.value,
        "standard_acceptance_status": acceptance["standard_acceptance_status"],
        "signoff_status": acceptance["signoff_status"],
        "signoff_blockers": acceptance["signoff_blockers"],
        "source_run_id": report.run_bundle.source_run_id,
        "source_run_dir": str(report.run_bundle.source_run_dir),
        "stage_id": stage_id,
        "stage_key": stage_key,
        "artifact_key": spec.artifact_key,
        "artifact_name": spec.top_level_name,
        "artifact_under_test": (
            standard.artifact_under_test if standard is not None else spec.artifact_key
        ),
        "artifact_path": str(artifact.path) if artifact and artifact.path else None,
        "artifact_sha256": artifact.sha256 if artifact is not None else None,
        "artifact_status": artifact.status.value if artifact is not None else Status.UNKNOWN.value,
        "artifact_source_kind": artifact.source_kind if artifact is not None else "missing",
        "artifact_hash_match": artifact.hash_match if artifact is not None else None,
        "standard_path": str(standard.path) if standard is not None else None,
        "standard_sha256": standard.sha256 if standard is not None else None,
        "standard_quality_status": (
            report.standard_quality.stage_statuses.get(spec.stage_key, Status.UNKNOWN.value)
            if spec.stage_key is not None
            else "not_applicable"
        ),
        "verifier_state": (
            check.verifier_state
            if check is not None
            else (standard.verifier_state if standard is not None else "not_applicable")
        ),
        "gate_enabled": (
            check.gate_enabled
            if check is not None
            else (standard.gate_enabled if standard is not None else None)
        ),
        "audit_status": check.audit_status if check is not None else None,
        "audit": check.audit if check is not None else {},
        "comparison_scope": (
            "stage_standard"
            if spec.stage_key is not None
            else (
                "run_bundle_binding_and_final_docx_observation"
                if spec.stage_id == "T6"
                else "run_bundle_artifact_binding"
            )
        ),
        "mismatches": mismatches,
        "stage_standard_diffs": stage_diffs,
        "root_causes": root_causes,
        "owner_assignments": owner_assignments,
        "fix_plan": fix_plan,
        "owner_summary": _owner_summary(root_causes),
        "top_blockers": _top_blockers(root_causes),
        "findings": [finding.to_dict() for finding in findings],
    }


def build_stage_standard_quality_markdown(report: dict[str, Any]) -> str:
    lines = [
        f"# {report['report_id']}",
        "",
        f"- Status: {report['status']}",
        f"- Standard acceptance: {report['standard_acceptance_status']}",
        f"- Sign-off: {report['signoff_status']}",
        f"- Stage: {report['stage_id']} / {report['stage_key']}",
        f"- Artifact: {report['artifact_name']}",
        f"- Artifact path: {report.get('artifact_path') or 'missing'}",
        f"- Artifact sha256: {report.get('artifact_sha256') or 'missing'}",
        f"- Artifact status: {report['artifact_status']}",
        f"- Standard path: {report.get('standard_path') or 'not_applicable'}",
        f"- Standard quality: {report['standard_quality_status']}",
        f"- Verifier state: {report['verifier_state']}",
        f"- Gate enabled: {report.get('gate_enabled')}",
        f"- Audit status: {report.get('audit_status') or 'not_applicable'}",
        "",
    ]
    if report.get("signoff_blockers"):
        lines.append("## Sign-off Blockers")
        for blocker in report["signoff_blockers"]:
            lines.append(f"- {blocker}")
        lines.append("")
    if report.get("top_blockers"):
        lines.append("## Top Blockers")
        for blocker in report["top_blockers"]:
            lines.append(
                "- "
                f"{blocker.get('stage_id')} {blocker.get('finding_type')}: "
                f"owner={blocker.get('owner')}, category={blocker.get('category')}, "
                f"next={blocker.get('next_action')}"
            )
        lines.append("")
    lines.append("## Diagnosis")
    if report.get("mismatches"):
        for mismatch in report["mismatches"]:
            lines.append(
                "- "
                f"{mismatch.get('id')} {mismatch.get('field')}: "
                f"{mismatch.get('problem')}"
            )
    else:
        lines.append("- mismatches: none")
    lines.append("")
    findings = report.get("findings") or []
    if findings:
        lines.append("## Findings")
        for finding in findings:
            lines.append(
                "- "
                f"[{finding.get('status')}] "
                f"{finding.get('stage')}/{finding.get('type')}: "
                f"{finding.get('message')}"
            )
    else:
        lines.append("No findings.")
    lines.append("")
    return "\n".join(lines)


def build_stage_standard_diagnosis_markdown(report: dict[str, Any]) -> str:
    lines = [
        f"# {report['report_id']}",
        "",
        f"- Status: {report['status']}",
        f"- Stage: {report['stage_id']} / {report['stage_key']}",
        f"- Artifact: {report['artifact_name']}",
        f"- Artifact path: {report.get('artifact_path') or 'missing'}",
        f"- Standard path: {report.get('standard_path') or 'not_applicable'}",
        f"- Verifier state: {report['verifier_state']}",
        f"- Gate enabled: {report.get('gate_enabled')}",
        "",
        "## Mismatches",
    ]
    if report["mismatches"]:
        for mismatch in report["mismatches"]:
            lines.append(
                "- "
                f"{mismatch['id']} {mismatch['field']}: {mismatch['problem']}"
            )
    else:
        lines.append("- none")
    lines.extend(["", "## Root Causes"])
    if report["root_causes"]:
        for root_cause in report["root_causes"]:
            lines.append(
                "- "
                f"{root_cause['mismatch_id']}: {root_cause['category']} "
                f"(first_bad_stage={root_cause['first_bad_stage']})"
            )
    else:
        lines.append("- none")
    lines.extend(["", "## Owner Assignments"])
    if report["owner_assignments"]:
        for assignment in report["owner_assignments"]:
            lines.append(
                "- "
                f"{assignment['mismatch_id']}: primary={assignment['primary']}, "
                f"secondary={assignment.get('secondary') or 'none'}"
            )
    else:
        lines.append("- none")
    lines.extend(["", "## Fix Plan"])
    if report["fix_plan"]:
        for plan in report["fix_plan"]:
            lines.append(f"- {plan['mismatch_id']}: {plan['action']}")
    else:
        lines.append("- none")
    lines.append("")
    return "\n".join(lines)


def build_template_generation_root_cause_markdown(report: dict[str, Any]) -> str:
    lines = [
        "# Template Generation Root Cause Report",
        "",
        f"- Status: {report['status']}",
        f"- Standard acceptance: {report['standard_acceptance_status']}",
        f"- Sign-off: {report['signoff_status']}",
        f"- First bad stage: {report.get('first_bad_stage') or 'none'}",
        f"- Source run id: {report['source_run_id']}",
        "",
        "## Root Causes",
    ]
    if report["root_causes"]:
        for root_cause in report["root_causes"]:
            lines.append(
                "- "
                f"{root_cause['mismatch_id']}: {root_cause['category']} "
                f"owner={root_cause['owner']} reason={root_cause['reason']}"
            )
    else:
        lines.append("- none")
    lines.extend(["", "## Fix Plan"])
    if report["fix_plan"]:
        for plan in report["fix_plan"]:
            lines.append(f"- {plan['mismatch_id']}: {plan['action']}")
    else:
        lines.append("- none")
    lines.append("")
    return "\n".join(lines)


def _stage_standard_quality_report_refs(
    report: TemplateGenerationJudgeReport,
) -> list[dict[str, Any]]:
    refs = []
    checks = {check.stage_key: check for check in report.stage_checks}
    for spec in RUN_ARTIFACT_SPECS:
        artifact = report.run_bundle.artifacts.get(spec.artifact_key)
        check = checks.get(spec.stage_key) if spec.stage_key else None
        status = check.status if check is not None else (
            artifact.status if artifact is not None else Status.UNKNOWN
        )
        report_id = _stage_quality_report_id(spec)
        refs.append(
            {
                "report_id": report_id,
                "json_report": f"{report_id}.json",
                "markdown_report": f"{report_id}.md",
                "stage_id": spec.stage_id or "T0",
                "stage_key": spec.stage_key or _implicit_stage_key(spec),
                "artifact_key": spec.artifact_key,
                "artifact_name": spec.top_level_name,
                "status": status.value,
            }
        )
    return refs


def _stage_standard_diff_report_refs(
    report: TemplateGenerationJudgeReport,
) -> list[dict[str, Any]]:
    refs = []
    checks = {check.stage_key: check for check in report.stage_checks}
    for spec in RUN_ARTIFACT_SPECS:
        artifact = report.run_bundle.artifacts.get(spec.artifact_key)
        check = checks.get(spec.stage_key) if spec.stage_key else None
        status = check.status if check is not None else (
            artifact.status if artifact is not None else Status.UNKNOWN
        )
        report_id = _stage_diff_report_id(spec)
        refs.append(
            {
                "report_id": report_id,
                "json_report": f"{report_id}.json",
                "markdown_report": f"{report_id}.md",
                "stage_id": spec.stage_id or "T0",
                "stage_key": spec.stage_key or _implicit_stage_key(spec),
                "artifact_key": spec.artifact_key,
                "artifact_name": spec.top_level_name,
                "status": status.value,
            }
        )
    return refs


def _artifact_acceptance_summary(
    *,
    status: Status,
    check: StageCheck | None,
    stage_diffs: list[dict[str, Any]],
    artifact_status: Status,
) -> dict[str, Any]:
    has_fail = (
        status == Status.FAIL
        or artifact_status == Status.FAIL
        or any(diff["status"] == Status.FAIL.value for diff in stage_diffs)
    )
    has_unknown = (
        status == Status.UNKNOWN
        or artifact_status == Status.UNKNOWN
        or any(diff["status"] == Status.UNKNOWN.value for diff in stage_diffs)
        or (
            check is not None
            and (check.verifier_state != "configured" or not check.gate_enabled)
        )
    )
    if has_fail:
        standard_acceptance_status = Status.FAIL.value
    elif has_unknown:
        standard_acceptance_status = Status.UNKNOWN.value
    else:
        standard_acceptance_status = Status.PASS.value
    signoff_blockers: list[str] = []
    if standard_acceptance_status != Status.PASS.value:
        signoff_blockers.append(
            f"standard_acceptance_status={standard_acceptance_status}"
        )
    if check is not None and check.verifier_state != "configured":
        signoff_blockers.append(f"verifier_state={check.verifier_state}")
    if check is not None and not check.gate_enabled:
        signoff_blockers.append("gate_enabled=false")
    signoff_status = "SIGNABLE" if not signoff_blockers else "NOT_SIGNABLE"
    return {
        "standard_acceptance_status": standard_acceptance_status,
        "signoff_status": signoff_status,
        "signoff_blockers": signoff_blockers,
    }


def _signoff_blockers(
    report: TemplateGenerationJudgeReport,
    *,
    standard_acceptance_status: str,
    stage_standard_diffs: list[dict[str, Any]],
) -> list[str]:
    blockers: list[str] = []
    if standard_acceptance_status != Status.PASS.value:
        blockers.append(f"standard_acceptance_status={standard_acceptance_status}")
    if report.standard_quality.status != Status.PASS:
        blockers.append(f"standard_quality={report.standard_quality.status.value}")
    if report.run_bundle.status != Status.PASS:
        blockers.append(f"run_bundle={report.run_bundle.status.value}")
    for check in report.stage_checks:
        if check.verifier_state != "configured":
            blockers.append(f"{check.stage_key}.verifier_state={check.verifier_state}")
        if not check.gate_enabled:
            blockers.append(f"{check.stage_key}.gate_enabled=false")
    if any(diff["status"] == Status.FAIL.value for diff in stage_standard_diffs):
        blockers.append("stage_standard_diffs contain FAIL")
    return sorted(set(blockers))


def _owner_summary(root_causes: list[dict[str, Any]]) -> dict[str, int]:
    summary = {
        "code": 0,
        "verifier": 0,
        "standard": 0,
        "ai_prompt": 0,
        "unknown": 0,
    }
    for root_cause in root_causes:
        owner = str(root_cause.get("owner") or "unknown")
        summary[owner if owner in summary else "unknown"] += 1
    return summary


def _top_blockers(root_causes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    priority = {
        Status.FAIL.value: 0,
        Status.UNKNOWN.value: 1,
        Status.PASS.value: 2,
    }
    ordered = sorted(
        root_causes,
        key=lambda item: (
            priority.get(str(item.get("status")), 9),
            str(item.get("stage_id")),
            str(item.get("finding_type")),
        ),
    )
    return ordered[:20]


def _attach_diagnosis_layers(
    mismatches: list[dict[str, Any]],
    root_causes: list[dict[str, Any]],
    owner_assignments: list[dict[str, Any]],
    fix_plan: list[dict[str, Any]],
) -> None:
    root_by_mismatch = {item["mismatch_id"]: item for item in root_causes}
    owner_by_mismatch = {item["mismatch_id"]: item for item in owner_assignments}
    fix_by_mismatch = {item["mismatch_id"]: item for item in fix_plan}
    for mismatch in mismatches:
        root_cause = root_by_mismatch.get(mismatch["id"])
        owner = owner_by_mismatch.get(mismatch["id"])
        plan = fix_by_mismatch.get(mismatch["id"])
        if root_cause is not None:
            mismatch["root_cause"] = {
                "category": root_cause["category"],
                "first_bad_stage": root_cause["first_bad_stage"],
                "reason": root_cause["reason"],
            }
        if owner is not None:
            mismatch["owner"] = {
                "primary": owner["primary"],
                "secondary": owner["secondary"],
                "rationale": owner["rationale"],
            }
        if plan is not None:
            mismatch["fix_plan"] = {
                "action": plan["action"],
                "likely_files": plan["likely_files"],
                "tests": plan["tests"],
                "acceptance": plan["acceptance"],
            }


def _normalized_mismatch_type(finding_type: str) -> str:
    mapping = {
        "t2_standard_unit_order_mismatch": "t2_unit_order_mismatch",
        "t2_standard_units_missing": "t2_expected_unit_missing",
        "t2_standard_units_unexpected": "t2_unexpected_unit_present",
        "t2_standard_custom_units_present": "t2_unexpected_unit_present",
        "t2_standard_anchor_owner_mismatch": "t2_anchor_owner_mismatch",
        "t5_required_input_hashes_missing": "t5_input_hash_missing",
        "t5_unit_section_profile_refs_missing": "t5_section_profile_refs_missing",
    }
    return mapping.get(finding_type, finding_type)


def _mismatch_field(finding_type: str, finding: Finding) -> str:
    field_by_type = {
        "t1_artifact_type_mismatch": "artifact_type",
        "t1_required_top_level_fields_missing": "required_top_level_fields",
        "t1_required_data_groups_missing": "data",
        "t1_visible_body_flow_locator_missing": "body_flow[].source_seq/source_ref",
        "t1_forbidden_semantic_fields_present": "forbidden_semantic_fields",
        "t2_unit_order_mismatch": "expected.unit_order",
        "t2_expected_unit_missing": "expected.units[].unit_id",
        "t2_unexpected_unit_present": "units[].unit_id",
        "t2_anchor_owner_mismatch": "expected.units[].anchors",
        "t2_source_range_mismatch": "units[].source_seq_refs",
        "t2_page_policy_mismatch": "units[].page_policy",
        "t2_standard_schema_invalid": "standard.expected",
        "t2_artifact_type_mismatch": "artifact_type",
        "t3_artifact_type_mismatch": "artifact_type",
        "t3_unit_order_mismatch": "expected.unit_order",
        "t3_policy_group_conflict": "expected.policy_groups",
        "t3_required_policy_fields_missing": (
            "expected.element_policy_contract.required_fields_by_policy"
        ),
        "t3_standard_element_expectations_missing": (
            "expected.element_expectations"
        ),
        "t3_standard_run_span_ledger_missing": "expected.run_span_ledger",
        "t4_artifact_type_mismatch": "artifact_type",
        "t4_global_layout_contract_missing": "expected.global_layout_contract",
        "t4_global_spec_evidence_fields_missing": "global_spec",
        "t4_section_profiles_missing": "section_profiles",
        "t4_page_numbering_missing": "page_numbering",
        "t5_artifact_type_mismatch": "artifact_type",
        "t5_unit_order_mismatch": "expected.unit_order",
        "t5_input_hash_missing": "input_hashes",
        "t5_review_flags_dropped": "review_flags",
        "t5_section_profile_refs_missing": "units[].section_profile_refs",
        "template_generation_stage_standard_missing": "stage_standard",
        "template_generation_stage_artifact_missing": "artifact",
        "template_generation_stage_standard_quality_not_pass": "standard_quality",
        "template_generation_run_bundle_not_pass": "run_bundle",
        "template_generation_stage_verifier_not_configured": "verifier_state",
        "template_generation_stage_gate_disabled": "gate_enabled",
    }
    return field_by_type.get(finding_type, finding.type)


def _mismatch_problem(finding_type: str, finding: Finding) -> str:
    problem_by_type = {
        "t2_unit_order_mismatch": (
            "T2 unit order differs from expected.unit_order in the signed standard."
        ),
        "t2_expected_unit_missing": "T2 is missing one or more units required by the standard.",
        "t2_unexpected_unit_present": "T2 emitted units that are not allowed by the standard.",
        "t2_anchor_owner_mismatch": (
            "T2 assigned a standard anchor to the wrong unit owner."
        ),
        "t2_source_range_mismatch": "T2 source ranges differ from the signed standard.",
        "t2_page_policy_mismatch": "T2 page policy differs from the signed standard.",
        "t2_standard_schema_invalid": "The signed T2 standard schema is invalid.",
        "t3_unit_order_mismatch": (
            "T3 element unit order differs from expected.unit_order in the signed standard."
        ),
        "t3_required_policy_fields_missing": (
            "T3 elements are missing fields required by their assigned policy."
        ),
        "t3_standard_element_expectations_missing": (
            "T3 standard lacks element/run-span expectations for the "
            "human-reviewed final_template element list."
        ),
        "t3_standard_run_span_ledger_missing": (
            "T3 standard lacks a run/span handling ledger for source runs."
        ),
        "t5_unit_order_mismatch": (
            "T5 template_spec unit order differs from expected.unit_order."
        ),
    }
    if finding_type == "t3_policy_group_conflict":
        return _t3_policy_conflict_problem(finding)
    if finding_type == "t5_review_flags_dropped":
        return "T5 dropped review_flags that upstream stages emitted."
    if finding_type == "t5_section_profile_refs_missing":
        return "T5 units are missing section_profile_refs required for layout traceability."
    if finding_type == "t5_input_hash_missing":
        return "T5 template_spec does not bind all required upstream input hashes."
    if finding_type in problem_by_type:
        return problem_by_type[finding_type]
    return finding.message


def _t3_policy_conflict_problem(finding: Finding) -> str:
    conflicts = _parse_finding_value(finding.actual)
    if not isinstance(conflicts, list) or not conflicts:
        return finding.message
    groups = sorted(
        {
            f"{item.get('expected')}->{item.get('actual')}"
            for item in conflicts
            if isinstance(item, dict)
            and item.get("expected") is not None
            and item.get("actual") is not None
        }
    )
    if not groups:
        return finding.message
    return f"Element policy group conflict: {', '.join(groups)}."


def _parse_finding_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    try:
        return ast.literal_eval(value)
    except (SyntaxError, ValueError):
        return value


def _source_seq_refs_from_finding(finding: Finding) -> list[int]:
    refs: list[int] = []
    for value in [_parse_finding_value(finding.actual), _parse_finding_value(finding.expected)]:
        refs.extend(_collect_source_seq_refs(value))
    return sorted(set(refs))


def _collect_source_seq_refs(value: Any) -> list[int]:
    refs: list[int] = []
    if isinstance(value, dict):
        for key, child in value.items():
            if key in {"source_seq", "source_seq_ref"}:
                parsed = _int_or_none(child)
                if parsed is not None:
                    refs.append(parsed)
            elif key == "source_seq_refs" and isinstance(child, list):
                refs.extend(
                    item
                    for item in (_int_or_none(item) for item in child)
                    if item is not None
                )
            else:
                refs.extend(_collect_source_seq_refs(child))
    elif isinstance(value, list):
        for child in value:
            refs.extend(_collect_source_seq_refs(child))
    return refs


def _int_or_none(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _root_cause_category(
    report: TemplateGenerationJudgeReport,
    mismatch: dict[str, Any],
) -> str:
    text = (
        f"{mismatch['stage_key']} {mismatch['finding_type']} "
        f"{mismatch['root_cause_bucket']} {mismatch['problem']}"
    ).lower()
    if "agent" in text or "ai" in text:
        return "ai_overlay"
    if "not_configured" in text or "gate_disabled" in text or "verifier" in text:
        return "comparator_issue"
    if _standard_document_problem(text):
        return "standard_issue"
    if "artifact_missing" in text or "run_bundle_missing" in text or "hash" in text:
        return "evidence_missing"
    if "input" in text and "hash" not in text:
        return "input_issue"
    if mismatch["stage_id"] == "T5" and report.first_bad_stage not in {None, "T5"}:
        return "downstream_symptom"
    if mismatch["status"] == Status.FAIL.value:
        return "generation_code"
    if "artifact_trace" in text or "missing" in text:
        return "evidence_missing"
    return "generation_code"


def _diagnosis_first_bad_stage(
    report: TemplateGenerationJudgeReport,
    mismatch: dict[str, Any],
) -> str:
    if mismatch["stage_id"] == "T5" and report.first_bad_stage not in {None, "T5"}:
        return str(report.first_bad_stage)
    return str(mismatch["stage_id"])


def _diagnosis_reason(category: str, mismatch: dict[str, Any]) -> str:
    if category == "generation_code":
        return (
            f"{mismatch['stage_id']} artifact field {mismatch['field']} differs from "
            "the signed stage standard."
        )
    if category == "ai_overlay":
        return (
            f"{mismatch['stage_id']} mismatch is tied to AI overlay or attribution "
            "evidence and needs deterministic reconciliation."
        )
    if category == "standard_issue":
        return (
            f"{mismatch['stage_id']} standard evidence is missing, invalid, or "
            "internally inconsistent."
        )
    if category == "comparator_issue":
        return (
            f"{mismatch['stage_id']} deterministic comparator/gate configuration "
            "cannot produce a signable judgement."
        )
    if category == "input_issue":
        return f"{mismatch['stage_id']} input data does not satisfy the stage contract."
    if category == "evidence_missing":
        return (
            f"{mismatch['stage_id']} lacks artifact, hash, or trace evidence needed "
            "to compare against the standard."
        )
    if category == "downstream_symptom":
        return (
            f"{mismatch['stage_id']} symptom likely propagates from an earlier stage; "
            "fix the first bad upstream stage before changing this output."
        )
    return f"{mismatch['stage_id']} mismatch needs human review."


def _legacy_owner_for_mismatch(mismatch: dict[str, Any]) -> tuple[str, str, str]:
    finding = Finding(
        finding_id=mismatch["id"],
        stage=mismatch["stage_key"],
        severity="blocking",
        status=Status(mismatch["status"]),
        type=mismatch["finding_type"],
        message=mismatch["message"],
        expected=repr(mismatch["expected"]),
        actual=repr(mismatch["observed"]),
        evidence_refs=mismatch["evidence"].get("evidence_refs") or [],
        affected_ids=mismatch["affected_ids"],
        root_cause_bucket=mismatch["root_cause_bucket"],
    )
    return _owner_for_finding(finding)


def _owner_assignment_for_category(category: str) -> tuple[str, str | None]:
    mapping = {
        "generation_code": ("template_generation_code_owner", "standard_judge_owner"),
        "ai_overlay": ("ai_integration_owner", "template_generation_code_owner"),
        "standard_issue": ("standard_owner", "standard_judge_owner"),
        "comparator_issue": ("standard_judge_owner", None),
        "input_issue": ("input_data_owner", "template_generation_code_owner"),
        "evidence_missing": ("input_data_owner", "standard_judge_owner"),
        "downstream_symptom": ("template_generation_code_owner", "standard_judge_owner"),
    }
    return mapping.get(category, ("human_review_owner", "standard_judge_owner"))


def _owner_assignment_rationale(
    root_cause: dict[str, Any],
    primary: str,
    secondary: str | None,
) -> str:
    secondary_text = f", with {secondary} confirming the diagnosis" if secondary else ""
    return (
        f"{root_cause['stage_id']} {root_cause['finding_type']} is classified as "
        f"{root_cause['category']}; {primary} should resolve it{secondary_text}."
    )


def _fix_plan_for_root_cause(root_cause: dict[str, Any]) -> dict[str, Any]:
    stage_id = root_cause["stage_id"]
    finding_type = root_cause["finding_type"]
    if root_cause["category"] == "standard_issue":
        return {
            "action": "Fix or complete the signed stage standard through standard review.",
            "likely_files": ["standards/targets/*/v1/template_generation/*.standard.yaml"],
            "tests": ["tests/contract/test_template_generation_standard_judge.py"],
            "acceptance": [
                "standard_quality_status becomes PASS",
                "standard_diff_report retains evidence for any remaining mismatch",
            ],
        }
    if root_cause["category"] == "comparator_issue":
        return {
            "action": "Implement or configure the deterministic stage comparator/gate.",
            "likely_files": [
                "src/docfit/harness/template_generation_stage_verifiers.py",
                "standards/targets/*/v1/template_generation/*.standard.yaml",
            ],
            "tests": ["tests/unit/test_template_generation_stage_verifiers.py"],
            "acceptance": [
                "verifier_state is configured where the comparator exists",
                "gate status no longer hides mismatches/root_causes/owner/fix_plan",
            ],
        }
    if stage_id == "T2":
        return {
            "action": _t2_fix_action(finding_type),
            "likely_files": [
                "src/docfit/template_generation/t2_ai.py",
                "src/docfit/template_generation/agent/prompt_templates/t2_rubric.txt",
                "src/docfit/template_generation/t2_standard.py",
            ],
            "tests": [
                "tests/unit/test_t2_standard.py",
                "tests/unit/test_template_generation_standard_diff_diagnosis.py",
            ],
            "acceptance": [
                "02_unit_map_standard_diff_report.json no longer contains this mismatch",
                "T2 audit_status changes from FAIL/UNKNOWN to PASS",
                "No new T3/T5 mismatch is introduced",
            ],
        }
    if stage_id == "T3":
        return {
            "action": _t3_fix_action(finding_type),
            "likely_files": [
                "src/docfit/template_generation/agent/t3_ai_materialize.py",
                "src/docfit/template_generation/artifacts.py",
                "src/docfit/harness/template_generation_stage_verifiers.py",
            ],
            "tests": [
                "tests/unit/test_template_generation_stage_verifiers.py",
                "tests/unit/test_template_generation_standard_diff_diagnosis.py",
            ],
            "acceptance": [
                "03_element_spec_standard_diff_report.json explains no policy conflict",
                "T3 audit_status changes from FAIL/UNKNOWN to PASS",
                "Fixed units only allow fill when the standard explicitly permits it",
            ],
        }
    if stage_id == "T5":
        return {
            "action": _t5_fix_action(finding_type, root_cause["category"]),
            "likely_files": [
                "src/docfit/template_generation/outputs.py",
                "src/docfit/template_generation/runner.py",
                "src/docfit/harness/template_generation_stage_verifiers.py",
            ],
            "tests": [
                "tests/contract/test_template_generate.py",
                "tests/unit/test_template_generation_standard_diff_diagnosis.py",
            ],
            "acceptance": [
                "05_template_spec_standard_diff_report.json no longer contains this mismatch",
                "T5 preserves upstream hashes, review_flags, and section_profile_refs",
                "If upstream caused the symptom, first_bad_stage points upstream",
            ],
        }
    return {
        "action": "Inspect the stage artifact and signed standard, then fix the owning stage.",
        "likely_files": ["src/docfit/template_generation/"],
        "tests": ["tests/contract/test_template_generation_standard_judge.py"],
        "acceptance": ["The mismatch disappears from the stage standard diff report"],
    }


def _t2_fix_action(finding_type: str) -> str:
    actions = {
        "t2_standard_unit_order_mismatch": "Fix T2 unit discovery ordering rules.",
        "t2_standard_units_missing": "Fix T2 unit discovery so all expected units are emitted.",
        "t2_standard_units_unexpected": "Remove unexpected non-standard units from T2 output.",
        "t2_standard_custom_units_present": "Map custom T2 units to signed standard units.",
        "t2_standard_anchor_owner_mismatch": "Fix T2 anchor ownership and source range assignment.",
    }
    return actions.get(finding_type, "Fix T2 unit pagination output against the signed standard.")


def _t3_fix_action(finding_type: str) -> str:
    actions = {
        "t3_unit_order_mismatch": "Fix T3 element ordering to follow expected.unit_order.",
        "t3_policy_group_conflict": (
            "Fix T3 policy assignment so fixed/manual/generated/fill groups do not conflict."
        ),
        "t3_required_policy_fields_missing": (
            "Emit all required fields for each T3 element policy."
        ),
    }
    return actions.get(finding_type, "Fix T3 element policy output against the signed standard.")


def _t5_fix_action(finding_type: str, category: str) -> str:
    if category == "downstream_symptom":
        return "Fix the upstream first_bad_stage, then rerun T5 to clear the propagated symptom."
    actions = {
        "t5_required_input_hashes_missing": "Bind all required upstream input hashes in T5.",
        "t5_review_flags_dropped": "Preserve upstream review_flags in template_spec.",
        "t5_unit_order_mismatch": "Fix T5 unit ordering to follow expected.unit_order.",
        "t5_unit_section_profile_refs_missing": (
            "Carry section_profile_refs from global layout into each T5 unit."
        ),
    }
    return actions.get(finding_type, "Fix T5 template_spec output against the signed standard.")


def _derived_id(mismatch_id: str, suffix: str) -> str:
    return mismatch_id.replace("MISMATCH", suffix)


def _diff_kind(finding: Finding) -> str:
    if finding.type in {
        "template_generation_stage_verifier_not_configured",
        "template_generation_stage_gate_disabled",
    }:
        return "gate_blocker"
    if finding.status == Status.FAIL:
        return "standard_mismatch"
    if "missing" in finding.type or "hash" in finding.type:
        return "evidence_gap"
    return "audit_observation"


def _owner_for_finding(finding: Finding) -> tuple[str, str, str]:
    finding_type = finding.type
    bucket = finding.root_cause_bucket
    text = f"{finding.stage} {finding_type} {bucket}".lower()
    if "agent" in text or "ai" in text:
        return (
            "ai_prompt",
            "AI observation or its materialization touched this path",
            "Review the AI evidence, output contract, and final materialization.",
        )
    if "verifier" in text or "gate_disabled" in text or "not_configured" in text:
        return (
            "verifier",
            "Stage standard verifier is not configured as a blocking gate",
            "Implement or configure the deterministic verifier before sign-off.",
        )
    if finding.status == Status.FAIL and not _standard_document_problem(text):
        return (
            "code",
            "Run artifact disagrees with signed standard under deterministic audit",
            "Fix template-generation deterministic logic for the first bad stage.",
        )
    if "standard" in text or "baseline" in text:
        return (
            "standard",
            "Signed standard or standard registry is missing or inconsistent",
            "Fix the signed standard through the review process, not from this run.",
        )
    if "run_bundle" in text or "hash" in text or "artifact" in text:
        return (
            "code",
            "Run evidence package is incomplete or internally inconsistent",
            "Fix artifact production or run bundle binding before sign-off.",
        )
    return (
        "unknown",
        "Current evidence is insufficient to assign ownership",
        "Add deterministic evidence or human review before changing code or standards.",
    )


def _standard_document_problem(text: str) -> bool:
    return any(
        marker in text
        for marker in [
            "standard_missing",
            "standard_invalid",
            "standard_quality_not_pass",
            "standard_incomplete",
            "standard_registry",
            "baseline_missing",
            "baseline_invalid",
        ]
    )


def _root_cause_reason(diff: dict[str, Any]) -> str:
    owner = diff["owner"]
    if owner == "verifier":
        return "The stage cannot be signed because the standard verifier is not an enabled gate."
    if owner == "standard":
        return "The signed standard or its registry is incomplete or inconsistent."
    if owner == "code":
        return "The generated stage artifact does not match the signed standard evidence."
    if owner == "ai_prompt":
        return "The mismatch is tied to AI proposal or attribution evidence."
    return "The report does not yet contain enough evidence to assign a precise owner."


def _stage_quality_report_id_for_stage(stage_key: str) -> str | None:
    for spec in RUN_ARTIFACT_SPECS:
        if spec.stage_key == stage_key:
            return _stage_quality_report_id(spec)
    return None


def _standard_quality_finding_stage_context(
    report: TemplateGenerationJudgeReport,
    finding: Finding,
) -> tuple[str, str, str | None]:
    stage_key = finding.stage
    stage = report.standard_set.stages.get(stage_key)
    if stage is not None:
        return stage.stage_id, stage_key, _stage_quality_report_id_for_stage(stage_key)
    return "STANDARD_QUALITY", stage_key, None


def _stage_check_for_spec(
    report: TemplateGenerationJudgeReport,
    spec: RunArtifactSpec,
) -> StageCheck | None:
    if spec.stage_key is None:
        return None
    for check in report.stage_checks:
        if check.stage_key == spec.stage_key:
            return check
    return None


def _run_bundle_findings_for_spec(
    report: TemplateGenerationJudgeReport,
    spec: RunArtifactSpec,
) -> list[Finding]:
    findings = []
    for finding in report.run_bundle.findings:
        if spec.artifact_key in finding.affected_ids:
            findings.append(finding)
        elif spec.artifact_key == "template_generation_request" and "source_hash" in finding.type:
            findings.append(finding)
        elif spec.artifact_key == "fillable_template_docx" and "fillable" in finding.type:
            findings.append(finding)
    if spec.stage_id == "T6":
        findings.extend(report.run_bundle.observed_t6_findings)
    return findings


def _stage_quality_report_id(spec: RunArtifactSpec) -> str:
    return f"{Path(spec.top_level_name).stem}_standard_quality_report"


def _stage_diff_report_id(spec: RunArtifactSpec) -> str:
    return f"{Path(spec.top_level_name).stem}_standard_diff_report"


def _implicit_stage_key(spec: RunArtifactSpec) -> str:
    if spec.artifact_key == "template_generation_request":
        return "template_generation_request"
    if spec.artifact_key == "fillable_template_docx":
        return "t6_fillable_template"
    if spec.artifact_key == "build_manifest":
        return "t6_build_manifest"
    if spec.artifact_key == "verification_report":
        return "t7_verification_report"
    return spec.artifact_key


def build_standard_quality_markdown(
    report: TemplateGenerationStandardQualityReport,
) -> str:
    lines = [
        "# Template Generation Stage Standard Quality",
        "",
        f"- Scope: {report.scope}",
        f"- Status: {report.status.value}",
        f"- Targets: {len(report.target_reports)}",
        f"- Findings: {len(report.findings)}",
        "",
        "## Stage Statuses",
    ]
    for stage_key, status in sorted(report.stage_statuses.items()):
        lines.append(f"- {stage_key}: {status}")
    if report.findings:
        lines.extend(["", "## Findings"])
        for finding in report.findings:
            lines.append(f"- [{finding.status.value}] {finding.type}: {finding.message}")
    else:
        lines.extend(["", "No findings."])
    lines.append("")
    return "\n".join(lines)


def build_judge_markdown(report: TemplateGenerationJudgeReport) -> str:
    report_dict = report.to_dict()
    lines = [
        "# Template Generation Judge Report",
        "",
        f"- Status: {report.status.value}",
        f"- Standard acceptance: {report_dict['standard_acceptance_status']}",
        f"- Sign-off: {report_dict['signoff_status']}",
        f"- First bad stage: {report.first_bad_stage or 'none'}",
        f"- Source run id: {report.run_bundle.source_run_id}",
        f"- Source run dir: {report.run_bundle.source_run_dir}",
        f"- Standard quality: {report.standard_quality.status.value}",
        f"- Run bundle: {report.run_bundle.status.value}",
        "",
    ]
    if report_dict["signoff_blockers"]:
        lines.append("## Sign-off Blockers")
        for blocker in report_dict["signoff_blockers"]:
            lines.append(f"- {blocker}")
        lines.append("")
    lines.append("## Owner Summary")
    for owner, count in report_dict["owner_summary"].items():
        lines.append(f"- {owner}: {count}")
    lines.extend(["", "## Top Blockers"])
    if report_dict["top_blockers"]:
        for blocker in report_dict["top_blockers"]:
            lines.append(
                "- "
                f"{blocker.get('stage_id')} {blocker.get('finding_type')}: "
                f"status={blocker.get('status')}, owner={blocker.get('owner')}, "
                f"next={blocker.get('next_action')}"
            )
    else:
        lines.append("- none")
    lines.extend([
        "",
        "## Stage Checks",
    ])
    for check in report.stage_checks:
        lines.append(
            "- "
            f"{check.stage_id} {check.stage_key}: "
            f"status={check.status.value}, audit={check.audit_status}, "
            f"verifier_state={check.verifier_state}, gate_enabled={check.gate_enabled}"
        )
    if report.findings:
        lines.extend(["", "## Findings"])
        for finding in report.findings[:100]:
            lines.append(
                f"- [{finding.status.value}] "
                f"{finding.stage}/{finding.type}: {finding.message}"
            )
        if len(report.findings) > 100:
            lines.append(f"- ... {len(report.findings) - 100} more findings")
    else:
        lines.extend(["", "No findings."])
    lines.append("")
    return "\n".join(lines)


_ROUTE_STAGE_FILES = {
    "T1": {
        "stage_key": "t1_document_facts",
        "routes": {
            "shared_input": {"filename": "01_document_facts.json", "payload_type": "json"},
        },
    },
    "L1": {
        "stage_key": "l1_input_contract",
        "routes": {
            "shared_input": {
                "filename": "01.5_l1_input_contract.json",
                "payload_type": "json",
            },
        },
    },
    "T2": {
        "stage_key": "t2_unit_pagination",
        "routes": {
            "ai": {"filename": "02_unit_map.yaml", "payload_type": "yaml"},
        },
    },
    "T3": {
        "stage_key": "t3_element_policy",
        "routes": {
            "ai": {"filename": "03_element_spec.yaml", "payload_type": "yaml"},
        },
    },
    "T4": {
        "stage_key": "t4_global_layout",
        "routes": {
            "ai": {"filename": "04_global_spec.yaml", "payload_type": "yaml"},
        },
    },
    "T5": {
        "stage_key": "t5_template_spec",
        "routes": {
            "canonical": {"filename": "05_template_spec.yaml", "payload_type": "yaml"},
        },
    },
    "T6": {
        "stage_key": "t6_fillable_template",
        "routes": {
            "canonical": {
                "filename": ["06.1_fillable_template.docx", "06.2_build_manifest.json"],
                "payload_type": "composite",
                "artifact_type": "t6_execution_bundle",
            },
        },
    },
    "T7": {
        "stage_key": "t7_verification_report",
        "routes": {
            "canonical": {"filename": "07_verification_report.json", "payload_type": "json"},
        },
    },
    "POST_T6": {
        "stage_key": "post_t6_template_gap",
        "routes": {
            "canonical": {"filename": "template_gap_report.json", "payload_type": "json"},
        },
    },
}

_OBSERVATION_POLICY_GROUPS = {
    "fixed_units": "fixed",
    "fill_units": "fill",
    "generated_units": "generated",
    "template_default_optional_units": "template_default",
}
_FORMAT_ANNOTATION_RE = re.compile(
    r"[（(][^（）()]{0,80}(?:"
    r"号|宋体|黑体|楷体|仿宋|华文|Times|Arial|加粗|居中|空[一二三四五六七八九十0-9]*行|"
    r"小四|小三|三号|四号|一号|二号|行距|磅"
    r")[^（）()]{0,80}[）)]"
)
_PLACEHOLDER_RE = re.compile(r"(?:□+|×{2,}|20×+|…{2,}|\.{6,}|_{4,})")


def build_template_generation_route_eval_report(
    report: TemplateGenerationJudgeReport,
    report_dict: dict[str, Any] | None = None,
) -> dict[str, Any]:
    routes: list[dict[str, Any]] = []
    stage_metrics: dict[str, dict[str, Any]] = {}
    mismatches: list[dict[str, Any]] = []
    run_dir = report.run_bundle.source_run_dir
    for stage_id, spec in _ROUTE_STAGE_FILES.items():
        stage_routes = {
            route_id: _route_candidate(
                run_dir,
                stage_id=stage_id,
                stage_key=str(spec["stage_key"]),
                route_id=route_id,
                filename=route_spec.get("filename"),
                payload_type=str(route_spec.get("payload_type") or "yaml"),
                artifact_type=route_spec.get("artifact_type"),
            )
            for route_id, route_spec in (spec.get("routes") or {}).items()
        }
        routes.extend(stage_routes.values())
        ai_route_id = "ai" if stage_id in {"T2", "T3", "T4"} else "canonical"
        ai_hash = (stage_routes.get(ai_route_id) or {}).get("payload_hash")
        ai_availability = (stage_routes.get(ai_route_id) or {}).get(
            "availability", "NOT_APPLICABLE"
        )
        stage_metrics[stage_id] = {
            "stage_key": spec["stage_key"],
            "ai_availability": ai_availability,
            "ai_hash": ai_hash,
            "route_availability": {
                route_id: route.get("availability")
                for route_id, route in stage_routes.items()
            },
            "route_reasons": {
                route_id: route.get("reason")
                for route_id, route in stage_routes.items()
                if route.get("reason")
            },
            "route_payload_summaries": {
                route_id: route.get("payload_summary")
                for route_id, route in stage_routes.items()
                if route.get("payload_summary")
            },
            "route_hashes": {
                route_id: route.get("payload_hash")
                for route_id, route in stage_routes.items()
            },
        }
        standard = report.standard_set.stages.get(str(spec["stage_key"]))
        expected = standard.expected if standard is not None else {}
        if stage_id in {"T2", "T3"}:
            _add_common_route_accuracy(
                stage_metrics[stage_id],
                stage_routes,
                stage_id=stage_id,
                expected=expected,
            )
        if stage_id == "T4":
            t4_contract, t4_contract_mismatches = _evaluate_t4_route_output_contract(
                stage_routes
            )
            stage_metrics[stage_id]["common_output_contract"] = t4_contract
            mismatches.extend(t4_contract_mismatches)
            if standard is not None:
                _add_t4_common_route_accuracy(
                    stage_metrics[stage_id],
                    stage_routes,
                    report=report,
                    standard=standard,
                )
        if stage_id == "L1":
            stage_metrics[stage_id]["coverage"] = _l1_route_coverage(
                stage_routes.get("shared_input")
            )
            mismatches.extend(
                _l1_route_mismatches(stage_routes.get("shared_input"))
            )
        if stage_id == "T3":
            residual_gate = _t3_residual_gate(report.run_bundle.source_run_dir)
            stage_metrics[stage_id]["residual_gate"] = residual_gate
            mismatches.extend(_t3_residual_mismatches(residual_gate))
        if ai_availability == "NOT_AVAILABLE":
            ai_reason = str((stage_routes.get(ai_route_id) or {}).get("reason") or "")
            mismatches.append(
                {
                    "id": f"{stage_id}-ROUTE-MISMATCH-001",
                    "stage_id": stage_id,
                    "stage_key": spec["stage_key"],
                    "type": (
                        "canonical_ai_not_available"
                        if stage_id in {"T2", "T3", "T4"}
                        else "canonical_route_not_available"
                    ),
                    "expected": (
                        "canonical AI route available"
                        if stage_id in {"T2", "T3", "T4"}
                        else "canonical route available"
                    ),
                    "observed": (
                        f"{ai_route_id} route is NOT_AVAILABLE"
                        + (f": {ai_reason}" if ai_reason else "")
                    ),
                    "route_ids": [ai_route_id],
                }
            )
    root_causes = [
        {
            "id": mismatch["id"].replace("MISMATCH", "ROOT-CAUSE"),
            "mismatch_id": mismatch["id"],
            "stage_id": mismatch["stage_id"],
            "stage_key": mismatch["stage_key"],
            "category": _route_mismatch_root_cause_category(mismatch),
            "reason": mismatch["observed"],
        }
        for mismatch in mismatches
    ]
    owner_assignments = [
        {
            "id": mismatch["id"].replace("MISMATCH", "OWNER"),
            "mismatch_id": mismatch["id"],
            "primary": _route_mismatch_owner(mismatch),
            "secondary": ["standard_judge_owner"],
        }
        for mismatch in mismatches
    ]
    fix_plan = [
        {
            "id": mismatch["id"].replace("MISMATCH", "FIX"),
            "mismatch_id": mismatch["id"],
            "action": _route_mismatch_fix_action(mismatch),
            "verification": "rerun template_generation_route_eval_report and confirm mismatch is gone",
        }
        for mismatch in mismatches
    ]
    ai_primary_gate_decision = _ai_primary_gate_decision(stage_metrics, mismatches)
    return {
        "artifact_type": "template_generation_route_eval_report",
        "artifact_version": "1.0",
        "profile_id": report.standard_set.target_standard.get("coverage_requirements", {}).get("profile"),
        "school_id": report.standard_set.school_id,
        "run_id": report.run_bundle.source_run_id,
        "routes": routes,
        "shared_inputs": {"t1_document_facts": "01_document_facts.json"},
        "stage_metrics": stage_metrics,
        "cross_route_summary": {
            "route_count": len(routes),
            "mismatch_count": len(mismatches),
            "availability_by_stage": {
                stage_id: metrics.get("route_availability", {})
                for stage_id, metrics in stage_metrics.items()
            },
        },
        "mismatches": mismatches,
        "root_causes": root_causes,
        "owner_assignments": owner_assignments,
        "fix_plan": fix_plan,
        "ai_primary_gate_decision": ai_primary_gate_decision,
    }


def _ai_primary_gate_decision(
    stage_metrics: dict[str, dict[str, Any]],
    mismatches: list[dict[str, Any]],
) -> dict[str, Any]:
    stage_decisions: dict[str, dict[str, Any]] = {}
    del mismatches
    for stage_id in ("T2", "T3", "T4"):
        metrics = stage_metrics.get(stage_id, {})
        availability = str(metrics.get("ai_availability") or "NOT_AVAILABLE")
        stage_decisions[stage_id] = {
            "eligible": availability == "AVAILABLE",
            "status": "CANONICAL_AI_ONLY" if availability == "AVAILABLE" else "BLOCKED",
            "authority_mode": "ai",
            "reason": (
                f"{stage_id} has one canonical AI route"
                if availability == "AVAILABLE"
                else f"canonical {stage_id} AI route is {availability}"
            ),
            "fallback": "safe_keep" if stage_id == "T3" else "none",
        }
    return {
        "artifact_type": "template_generation_ai_primary_gate_decision",
        "artifact_version": "2.0",
        "default_authority_mode": "ai",
        "allowed_ai_primary_stages": ["T2", "T3", "T4"],
        "stages": stage_decisions,
    }


def _route_mismatch_root_cause_category(mismatch: dict[str, Any]) -> str:
    mismatch_type = str(mismatch.get("type") or "")
    if mismatch_type.startswith("l1_"):
        return "l1_input_contract_gap"
    if mismatch_type.startswith("t4_"):
        return "t4_ai_layout_contract_gap"
    return "ai_route_missing"


def _route_mismatch_owner(mismatch: dict[str, Any]) -> str:
    mismatch_type = str(mismatch.get("type") or "")
    if mismatch_type.startswith("l1_"):
        return "template_generation_input_contract_owner"
    if mismatch_type.startswith("t4_"):
        return "template_generation_t4_layout_owner"
    return "template_generation_observation_owner"


def _route_mismatch_fix_action(mismatch: dict[str, Any]) -> str:
    mismatch_type = str(mismatch.get("type") or "")
    if mismatch_type.startswith("l1_"):
        return (
            "Patch L1 input projection, render binding, or bundle gate coverage until "
            "the contract is explicit."
        )
    if mismatch_type.startswith("t4_"):
        return (
            "Fix the T4 AI observation or its identity/materialization contract, then "
            "republish the canonical global_spec."
        )
    return "Run the canonical AI observation stage and republish its final artifact."


def _route_candidate(
    run_dir: Path,
    *,
    stage_id: str,
    stage_key: str,
    route_id: str,
    filename: Any,
    payload_type: str,
    artifact_type: Any = None,
) -> dict[str, Any]:
    filenames = (
        [str(item) for item in filename]
        if isinstance(filename, list)
        else [str(filename)]
    )
    paths = [
        _route_artifact_path(
            run_dir,
            item,
        )
        for item in filenames
    ]
    found_paths = [path for path in paths if path is not None]
    path = found_paths[0] if found_paths else None
    payload: dict[str, Any] = {}
    payload_hash = None
    availability = "NOT_AVAILABLE"
    if found_paths and len(found_paths) == len(filenames):
        payload_hash = (
            sha256_file(path)
            if len(found_paths) == 1 and path is not None
            else sha256_json({item.name: sha256_file(item) for item in found_paths})
        )
        payload = _read_route_payload(found_paths, payload_type)
        final_availability = payload.get("availability")
        final_availability = (
            final_availability if isinstance(final_availability, dict) else {}
        )
        availability = str(
            final_availability.get("status")
            or (payload.get("route") or {}).get("availability")
            or ("AVAILABLE" if payload else "UNKNOWN")
        )
    route = payload.get("route") or {}
    final_availability = payload.get("availability")
    final_availability = (
        final_availability if isinstance(final_availability, dict) else {}
    )
    lineage = payload.get("lineage")
    lineage = lineage if isinstance(lineage, dict) else {}
    return {
        "route_id": route_id,
        "stage_key": stage_key,
        "stage_id": stage_id,
        "artifact_type": artifact_type or payload.get("artifact_type"),
        "payload_path": str(path) if path is not None else None,
        "payload_paths": [str(item) for item in found_paths],
        "payload_hash": payload_hash,
        "availability": availability,
        "origin": route.get("origin") or lineage.get("producer_mode"),
        "reason": final_availability.get("reason") or route.get("reason"),
        "payload_summary": _route_payload_summary(payload),
    }


def _route_payload_summary(payload: dict[str, Any]) -> dict[str, Any]:
    if not payload:
        return {}
    items = payload.get("items")
    unknown_items = payload.get("unknown_items")
    summary: dict[str, Any] = {
        "model": payload.get("model"),
        "schema_version": payload.get("schema_version"),
    }
    if isinstance(items, list):
        summary["item_count"] = len(items)
    if isinstance(unknown_items, list):
        summary["unknown_item_count"] = len(unknown_items)
    coverage = payload.get("coverage")
    if isinstance(coverage, dict):
        summary["coverage"] = {
            key: value
            for key, value in coverage.items()
            if key in {"total", "owned_source_seq", "unknown_source_seq"}
        }
        for key in ("owned_source_seq", "unknown_source_seq"):
            value = summary["coverage"].get(key)
            if isinstance(value, list):
                summary["coverage"][f"{key}_count"] = len(value)
                summary["coverage"].pop(key, None)
    return {key: value for key, value in summary.items() if value is not None}


def _route_candidate_payload(route: dict[str, Any]) -> dict[str, Any]:
    payload_path = route.get("payload_path")
    if not payload_path:
        return {}
    path = Path(str(payload_path))
    if not path.exists():
        return {}
    try:
        loaded = read_json(path) if path.suffix == ".json" else read_yaml(path)
    except Exception:
        return {}
    return loaded if isinstance(loaded, dict) else {}


def _add_common_route_accuracy(
    metrics: dict[str, Any],
    routes: dict[str, dict[str, Any]],
    *,
    stage_id: str,
    expected: dict[str, Any],
) -> None:
    """Score every available route against one stage gold and one evaluator."""
    for route_id in ("ai",):
        route = routes.get(route_id) or {}
        if route.get("availability") != "AVAILABLE":
            continue
        payload = _route_candidate_payload(route)
        if not payload:
            continue
        if stage_id == "T2":
            result = _evaluate_t2_route_accuracy(payload, expected)
        else:
            result = _evaluate_t3_route_accuracy(payload, expected)
        metrics[f"{route_id}_accuracy"] = result
    metrics["primary_metric"] = "unit_f1" if stage_id == "T2" else "exact_action_accuracy"


def _evaluate_t4_route_output_contract(
    routes: dict[str, dict[str, Any]],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    required_fields = (
        "section_profiles",
        "default_font",
        "page_numbering",
        "header_footer",
        "numbering_rules",
        "flags",
    )
    route_results: dict[str, Any] = {}
    mismatches: list[dict[str, Any]] = []
    for route_id in ("ai",):
        route = routes.get(route_id) or {}
        if route.get("availability") != "AVAILABLE":
            route_results[route_id] = {
                "status": "NOT_AVAILABLE",
                "artifact_type": None,
                "missing_fields": list(required_fields),
            }
            continue
        payload = _route_candidate_payload(route)
        missing_fields = [field for field in required_fields if field not in payload]
        artifact_type = payload.get("artifact_type")
        status = (
            "PASS"
            if artifact_type == "global_spec" and not missing_fields
            else "FAIL"
        )
        route_results[route_id] = {
            "status": status,
            "artifact_type": artifact_type,
            "missing_fields": missing_fields,
        }
        if status == "FAIL":
            mismatches.append(
                {
                    "id": f"T4-{route_id.upper()}-OUTPUT-CONTRACT-MISMATCH-001",
                    "stage_id": "T4",
                    "stage_key": "t4_global_layout",
                    "type": "t4_route_output_contract_mismatch",
                    "expected": {
                        "artifact_type": "global_spec",
                        "required_fields": list(required_fields),
                    },
                    "observed": {
                        "artifact_type": artifact_type,
                        "missing_fields": missing_fields,
                    },
                    "route_ids": [route_id],
                }
            )
    available_results = [
        result
        for result in route_results.values()
        if result.get("status") != "NOT_AVAILABLE"
    ]
    return {
        "artifact_type": "global_spec",
        "required_fields": list(required_fields),
        "status": (
            "PASS"
            if available_results
            and all(result.get("status") == "PASS" for result in available_results)
            else "UNKNOWN" if not available_results else "FAIL"
        ),
        "routes": route_results,
    }, mismatches


def _add_t4_common_route_accuracy(
    metrics: dict[str, Any],
    routes: dict[str, dict[str, Any]],
    *,
    report: TemplateGenerationJudgeReport,
    standard: StageStandardSpec,
) -> None:
    """Run the canonical T4 AI final through the stage judge and signed gold."""
    for route_id in ("ai",):
        route = routes.get(route_id) or {}
        if route.get("availability") != "AVAILABLE":
            continue
        payload = _route_candidate_payload(route)
        payload_path = Path(str(route.get("payload_path") or ""))
        if not payload or not payload_path.exists():
            continue
        artifact_hash = str(route.get("payload_hash") or sha256_file(payload_path))
        artifact = BoundArtifact(
            artifact_key="global_spec",
            stage_key="t4_global_layout",
            stage_id="T4",
            path=payload_path,
            sha256=artifact_hash,
            declared_sha256=artifact_hash,
            source_kind=f"route_eval.{route_id}",
            status=Status.PASS,
            payload=payload,
            hash_match=True,
        )
        check = judge_template_generation_stage(
            "t4_global_layout",
            standard=standard,
            artifact=artifact,
            standard_quality=report.standard_quality,
            run_bundle=report.run_bundle,
        )
        result = _t4_accuracy_metric(report, check)
        result.update(
            {
                "scoring_contract": "t4_stage_standard_v1",
                "gold_path": str(standard.path),
                "gold_sha256": standard.sha256,
                "route_id": route_id,
            }
        )
        metrics[f"{route_id}_accuracy"] = result
    metrics["primary_metric"] = "layout_contract_completeness"


def _evaluate_t2_route_accuracy(
    payload: dict[str, Any],
    expected: dict[str, Any],
) -> dict[str, Any]:
    units = payload.get("units")
    projection = {
        "units": [unit for unit in (units or []) if isinstance(unit, dict)],
    }
    result = _evaluate_ai_unit_accuracy(projection, expected)
    result["scoring_contract"] = "t2_common_route_v1"
    result["gold_universe"] = "expected.unit_order"
    return result


def _evaluate_t3_route_accuracy(
    payload: dict[str, Any],
    expected: dict[str, Any],
) -> dict[str, Any]:
    items = payload.get("items")
    if not isinstance(items, list):
        items = payload.get("elements")
    projected = [item for item in (items or []) if isinstance(item, dict)]
    ledger = [
        row
        for row in (expected.get("run_span_ledger") or [])
        if isinstance(row, dict) and row.get("expected_action") in {"keep", "fill", "delete"}
    ]
    if not ledger:
        legacy = _evaluate_ai_element_accuracy({"items": projected}, expected)
        legacy.update(
            {
                "exact_action_accuracy": None,
                "scoring_contract": "t3_common_route_v2",
                "gold_universe": "unavailable",
            }
        )
        return legacy

    matches = 0
    missing: list[dict[str, Any]] = []
    mismatches: list[dict[str, Any]] = []
    mixed: list[dict[str, Any]] = []
    for row in ledger:
        identity = {
            field: str(row[field])
            for field in ("raw_run_id", "logical_run_id")
            if row.get(field)
        }
        identity["target_kind"] = str(row.get("target_kind") or "run")
        if identity["target_kind"] == "span":
            identity["start"] = row.get("start")
            identity["end"] = row.get("end")
        item_projection = project_t3_gold_item(projected, row)
        actions = item_projection.actions
        expected_action = str(row["expected_action"])
        if not item_projection.coverage_complete or not actions:
            missing.append(
                {
                    "expected_action": expected_action,
                    "actual_actions": sorted(actions),
                    "identity_covered": item_projection.identity_covered,
                    **identity,
                }
            )
        elif actions == {expected_action}:
            matches += 1
        else:
            mismatch = {
                "expected_action": expected_action,
                "actual_actions": sorted(actions),
                "status": "mixed" if len(actions) > 1 else "mismatch",
                **identity,
            }
            mismatches.append(mismatch)
            if len(actions) > 1:
                mixed.append(mismatch)
    total = len(ledger)
    return {
        "scoring_contract": "t3_common_route_v2",
        "gold_universe": "expected.run_span_ledger.adaptive_scored_actions",
        "gold_count": total,
        "match_count": matches,
        "mismatch_count": len(mismatches),
        "mixed_action_count": len(mixed),
        "conflicted_run_count": len(mixed),
        "missing_count": len(missing),
        "coverage": round((total - len(missing)) / total, 4) if total else 0.0,
        "exact_action_accuracy": round(matches / total, 4) if total else 0.0,
        "mismatch_samples": mismatches[:10],
        "mixed_action_samples": mixed[:10],
        "conflict_samples": mixed[:10],
        "missing_samples": missing[:10],
    }


def _evaluate_ai_unit_accuracy(
    observation: dict[str, Any],
    expected: dict[str, Any],
) -> dict[str, Any]:
    gold_order = [str(unit_id) for unit_id in expected.get("unit_order", []) or []]
    gold_set = set(gold_order)
    observed_order = [
        str(unit.get("unit_id"))
        for unit in observation.get("units", [])
        if isinstance(unit, dict) and unit.get("unit_id")
    ]
    observed_set = set(observed_order)
    hits = observed_set & gold_set
    precision = len(hits) / len(observed_set) if observed_set else 0.0
    recall = len(hits) / len(gold_set) if gold_set else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    result = {
        "gold_units": len(gold_set),
        "observed_units": len(observed_set),
        "precision": round(precision, 4),
        "recall": round(recall, 4),
        "f1": round(f1, 4),
        "order_exact_match": observed_order == gold_order,
        "missing_units": sorted(gold_set - observed_set),
        "extra_units": sorted(observed_set - gold_set),
    }
    result["page_boundary_accuracy"] = _evaluate_ai_unit_page_boundary_accuracy(
        observation,
        expected,
    )
    return result


def _evaluate_ai_unit_page_boundary_accuracy(
    observation: dict[str, Any],
    expected: dict[str, Any],
) -> dict[str, Any]:
    expected_boundaries: dict[str, dict[str, int]] = {}
    for unit in expected.get("units", []) or []:
        if not isinstance(unit, dict):
            continue
        unit_id = str(unit.get("unit_id") or "")
        boundary = _page_boundary(unit.get("boundary"))
        if unit_id and boundary is not None:
            expected_boundaries[unit_id] = boundary
    if not expected_boundaries:
        return {
            "page_boundary_evaluable": False,
            "reason": "T2 standard does not provide page-native unit boundaries",
        }

    observed_boundaries: dict[str, dict[str, int]] = {}
    for unit in observation.get("units", []) or []:
        if not isinstance(unit, dict):
            continue
        unit_id = str(unit.get("unit_id") or "")
        boundary = _page_boundary(unit.get("boundary"))
        if unit_id and boundary is not None and unit_id not in observed_boundaries:
            observed_boundaries[unit_id] = boundary

    exact_match_count = 0
    missing: list[dict[str, Any]] = []
    mismatches: list[dict[str, Any]] = []
    for unit_id, expected_boundary in expected_boundaries.items():
        actual_boundary = observed_boundaries.get(unit_id)
        if actual_boundary is None:
            missing.append(
                {
                    "unit_id": unit_id,
                    "expected": expected_boundary,
                    "reason": "T2 observation does not provide this unit boundary",
                }
            )
            continue
        if actual_boundary == expected_boundary:
            exact_match_count += 1
        else:
            mismatches.append(
                {
                    "unit_id": unit_id,
                    "expected": expected_boundary,
                    "actual": actual_boundary,
                }
            )
    expected_count = len(expected_boundaries)
    compared_count = expected_count - len(missing)
    return {
        "page_boundary_evaluable": True,
        "expected_units": expected_count,
        "observed_units": compared_count,
        "coverage": round(compared_count / expected_count, 4),
        "exact_match_count": exact_match_count,
        "exact_match_accuracy": round(exact_match_count / expected_count, 4),
        "missing_count": len(missing),
        "mismatch_count": len(mismatches),
        "missing_samples": missing[:10],
        "mismatch_samples": mismatches[:10],
    }


def _page_boundary(value: Any) -> dict[str, int] | None:
    if not isinstance(value, dict):
        return None
    start_page = value.get("start_page")
    end_page = value.get("end_page")
    if (
        isinstance(start_page, bool)
        or isinstance(end_page, bool)
        or not isinstance(start_page, int)
        or not isinstance(end_page, int)
    ):
        return None
    return {"start_page": start_page, "end_page": end_page}


def _evaluate_ai_element_accuracy(
    observation: dict[str, Any],
    expected: dict[str, Any],
) -> dict[str, Any]:
    groups = expected.get("policy_groups", {}) or {}
    expected_policy: dict[str, str] = {}
    for group_name, policy in _OBSERVATION_POLICY_GROUPS.items():
        for unit_id in groups.get(group_name, []) or []:
            expected_policy[str(unit_id)] = policy

    by_unit: dict[str, list[str]] = {}
    items = [item for item in observation.get("items", []) if isinstance(item, dict)]
    for item in items:
        unit_id = str(item.get("unit_id") or "")
        policy = str(item.get("policy") or "")
        if unit_id and policy:
            by_unit.setdefault(unit_id, []).append(policy)
    dominant = {
        unit_id: Counter(policies).most_common(1)[0][0]
        for unit_id, policies in by_unit.items()
    }
    policy_sets = {unit_id: set(policies) for unit_id, policies in by_unit.items()}
    evaluated = [unit_id for unit_id in expected_policy if unit_id in dominant]
    dominant_hits = [
        unit_id
        for unit_id in evaluated
        if dominant[unit_id] == expected_policy[unit_id]
    ]
    present_hits = [
        unit_id
        for unit_id in evaluated
        if expected_policy[unit_id] in policy_sets[unit_id]
    ]
    mismatches = [
        {
            "unit_id": unit_id,
            "expected": expected_policy[unit_id],
            "ai_dominant": dominant[unit_id],
            "expected_present": expected_policy[unit_id] in policy_sets[unit_id],
        }
        for unit_id in evaluated
        if dominant[unit_id] != expected_policy[unit_id]
    ]
    demotions = observation.get("quality_report", {}).get("demotions", []) or []
    required_field_fails = sum(
        1
        for demotion in demotions
        if isinstance(demotion, dict)
        and demotion.get("check_id") == "C-REQUIRED-FIELD"
    )
    total_items = len(items) + required_field_fails
    result = {
        "units_evaluated": len(evaluated),
        "unit_dominant_policy_accuracy": (
            round(len(dominant_hits) / len(evaluated), 4) if evaluated else 0.0
        ),
        "distinguishing_policy_recall": (
            round(len(present_hits) / len(evaluated), 4) if evaluated else 0.0
        ),
        "policy_mismatches": mismatches,
        "required_field_compliance": (
            round(1 - required_field_fails / total_items, 4) if total_items else 1.0
        ),
    }
    expectations = expected.get("element_expectations") or []
    result["element_expectation_eval"] = (
        _evaluate_ai_element_expectations(items, expectations)
        if isinstance(expectations, list) and expectations
        else {
            "status": "NOT_AVAILABLE",
            "reason": "t3 standard does not provide expected.element_expectations",
        }
    )
    return result


def _evaluate_ai_element_expectations(
    items: list[dict[str, Any]],
    expectations: list[Any],
) -> dict[str, Any]:
    expected_items = [item for item in expectations if isinstance(item, dict)]
    matched: list[dict[str, Any]] = []
    policy_mismatches: list[dict[str, Any]] = []
    missing: list[dict[str, Any]] = []
    for expected in expected_items:
        unit_id = str(expected.get("unit_id") or "")
        policy = str(expected.get("policy") or "")
        expected_refs = _observation_int_set(expected.get("source_seq_refs"))
        if not unit_id or not policy or not expected_refs:
            continue
        candidates = [
            item
            for item in items
            if str(item.get("unit_id") or "") == unit_id
            and expected_refs & _observation_int_set(item.get("source_seq_refs"))
        ]
        exact = [item for item in candidates if str(item.get("policy") or "") == policy]
        if exact:
            matched.append(
                {
                    "unit_id": unit_id,
                    "expected_policy": policy,
                    "source_seq_refs": sorted(expected_refs),
                    "observed_policy": exact[0].get("policy"),
                    "observed_element_id": exact[0].get("element_id"),
                }
            )
        elif candidates:
            policy_mismatches.append(
                {
                    "unit_id": unit_id,
                    "expected_policy": policy,
                    "source_seq_refs": sorted(expected_refs),
                    "observed_policies": sorted(
                        {
                            str(item.get("policy") or "")
                            for item in candidates
                            if item.get("policy")
                        }
                    ),
                    "observed_element_ids": [
                        str(item.get("element_id") or "")
                        for item in candidates[:3]
                        if item.get("element_id")
                    ],
                }
            )
        else:
            missing.append(
                {
                    "unit_id": unit_id,
                    "expected_policy": policy,
                    "source_seq_refs": sorted(expected_refs),
                    "name": expected.get("name"),
                    "content_contains": expected.get("content_contains"),
                }
            )
    evaluable = len(matched) + len(policy_mismatches) + len(missing)
    source_overlap = len(matched) + len(policy_mismatches)
    return {
        "status": "AVAILABLE",
        "expectation_count": len(expected_items),
        "evaluable_count": evaluable,
        "source_overlap_count": source_overlap,
        "source_overlap_recall": round(source_overlap / evaluable, 4) if evaluable else 0.0,
        "policy_match_count": len(matched),
        "source_policy_overlap_accuracy": (
            round(len(matched) / evaluable, 4) if evaluable else 0.0
        ),
        "policy_mismatch_count": len(policy_mismatches),
        "missing_count": len(missing),
        "policy_mismatch_samples": policy_mismatches[:10],
        "missing_samples": missing[:10],
    }


def _evaluate_ai_layout_accuracy(
    observation: dict[str, Any],
    unit_observation: dict[str, Any],
    expected: dict[str, Any],
) -> dict[str, Any]:
    global_profile_present = any(
        isinstance(item, dict) and item.get("source") == "deterministic_facts"
        for item in observation.get("items", [])
    )
    page_observations = observation.get("page_observations") or []
    return {
        "global_profile_present": global_profile_present,
        "page_policy_evaluable": False,
        "page_policy_owner": "T2",
        "reason": "T4 layout accuracy does not evaluate or generate unit page policy",
        "page_count": observation.get("page_count"),
        "page_structure_source": observation.get("page_structure_source"),
        "vision_page_observation_count": len(page_observations)
        if isinstance(page_observations, list)
        else 0,
    }


def _observation_int_set(values: Any) -> set[int]:
    if not isinstance(values, list):
        return set()
    result: set[int] = set()
    for value in values:
        try:
            if value not in (None, ""):
                result.add(int(value))
        except (TypeError, ValueError):
            continue
    return result


def _read_route_payload(paths: list[Path], payload_type: str) -> dict[str, Any]:
    if payload_type == "json":
        loaded = read_json(paths[0])
        return loaded if isinstance(loaded, dict) else {}
    if payload_type == "yaml":
        loaded = read_yaml(paths[0])
        return loaded if isinstance(loaded, dict) else {}
    if payload_type == "composite":
        payload: dict[str, Any] = {"artifact_type": "composite_route_artifact", "parts": []}
        for path in paths:
            part: dict[str, Any] = {
                "name": path.name,
                "path": str(path),
                "sha256": sha256_file(path),
            }
            if path.suffix == ".json":
                loaded = read_json(path)
                if isinstance(loaded, dict):
                    part["artifact_type"] = loaded.get("artifact_type")
                    part["status"] = loaded.get("status")
            payload["parts"].append(part)
        return payload
    return {"artifact_type": payload_type}


def _route_artifact_path(
    run_dir: Path,
    filename: str,
) -> Path | None:
    candidates = [run_dir / filename]
    if filename == "template_gap_report.json":
        candidates.extend(
            [
                run_dir.parent / "template_gap" / filename,
                run_dir.parent / "template_gap" / "artifacts" / filename,
                run_dir.parent.parent / "template_gap" / "artifacts" / filename,
            ]
        )
    for path in candidates:
        if path.exists():
            return path
    return None


def _l1_route_coverage(route: dict[str, Any] | None) -> dict[str, Any]:
    if not route or route.get("availability") != "AVAILABLE" or not route.get("payload_path"):
        return {"available": False}
    payload = read_json(Path(str(route["payload_path"])))
    coverage = payload.get("coverage", {}) if isinstance(payload, dict) else {}
    return {"available": True, **coverage}


def _l1_route_mismatches(route: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not route or route.get("availability") != "AVAILABLE" or not route.get("payload_path"):
        return [
            {
                "id": "L1-ROUTE-MISMATCH-001",
                "stage_id": "L1",
                "stage_key": "l1_input_contract",
                "type": "l1_input_contract_not_available",
                "expected": "L1 input contract artifact exists",
                "observed": "template_generation_l1_input_contract route is not AVAILABLE",
                "route_ids": ["shared_input"],
            }
        ]
    payload = read_json(Path(str(route["payload_path"])))
    coverage = payload.get("coverage", {}) if isinstance(payload, dict) else {}
    mismatches: list[dict[str, Any]] = []
    if coverage.get("render_status") in {None, "", "not_available", "projection_fallback"}:
        mismatches.append(
            {
                "id": "L1-ROUTE-MISMATCH-002",
                "stage_id": "L1",
                "stage_key": "l1_input_contract",
                "type": "l1_visual_render_not_real",
                "expected": "visual_page_index binds a real render or records why it cannot",
                "observed": (
                    f"render_status={coverage.get('render_status')}; "
                    f"render_error={coverage.get('render_error') or 'none'}"
                ),
                "route_ids": ["shared_input"],
            }
        )
    if int(coverage.get("source_object_unbound_count") or 0) > 0:
        mismatches.append(
            {
                "id": "L1-ROUTE-MISMATCH-003",
                "stage_id": "L1",
                "stage_key": "l1_input_contract",
                "type": "l1_object_binding_gaps",
                "expected": "source_object_index entries have explicit page/bbox binding or reason",
                "observed": f"{coverage.get('source_object_unbound_count')} object(s) unbound",
                "route_ids": ["shared_input"],
            }
        )
    invalid = coverage.get("bundle_gate_invalid_stages") or []
    if invalid:
        mismatches.append(
            {
                "id": "L1-ROUTE-MISMATCH-004",
                "stage_id": "L1",
                "stage_key": "l1_input_contract",
                "type": "l1_bundle_gate_invalid",
                "expected": "AI observation bundle passes hash/schema/coverage gate",
                "observed": f"invalid bundle stages: {invalid}",
                "route_ids": ["shared_input"],
            }
        )
    return mismatches


def _t3_residual_gate(run_dir: Path) -> dict[str, Any]:
    element_spec = _load_run_yaml(
        run_dir,
        "03_element_spec.yaml",
    ) or {}
    format_element_hits = []
    placeholder_element_hits = []
    for element in _iter_element_dicts(element_spec):
        policy = str(element.get("policy") or "")
        content = str(
            element.get("content")
            or element.get("text")
            or element.get("name")
            or ""
        )
        element_id = str(element.get("element_id") or element.get("id") or "")
        unit_id = str(element.get("unit_id") or "")
        if not _is_instruction_policy(policy) and _FORMAT_ANNOTATION_RE.search(content):
            format_element_hits.append(
                {
                    "unit_id": unit_id,
                    "element_id": element_id,
                    "policy": policy,
                    "content": content[:160],
                }
            )
        if not _is_instruction_policy(policy) and _PLACEHOLDER_RE.search(content):
            placeholder_element_hits.append(
                {
                    "unit_id": unit_id,
                    "element_id": element_id,
                    "policy": policy,
                    "content": content[:160],
                }
            )
    visible_text = _visible_docx_text(_generated_docx_for_run(run_dir))
    format_docx_hits = _regex_samples(_FORMAT_ANNOTATION_RE, visible_text)
    placeholder_docx_hits = _regex_samples(_PLACEHOLDER_RE, visible_text)
    return {
        "artifact_type": "template_generation_t3_residual_gate",
        "artifact_version": "1.0",
        "format_annotation_non_instruction_count": len(format_element_hits),
        "format_annotation_final_docx_count": len(format_docx_hits),
        "placeholder_non_instruction_count": len(placeholder_element_hits),
        "placeholder_final_docx_count": len(placeholder_docx_hits),
        "format_annotation_non_instruction_samples": format_element_hits[:20],
        "format_annotation_final_docx_samples": format_docx_hits[:20],
        "placeholder_non_instruction_samples": placeholder_element_hits[:20],
        "placeholder_final_docx_samples": placeholder_docx_hits[:20],
        "status": (
            "PASS"
            if not (
                format_element_hits
                or placeholder_element_hits
                or format_docx_hits
                or placeholder_docx_hits
            )
            else "FAIL"
        ),
    }


def _t3_residual_mismatches(gate: dict[str, Any]) -> list[dict[str, Any]]:
    mismatches: list[dict[str, Any]] = []
    if int(gate.get("format_annotation_non_instruction_count") or 0) or int(
        gate.get("format_annotation_final_docx_count") or 0
    ):
        mismatches.append(
            {
                "id": "T3-RESIDUAL-MISMATCH-001",
                "stage_id": "T3",
                "stage_key": "t3_element_policy",
                "type": "t3_inline_format_instruction_residual",
                "expected": "format annotations are instruction_remove or absent from final DOCX",
                "observed": (
                    f"{gate.get('format_annotation_non_instruction_count')} element residual(s), "
                    f"{gate.get('format_annotation_final_docx_count')} final DOCX residual(s)"
                ),
                "route_ids": ["ai"],
            }
        )
    if int(gate.get("placeholder_non_instruction_count") or 0) or int(
        gate.get("placeholder_final_docx_count") or 0
    ):
        mismatches.append(
            {
                "id": "T3-RESIDUAL-MISMATCH-002",
                "stage_id": "T3",
                "stage_key": "t3_element_policy",
                "type": "t3_placeholder_like_residual",
                "expected": "placeholder-like sample spans are replaced with slots or justified",
                "observed": (
                    f"{gate.get('placeholder_non_instruction_count')} element residual(s), "
                    f"{gate.get('placeholder_final_docx_count')} final DOCX residual(s)"
                ),
                "route_ids": ["ai"],
            }
        )
    return mismatches


def _is_instruction_policy(policy: str) -> bool:
    return str(policy or "") in {"instruction_remove", "remove_instruction"}


def _iter_element_dicts(payload: Any) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    if isinstance(payload, dict):
        if "policy" in payload and any(key in payload for key in ("element_id", "content", "text", "name")):
            result.append(payload)
        for value in payload.values():
            result.extend(_iter_element_dicts(value))
    elif isinstance(payload, list):
        for item in payload:
            result.extend(_iter_element_dicts(item))
    return result


def _generated_docx_for_run(run_dir: Path) -> Path | None:
    path = run_dir / "06.1_fillable_template.docx"
    return path if path.exists() else None


def _visible_docx_text(path: Path | None) -> str:
    if path is None or not path.exists():
        return ""
    try:
        doc = Document(path)
    except Exception:
        return ""
    chunks: list[str] = []
    chunks.extend(paragraph.text for paragraph in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(chunks)


def _regex_samples(pattern: re.Pattern[str], text: str) -> list[str]:
    samples: list[str] = []
    for match in pattern.finditer(text):
        value = match.group(0).strip()
        if value and value not in samples:
            samples.append(value[:160])
        if len(samples) >= 50:
            break
    return samples


def build_template_generation_route_eval_markdown(report: dict[str, Any]) -> str:
    lines = [
        "# Template Generation Route Eval Report",
        "",
        f"- Run: {report.get('run_id')}",
        f"- Routes: {report.get('cross_route_summary', {}).get('route_count', 0)}",
        f"- Mismatches: {report.get('cross_route_summary', {}).get('mismatch_count', 0)}",
        "",
        "## Stage Metrics",
    ]
    for stage_id, metrics in (report.get("stage_metrics") or {}).items():
        lines.append(
            "- "
            f"{stage_id} {metrics.get('stage_key')}: "
            f"ai={metrics.get('ai_availability')}, "
            f"changed_from_code={metrics.get('changed_from_code')}"
        )
    lines.extend(["", "## Mismatches"])
    if report.get("mismatches"):
        for mismatch in report["mismatches"]:
            lines.append(
                f"- {mismatch['id']} {mismatch['type']}: {mismatch['observed']}"
            )
    else:
        lines.append("- none")
    lines.append("")
    return "\n".join(lines)


def _first_bad_stage(stage_checks: list[StageCheck]) -> str | None:
    for check in stage_checks:
        if check.status in {Status.FAIL, Status.UNKNOWN}:
            return check.stage_id
    return None


def _blocked_at_from_findings(findings: list[Finding]) -> str | None:
    for finding in findings:
        if finding.severity == "blocking":
            return finding.stage
    return None
