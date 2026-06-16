from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document

from docfit.core.io import read_json
from docfit.core.models import Finding, make_finding
from docfit.core.status import Status
from docfit.harness import template_units


DONOR_FRONT_MATTER_MARKERS = (
    "湖 南 农 业 大 学",
    "全日制普通本科生毕业论文",
    "学生姓名：",
    "学    号：",
    "年级专业及班级：",
    "指导老师及职称：",
    "学    院：",
)

BUSINESS_ACCEPTANCE_STAGES = ("template", "content", "placement", "render")

REQUIRED_E2E_AUDIT_FILES = (
    Path("artifacts/template_artifact.json"),
    Path("artifacts/student_content_artifact.json"),
    Path("artifacts/placement_plan.json"),
    Path("artifacts/render_manifest.json"),
    Path("final.docx"),
)


def audit_e2e_case(case_dir: Path) -> list[Finding]:
    """Return deterministic product-quality findings for a rendered e2e case."""

    artifacts = case_dir / "artifacts"
    template_artifact = read_json(artifacts / "template_artifact.json")
    content_artifact = read_json(artifacts / "student_content_artifact.json")
    placement_plan = read_json(artifacts / "placement_plan.json")
    render_manifest = read_json(artifacts / "render_manifest.json")
    final_docx = case_dir / "final.docx"

    findings: list[Finding] = []
    for stage_findings in (
        audit_template_artifact(template_artifact),
        audit_content_artifact(content_artifact),
        audit_placement_plan(placement_plan, content_artifact),
        audit_render_output(final_docx, template_artifact, placement_plan, render_manifest),
    ):
        findings.extend(_renumber(stage_findings, start_index=len(findings) + 1))
    return findings


def missing_e2e_audit_inputs(case_dir: Path) -> list[Path]:
    return [
        case_dir / relative_path
        for relative_path in REQUIRED_E2E_AUDIT_FILES
        if not (case_dir / relative_path).exists()
    ]


def business_acceptance_coverage(findings: list[Finding]) -> dict[str, bool]:
    return {
        f"business.{stage}_acceptance": not any(
            finding.stage == stage and finding.severity == "blocking"
            for finding in findings
        )
        for stage in BUSINESS_ACCEPTANCE_STAGES
    }


def audit_template_artifact(template_artifact: dict[str, Any]) -> list[Finding]:
    data = template_artifact.get("data", {})
    paragraphs = data.get("paragraphs", [])
    units = data.get("units", [])
    slots = data.get("slots", [])
    protected_zones = data.get("protected_zones", [])
    required_fields = data.get("required_fields", [])
    findings: list[Finding] = []
    next_index = 1

    if paragraphs and (
        len(units) < 5
        or _units_lack_elements(units)
        or not _has_non_virtual_slot(slots)
    ):
        findings.append(
            make_finding(
                next_index,
                "template",
                Status.UNKNOWN,
                "template_unit_tree_missing",
                "模板解析没有列出目标学校模板的主要位置",
                (
                    "应能分出封面、声明、目录、题名、摘要、正文、参考文献、"
                    "附录/致谢、手工表单，并说明哪些位置需要填写"
                ),
                (
                    f"{template_artifact.get('school_id')}/{template_artifact.get('template_version')} "
                    f"有 {len(paragraphs)} 个非空模板段落，"
                    f"{len(units)} 个单元，{_element_count(units)} 个元素，"
                    f"{len(slots)} 个位置={_slot_ids(slots)}，"
                    f"{len(protected_zones)} 个受保护区域，"
                    f"{len(required_fields)} 个必填字段；第一个段落："
                    f"{_preview(paragraphs[0].get('text', ''))}"
                ),
                evidence_refs=[template_artifact.get("provenance", {}).get("template_docx", "")],
                affected_ids=[slot.get("slot_id", "") for slot in slots],
                root_cause_bucket="template_unit_model_gap",
            )
        )
        next_index += 1

    instruction_examples = [
        paragraph
        for paragraph in paragraphs
        if _contains_instruction_marker(str(paragraph.get("text", "")))
        and not _has_output_policy(paragraph)
    ][:3]
    if instruction_examples:
        findings.append(
            make_finding(
                next_index,
                "template",
                Status.UNKNOWN,
                "template_instruction_paragraph_unclassified",
                "模板说明/示例文字仍被当成普通模板段落",
                "说明/示例段落应标成固定保留、需要填写，或不能写入最终成品",
                "; ".join(
                    f"p[{item.get('index')}]: {_preview(item.get('text', ''))}"
                    for item in instruction_examples
                ),
                evidence_refs=[
                    f"template_artifact.data.paragraphs[{item.get('index')}]"
                    for item in instruction_examples
                ],
                root_cause_bucket="template_instruction_policy_gap",
            )
        )
        next_index += 1
    findings.extend(
        template_units.verify_template_units_against_source_facts(
            template_artifact,
            start_index=next_index,
        )
    )
    return findings


def audit_content_artifact(content_artifact: dict[str, Any]) -> list[Finding]:
    ledger = content_artifact.get("data", {}).get("visible_content_ledger", [])
    findings: list[Finding] = []
    next_index = 1

    heading_like = [
        item
        for item in ledger
        if item.get("kind") == "paragraph"
        and _looks_like_heading(str(item.get("text", "")))
        and not _has_candidate(item, "heading")
    ][:5]
    if heading_like:
        findings.append(
            make_finding(
                next_index,
                "content",
                Status.UNKNOWN,
                "content_heading_semantics_unclassified",
                "内容抽取把正文标题当成普通段落",
                "标题、摘要、参考文献、附录、致谢等内容应有可检查的类型标记",
                "; ".join(_content_example(item) for item in heading_like),
                evidence_refs=[item.get("source_ref", "") for item in heading_like],
                affected_ids=[item.get("content_id", "") for item in heading_like],
                root_cause_bucket="content_semantic_gap",
            )
        )
        next_index += 1

    donor_front_matter = [
        item
        for item in ledger
        if item.get("kind") != "source_format"
        and any(marker in str(item.get("text", "")) for marker in DONOR_FRONT_MATTER_MARKERS)
    ][:5]
    if donor_front_matter:
        findings.append(
            make_finding(
                next_index,
                "content",
                Status.UNKNOWN,
                "content_donor_front_matter_not_disposed",
                "内容抽取把旧封面/源文档前置页当成学生正文",
                "旧学校封面和源模板前置页应标为源文档格式内容，或按已 review 规则忽略",
                "; ".join(_content_example(item) for item in donor_front_matter),
                evidence_refs=[item.get("source_ref", "") for item in donor_front_matter],
                affected_ids=[item.get("content_id", "") for item in donor_front_matter],
                root_cause_bucket="content_source_format_gap",
            )
        )
    return findings


def audit_placement_plan(
    placement_plan: dict[str, Any],
    content_artifact: dict[str, Any] | None = None,
) -> list[Finding]:
    actions = placement_plan.get("data", {}).get("actions", [])
    place_actions = [
        action for action in actions if action.get("disposition") == "place"
    ]
    if not place_actions:
        return []

    target_counts: dict[str, int] = {}
    for action in place_actions:
        target = str(action.get("target_slot_id") or "missing")
        target_counts[target] = target_counts.get(target, 0) + 1
    dominant_target, dominant_count = max(target_counts.items(), key=lambda item: item[1])
    if dominant_target != "slot_body_start" or dominant_count / len(place_actions) < 0.8:
        return []

    content_by_id = _content_by_id(content_artifact or {})
    examples = [
        _placement_example(action, content_by_id)
        for action in place_actions[:5]
    ]
    return [
        make_finding(
            1,
            "placement",
            Status.UNKNOWN,
            "placement_actions_collapsed_to_virtual_body_slot",
            "内容放置把学生内容都放到同一个兜底位置",
            (
                "每段内容应写到目标学校的具体位置，例如题名区、摘要区、"
                "正文标题、参考文献、附录或致谢"
            ),
            (
                f"{dominant_count}/{len(place_actions)} 个写入动作都指向 "
                f"{dominant_target}；例子：{'；'.join(examples)}"
            ),
            affected_ids=[
                content_id
                for action in place_actions[:5]
                for content_id in action.get("content_ids", [])
            ],
            root_cause_bucket="placement_unit_mapping_gap",
        )
    ]


def audit_render_output(
    final_docx: Path,
    template_artifact: dict[str, Any],
    placement_plan: dict[str, Any],
    render_manifest: dict[str, Any],
) -> list[Finding]:
    if not final_docx.exists():
        return [
            make_finding(
                1,
                "render",
                Status.UNKNOWN,
                "render_output_missing_for_product_audit",
                "成品质量检查需要生成的 final.docx",
                str(final_docx),
                "文件不存在",
                root_cause_bucket="render_output_missing",
            )
        ]

    doc = Document(final_docx)
    paragraphs = [
        {"index": index, "text": paragraph.text.strip()}
        for index, paragraph in enumerate(doc.paragraphs, start=1)
        if paragraph.text.strip()
    ]
    findings: list[Finding] = []
    next_index = 1

    leaked = [
        paragraph
        for paragraph in paragraphs
        if _contains_instruction_marker(paragraph["text"])
    ][:3]
    if leaked:
        findings.append(
            make_finding(
                next_index,
                "render",
                Status.FAIL,
                "render_template_instruction_text_leaked",
                "最终 Word 仍包含目标模板说明/示例文字",
                "最终 Word 应只包含目标学校固定内容、生成字段和已接受的学生内容",
                "; ".join(
                    f"docx paragraph {item['index']}: {_preview(item['text'])}"
                    for item in leaked
                ),
                evidence_refs=[f"{final_docx}:paragraph[{item['index']}]" for item in leaked],
                root_cause_bucket="render_template_leak",
            )
        )
        next_index += 1

    append_finding = _append_only_finding(
        template_artifact,
        placement_plan,
        render_manifest,
        paragraphs,
    )
    if append_finding is not None:
        findings.append(_renumber([append_finding], start_index=next_index)[0])
    return findings


def _append_only_finding(
    template_artifact: dict[str, Any],
    placement_plan: dict[str, Any],
    render_manifest: dict[str, Any],
    output_paragraphs: list[dict[str, Any]],
) -> Finding | None:
    template_paragraph_count = len(template_artifact.get("data", {}).get("paragraphs", []))
    writing_action_ids = {
        action.get("action_id")
        for action in placement_plan.get("data", {}).get("actions", [])
        if action.get("disposition") != "discard_as_source_format"
    }
    executed = [
        action
        for action in render_manifest.get("actions_executed", [])
        if action.get("action_id") in writing_action_ids
    ]
    first_ref = next(
        (
            str(action.get("actual_ooxml_ref", ""))
            for action in executed
            if str(action.get("actual_ooxml_ref", "")).startswith("word/document.xml:p[")
        ),
        "",
    )
    match = re.search(r"p\[(\d+)\]", first_ref)
    if not match:
        return None

    first_rendered_paragraph_index = int(match.group(1))
    if first_rendered_paragraph_index <= max(1, template_paragraph_count):
        return None

    first_output = _preview(output_paragraphs[0]["text"]) if output_paragraphs else "no text"
    return make_finding(
        1,
        "render",
        Status.FAIL,
        "render_append_only_insertion",
        "Word 生成把学生内容追加到复制模板之后",
        "第一个写入动作应进入目标位置，而不是整份模板正文之后",
        (
            f"第一个写入位置={first_ref}；模板解析结果有 "
            f"{template_paragraph_count} 个非空段落；输出第一个非空段落仍是 "
            f"{_preview(first_output)}"
        ),
        evidence_refs=[first_ref],
        root_cause_bucket="render_append_only",
    )


def _contains_instruction_marker(text: str) -> bool:
    return template_units.contains_instruction_marker(text)


def _has_non_virtual_slot(slots: list[dict[str, Any]]) -> bool:
    return any(
        slot.get("slot_id") != "slot_body_start"
        and not str(slot.get("source_ref", "")).startswith("real-core:virtual")
        for slot in slots
    )


def _units_lack_elements(units: list[dict[str, Any]]) -> bool:
    return any(not unit.get("elements") for unit in units)


def _element_count(units: list[dict[str, Any]]) -> int:
    return sum(len(unit.get("elements", [])) for unit in units)


def _has_output_policy(paragraph: dict[str, Any]) -> bool:
    return paragraph.get("template_policy") in {
        "fixed",
        "fill",
        "generated",
        "manual_only",
        "strip",
        "template_default_optional",
    }


def _slot_ids(slots: list[dict[str, Any]]) -> list[str]:
    return [str(slot.get("slot_id", "missing")) for slot in slots]


def _looks_like_heading(text: str) -> bool:
    stripped = text.strip()
    if not stripped or len(stripped) > 80:
        return False
    if stripped in {"摘要", "摘 要", "ABSTRACT", "Abstract", "参考文献", "致谢", "附录"}:
        return True
    if stripped.startswith(("关键词", "KEY WORDS", "KEYWORDS")):
        return True
    if re.match(r"^第[一二三四五六七八九十0-9]+[章节篇]\s*", stripped):
        return True
    if re.match(r"^\d+(?:\.\d+)*\s*[、.．]?\s*[\u4e00-\u9fffA-Za-z]", stripped):
        return True
    return False


def _has_candidate(item: dict[str, Any], kind: str) -> bool:
    return any(
        candidate.get("kind") == kind for candidate in item.get("semantic_candidates", [])
    )


def _content_example(item: dict[str, Any]) -> str:
    signals = item.get("style_signals", {})
    return (
        f"{item.get('content_id')}: {_preview(item.get('text', ''))} "
        f"(kind={item.get('kind')}, style={signals.get('style_name')}, "
        f"alignment={signals.get('alignment')}, font_size={signals.get('font_size')})"
    )


def _content_by_id(content_artifact: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        str(item.get("content_id")): item
        for item in content_artifact.get("data", {}).get("visible_content_ledger", [])
    }


def _placement_example(
    action: dict[str, Any],
    content_by_id: dict[str, dict[str, Any]],
) -> str:
    content_ids = [str(content_id) for content_id in action.get("content_ids", [])]
    content = content_by_id.get(content_ids[0], {}) if content_ids else {}
    return (
        f"{action.get('action_id')} {content_ids} -> {action.get('target_slot_id')} "
        f"text={_preview(content.get('text', action.get('render_kind', '')))}"
    )


def _preview(value: Any, *, limit: int = 96) -> str:
    text = str(value).replace("\n", " ").replace("\t", "<TAB>").strip()
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _renumber(findings: list[Finding], *, start_index: int) -> list[Finding]:
    for offset, finding in enumerate(findings):
        finding.finding_id = f"f_{start_index + offset:03d}"
    return findings
