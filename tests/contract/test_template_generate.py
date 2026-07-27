from __future__ import annotations

from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from docx import Document
from typer.testing import CliRunner

from docfit.cli.main import app
from docfit.convert.orchestrator import run_template_generate_eval
from docfit.core.io import read_json, read_yaml, sha256_file, sha256_json, write_json
from docfit.core.status import Status
from docfit.template_generation.constants import BODY_SLOT_MARKER
from docfit.template_generation.agent.config import AgentConfig
from docfit.template_generation.source_tree import inspect_document_facts_docx


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def write_source_docx(path: Path, paragraphs: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)


def write_source_docx_with_runs(
    path: Path,
    paragraphs: list[str | list[tuple[str, dict[str, object]]]],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for paragraph_spec in paragraphs:
        if isinstance(paragraph_spec, str):
            doc.add_paragraph(paragraph_spec)
            continue
        paragraph = doc.add_paragraph()
        for text, attrs in paragraph_spec:
            run = paragraph.add_run(text)
            if "bold" in attrs:
                run.bold = bool(attrs["bold"])
            if "italic" in attrs:
                run.italic = bool(attrs["italic"])
            if "underline" in attrs:
                run.underline = bool(attrs["underline"])
    doc.save(path)


def docx_texts(path: Path) -> list[str]:
    return [paragraph.text for paragraph in Document(path).paragraphs]


def table_texts(path: Path) -> list[str]:
    doc = Document(path)
    return [
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    ]


def docx_sdt_tags(path: Path) -> set[str]:
    tags: set[str] = set()
    with ZipFile(path) as package:
        root = ET.fromstring(package.read("word/document.xml"))
    for tag in root.iter(f"{W_NS}tag"):
        value = tag.attrib.get(f"{W_NS}val")
        if value:
            tags.add(value)
    return tags


def ai_replay_config(
    tmp_path: Path,
    *,
    end_page: int = 1,
    end_source_seq: int = 1,
) -> AgentConfig:
    replay = tmp_path / "t2-ai-replay.json"
    write_json(
        replay,
        {
            "t2": [
                {
                    "units": [
                        {
                            "unit_id": "template_pages",
                            "unit_name": "模板页面",
                            "boundary": {
                                "start_page": 1,
                                "end_page": end_page,
                            },
                        }
                    ]
                }
            ],
            "t3": {},
            "t4": {
                "section_profiles": [
                    {
                        "section_profile_id": "section_001",
                        "source_ref": "word/document.xml:body/sectPr",
                        "boundary": {
                            "start_source_seq": 1,
                            "end_source_seq": end_source_seq,
                            "confidence": "high",
                        },
                        "page_setup": {},
                        "header_footer": [],
                        "page_numbering": {
                            "declared": {"status": "none"},
                            "fields": [],
                            "display": {
                                "status": "no_page_field",
                                "has_page_field": False,
                                "checked_scopes": {
                                    "body_source_seq_range": {
                                        "start": 1,
                                        "end": end_source_seq,
                                    },
                                    "header_footer_parts": [],
                                },
                            },
                            "flags": [],
                        },
                        "evidence_refs": [
                            {"page_no": 1, "render_target_id": "page:1"}
                        ],
                    }
                ],
                "default_font": None,
                "page_numbering": {"status": "none"},
                "header_footer": [],
                "numbering_rules": [],
            },
        },
    )
    return AgentConfig(
        enabled=True,
        observation_mode="replay",
        observation_transcript_path=replay,
    )


def run_template_generate_replay(
    tmp_path: Path,
    source: Path,
    out_dir: Path,
):
    source_seq_count = len(inspect_document_facts_docx(source).get("body_flow", []))
    return run_template_generate_eval(
        tmp_path,
        source,
        out_dir,
        agent_config=ai_replay_config(
            tmp_path,
            end_source_seq=max(1, source_seq_count),
        ),
    )


def test_template_generate_writes_full_stage_artifact_chain(tmp_path) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template.docx"
    write_source_docx(source, ["学校固定封面", "目录", "正文开始", "格式说明：小四宋体"])

    bundle_root = (
        tmp_path
        / "test_outputs/debug/template_generation/20260621T160704467333+0800"
    )
    out_dir = (
        bundle_root
        / "eval_runs/template_generate"
    )
    result = run_template_generate_eval(
        tmp_path,
        source,
        out_dir,
        agent_config=ai_replay_config(tmp_path, end_source_seq=4),
    )

    fillable = out_dir / "06.1_fillable_template.docx"
    manifest_path = out_dir / "06.2_build_manifest.json"
    summary = read_json(out_dir / "summary.json")
    manifest = read_json(manifest_path)
    plan = result.artifacts["template_generation_plan"]
    document_facts = read_json(out_dir / "01_document_facts.json")
    l1_input_contract = read_json(out_dir / "01.5_l1_input_contract.json")
    unit_map = read_yaml(out_dir / "02_unit_map.yaml")
    t2_ai = read_yaml(out_dir / "02.2_t2_ai_unit_observation.yaml")
    element_spec = read_yaml(out_dir / "03_element_spec.yaml")
    t3_ai = read_yaml(out_dir / "03.1_t3_ai_element_observation.yaml")
    global_spec = read_yaml(out_dir / "04_global_spec.yaml")
    t4_ai = read_yaml(out_dir / "04.1_t4_ai_layout_observation.yaml")
    template_spec = read_yaml(out_dir / "05_template_spec.yaml")
    verification_report = read_json(out_dir / "07_verification_report.json")
    source_tree = result.artifacts["source_template_tree"]
    t2_input = read_json(out_dir / "02.1_t2_input.json")
    generation_model = result.artifacts["template_generation_model"]
    debug_index = read_json(out_dir / "99_template_generation_debug_index.json")
    issue_clusters = read_json(out_dir / "issue_clusters.json")

    assert result.status == Status.UNKNOWN
    assert result.run_status in {Status.PASS, Status.UNKNOWN, Status.FAIL}
    assert result.quality_status == Status.UNKNOWN
    assert out_dir.parent.name == "eval_runs"
    assert fillable.exists()
    assert (out_dir / "00_template_generation_request.json").exists()
    assert document_facts["artifact_type"] == "document_facts"
    assert l1_input_contract["input_hashes"]["document_facts"] == sha256_json(
        document_facts
    )
    serialized_document_facts = repr(document_facts)
    for forbidden in ("template_policy", "final_disposition", "policy_reason"):
        assert forbidden not in serialized_document_facts
    assert unit_map["artifact_type"] == "unit_map"
    assert t2_ai["artifact_type"] == "ai_unit_observation"
    assert t2_ai["units"] == [
        {
            "unit_id": "template_pages",
            "unit_name": "模板页面",
            "boundary": {"start_page": 1, "end_page": 1},
        }
    ]
    assert unit_map["lineage"]["producer_mode"] == "ai"
    assert unit_map["units"][0]["page_policy"] == {
        "start": "document_start",
        "scope": "page_range_exclusive",
    }
    assert element_spec["artifact_type"] == "element_spec"
    assert element_spec["route"]["route_id"] == "ai"
    assert element_spec["route"]["availability"] == "AVAILABLE"
    assert all(element["policy"] == "fixed" for element in element_spec["elements"])
    assert t3_ai["artifact_type"] == "ai_element_observation"
    assert t3_ai["route"]["route_id"] == "ai_raw"
    assert t3_ai["route"]["availability"] == "AVAILABLE"
    assert t3_ai["abstain"] is False
    assert t3_ai["coverage"]["unknown_source_seq"] == []
    assert t3_ai["coverage"]["total"] == len(document_facts["body_flow"])
    assert global_spec["artifact_type"] == "global_spec"
    assert global_spec["availability"]["status"] == "AVAILABLE"
    assert global_spec["lineage"]["producer_mode"] == "ai"
    assert len(global_spec["lineage"]["selected_from"]) == 1
    assert global_spec["lineage"]["selected_from"][0]["artifact"] == (
        "04.1_t4_ai_layout_observation.yaml"
    )
    assert t4_ai["artifact_type"] == "ai_layout_observation"
    assert "route" not in t4_ai
    assert t4_ai["abstain"] is False
    assert t4_ai["coverage"]["unknown_source_seq"] == []
    assert t4_ai["coverage"]["total"] == len(document_facts["body_flow"])
    assert set(global_spec) >= {
        "section_profiles",
        "default_font",
        "page_numbering",
        "header_footer",
        "numbering_rules",
        "flags",
    }
    assert template_spec["artifact_type"] == "template_spec"
    assert manifest["artifact_type"] == "build_manifest"
    assert manifest["observed_layout_effects"]["status"] == Status.PASS.value
    assert manifest["observed_layout_effects"]["output_docx_hash"] == sha256_file(fillable)
    assert verification_report["status"] == Status.UNKNOWN.value
    assert verification_report["first_bad_stage"] is None
    assert unit_map["flags"] == []
    # Element confidence is graded by final policy/evidence (not blanket medium):
    # unambiguous fixed/instruction/generated elements grade `high` and raise no
    # review flag; only genuinely-ambiguous elements stay medium/low and get one.
    assert all(
        element["confidence"] in {"high", "medium", "low"}
        for element in element_spec["elements"]
    )
    assert any(element["confidence"] == "high" for element in element_spec["elements"])
    assert all(
        flag["confidence"] in {"medium", "low"}
        for flag in element_spec["flags"]
        if flag["type"] == "element_confidence_needs_review"
    )
    assert global_spec["section_profiles"][0]["boundary"]["status"] == "detected"
    assert global_spec["section_profiles"][0]["page_numbering"]["display"]["status"] == (
        "no_page_field"
    )
    assert global_spec["page_numbering"]["status"] == "none"
    assert not any(flag["type"] == "page_numbering_unknown" for flag in global_spec["flags"])
    assert template_spec["review_flags"] == []
    assert not any(
        finding["type"] == "t2_unit_confidence_needs_review"
        for finding in verification_report["findings"]
    )
    unavailable_stages = {
        finding["affected_ids"][0]
        for finding in verification_report["findings"]
        if finding["type"] == "stage_final_not_available"
    }
    assert not unavailable_stages
    assert not any(
        finding["type"] == "t4_page_numbering_unknown"
        for finding in verification_report["findings"]
    )
    assert issue_clusters
    assert source_tree["artifact_type"] == "source_template_tree"
    assert t2_input["scope"] == "t2_page_groups"
    assert generation_model["artifact_type"] == "template_generation_model"
    assert unit_map["open_questions"] == []
    assert manifest_path.exists()
    source_seq_refs = [
        item["source_seq"] for item in document_facts["body_flow"]
    ]
    assert source_seq_refs == list(range(1, len(source_seq_refs) + 1))
    assert document_facts["indexes"]["by_source_seq"]["1"]["node_id"] == "body_0001"
    assert document_facts["runs"]
    assert all(run["raw_run_id"] and run["logical_run_id"] for run in document_facts["runs"])
    assert not (tmp_path / "runs/template_generation").exists()
    assert not (out_dir / "artifacts").exists()
    assert not (bundle_root / "human").exists()
    assert not (out_dir / "fillable_template.docx").exists()
    assert (out_dir / "00_input_source_template.docx").exists()
    assert (out_dir / "00_template_generation_request.json").exists()
    assert (out_dir / "01_document_facts.json").exists()
    assert not (out_dir / "02.0_t2_code_unit_map.yaml").exists()
    assert (out_dir / "02_unit_map.yaml").exists()
    assert (out_dir / "02.1_t2_input.json").exists()
    assert (out_dir / "02.2_t2_ai_unit_observation.yaml").exists()
    assert not (out_dir / "02.3_t2_merged_unit_map.yaml").exists()
    assert not (out_dir / "03.0_t3_code_element_spec.yaml").exists()
    assert (out_dir / "03.1_t3_ai_element_observation.yaml").exists()
    assert not (out_dir / "03.2_t3_merged_element_spec.yaml").exists()
    assert (out_dir / "03_element_spec.yaml").exists()
    assert not (out_dir / "04.0_t4_code_global_spec.yaml").exists()
    assert (out_dir / "04.1_t4_ai_layout_observation.yaml").exists()
    assert not (out_dir / "04.1.5_t4_ai_global_spec.yaml").exists()
    assert not (out_dir / "04.2_t4_merged_global_spec.yaml").exists()
    assert (out_dir / "04_global_spec.yaml").exists()
    assert (out_dir / "05_template_spec.yaml").exists()
    assert (out_dir / "06.0_copy_source_docx.docx").exists()
    assert (out_dir / "06.1_fillable_template.docx").exists()
    assert (out_dir / "06.2_build_manifest.json").exists()
    assert (out_dir / "07_verification_report.json").exists()
    assert (out_dir / "99_template_generation_debug_index.json").exists()
    assert summary["status"] == Status.UNKNOWN.value
    assert summary["run_status"] == result.run_status.value
    assert summary["quality_status"] == Status.UNKNOWN.value
    assert summary["unknown_findings"] > 0
    assert summary["artifacts"]["fillable_template_docx"] == str(fillable)
    assert "slot_body_start" in docx_sdt_tags(fillable)
    assert template_spec["availability"]["status"] == "AVAILABLE"
    assert manifest["availability"]["status"] == "AVAILABLE"
    assert manifest["input_refs"]["t5_final"]["sha256"] == sha256_json(template_spec)
    assert not any("[[DOCFIT_" in text for text in docx_texts(fillable))
    assert BODY_SLOT_MARKER not in docx_texts(out_dir / "06.0_copy_source_docx.docx")
    assert "格式说明：小四宋体" in docx_texts(out_dir / "06.0_copy_source_docx.docx")
    assert manifest["strategy"] == "source_copy_scaffold"
    assert "debug_snapshot" not in manifest
    assert manifest["output"]["fillable_template_docx"] == str(fillable)
    assert manifest["output"]["fillable_template_docx_hash"] == sha256_file(fillable)
    assert {
        item["name"] for item in debug_index["files"]
    } >= {
        "00_input_source_template.docx",
        "06.0_copy_source_docx.docx",
        "06.1_fillable_template.docx",
    }
    assert {slot["slot_id"] for slot in manifest["slots"]} >= {"slot_body_start"}
    assert manifest["actions_executed"]
    assert "actions_deferred" not in manifest
    assert not any(item["policy"] == "strip" for item in generation_model["cleanup"])
    assert "格式说明：小四宋体" in docx_texts(fillable)
    assert all("affected_source_seq_refs" in action for action in plan["actions"])


def test_template_generate_preserves_existing_body_slot(tmp_path) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template-with-slot.docx"
    write_source_docx(source, ["学校固定封面", BODY_SLOT_MARKER])

    result = run_template_generate_replay(
        tmp_path, source, tmp_path / "template_generate"
    )
    fillable = tmp_path / "template_generate/06.1_fillable_template.docx"
    manifest = read_json(
        tmp_path / "template_generate/06.2_build_manifest.json"
    )

    assert result.status == Status.UNKNOWN
    assert BODY_SLOT_MARKER not in docx_texts(fillable)
    assert "slot_body_start" in docx_sdt_tags(fillable)
    assert any(
        action["action_type"] == "ensure_body_slot"
        for action in manifest["actions_executed"]
    )
    assert any(slot["slot_id"] == "slot_body_start" for slot in manifest["slots"])


def test_template_generate_without_ai_preserves_instruction_text_inside_table_cells(tmp_path) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template-table.docx"
    source.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("摘要")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "学生姓名：×××"
    table.cell(0, 1).text = "格式说明：此处用小四宋体"
    doc.add_paragraph("正文")
    doc.save(source)

    result = run_template_generate_replay(
        tmp_path, source, tmp_path / "template_generate"
    )
    fillable = tmp_path / "template_generate/06.1_fillable_template.docx"
    manifest = read_json(
        tmp_path / "template_generate/06.2_build_manifest.json"
    )

    assert result.status == Status.UNKNOWN
    assert not manifest["actions_requiring_review"]
    assert any("格式说明" in text for text in table_texts(fillable))
    assert docx_sdt_tags(fillable) == {"slot_body_start"}


def test_template_generate_without_ai_safely_keeps_form_usage_notes(tmp_path) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template-form-notes.docx"
    out_dir = tmp_path / "template_generate"
    notes = [
        "一、毕业设计任务书是学校根据已经确定的毕业设计题目下达给学生的一种教学文件，是学生在指导教师指导下独立从事毕业设计工作的依据。此表由指导教师填写。",
        "四、任务书一经下达，不得随意更改。",
        "请在合适的对应选项前的“□”内打“√”，科研课题请注明课题项目和名称，项目指“国家青年基金”等。",
        "六、本表可从毕业论文管理系统填写打印或教务处网站下载中心下载填写打印，但签名栏必须相应责任人亲笔签名，且应用黑色签字笔填写。",
        "注：此表如不够填写，可另加附页。",
        "注：此表可从毕业论文管理系统或教务处网站下载中心下载。记录、签名栏必须用黑色笔手工填写。",
    ]
    write_source_docx(
        source,
        [
            "毕业设计任务书",
            *notes,
            "指导教师签名：",
            "评阅教师意见：",
            "正文",
        ],
    )

    result = run_template_generate_replay(tmp_path, source, out_dir)
    fillable = out_dir / "06.1_fillable_template.docx"
    plan = result.artifacts["template_generation_plan"]
    output_text = "\n".join(docx_texts(fillable))

    assert result.status == Status.UNKNOWN
    for note in notes:
        assert note in output_text
    assert "指导教师签名：" in output_text
    assert "评阅教师意见：" in output_text
    assert not any(
        action["action_type"] == "create_manual_placeholder"
        for action in plan["actions"]
    )
    assert not any(slot.get("kind") == "manual_only" for slot in result.artifacts["build_manifest"]["slots"])
    assert not any(
        action["action_type"] == "remove_instruction_text"
        for action in plan["actions"]
    )


def test_template_generate_keeps_table_cells_neutral_without_code_side_t2_merging(
    tmp_path,
) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template-table-label.docx"
    source.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("摘要")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "学生姓名："
    table.cell(0, 1).text = "____"
    doc.add_paragraph("正文")
    doc.save(source)

    result = run_template_generate_replay(
        tmp_path, source, tmp_path / "template_generate"
    )
    generation_model = result.artifacts["template_generation_model"]
    elements = generation_model["units"][0]["elements"]
    table_elements = [
        element
        for element in elements
        if element.get("source_refs")
        in (
            ["word/document.xml:tbl[1]/tr[1]/tc[1]"],
            ["word/document.xml:tbl[1]/tr[1]/tc[2]"],
        )
    ]

    assert result.status == Status.UNKNOWN
    assert len(table_elements) == 2
    assert all(element["candidate_policy"] == "fixed" for element in table_elements)
    assert all(
        element["role_hint"] == "unclassified_source_content"
        for element in table_elements
    )
    assert all("merge" not in element for element in table_elements)


def test_template_generate_keeps_paragraphs_neutral_without_code_side_t2_merging(
    tmp_path,
) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template-continuation.docx"
    out_dir = tmp_path / "template_generate"
    write_source_docx(
        source,
        [
            "摘要",
            "摘要正文：本研究围绕生成流程，",
            "重点讨论阶段证据链。",
            "正文",
        ],
    )

    result = run_template_generate_replay(tmp_path, source, out_dir)
    generation_model = result.artifacts["template_generation_model"]
    unit = generation_model["units"][0]
    continuation_elements = [
        element
        for element in unit["elements"]
        if element.get("source_refs")
        in (["word/document.xml:p[2]"], ["word/document.xml:p[3]"])
    ]

    assert result.status == Status.UNKNOWN
    assert unit["unit_id"] == "template_pages"
    assert len(continuation_elements) == 2
    assert [element["source_seq_refs"] for element in continuation_elements] == [[2], [3]]
    assert all(
        element["role_hint"] == "unclassified_source_content"
        for element in continuation_elements
    )
    assert all("merge" not in element for element in continuation_elements)


def test_template_generate_invalid_docx_fails_without_output(tmp_path) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/not-a-docx.docx"
    source.parent.mkdir(parents=True, exist_ok=True)
    source.write_text("not a zip package", encoding="utf-8")

    result = run_template_generate_replay(
        tmp_path, source, tmp_path / "template_generate"
    )
    summary = read_json(tmp_path / "template_generate/summary.json")

    assert result.status == Status.FAIL
    assert summary["status"] == Status.FAIL.value
    assert summary["blocked_at"] == "template_generate"
    assert not (tmp_path / "template_generate/generated_template.docx").exists()
    assert any(
        finding.type == "template_generation_source_invalid"
        for finding in result.findings
    )


def test_template_generate_cli_writes_public_outputs(tmp_path) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template.docx"
    out_dir = tmp_path / "cli_template_generate"
    write_source_docx(source, ["学校固定封面"])
    replay = ai_replay_config(tmp_path).observation_transcript_path
    assert replay is not None

    result = CliRunner().invoke(
        app,
        [
            "template",
            "generate",
            "--template",
            str(source),
            "--out",
            str(out_dir),
            "--ai",
            "replay",
            "--replay",
            str(replay),
        ],
    )

    assert result.exit_code == 0
    assert "status = UNKNOWN" in result.stdout
    assert (out_dir / "06.1_fillable_template.docx").exists()
    assert (out_dir / "06.2_build_manifest.json").exists()
    assert (out_dir / "05_template_spec.yaml").exists()
    assert (out_dir / "summary.json").exists()


def test_template_generate_cli_rejects_school_standard_input(tmp_path) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template.docx"
    out_dir = tmp_path / "cli_template_generate"
    write_source_docx(source, ["学校固定封面"])

    result = CliRunner().invoke(
        app,
        [
            "eval",
            "template-generate",
            "--school",
            "hunannongye",
            "--template",
            str(source),
            "--out",
            str(out_dir),
        ],
    )

    assert result.exit_code != 0
    assert "No such option" in result.output or "No such option" in result.stderr


def test_template_generate_uses_copy_then_patch_actions(tmp_path) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template.docx"
    out_dir = tmp_path / "template_generate"
    write_source_docx(source, ["封面", "参考文献"])

    result = run_template_generate_replay(tmp_path, source, out_dir)
    generation_model = result.artifacts["template_generation_model"]
    plan = result.artifacts["template_generation_plan"]
    strategies = generation_model["unit_strategies"]

    assert result.status == Status.UNKNOWN
    assert [strategy["unit_id"] for strategy in strategies] == ["template_pages"]
    assert strategies[0]["generation_mode"] == "copy_then_patch"
    assert strategies[0]["generation_policy"] == "unit_actions"


def test_template_generate_does_not_invent_custom_units_outside_ai_output(tmp_path) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template.docx"
    out_dir = tmp_path / "template_generate"
    source.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("封面")
    doc.add_paragraph("研究方案设计", style="Heading 1")
    doc.add_paragraph("论文题目：____")
    doc.add_paragraph("正文", style="Heading 1")
    doc.save(source)

    result = run_template_generate_replay(tmp_path, source, out_dir)
    generation_model = result.artifacts["template_generation_model"]
    plan = result.artifacts["template_generation_plan"]
    strategies = generation_model["unit_strategies"]

    assert result.status == Status.UNKNOWN
    assert [strategy["unit_id"] for strategy in strategies] == ["template_pages"]
    assert not any(
        str(strategy["unit_id"]).startswith("custom")
        for strategy in strategies
    )
    assert not any(
        action["action_type"] in {"create_fillable_slot", "replace_span_with_slot"}
        for action in plan["actions"]
    )


def test_template_generate_cover_uses_patch_analysis(
    tmp_path,
) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template.docx"
    out_dir = tmp_path / "template_generate"
    write_source_docx(
        source,
        [
            "封面",
            "论文题目：____",
            "格式说明：小四宋体",
            "目录",
            "正文",
        ],
    )

    result = run_template_generate_replay(tmp_path, source, out_dir)
    fillable = out_dir / "06.1_fillable_template.docx"
    generation_model = result.artifacts["template_generation_model"]
    plan = result.artifacts["template_generation_plan"]
    unit_model = generation_model["units"][0]
    cover_strategy = next(
        unit
        for unit in generation_model["unit_strategies"]
        if unit["unit_id"] == "template_pages"
    )
    unit_elements = unit_model["elements"]
    model_title = next(
        element
        for element in unit_elements
        if element.get("source_refs") == ["word/document.xml:p[2]"]
    )
    model_instruction = next(
        element
        for element in unit_elements
        if element.get("source_refs") == ["word/document.xml:p[3]"]
    )

    assert result.status == Status.UNKNOWN
    assert unit_model["unit_id"] == "template_pages"
    assert model_title["role_hint"] == "unclassified_source_content"
    assert model_title["candidate_policy"] == "fixed"
    assert model_title["policy"] == "fixed"
    assert model_instruction["role_hint"] == "unclassified_source_content"
    assert model_instruction["candidate_policy"] == "fixed"
    assert cover_strategy["generation_mode"] == "copy_then_patch"
    assert not any(
        decision["decision_type"] == "remove_instruction_text"
        and decision["source_ref"] == "word/document.xml:p[3]"
        for decision in cover_strategy["decisions"]
    )
    assert not any(
        action.get("unit_id") == "template_pages"
        and action["action_type"] in {"create_fillable_slot", "replace_span_with_slot"}
        and action["affected_source_seq_refs"] == [2]
        for action in plan["actions"]
    )
    assert not any(
        action["action_type"] == "remove_instruction_text"
        and action.get("source_ref") == "word/document.xml:p[3]"
        and action["affected_source_seq_refs"] == [3]
        for action in plan["actions"]
    )
    assert "格式说明：小四宋体" in docx_texts(fillable)
    assert not any(
        tag.startswith("template_pages.") for tag in docx_sdt_tags(fillable)
    )


def test_template_generate_splits_within_paragraph_format_instruction_runs(
    tmp_path,
) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template.docx"
    out_dir = tmp_path / "template_generate"
    write_source_docx_with_runs(
        source,
        [
            "封面",
            [
                ("毕业论文（设计）中文题目  ", {"bold": False}),
                ("（小二黑体加粗）", {"bold": True}),
            ],
            "目录",
            "正文",
            [
                ("这是", {"bold": False}),
                ("普通", {"bold": True}),
                ("正文段落", {"bold": False}),
            ],
        ],
    )

    result = run_template_generate_replay(tmp_path, source, out_dir)
    fillable = out_dir / "06.1_fillable_template.docx"
    generation_model = result.artifacts["template_generation_model"]
    element_spec = read_yaml(out_dir / "03_element_spec.yaml")
    plan = result.artifacts["template_generation_plan"]
    manifest = read_json(out_dir / "06.2_build_manifest.json")

    unit_model = generation_model["units"][0]
    model_cover_elements = [
        element
        for element in unit_model["elements"]
        if element.get("source_refs") == ["word/document.xml:p[2]"]
    ]
    assert result.status == Status.UNKNOWN
    assert unit_model["unit_id"] == "template_pages"
    assert model_cover_elements
    assert all(element["policy"] == "fixed" for element in model_cover_elements)
    assert all(element["spans"] == [] for element in model_cover_elements)
    assert all(element["policy"] == "fixed" for element in element_spec["elements"])
    assert not any(
        action["action_type"] in {"remove_instruction_text", "replace_span_with_slot"}
        for action in plan["actions"]
    )
    assert not any(
        action["action_type"] in {"remove_instruction_text", "replace_span_with_slot"}
        for action in manifest["actions_executed"]
    )
    assert any("毕业论文（设计）中文题目" in text for text in docx_texts(fillable))
    assert any("小二黑体加粗" in text for text in docx_texts(fillable))


def test_template_generate_ai_replay_preserves_field_line_placeholder_spans(tmp_path) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template.docx"
    out_dir = tmp_path / "template_generate"
    write_source_docx_with_runs(
        source,
        [
            "封面",
            [
                ("□□□□□□", {"bold": True}),
                ("学□□号：", {"bold": True}),
                ("20××××××××××", {"bold": True}),
            ],
            "目录",
            "正文",
        ],
    )

    result = run_template_generate_replay(tmp_path, source, out_dir)
    fillable = out_dir / "06.1_fillable_template.docx"
    element_spec = read_yaml(out_dir / "03_element_spec.yaml")
    plan = result.artifacts["template_generation_plan"]
    manifest = read_json(out_dir / "06.2_build_manifest.json")
    cover_field = next(
        element
        for element in element_spec["elements"]
        if element["unit_id"] == "template_pages"
        and "学□□号" in element.get("content", "")
    )

    assert result.status == Status.UNKNOWN
    assert cover_field["policy"] == "fixed"
    assert cover_field["spans"] == []
    assert not any(
        action["action_type"] == "replace_span_with_slot"
        and action["unit_id"] == "template_pages"
        and action["element_id"] == cover_field["element_id"]
        for action in plan["actions"]
    )
    assert not any(
        action["action_type"] == "replace_span_with_slot"
        for action in manifest["actions_executed"]
    )
    joined_text = "\n".join(docx_texts(fillable))
    assert "□□□□□□" in joined_text
    assert "学□□号" in joined_text
    assert "20××" in joined_text
    assert not any(
        tag.startswith("template_pages.") for tag in docx_sdt_tags(fillable)
    )


def test_template_generate_does_not_invent_references_unit_outside_ai_output(
    tmp_path,
) -> None:
    source = tmp_path / "inputs/targets/demo-school/raw/school-template.docx"
    out_dir = tmp_path / "template_generate"
    write_source_docx(source, ["封面", "正文", "参考文献", "学生文献内容占位"])

    result = run_template_generate_replay(tmp_path, source, out_dir)
    generation_model = result.artifacts["template_generation_model"]
    plan = result.artifacts["template_generation_plan"]
    strategies = generation_model["unit_strategies"]

    assert result.status == Status.UNKNOWN
    assert [strategy["unit_id"] for strategy in strategies] == ["template_pages"]
    assert not any(
        action["action_type"] in {"create_fillable_slot", "replace_span_with_slot"}
        for action in plan["actions"]
    )
    assert not any(
        tag.startswith("references.")
        for tag in docx_sdt_tags(out_dir / "06.1_fillable_template.docx")
    )
