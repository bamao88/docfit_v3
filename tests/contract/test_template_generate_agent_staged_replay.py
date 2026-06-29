from __future__ import annotations

from docx import Document
from typer.testing import CliRunner

from docfit.cli.main import app
from docfit.core.io import read_json, read_yaml, write_json
from docfit.template_generation.agent.packet import build_template_agent_render_packet
from docfit.template_generation.artifacts import source_tree_from_document_facts
from docfit.template_generation.source_tree import inspect_document_facts_docx
from docfit.template_generation.structure_candidates import build_template_structure_candidates


def test_template_generate_cli_staged_replay_writes_pass_plan(tmp_path) -> None:
    source = tmp_path / "school-template.docx"
    doc = Document()
    doc.add_paragraph("封面")
    doc.add_paragraph("学生姓名：____")
    doc.add_paragraph("正文")
    doc.add_paragraph("阶段二新增单元")
    doc.add_paragraph("正文续段")
    doc.save(source)
    facts = inspect_document_facts_docx(source)
    candidates = build_template_structure_candidates(source_tree_from_document_facts(facts))
    packet = build_template_agent_render_packet(
        document_facts=facts,
        structure_candidates=candidates,
        source_template_docx=source,
    )
    packet_path = tmp_path / "packet.json"
    transcript_path = tmp_path / "transcript.json"
    write_json(packet_path, packet)
    write_json(
        transcript_path,
        {
            "artifact_type": "template_agent_transcript",
            "rounds": [
                {
                    "round_id": "round_001",
                    "pass_id": "t2_unit_scan",
                    "pass_kind": "t2_unit_scan",
                    "window_id": "full_document",
                    "allowed_layers": ["t2"],
                    "submission": {
                        "source_render_hash": packet["source_render_hash"],
                        "round_id": "round_001",
                        "model": "fixture",
                        "layers": {
                            "t2": {
                                "unit_candidates": [
                                    {
                                        "proposal_id": "t2_add_001",
                                        "kind": "unit_candidate",
                                        "operation": "add_unit",
                                        "unit_id": "agent_added_unit",
                                        "display_name": "阶段二新增单元",
                                        "source_seq_refs": [4],
                                    }
                                ],
                                "block_candidates": [],
                                "boundary_adjustments": [],
                                "open_questions": [],
                            },
                            "t3": {"element_policy_candidates": [], "open_questions": []},
                            "t4": {
                                "page_policy_hints": [],
                                "section_profile_hints": [],
                                "page_numbering_hints": [],
                                "open_questions": [],
                            },
                        },
                    },
                },
                {
                    "round_id": "round_002",
                    "pass_id": "t3_unit_elements_cover",
                    "pass_kind": "t3_unit_elements",
                    "window_id": "unit:cover",
                    "unit_id": "cover",
                    "allowed_layers": ["t3"],
                    "submission": {
                        "source_render_hash": packet["source_render_hash"],
                        "round_id": "round_002",
                        "model": "fixture",
                        "layers": {
                            "t2": {
                                "unit_candidates": [],
                                "block_candidates": [],
                                "boundary_adjustments": [],
                                "open_questions": [],
                            },
                            "t3": {
                                "element_policy_candidates": [
                                    {
                                        "proposal_id": "t3_fill_001",
                                        "kind": "element_policy_candidate",
                                        "policy": "fill",
                                        "source_seq_refs": [2],
                                    }
                                ],
                                "open_questions": [],
                            },
                            "t4": {
                                "page_policy_hints": [],
                                "section_profile_hints": [],
                                "page_numbering_hints": [],
                                "open_questions": [],
                            },
                        },
                    },
                },
            ],
        },
    )
    out_dir = tmp_path / "template_generate"

    result = CliRunner().invoke(
        app,
        [
            "eval",
            "template-generate",
            "--template",
            str(source),
            "--out",
            str(out_dir),
            "--agent-replay",
            str(transcript_path),
            "--agent-render-packet",
            str(packet_path),
        ],
    )

    assert result.exit_code == 0, result.output
    pass_plan = read_json(out_dir / "artifacts/template_agent_pass_plan.json")
    checkpoint = read_json(out_dir / "artifacts/template_agent_post_t2_checkpoint.json")
    post_t2_input = read_json(out_dir / "artifacts/template_agent_post_t2_input.json")
    unit_windows = read_json(out_dir / "artifacts/template_agent_unit_windows.json")
    decisions = read_json(out_dir / "artifacts/template_agent_decisions.json")
    comparison = read_json(
        out_dir / "artifacts/template_agent_submission_comparison.json"
    )
    manual_review = read_json(
        out_dir / "artifacts/template_agent_manual_review_items.json"
    )
    unit_map = read_yaml(out_dir / "artifacts/unit_map.yaml")
    element_spec = read_yaml(out_dir / "artifacts/element_spec.yaml")
    assert [item["pass_kind"] for item in pass_plan["passes"]] == [
        "t2_unit_scan",
        "t3_unit_elements",
    ]
    assert decisions["accepted_proposal_ids"] == ["t2_add_001", "t3_fill_001"]
    assert comparison["summary"]["auto_executable"] == 2
    assert manual_review["summary"]["total"] == 0
    assert checkpoint["changed"] is True
    assert post_t2_input["artifact_type"] == "template_agent_post_t2_input"
    assert post_t2_input["source_seq_ownership"]["4"] == "agent_added_unit"
    assert any(
        unit["unit_id"] == "agent_added_unit"
        for unit in checkpoint["unit_map"]["units"]
    )
    cover_window = next(
        window
        for window in unit_windows["windows"]
        if window["unit_id"] == "cover"
    )
    assert cover_window["window_id"] == "unit:cover"
    assert 2 in cover_window["source_seq_refs"]
    assert any(unit["unit_id"] == "agent_added_unit" for unit in unit_map["units"])
    assert any(
        element["source_seq_refs"] == [2] and element["policy"] == "fill"
        for element in element_spec["elements"]
    )
    assert (out_dir / "08.5_agent_pass_plan.json").exists()
    assert (out_dir / "08.6_agent_post_t2_checkpoint.json").exists()
    assert (out_dir / "08.65_agent_post_t2_input.json").exists()
    assert (out_dir / "08.7_agent_unit_windows.json").exists()
    assert (out_dir / "09.5_agent_submission_comparison.json").exists()
    assert (out_dir / "10.5_agent_manual_review_items.json").exists()
