from __future__ import annotations

from docx import Document
from typer.testing import CliRunner

from docfit.cli.main import app
from docfit.core.io import read_json, read_yaml, write_json
from docfit.template_generation.agent.packet import build_template_agent_render_packet
from docfit.template_generation.input_contract import build_l1_input_contract
from docfit.template_generation.stage_inputs import build_agent_stage_packet
from docfit.template_generation.source_tree import inspect_document_facts_docx


def test_template_generate_cli_rejects_cross_run_t4_bundle_without_code_fallback(
    tmp_path,
) -> None:
    source = tmp_path / "school-template.docx"
    doc = Document()
    doc.add_paragraph("封面")
    doc.add_paragraph("承诺书")
    doc.add_paragraph("正文")
    doc.save(source)
    facts = inspect_document_facts_docx(source)
    packet = build_template_agent_render_packet(
        document_facts=facts,
        source_template_docx=source,
        render_artifacts_dir=tmp_path / "agent_render",
    )
    stage_packet = build_agent_stage_packet(
        build_l1_input_contract(document_facts=facts, render_packet=packet)
    )
    page_count = packet["render_artifacts"]["page_count"]
    packet_path = tmp_path / "packet.json"
    bundle_path = tmp_path / "observation_bundle.json"
    write_json(packet_path, packet)
    write_json(
        bundle_path,
        {
            "artifact_type": "ai_observation_bundle",
            "source_render_hash": packet["source_render_hash"],
            "model": "fixture",
            "ai_unit_observation": {
                "artifact_type": "ai_unit_observation",
                "units": [
                    {
                        "unit_id": "template_pages",
                        "unit_name": "模板页面",
                        "boundary": {"start_page": 1, "end_page": page_count},
                    }
                ],
            },
            "ai_element_observation": {
                "artifact_type": "ai_element_observation",
                "items": [],
                "quality_report": {"demotions": []},
            },
            "ai_layout_observation": {
                "artifact_type": "ai_layout_observation",
                "source_render_hash": stage_packet["source_render_hash"],
                "input_contract_hash": stage_packet["input_contract_hash"],
                "items": [
                        {
                            "section_profile_id": "section_001",
                            "source_ref": "word/document.xml:body/sectPr",
                        "boundary": {
                            "start_source_seq": 1,
                            "end_source_seq": 3,
                            "confidence": "high",
                        },
                        "source_seq_refs": [1, 2, 3],
                        "page_setup": {},
                        "header_footer": [],
                        "page_numbering": {},
                        "evidence_refs": [
                            {"page_no": 1, "render_target_id": "page:1"}
                        ],
                    }
                ],
                "abstain": False,
                "coverage": {
                    "total": 3,
                    "owned_source_seq": [1, 2, 3],
                    "unknown_source_seq": [],
                },
                "page_numbering": {"status": "unknown"},
                "header_footer": [],
                "numbering_rules": [],
                "quality_report": {"demotions": []},
            },
        },
    )
    out_dir = tmp_path / "template_generate_observation"

    result = CliRunner().invoke(
        app,
        [
            "eval",
            "template-generate",
            "--template",
            str(source),
            "--out",
            str(out_dir),
            "--agent-render-packet",
            str(packet_path),
            "--agent-observation-bundle",
            str(bundle_path),
        ],
    )

    assert result.exit_code == 0, result.output
    ai_t2 = read_yaml(out_dir / "02.2_t2_ai_unit_observation.yaml")
    ai_t3 = read_yaml(out_dir / "03.1_t3_ai_element_observation.yaml")
    ai_t4 = read_yaml(out_dir / "04.1_t4_ai_layout_observation.yaml")
    t4_final = read_yaml(out_dir / "04_global_spec.yaml")
    assert ai_t2["artifact_type"] == "ai_unit_observation"
    assert "route" not in ai_t2
    assert ai_t2["units"][0]["unit_id"] == "template_pages"
    assert ai_t3["route"]["route_id"] == "ai_raw"
    assert ai_t3["route"]["availability"] == "AVAILABLE"
    assert ai_t3["artifact_type"] == "ai_element_observation"
    assert "route" not in ai_t4
    assert ai_t4["artifact_type"] == "ai_layout_observation"
    assert t4_final["artifact_type"] == "global_spec"
    assert t4_final["availability"]["status"] == "NOT_AVAILABLE"
    assert t4_final["section_profiles"] == []
    assert t4_final["lineage"]["producer_mode"] == "ai"
    assert not (out_dir / "04.0_t4_code_global_spec.yaml").exists()
    assert not (out_dir / "04.1.5_t4_ai_global_spec.yaml").exists()
    assert not (out_dir / "04.2_t4_merged_global_spec.yaml").exists()
    assert not (out_dir / "09.25_agent_observation_bridge.json").exists()
    assert not (out_dir / "09.5_agent_submission_comparison.json").exists()
    assert not (out_dir / "14_agent_attribution.json").exists()


def test_template_generate_cli_observation_replay_runs_module1_in_same_run(tmp_path) -> None:
    source = tmp_path / "school-template.docx"
    doc = Document()
    doc.add_paragraph("封面")
    doc.add_paragraph("承诺书")
    doc.add_paragraph("正文")
    doc.save(source)
    facts = inspect_document_facts_docx(source)
    packet = build_template_agent_render_packet(
        document_facts=facts,
        source_template_docx=source,
        render_artifacts_dir=tmp_path / "agent_render",
    )
    page_count = packet["render_artifacts"]["page_count"]
    packet_path = tmp_path / "packet.json"
    observation_replay_path = tmp_path / "observation_replay.json"
    write_json(packet_path, packet)
    write_json(
        observation_replay_path,
        {
            "t2": [
                {
                    "units": [
                        {
                            "unit_id": "template_pages",
                            "unit_name": "模板页面",
                            "boundary": {
                                "start_page": 1,
                                "end_page": page_count,
                            },
                        },
                    ]
                }
            ],
            "t3": {
                "template_pages": {
                    "items": [
                        {
                            "element_id": "template_pages.001",
                            "policy": "fixed",
                            "role": "template_fixed",
                            "content": "封面",
                            "source_seq_refs": [1],
                            "confidence": "high",
                        },
                        {
                            "element_id": "template_pages.002",
                            "policy": "fixed",
                            "role": "template_fixed",
                            "content": "承诺书",
                            "source_seq_refs": [2],
                            "confidence": "high",
                        },
                        {
                            "element_id": "template_pages.003",
                            "policy": "fill",
                            "fill_source": "student_content",
                            "source_seq_refs": [3],
                            "confidence": "high",
                        }
                    ]
                },
            },
            "t4": {
                "section_profiles": [
                    {
                        "section_profile_id": "section_001",
                        "source_ref": "word/document.xml:body/sectPr",
                        "boundary": {
                            "start_source_seq": 1,
                            "end_source_seq": 3,
                            "confidence": "high",
                        },
                        "page_setup": {},
                        "header_footer": [],
                        "page_numbering": {},
                        "evidence_refs": [
                            {"page_no": 1, "render_target_id": "page:1"}
                        ],
                    }
                ],
                "page_numbering": {"status": "unknown"},
                "header_footer": [],
                "numbering_rules": [],
            },
        },
    )
    out_dir = tmp_path / "template_generate_observation_replay"

    result = CliRunner().invoke(
        app,
        [
            "eval",
            "template-generate",
            "--template",
            str(source),
            "--out",
            str(out_dir),
            "--agent-render-packet",
            str(packet_path),
            "--agent-observation-replay",
            str(observation_replay_path),
        ],
    )

    assert result.exit_code == 0, result.output
    bundle = read_json(out_dir / "09.1_ai_observation_bundle.json")
    l1_input_contract = read_json(out_dir / "01.5_l1_input_contract.json")
    ai_t3 = read_yaml(out_dir / "03.1_t3_ai_element_observation.yaml")
    t4_final = read_yaml(out_dir / "04_global_spec.yaml")
    element_spec = read_yaml(out_dir / "03_element_spec.yaml")
    hierarchical_input = read_json(
        out_dir / "03.0_t3_hierarchical_stage_input.json"
    )
    sparse_trace = read_json(out_dir / "03.1.5_t3_sparse_decision_trace.json")
    assert bundle["source_render_hash"] == (
        l1_input_contract["visual_page_index"]["source_render_hash"]
    )
    assert ai_t3["route"]["route_id"] == "ai_raw"
    assert ai_t3["route"]["availability"] == "AVAILABLE"
    assert ai_t3["items"]
    assert element_spec["route"]["route_id"] == "ai"
    assert element_spec["route"]["availability"] == "AVAILABLE"
    assert hierarchical_input["artifact_type"] == "t3_hierarchical_stage_input"
    assert sparse_trace["artifact_type"] == "t3_sparse_decision_trace"
    assert sparse_trace["stage_input_ref"]["tree_hash"] == hierarchical_input["tree_hash"]
    assert sparse_trace["validation"] == {"valid": True, "errors": []}
    assert len({row["member_ref"] for row in sparse_trace["atomic_coverage"]}) == len(
        sparse_trace["atomic_coverage"]
    )
    assert t4_final["availability"]["status"] == "AVAILABLE"
    assert t4_final["lineage"]["producer_mode"] == "ai"
    assert (out_dir / "09.1_ai_observation_bundle.json").exists()
    assert (out_dir / "03.0_t3_hierarchical_stage_input.json").exists()
    assert (out_dir / "03.1.5_t3_sparse_decision_trace.json").exists()
    assert not (out_dir / "03.0_t3_code_element_spec.yaml").exists()
    assert not (out_dir / "03.2_t3_merged_element_spec.yaml").exists()
    assert not (out_dir / "03.2.5_t3_atomic_route_comparison.json").exists()
    assert not (out_dir / "09.25_agent_observation_bridge.json").exists()
