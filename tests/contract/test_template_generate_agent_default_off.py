from __future__ import annotations

from docx import Document

from docfit.convert.orchestrator import run_template_generate_eval
from docfit.core.io import read_json


def test_template_generate_default_off_fails_without_code_fallback(tmp_path) -> None:
    source = tmp_path / "school-template.docx"
    doc = Document()
    doc.add_paragraph("封面")
    doc.add_paragraph("正文")
    doc.save(source)

    out_dir = tmp_path / "template_generate"
    result = run_template_generate_eval(tmp_path, source, out_dir)

    assert result.status.value == "FAIL"
    assert result.blocked_at == "template_generate"
    assert "template_agent_render_packet" not in result.artifacts
    assert not (out_dir / "08_agent_render_packet.json").exists()
    assert not (out_dir / "01.6_t2_l1_stage_input.json").exists()
    assert not (out_dir / "02.0_t2_code_unit_map.yaml").exists()
    assert not (out_dir / "02_unit_map.yaml").exists()
    summary = read_json(out_dir / "summary.json")
    assert summary["status"] == "FAIL"
    assert summary["primary_failure_bucket"] == "agent_config_invalid"
    assert "ordered.template_agent_render_packet" not in summary["artifacts"]
    run_manifest = read_json(out_dir / "run_manifest.json")
    assert run_manifest["entrypoint"] == "python.run_template_generate_eval"
    assert run_manifest["ai_mode"] == "off"
    assert run_manifest["api_call_count"] == 0
    assert run_manifest["source_template_hash"]
    assert run_manifest["source_render_hash"] is None
    assert "ordered.template_agent_render_packet" not in run_manifest["output_artifacts"]
