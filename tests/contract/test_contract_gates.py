from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from docfit.convert.orchestrator import run_content_eval
from docfit.core.io import write_json
from docfit.core.status import Status
from docfit.harness.profiles import BOOTSTRAP_PROFILE
from docfit.harness.coverage import evaluate_bootstrap_coverage
from docfit.harness.standards import load_standard_bundle
from docfit.stages.content_extract.runner import extract_student_content
from docfit.stages.placement.runner import build_placement_plan
from docfit.stages.render.runner import render_docx, verify_render_outputs
from docfit.stages.template_parse.runner import parse_template


ROOT = Path.cwd()


def _bundle():
    bundle, findings = load_standard_bundle(ROOT, "demo-school")
    assert bundle is not None
    assert findings == []
    return bundle


def test_unknown_when_standard_missing() -> None:
    bundle, findings = load_standard_bundle(ROOT, "missing-school")
    assert bundle is None
    assert findings[0].status == Status.UNKNOWN
    assert findings[0].type == "missing_signed_standard"


def test_school_standard_tree_contains_only_runnable_standards() -> None:
    version_dirs = [
        path
        for school_dir in (ROOT / "standards/schools").iterdir()
        if school_dir.is_dir()
        for path in school_dir.iterdir()
        if path.is_dir()
    ]

    assert version_dirs
    assert all((path / "signed_standard.yaml").exists() for path in version_dirs)


def test_unknown_when_signed_standard_capability_profile_drifts(tmp_path) -> None:
    copied_school_root = tmp_path / "standards/schools/demo-school"
    copied_school_root.parent.mkdir(parents=True)
    shutil.copytree(ROOT / "standards/schools/demo-school", copied_school_root)
    signed_standard = copied_school_root / "v1/signed_standard.yaml"
    signed_standard.write_text(
        signed_standard.read_text(encoding="utf-8").replace(
            "  - content.visible_tables\n",
            "  - content.visible_text_blocks\n",
        ),
        encoding="utf-8",
    )

    _, findings = load_standard_bundle(tmp_path, "demo-school")

    assert any(finding.status == Status.UNKNOWN for finding in findings)
    assert any(finding.type == "coverage_requirements_drift" for finding in findings)


def test_fail_when_template_slot_missing(tmp_path) -> None:
    template = tmp_path / "missing-slot.docx"
    doc = Document()
    doc.add_paragraph("No slot marker here")
    doc.save(template)

    result = parse_template(template, _bundle())

    assert result.status == Status.FAIL
    assert any(finding.type == "required_slot_missing" for finding in result.findings)


def test_fail_when_visible_content_missing_from_ledger() -> None:
    result = extract_student_content(
        ROOT / "inputs/bootstrap-demo-student-pass.docx",
        omit_content_id_for_test="c_002",
    )

    assert result.status == Status.FAIL
    assert any(finding.type == "visible_content_missing_from_ledger" for finding in result.findings)


def test_unknown_when_unsupported_visible_object() -> None:
    result = extract_student_content(
        ROOT / "inputs/bootstrap-demo-student-unsupported-textbox.docx"
    )

    assert result.status == Status.UNKNOWN
    assert any(finding.type == "unsupported_visible_object" for finding in result.findings)


def test_unknown_when_content_coverage_insufficient(tmp_path) -> None:
    student = tmp_path / "no-table.docx"
    doc = Document()
    doc.add_heading("No Table Thesis", level=1)
    doc.add_paragraph("This document intentionally has no table.")
    doc.save(student)

    result = run_content_eval(student, tmp_path / "content")

    assert result.status == Status.UNKNOWN
    assert any(finding.type == "coverage_insufficient" for finding in result.findings)


def test_bootstrap_coverage_checks_fixture_content_not_just_paths(tmp_path) -> None:
    template = tmp_path / "inputs/bootstrap-demo-school-template.docx"
    student = tmp_path / "inputs/bootstrap-demo-student-pass.docx"
    expected_paths = BOOTSTRAP_PROFILE.expected_paths(tmp_path)
    expected_snapshot = expected_paths["feature_snapshot"]
    expected_placement = expected_paths["placement_plan"]
    template.parent.mkdir(parents=True, exist_ok=True)
    student.parent.mkdir(parents=True, exist_ok=True)

    template_doc = Document()
    template_doc.add_paragraph("[[DOCFIT_SLOT:body]]")
    template_doc.save(template)

    student_doc = Document()
    student_doc.add_paragraph("This student document has no table.")
    student_doc.save(student)

    write_json(expected_placement, {"data": {"actions": []}})
    write_json(
        expected_snapshot,
        {"required_content_hashes": ["sha256:demo"]},
    )

    report, findings = evaluate_bootstrap_coverage(tmp_path)

    assert report["status"] == Status.UNKNOWN.value
    assert "content.visible_tables" in report["missing"]
    assert any(finding.type == "coverage_insufficient" for finding in findings)


def test_fail_when_content_unplaced() -> None:
    bundle = _bundle()
    template = parse_template(bundle.template_docx, bundle)
    content = extract_student_content(ROOT / "inputs/bootstrap-demo-student-pass.docx")

    result = build_placement_plan(
        template.artifacts["template_artifact"],
        content.artifacts["student_content_artifact"],
        drop_content_id_for_test="c_002",
    )

    assert result.status == Status.FAIL
    assert any(finding.type == "unplaced_content" for finding in result.findings)


def test_source_toc_entry_is_discarded_as_source_format(tmp_path) -> None:
    student = tmp_path / "student-with-old-toc.docx"
    doc = Document()
    doc.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("摘要\tI", style="toc 1")
    doc.add_paragraph("正文第一段。")
    doc.save(student)

    bundle = _bundle()
    template = parse_template(bundle.template_docx, bundle)
    content = extract_student_content(student)

    first_item = content.artifacts["student_content_artifact"]["data"][
        "visible_content_ledger"
    ][0]
    assert first_item["kind"] == "source_format"
    assert first_item["semantic_candidates"][0]["kind"] == "source_toc_entry"

    placement = build_placement_plan(
        template.artifacts["template_artifact"],
        content.artifacts["student_content_artifact"],
    )

    discard_action = placement.artifacts["placement_plan"]["data"]["actions"][0]
    assert discard_action["disposition"] == "discard_as_source_format"
    assert placement.status == Status.PASS

    render = render_docx(
        template.artifacts["template_artifact"],
        placement.artifacts["placement_plan"],
        bundle,
        tmp_path / "rendered",
    )
    rendered_doc = Document(render.artifact_paths["final_docx"])

    assert all("摘要\tI" not in paragraph.text for paragraph in rendered_doc.paragraphs)
    assert render.artifacts["render_manifest"]["actions_executed"][0][
        "actual_ooxml_ref"
    ] == "discarded:source_format"
    assert first_item["text_hash"] not in render.artifacts["feature_snapshot"][
        "expected_content_hashes"
    ]


def test_unknown_when_golden_missing(tmp_path) -> None:
    plan = {"data": {"actions": [{"action_id": "a_001"}]}}
    manifest = {"actions_executed": [{"action_id": "a_001"}]}
    feature_diff = {"status": "UNKNOWN", "diffs": []}
    oracle = {"valid_docx_package": True}

    findings = verify_render_outputs(plan, manifest, feature_diff, oracle, tmp_path / "missing.json")

    assert any(finding.status == Status.UNKNOWN for finding in findings)
    assert any(finding.type == "missing_signed_golden" for finding in findings)
