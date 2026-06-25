from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from docfit.convert.orchestrator import run_template_gap_eval
from docfit.core.io import read_json, sha256_file
from docfit.core.status import Status
from docfit.template_gap.gap import summarize_template_gap_report
from docfit.template_gap.inspector import (
    inspect_generated_template_docx,
    iter_visible_text_entries,
)
from docfit.harness.profiles import (
    BOOTSTRAP_PROFILE,
    REAL_CORE_SCHOOLS,
)


ROOT = Path.cwd()


def unit_by_id(report: dict, unit_id: str) -> dict:
    return next(unit for unit in report["units"] if unit["unit_id"] == unit_id)


def element_by_id(unit: dict, element_id: str) -> dict:
    return next(element for element in unit["elements"] if element["element_id"] == element_id)


def collect_checks(
    report: dict,
    *,
    type: str | None = None,
    category: str | None = None,
    status: str | None = None,
) -> list[dict]:
    checks: list[dict] = []
    if isinstance(report.get("input"), dict):
        checks.append(report["input"])
    checks.extend(report.get("global_checks", []))
    for unit in report.get("units", []):
        if isinstance(unit.get("presence"), dict):
            checks.append(unit["presence"])
        for element in unit.get("elements", []):
            for key in ("presence", "style"):
                if isinstance(element.get(key), dict):
                    checks.append(element[key])
        for dimension_checks in unit.get("dimensions", {}).values():
            checks.extend(dimension_checks)
    checks.extend(report.get("unmodeled_objects", []))
    if type is not None:
        checks = [check for check in checks if check.get("type") == type]
    if category is not None:
        checks = [check for check in checks if check.get("category") == category]
    if status is not None:
        checks = [check for check in checks if check.get("status") == status]
    return checks


def write_minimal_gap_standard(
    root: Path,
    school_id: str,
    expected_units: list[dict],
) -> None:
    school_dir = root / "standards/targets" / school_id / "v1"
    school_dir.mkdir(parents=True)
    required = list(BOOTSTRAP_PROFILE.all_required_capabilities())
    signed_standard = {
        "standard_id": f"{school_id}-v1",
        "school_id": school_id,
        "template_version": "v1",
        "status": "test_fixture",
        "owner": "docfit-tests",
        "source": {
            "template_docx": "inputs/targets/demo-school/raw/source_template.docx",
            "template_docx_sha256": "sha256:test",
        },
        "contracts": {
            "template_contract": "template_contract.json",
            "template_generation_contract": "template_generation_contract.json",
            "template_quality_contract": "template_quality_contract.json",
            "student_content_contract": "student_content_contract.json",
            "placement_contract": "placement_contract.json",
            "render_contract": "render_contract.json",
        },
        "evidence_baselines": {
            "template_generation_final": "template_quality/final_template.expected.yaml",
        },
        "coverage_requirements": {
            "profile": BOOTSTRAP_PROFILE.profile_id,
            "required_capabilities": required,
        },
    }
    (school_dir / "target.standard.yaml").write_text(
        _simple_yaml(signed_standard),
        encoding="utf-8",
    )
    contract = {
        "contract_type": "test",
        "contract_version": "1.0",
        "owner": "docfit-tests",
        "required_invariants": [],
        "required_capabilities": [],
        "unsupported_policy": {},
        "coverage_requirements": {},
        "verifier_refs": [],
    }
    for name in signed_standard["contracts"].values():
        (school_dir / name).write_text(json.dumps(contract), encoding="utf-8")
    baseline = {
        "baseline_type": "template_generation_final",
        "profile_id": "test",
        "school_id": school_id,
        "review_metadata": {
            "reviewed_by": "test",
            "review_source": "tests",
            "source_docx_sha256": "sha256:test",
            "change_reason": "test fixture",
            "auto_update_allowed": False,
        },
        "dimensions": [
            {"dimension_id": "expected.units", "comparator_mode": "exact"},
        ],
        "expected": {"units": expected_units},
    }
    final_path = school_dir / "template_quality/final_template.expected.yaml"
    final_path.parent.mkdir(parents=True, exist_ok=True)
    final_path.write_text(
        _simple_yaml(baseline),
        encoding="utf-8",
    )


def write_docx(path: Path, paragraphs: list[dict | str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for item in paragraphs:
        if isinstance(item, str):
            doc.add_paragraph(item)
            continue
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(str(item["text"]))
        if item.get("style"):
            paragraph.style = str(item["style"])
        if item.get("font"):
            run.font.name = str(item["font"])
        if item.get("size"):
            run.font.size = Pt(float(item["size"]))
        if "bold" in item:
            run.bold = bool(item["bold"])
        if item.get("alignment") == "center":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


def write_docx_with_header(path: Path, paragraphs: list[str], header_text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = header_text
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)


def _simple_yaml(value, indent: int = 0) -> str:
    if isinstance(value, dict):
        lines = []
        for key, item in value.items():
            if isinstance(item, (dict, list)):
                lines.append(" " * indent + f"{key}:")
                lines.append(_simple_yaml(item, indent + 2).rstrip())
            else:
                lines.append(" " * indent + f"{key}: {json.dumps(item, ensure_ascii=False)}")
        return "\n".join(lines) + "\n"
    if isinstance(value, list):
        lines = []
        for item in value:
            if isinstance(item, (dict, list)):
                lines.append(" " * indent + "-")
                lines.append(_simple_yaml(item, indent + 2).rstrip())
            else:
                lines.append(" " * indent + f"- {json.dumps(item, ensure_ascii=False)}")
        return "\n".join(lines) + "\n"
    return " " * indent + json.dumps(value, ensure_ascii=False) + "\n"


def minimal_expected_units(*, style: str = "") -> list[dict]:
    return [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "学校名称",
                    "policy": "fixed",
                    "content": "测试大学",
                    "style": style,
                },
                {
                    "element_id": "e_002",
                    "name": "学生姓名",
                    "policy": "manual_only",
                    "content": "学生姓名：",
                    "style": "",
                },
            ],
        },
        {
            "unit_id": "body",
            "name": "正文",
            "order": 20,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "正文标题",
                    "policy": "fixed",
                    "content": "第一章 绪论",
                    "style": "",
                }
            ],
        },
    ]


def test_template_gap_minimal_fixture_can_pass_cleanly(tmp_path) -> None:
    school_id = "minimal-school"
    write_minimal_gap_standard(tmp_path, school_id, minimal_expected_units())
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        ["测试大学", "学生姓名：", "第一章 绪论"],
    )

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "minimal_gap",
    )
    report = read_json(tmp_path / "minimal_gap/artifacts/template_gap_report.json")

    assert result.status == Status.PASS
    assert report["artifact_version"] == "2.0"
    assert "check_items" not in report
    assert report["summary"]["blocking_status"] == Status.PASS.value
    assert report["summary"]["failed_count"] == 0
    assert report["summary"]["unknown_count"] == 0
    assert unit_by_id(report, "cover")["verdict"] == Status.PASS.value


def test_template_gap_treats_document_start_as_first_page_break(tmp_path) -> None:
    school_id = "document-start-page-school"
    expected = minimal_expected_units()
    expected[0]["page"] = {"page_break": "是"}
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        ["测试大学", "学生姓名：", "第一章 绪论"],
    )

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "document_start_page_gap",
    )
    report = read_json(
        tmp_path / "document_start_page_gap/artifacts/template_gap_report.json"
    )
    page_items = collect_checks(report, category="page_rule")

    assert result.status == Status.PASS
    assert any(
        item["type"] == "template_generation_page_rule_match"
        and item["affected_ids"] == ["cover.page.page_break"]
        and item["actual"] == "document starts at p[1]"
        for item in page_items
    )


def test_template_gap_accepts_docfit_slot_and_generated_markers(tmp_path) -> None:
    school_id = "marker-school"
    expected = [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "学校名称",
                    "policy": "fixed",
                    "content": "测试大学",
                    "style": "",
                },
                {
                    "element_id": "e_002",
                    "name": "学生姓名",
                    "policy": "fill",
                    "content": "学生姓名",
                    "style": "",
                },
            ],
        },
        {
            "unit_id": "auto_block",
            "name": "自动内容块",
            "order": 20,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "系统生成内容",
                    "policy": "generated",
                    "content": "系统生成内容",
                    "style": "",
                }
            ],
        },
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        [
            "测试大学",
            "[[DOCFIT_SLOT:cover.e_002]]",
            "[[DOCFIT_GENERATED:auto_block.e_001]]",
        ],
    )

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "marker_gap",
    )
    report = read_json(tmp_path / "marker_gap/artifacts/template_gap_report.json")
    cover_slot = element_by_id(unit_by_id(report, "cover"), "e_002")["presence"]
    generated = element_by_id(unit_by_id(report, "auto_block"), "e_001")["presence"]

    assert result.status == Status.PASS
    assert cover_slot["type"] == "template_generation_slot_marker_found"
    assert generated["type"] == "template_generation_generated_marker_found"


def test_template_gap_accepts_explicit_marker_outside_located_unit_range(
    tmp_path,
) -> None:
    school_id = "marker-range-school"
    expected = [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "学校名称",
                    "policy": "fixed",
                    "content": "测试大学",
                    "style": "",
                }
            ],
        },
        {
            "unit_id": "auto_block",
            "name": "自动内容块",
            "order": 20,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "系统生成内容",
                    "policy": "generated",
                    "content": "系统生成内容",
                    "style": "",
                }
            ],
        },
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        [
            "测试大学",
            "[[DOCFIT_GENERATED:auto_block.e_001]]",
            "自动内容块",
        ],
    )

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "marker_range_gap",
    )
    report = read_json(tmp_path / "marker_range_gap/artifacts/template_gap_report.json")
    generated = element_by_id(unit_by_id(report, "auto_block"), "e_001")[
        "presence"
    ]

    assert result.status == Status.PASS
    assert generated["type"] == "template_generation_generated_marker_found"
    assert generated["evidence_refs"] == ["word/document.xml:p[2]"]


def test_template_gap_starts_first_unit_at_first_body_entry_when_anchor_is_ambiguous(
    tmp_path,
) -> None:
    school_id = "first-unit-school"
    expected = [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "封面本体",
                    "policy": "fixed",
                    "content": "保留学校封面本体",
                    "style": "",
                }
            ],
        },
        {
            "unit_id": "copyright_notice",
            "name": "版权声明",
            "order": 20,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "版权声明",
                    "policy": "fixed",
                    "content": "版权声明",
                    "style": "",
                }
            ],
        },
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        [
            "博士研究生学位论文",
            "版权声明",
            "论文打印装订时的注意事项",
            "封面—实名评审专家名单—版权声明—论文本体—原创性声明—封底",
        ],
    )

    run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "first_unit_gap",
    )
    report = read_json(tmp_path / "first_unit_gap/artifacts/template_gap_report.json")

    assert unit_by_id(report, "cover")["located"]["source_ref"] == "word/document.xml:p[1]"
    assert (
        unit_by_id(report, "copyright_notice")["located"]["source_ref"]
        == "word/document.xml:p[2]"
    )


def test_template_gap_does_not_use_header_footer_as_unit_anchor(tmp_path) -> None:
    school_id = "header-anchor-school"
    expected = [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "学校名称",
                    "policy": "fixed",
                    "content": "测试大学",
                    "style": "",
                }
            ],
        },
        {
            "unit_id": "acknowledgement",
            "name": "致谢",
            "order": 20,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "致谢标题",
                    "policy": "fixed",
                    "content": "致谢",
                    "style": "",
                }
            ],
        },
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx_with_header(generated_template, ["测试大学"], "致谢")

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "header_anchor_gap",
    )
    report = read_json(tmp_path / "header_anchor_gap/artifacts/template_gap_report.json")
    acknowledgement = unit_by_id(report, "acknowledgement")

    assert result.status == Status.FAIL
    assert acknowledgement["located"]["found"] is False
    assert acknowledgement["presence"]["type"] == "template_generation_unit_missing"


def test_template_gap_keeps_body_anchor_at_earlier_slot_before_placeholder_chapter(
    tmp_path,
) -> None:
    school_id = "body-anchor-school"
    expected = [
        {
            "unit_id": "toc",
            "name": "目录",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "目录",
                    "policy": "fixed",
                    "content": "目录",
                    "style": "",
                }
            ],
        },
        {
            "unit_id": "body_main",
            "name": "正文主体",
            "order": 20,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "正文章标题",
                    "policy": "fill",
                    "content": "学生章节标题",
                    "style": "",
                },
                {
                    "element_id": "e_002",
                    "name": "固定类型结尾章",
                    "policy": "fill",
                    "content": "第X章 结论与展望",
                    "style": "",
                },
            ],
        },
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        [
            "目录",
            "[[DOCFIT_SLOT:body_main.e_001]]",
            {"text": "第X章 结论与展望", "style": "Heading 1"},
        ],
    )

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "body_anchor_gap",
    )
    report = read_json(tmp_path / "body_anchor_gap/artifacts/template_gap_report.json")
    body_main = unit_by_id(report, "body_main")

    assert result.status == Status.PASS
    assert body_main["located"]["source_ref"] == "word/document.xml:p[2]"


def test_template_gap_reports_out_of_order_toc_marker_instead_of_missing(
    tmp_path,
) -> None:
    school_id = "toc-marker-school"
    expected = [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "封面",
                    "policy": "fixed",
                    "content": "封面",
                    "style": "",
                }
            ],
        },
        {
            "unit_id": "abstract_en",
            "name": "英文摘要",
            "order": 20,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "ABSTRACT",
                    "policy": "fixed",
                    "content": "ABSTRACT",
                    "style": "",
                }
            ],
        },
        {
            "unit_id": "toc",
            "name": "目录",
            "order": 30,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "目录标题",
                    "policy": "generated",
                    "content": "目录",
                    "style": "",
                },
                {
                    "element_id": "e_003",
                    "name": "目录结果条目",
                    "policy": "generated",
                    "content": "一级标题文字、点引导线、页码。",
                    "style": "",
                },
            ],
        },
        {
            "unit_id": "figure_list",
            "name": "图目录",
            "order": 40,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "图目录",
                    "policy": "fixed",
                    "content": "图目录",
                    "style": "",
                }
            ],
        },
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        [
            "封面",
            "[[DOCFIT_GENERATED:toc.e_001]]",
            "ABSTRACT",
            "图目录",
            "[[DOCFIT_GENERATED:toc.e_003]]",
        ],
    )

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "toc_marker_gap",
    )
    report = read_json(tmp_path / "toc_marker_gap/artifacts/template_gap_report.json")
    toc = unit_by_id(report, "toc")

    assert result.status == Status.FAIL
    assert toc["located"]["source_ref"] == "word/document.xml:p[2]"
    assert any(
        item["type"] == "template_generation_unit_order_mismatch"
        for item in collect_checks(report)
    )


def test_template_gap_single_style_mutation_fails_only_that_element(tmp_path) -> None:
    school_id = "style-school"
    expected = minimal_expected_units(style="宋体；12pt；加粗；居中。")
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        [
            {"text": "测试大学", "font": "宋体", "size": 14, "bold": True, "alignment": "center"},
            "学生姓名：",
            "第一章 绪论",
        ],
    )

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "style_gap",
    )
    report = read_json(tmp_path / "style_gap/artifacts/template_gap_report.json")
    mismatches = collect_checks(
        report,
        type="template_generation_style_mismatch",
    )

    assert result.status == Status.FAIL
    assert len(mismatches) == 1
    assert mismatches[0]["path"] == ["units", "cover", "elements", "e_001", "style"]
    assert "font_size expected 12pt" in mismatches[0]["actual"]


def test_template_gap_missing_required_element_is_presence_fail(tmp_path) -> None:
    school_id = "missing-element-school"
    write_minimal_gap_standard(tmp_path, school_id, minimal_expected_units())
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(generated_template, ["测试大学", "第一章 绪论"])

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "missing_element_gap",
    )
    report = read_json(tmp_path / "missing_element_gap/artifacts/template_gap_report.json")
    cover = unit_by_id(report, "cover")
    missing = element_by_id(cover, "e_002")["presence"]

    assert result.status == Status.FAIL
    assert missing["status"] == Status.FAIL.value
    assert missing["type"] == "template_generation_element_missing"
    assert missing["path"] == ["units", "cover", "elements", "e_002", "presence"]


def test_template_gap_unsearchable_fixed_element_is_unknown_not_fail(tmp_path) -> None:
    school_id = "unsearchable-school"
    expected = [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "学校名称",
                    "policy": "fixed",
                    "content": "测试大学",
                    "style": "",
                },
                {
                    "element_id": "e_002",
                    "name": "内容",
                    "policy": "fixed",
                    "content": "正文内容",
                    "style": "",
                },
            ],
        }
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(generated_template, ["测试大学"])

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "unsearchable_gap",
    )
    report = read_json(tmp_path / "unsearchable_gap/artifacts/template_gap_report.json")
    uncheckable = collect_checks(
        report,
        type="template_generation_element_uncheckable",
    )

    assert result.status == Status.UNKNOWN
    assert len(uncheckable) == 1
    assert uncheckable[0]["status"] == Status.UNKNOWN.value


def test_template_gap_checks_page_rule_text_against_ooxml(
    tmp_path,
) -> None:
    school_id = "nonvisible-rule-school"
    expected = [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "学校名称",
                    "policy": "fixed",
                    "content": "测试大学",
                    "style": "",
                },
                {
                    "element_id": "e_002",
                    "name": "纸张：A4。",
                    "policy": "fixed",
                    "content": "纸张：A4。",
                    "style": "",
                },
            ],
        }
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(generated_template, ["测试大学"])

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "nonvisible_rule_gap",
    )
    report = read_json(tmp_path / "nonvisible_rule_gap/artifacts/template_gap_report.json")
    paper_rule = element_by_id(unit_by_id(report, "cover"), "e_002")["presence"]

    assert result.status == Status.FAIL
    assert paper_rule["status"] == Status.FAIL.value
    assert paper_rule["type"] == "template_generation_page_setup_mismatch"


def test_template_gap_treats_word_paragraph_options_as_uncheckable(
    tmp_path,
) -> None:
    school_id = "word-option-school"
    expected = [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "学校名称",
                    "policy": "fixed",
                    "content": "测试大学",
                    "style": "",
                },
                {
                    "element_id": "e_002",
                    "name": "孤行控制",
                    "policy": "manual_only",
                    "content": "孤行控制",
                    "style": "",
                },
            ],
        }
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(generated_template, ["测试大学"])

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "word_option_gap",
    )
    report = read_json(tmp_path / "word_option_gap/artifacts/template_gap_report.json")
    option = element_by_id(unit_by_id(report, "cover"), "e_002")["presence"]

    assert result.status == Status.UNKNOWN
    assert option["status"] == Status.UNKNOWN.value
    assert option["type"] == "template_generation_element_uncheckable"


def test_template_gap_still_fails_missing_visible_fixed_text(tmp_path) -> None:
    school_id = "visible-missing-school"
    expected = [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "学校名称",
                    "policy": "fixed",
                    "content": "测试大学",
                    "style": "",
                },
                {
                    "element_id": "e_002",
                    "name": "地点",
                    "policy": "fixed",
                    "content": "湖南·长沙",
                    "style": "",
                },
            ],
        }
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(generated_template, ["测试大学"])

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "visible_missing_gap",
    )
    report = read_json(tmp_path / "visible_missing_gap/artifacts/template_gap_report.json")
    place = element_by_id(unit_by_id(report, "cover"), "e_002")["presence"]

    assert result.status == Status.FAIL
    assert place["status"] == Status.FAIL.value
    assert place["type"] == "template_generation_element_missing"


def test_template_gap_normalizes_template_noise(tmp_path) -> None:
    school_id = "noise-school"
    expected = [
        {
            "unit_id": "toc",
            "name": "目录",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "目录标签",
                    "policy": "fixed",
                    "content": "目 录",
                    "style": "",
                },
                {
                    "element_id": "e_002",
                    "name": "学生信息",
                    "policy": "manual_only",
                    "content": "学生姓名：；学号：；年级专业及班级：",
                    "style": "",
                },
            ],
        }
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        [
            "目□□录   (二号黑体，居中)",
            "□□□□□□学生姓名（三号黑体加粗）：×××；学号：×××；年级专业及班级：×××",
        ],
    )

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "noise_gap",
    )
    report = read_json(tmp_path / "noise_gap/artifacts/template_gap_report.json")

    assert result.status == Status.PASS
    assert not collect_checks(report, type="template_generation_element_missing")


def test_template_gap_matches_repeated_text_inside_unit_range(tmp_path) -> None:
    school_id = "range-school"
    expected = [
        {
            "unit_id": "first",
            "name": "第一单元",
            "order": 10,
            "status": "required",
            "elements": [
                {"element_id": "e_001", "name": "第一锚点", "policy": "fixed", "content": "第一单元", "style": ""},
                {"element_id": "e_002", "name": "重复标签", "policy": "fixed", "content": "重复标签", "style": ""},
            ],
        },
        {
            "unit_id": "second",
            "name": "第二单元",
            "order": 20,
            "status": "required",
            "elements": [
                {"element_id": "e_001", "name": "第二锚点", "policy": "fixed", "content": "第二单元", "style": ""},
                {"element_id": "e_002", "name": "重复标签", "policy": "fixed", "content": "重复标签", "style": ""},
            ],
        },
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        ["第一单元", "重复标签", "第二单元", "重复标签"],
    )

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "range_gap",
    )
    report = read_json(tmp_path / "range_gap/artifacts/template_gap_report.json")
    first_repeat = element_by_id(unit_by_id(report, "first"), "e_002")["presence"]
    second_repeat = element_by_id(unit_by_id(report, "second"), "e_002")["presence"]

    assert result.status == Status.PASS
    assert first_repeat["evidence_refs"] == ["word/document.xml:p[2]"]
    assert second_repeat["evidence_refs"] == ["word/document.xml:p[4]"]


def test_template_gap_checks_out_of_order_units_in_their_actual_regions(
    tmp_path,
) -> None:
    school_id = "out-of-order-school"
    expected = [
        {
            "unit_id": "first",
            "name": "第一单元",
            "order": 10,
            "status": "required",
            "elements": [
                {"element_id": "e_001", "name": "第一锚点", "policy": "fixed", "content": "第一单元", "style": ""},
            ],
        },
        {
            "unit_id": "second",
            "name": "第二单元",
            "order": 20,
            "status": "required",
            "elements": [
                {"element_id": "e_001", "name": "第二锚点", "policy": "fixed", "content": "第二单元", "style": ""},
            ],
        },
        {
            "unit_id": "third",
            "name": "第三单元",
            "order": 30,
            "status": "required",
            "elements": [
                {"element_id": "e_001", "name": "第三锚点", "policy": "fixed", "content": "第三单元", "style": ""},
            ],
        },
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(generated_template, ["第一单元", "第三单元", "第二单元"])

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "out_of_order_gap",
    )
    report = read_json(tmp_path / "out_of_order_gap/artifacts/template_gap_report.json")
    third_anchor = element_by_id(unit_by_id(report, "third"), "e_001")["presence"]

    assert result.status == Status.FAIL
    assert unit_by_id(report, "third")["located"]["source_ref"] == "word/document.xml:p[2]"
    assert third_anchor["evidence_refs"] == ["word/document.xml:p[2]"]
    assert any(
        item["type"] == "template_generation_unit_order_mismatch"
        for item in collect_checks(report)
    )


def test_template_gap_matches_form_element_across_multiple_nodes(tmp_path) -> None:
    school_id = "form-school"
    expected = [
        {
            "unit_id": "cover",
            "name": "封面",
            "order": 10,
            "status": "required",
            "elements": [
                {
                    "element_id": "e_001",
                    "name": "学校名称",
                    "policy": "fixed",
                    "content": "测试大学",
                    "style": "",
                },
                {
                    "element_id": "e_002",
                    "name": "学生基本信息",
                    "policy": "manual_only",
                    "content": "学生姓名；学号；年级专业及班级",
                    "style": "",
                },
            ],
        }
    ]
    write_minimal_gap_standard(tmp_path, school_id, expected)
    generated_template = tmp_path / "inputs/targets/demo-school/fixtures/template_gap/generated_template.input.docx"
    write_docx(
        generated_template,
        ["测试大学", "学生姓名：", "学号：", "年级专业及班级："],
    )

    result = run_template_gap_eval(
        tmp_path,
        school_id,
        generated_template,
        tmp_path / "form_gap",
    )
    report = read_json(tmp_path / "form_gap/artifacts/template_gap_report.json")
    form_presence = element_by_id(unit_by_id(report, "cover"), "e_002")["presence"]

    assert result.status == Status.PASS
    assert form_presence["type"] == "template_generation_element_found_across_sources"
    assert form_presence["evidence_refs"] == [
        "word/document.xml:p[2]",
        "word/document.xml:p[3]",
        "word/document.xml:p[4]",
    ]


def test_real_core_template_gap_outputs_tree_and_reports_for_all_schools(
    tmp_path,
) -> None:
    for school in REAL_CORE_SCHOOLS:
        school_id = str(school["school_id"])
        out_dir = tmp_path / school_id

        result = run_template_gap_eval(
            ROOT,
            school_id,
            ROOT / school["generated_template_docx"],
            out_dir,
        )

        artifacts = out_dir / "artifacts"
        generated_template = artifacts / "generated_template.docx"
        tree_path = artifacts / "generated_template_tree.json"
        report_json = artifacts / "template_gap_report.json"
        report_md = artifacts / "template_gap_report.md"
        report_docx = artifacts / "template_gap_report.docx"

        assert result.status in {Status.PASS, Status.FAIL, Status.UNKNOWN}
        assert generated_template.exists()
        assert tree_path.exists()
        assert report_json.exists()
        assert report_md.exists()
        assert report_docx.exists()

        tree = read_json(tree_path)
        report = read_json(report_json)
        assert tree["input_valid_docx"] is True
        assert tree["data"]["paragraphs"] or tree["data"]["tables"]
        assert report["generated_template"]["sha256"] == sha256_file(generated_template)
        assert {
            "known_status",
            "display_status",
            "passed_count",
            "failed_count",
            "unknown_count",
            "blocking_status",
            "per_unit",
        } <= set(report["summary"])
        assert report["generated_template"]["source_path"] == str(
            ROOT / school["generated_template_docx"]
        )
        assert "/fixtures/template_gap/generated_template.input.docx" in report[
            "generated_template"
        ]["source_path"]
        assert "check_items" not in report
        assert collect_checks(report)
        assert any(
            item["check_id"] == "template_generation.output_docx"
            for item in collect_checks(report)
        )


def test_generated_template_gap_bad_docx_does_not_pass(tmp_path) -> None:
    bad_docx = tmp_path / "bad_generated_template.docx"
    doc = Document()
    doc.add_paragraph("这不是湖南农业大学模板。")
    doc.save(bad_docx)

    result = run_template_gap_eval(
        ROOT,
        "hunannongye",
        bad_docx,
        tmp_path / "bad_gap",
    )
    report = read_json(tmp_path / "bad_gap/artifacts/template_gap_report.json")

    assert result.status == Status.FAIL
    assert report["summary"]["blocking_status"] == Status.FAIL.value
    assert report["summary"]["failed_count"] > 0
    assert any(
        item["type"] == "template_generation_unit_missing"
        for item in collect_checks(report)
    )


def test_generated_template_gap_reports_ooxml_style_details(tmp_path) -> None:
    result = run_template_gap_eval(
        ROOT,
        "hunannongye",
        ROOT / "inputs/targets/hunannongye/raw/source_template.docx",
        tmp_path / "style_gap",
    )
    tree = read_json(tmp_path / "style_gap/artifacts/generated_template_tree.json")
    report = read_json(tmp_path / "style_gap/artifacts/template_gap_report.json")

    assert result.status == Status.FAIL
    paragraph = next(
        item
        for item in tree["data"]["paragraphs"]
        if item["text"].startswith("湖 南 农 业 大 学")
    )
    assert paragraph["style_details"]["dominant_run"]["font_names"] == ["华文行楷"]
    assert paragraph["style_details"]["dominant_run"]["font_size_pt"] == 26.0
    assert paragraph["style_details"]["paragraph"]["alignment"] == "center"

    style_mismatches = [
        item
        for item in collect_checks(report)
        if item["type"] == "template_generation_style_mismatch"
    ]
    assert style_mismatches
    assert any("font_size=" in item["actual"] for item in style_mismatches)


def test_generated_template_gap_resolves_ooxml_style_inheritance(tmp_path) -> None:
    result = run_template_gap_eval(
        ROOT,
        "nannong-undergraduate",
        ROOT / "inputs/targets/nannong-undergraduate/raw/source_template.docx",
        tmp_path / "style_inheritance_gap",
    )
    tree = read_json(
        tmp_path / "style_inheritance_gap/artifacts/generated_template_tree.json"
    )
    report = read_json(
        tmp_path / "style_inheritance_gap/artifacts/template_gap_report.json"
    )

    assert result.status == Status.FAIL
    signature_line = next(
        item
        for item in tree["data"]["paragraphs"]
        if item["text"].startswith("论文作者签名：")
    )
    assert signature_line["xml_index"] == 29
    assert signature_line["style_details"]["paragraph"]["spacing"][
        "line_spacing"
    ] == "exact:20pt"
    assert signature_line["style_details"]["style_inheritance"]["source_refs"] == [
        "word/styles.xml:style[a]"
    ]
    assert not any(
        item["type"] == "template_generation_style_mismatch"
        and "line_spacing expected exact:20pt, actual exact:35pt" in item["actual"]
        for item in collect_checks(report)
    )


def test_generated_template_gap_binds_page_rules_to_ooxml_sources(tmp_path) -> None:
    result = run_template_gap_eval(
        ROOT,
        "hunannongye",
        ROOT / "inputs/targets/hunannongye/raw/source_template.docx",
        tmp_path / "page_gap",
    )
    report = read_json(tmp_path / "page_gap/artifacts/template_gap_report.json")

    assert result.status == Status.FAIL
    page_items = [
        item
        for item in collect_checks(report)
        if item["category"] == "page_rule"
    ]
    assert any(
        item["type"] == "template_generation_page_rule_match"
        and item["affected_ids"] == ["cover.page.page_break"]
        and "word/document.xml:p[" in item["evidence_refs"][0]
        for item in page_items
    )
    assert any(
        item["type"] == "template_generation_page_rule_mismatch"
        and item["affected_ids"] == ["integrity_statement.page.page_break"]
        and "no explicit page/section break" in item["actual"]
        for item in page_items
    )


def test_generated_template_gap_binds_keep_together_to_table_sources(tmp_path) -> None:
    result = run_template_gap_eval(
        ROOT,
        "nannong-undergraduate",
        ROOT / "inputs/targets/nannong-undergraduate/raw/source_template.docx",
        tmp_path / "table_keep_gap",
    )
    tree = read_json(tmp_path / "table_keep_gap/artifacts/generated_template_tree.json")
    report = read_json(tmp_path / "table_keep_gap/artifacts/template_gap_report.json")

    assert result.status == Status.FAIL
    cover_table = tree["data"]["tables"][0]
    assert cover_table["first_paragraph_index"] == 1
    assert cover_table["last_paragraph_index"] == 21
    assert cover_table["cant_split_row_refs"] == [
        "word/document.xml:tbl[1]/tr[1]/cantSplit"
    ]
    assert cover_table["cells"][0]["first_paragraph_index"] == 1

    assert any(
        item["type"] == "template_generation_element_found"
        and item["affected_ids"] == ["cover.e_001"]
        and item["evidence_refs"][0] == "word/document.xml:tbl[1]/tr[1]/tc[1]"
        for item in collect_checks(report)
    )
    assert any(
        item["type"] == "template_generation_page_rule_match"
        and item["affected_ids"] == ["cover.page.keep_together"]
        and item["evidence_refs"] == ["word/document.xml:tbl[1]/tr[1]/cantSplit"]
        for item in collect_checks(report)
    )


def test_generated_template_inspector_keeps_merged_cell_copies_in_row_order(
    tmp_path,
) -> None:
    path = tmp_path / "merged-table.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "合并标题"
    table.cell(1, 0).text = "下一行"
    doc.save(path)

    tree = inspect_generated_template_docx(path)
    entries = [
        entry
        for entry in iter_visible_text_entries(tree)
        if entry["kind"] == "table_cell" and entry["text"] == "合并标题"
    ]

    assert len(entries) == 3
    assert {entry["order"] for entry in entries} == {1}
    assert all(entry["order"] < 10_000 for entry in entries)
    merged_cells = [
        cell
        for cell in tree["data"]["tables"][0]["cells"]
        if cell["text"] == "合并标题"
    ]
    assert all(cell["first_paragraph_index"] == 1 for cell in merged_cells)
    assert all(
        cell["cell_paragraph_refs"] == ["word/document.xml:p[1]"]
        for cell in merged_cells
    )
    assert all(cell["row_first_paragraph_index"] == 1 for cell in merged_cells)


def test_generated_template_inspector_models_footnotes_in_generated_tree() -> None:
    tree = inspect_generated_template_docx(
        ROOT / "inputs/targets/pku-graduate/raw/source_template.docx"
    )

    assert tree["data"]["footnotes"]
    assert all(
        item.get("object_type") != "footnote"
        for item in tree["data"]["unknown_visible_objects"]
    )


def test_generated_template_inspector_models_text_boxes_in_generated_tree() -> None:
    tree = inspect_generated_template_docx(
        ROOT / "inputs/targets/nannong-undergraduate/raw/source_template.docx"
    )

    assert any("规范化要求" in item["text"] for item in tree["data"]["text_boxes"])
    assert all(
        item.get("object_type") != "text_box"
        for item in tree["data"]["unknown_visible_objects"]
    )


def test_generated_template_inspector_models_images_in_generated_tree() -> None:
    tree = inspect_generated_template_docx(
        ROOT / "inputs/targets/nannong-undergraduate/raw/source_template.docx"
    )

    assert {
        item["target"] for item in tree["data"]["images"]
    } >= {"word/media/image1.jpeg", "word/media/image2.png"}
    assert all(item.get("sha256") for item in tree["data"]["images"])
    assert all(
        item.get("object_type") != "image"
        for item in tree["data"]["unknown_visible_objects"]
    )


def test_generated_template_gap_binds_header_footer_rules_to_sections(tmp_path) -> None:
    result = run_template_gap_eval(
        ROOT,
        "nannong-undergraduate",
        ROOT / "inputs/targets/nannong-undergraduate/raw/source_template.docx",
        tmp_path / "header_footer_gap",
    )
    tree = read_json(tmp_path / "header_footer_gap/artifacts/generated_template_tree.json")
    report = read_json(
        tmp_path / "header_footer_gap/artifacts/template_gap_report.json"
    )

    assert result.status == Status.FAIL
    assert tree["data"]["sections"]
    toc_section = next(
        item for item in tree["data"]["sections"] if item["paragraph_index"] == 45
    )
    assert {
        (item["kind"], item["type"], item["part_name"])
        for item in toc_section["effective_references"]
    } >= {
        ("header", "default", "word/header2.xml"),
        ("footer", "default", "word/footer4.xml"),
    }
    assert toc_section["page_numbering"] == {"format": "upperRoman", "start": 1}

    header_footer_items = [
        item
        for item in collect_checks(report)
        if item["category"] == "header_footer"
    ]
    assert any(
        item["type"] == "template_generation_header_footer_match"
        and item["affected_ids"] == ["toc.header_footer.header"]
        and "word/document.xml:p[74]/sectPr" in item["evidence_refs"][0]
        for item in header_footer_items
    )
    toc_page_number = next(
        item
        for item in header_footer_items
        if item["affected_ids"] == ["toc.header_footer.page_number"]
    )
    assert toc_page_number["type"] == "template_generation_page_number_rule_match"
    assert "page_numbering={'format': 'lowerRoman', 'start': 1}" in toc_page_number["actual"]
    assert any("sectPr" in ref for ref in toc_page_number["evidence_refs"])


def test_generated_template_gap_binds_word_fields_to_units(tmp_path) -> None:
    nannong_result = run_template_gap_eval(
        ROOT,
        "nannong-undergraduate",
        ROOT / "inputs/targets/nannong-undergraduate/raw/source_template.docx",
        tmp_path / "nannong_field_gap",
    )
    nannong_tree = read_json(
        tmp_path / "nannong_field_gap/artifacts/generated_template_tree.json"
    )
    nannong_report = read_json(
        tmp_path / "nannong_field_gap/artifacts/template_gap_report.json"
    )

    assert nannong_result.status == Status.FAIL
    nannong_toc_field = next(
        item
        for item in nannong_tree["data"]["fields"]
        if item["field_type"] == "TOC"
    )
    assert nannong_toc_field["kind"] == "complexField"
    assert nannong_toc_field["instruction"] == 'TOC \\o "1-3" \\h \\z \\u'
    assert nannong_toc_field["paragraph_index"] == 49
    assert nannong_toc_field["end_paragraph_index"] == 74
    assert any(
        item["type"] == "template_generation_field_match"
        and item["affected_ids"] == ["toc.e_002.field"]
        and item["evidence_refs"] == [
            "word/document.xml:p[49]/field[1]",
            "word/document.xml:p[74]",
        ]
        for item in collect_checks(nannong_report)
    )

    pku_result = run_template_gap_eval(
        ROOT,
        "pku-graduate",
        ROOT / "inputs/targets/pku-graduate/raw/source_template.docx",
        tmp_path / "pku_field_gap",
    )
    pku_tree = read_json(
        tmp_path / "pku_field_gap/artifacts/generated_template_tree.json"
    )
    pku_report = read_json(tmp_path / "pku_field_gap/artifacts/template_gap_report.json")
    pku_field_items = [
        item
        for item in collect_checks(pku_report)
        if item["category"] == "field"
    ]
    pku_seq_instructions = {
        item["instruction"]
        for item in pku_tree["data"]["fields"]
        if item["field_type"] == "SEQ"
    }

    assert pku_result.status == Status.FAIL
    assert {
        "SEQ 图 \\* ARABIC \\S 1",
        "SEQ 表 \\* ARABIC \\s 1",
        "SEQ 公式 \\* ARABIC \\s 1",
    } <= pku_seq_instructions
    assert unit_by_id(pku_report, "figure_list")["located"]["source_ref"] == (
        "word/document.xml:p[104]"
    )
    assert unit_by_id(pku_report, "table_list")["located"]["source_ref"] == (
        "word/document.xml:p[121]"
    )
    assert unit_by_id(pku_report, "body_main")["located"]["source_ref"] == (
        "word/document.xml:p[127]"
    )
    assert unit_by_id(pku_report, "references")["located"]["source_ref"] == (
        "word/document.xml:p[406]"
    )
    assert unit_by_id(pku_report, "acknowledgement")["located"]["source_ref"] == (
        "word/document.xml:p[428]"
    )
    assert any(
        item["type"] == "template_generation_field_unverified"
        and item["affected_ids"] == ["toc.field"]
        and item["status"] == Status.UNKNOWN.value
        for item in pku_field_items
    )
    assert any(
        item["type"] == "template_generation_field_match"
        and item["affected_ids"] == ["figure_list.field"]
        and item["actual"] == 'TOC \\h \\z \\t "PKU图题" \\c'
        for item in pku_field_items
    )
    assert any(
        item["type"] == "template_generation_field_match"
        and item["affected_ids"] == ["table_list.field"]
        and item["actual"] == 'TOC \\h \\z \\t "PKU表题" \\c'
        for item in pku_field_items
    )
    assert any(
        item["type"] == "template_generation_field_match"
        and item["affected_ids"] == ["body_main.e_014.field"]
        and "SEQ 图" in item["actual"]
        for item in pku_field_items
    )
    assert any(
        item["type"] == "template_generation_field_match"
        and item["affected_ids"] == ["body_main.e_017.field"]
        and "SEQ 表" in item["actual"]
        for item in pku_field_items
    )
    assert any(
        item["type"] == "template_generation_field_match"
        and item["affected_ids"] == ["body_main.e_024.field"]
        and "SEQ 公式" in item["actual"]
        for item in pku_field_items
    )


def test_generated_template_gap_binds_numbering_rules_to_units(tmp_path) -> None:
    result = run_template_gap_eval(
        ROOT,
        "pku-graduate",
        ROOT / "inputs/targets/pku-graduate/raw/source_template.docx",
        tmp_path / "pku_numbering_gap",
    )
    tree = read_json(
        tmp_path / "pku_numbering_gap/artifacts/generated_template_tree.json"
    )
    report = read_json(
        tmp_path / "pku_numbering_gap/artifacts/template_gap_report.json"
    )

    assert result.status == Status.FAIL
    heading_definition = next(
        item
        for item in tree["data"]["numbering_definitions"]
        if item["paragraph_style_id"] == "1"
        and item["lvl_text"] == "第%1章"
    )
    assert heading_definition["num_fmt"] == "chineseCountingThousand"

    heading_ref = next(
        item
        for item in tree["data"]["numbering_refs"]
        if item["paragraph_index"] == 127
    )
    assert heading_ref["source_kind"] == "paragraph_style"
    assert heading_ref["num_id"] == "1"
    assert heading_ref["ilvl"] == "0"
    assert heading_ref["lvl_text"] == "第%1章"
    assert heading_ref["source_refs"] == [
        "word/styles.xml:style[1]/numPr",
        "word/numbering.xml:abstractNum[0]/lvl[0]",
    ]

    numbering_items = [
        item
        for item in collect_checks(report)
        if item["category"] == "numbering"
    ]
    assert any(
        item["type"] == "template_generation_numbering_match"
        and item["affected_ids"] == ["body_main.e_001.numbering"]
        and "text=第%1章" in item["actual"]
        and "word/document.xml:p[175]/pStyle" in item["evidence_refs"]
        for item in numbering_items
    )
    assert any(
        item["type"] == "template_generation_numbering_match"
        and item["affected_ids"] == ["body_main.e_004.numbering"]
        and "%1.%2.%3.%4.%5" in item["actual"]
        and "word/document.xml:p[165]/numPr" in item["evidence_refs"]
        for item in numbering_items
    )


def test_generated_template_gap_missing_docx_is_unknown(tmp_path) -> None:
    result = run_template_gap_eval(
        ROOT,
        "hunannongye",
        tmp_path / "missing_generated_template.docx",
        tmp_path / "missing_gap",
    )
    report = read_json(tmp_path / "missing_gap/artifacts/template_gap_report.json")

    assert result.status == Status.UNKNOWN
    assert report["summary"]["blocking_status"] == Status.UNKNOWN.value
    assert report["summary"]["display_status"] == Status.UNKNOWN.value
    assert any(
        item["type"] == "template_generation_output_missing"
        for item in collect_checks(report)
    )


def test_template_gap_status_combination_rules() -> None:
    def summary_for(*statuses: str) -> dict:
        report = {
            "input": {
                "status": statuses[0],
                "check_id": "test",
                "type": "test",
                "category": "test",
            },
            "units": [],
            "global_checks": [
                {
                    "status": status,
                    "check_id": "test",
                    "type": "test",
                    "category": "test",
                }
                for status in statuses[1:]
            ],
            "unmodeled_objects": [],
        }
        return summarize_template_gap_report(report)

    assert summary_for("PASS") == {
        "known_status": "PASS",
        "display_status": "PASS",
        "passed_count": 1,
        "failed_count": 0,
        "unknown_count": 0,
        "blocking_status": "PASS",
        "per_unit": [],
    }
    assert summary_for("PASS", "UNKNOWN")["display_status"] == "PASS + UNKNOWN"
    assert summary_for("FAIL")["blocking_status"] == "FAIL"
    assert summary_for("FAIL", "UNKNOWN")["display_status"] == "FAIL + UNKNOWN"
    assert summary_for("UNKNOWN") == {
        "known_status": "UNKNOWN",
        "display_status": "UNKNOWN",
        "passed_count": 0,
        "failed_count": 0,
        "unknown_count": 1,
        "blocking_status": "UNKNOWN",
        "per_unit": [],
    }
