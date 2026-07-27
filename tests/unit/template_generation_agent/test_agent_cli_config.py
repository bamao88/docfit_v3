from __future__ import annotations

from docx import Document
from typer.testing import CliRunner

from docfit.cli import main as cli_main
from docfit.core.models import StageResult
from docfit.core.status import Status


def test_cli_live_without_packet_requests_auto_render_packet(monkeypatch, tmp_path) -> None:
    source = tmp_path / "template.docx"
    Document().save(source)
    captured = {}

    def fake_run_template_generate_eval(
        _root,
        _template,
        _out,
        *,
        agent_config=None,
    ):
        captured["agent_config"] = agent_config
        return StageResult("template_generate", Status.PASS)

    monkeypatch.setattr(
        cli_main,
        "run_template_generate_eval",
        fake_run_template_generate_eval,
    )

    result = CliRunner().invoke(
        cli_main.app,
        [
            "eval",
            "template-generate",
            "--template",
            str(source),
            "--out",
            str(tmp_path / "out"),
            "--agent-observe-live",
        ],
    )

    assert result.exit_code == 0, result.output
    agent_config = captured["agent_config"]
    assert agent_config.observation_mode == "live"
    assert agent_config.render_packet_path is None


def test_cli_llm_is_simple_live_observation_opt_in(monkeypatch, tmp_path) -> None:
    source = tmp_path / "template.docx"
    Document().save(source)
    captured = {}

    def fake_run_template_generate_eval(
        _root,
        _template,
        _out,
        *,
        agent_config=None,
    ):
        captured["agent_config"] = agent_config
        return StageResult("template_generate", Status.PASS)

    monkeypatch.setattr(
        cli_main,
        "run_template_generate_eval",
        fake_run_template_generate_eval,
    )

    result = CliRunner().invoke(
        cli_main.app,
        [
            "eval",
            "template-generate",
            "--template",
            str(source),
            "--out",
            str(tmp_path / "out"),
            "--llm",
        ],
    )

    assert result.exit_code == 0, result.output
    agent_config = captured["agent_config"]
    assert agent_config.enabled is True
    assert agent_config.observation_mode == "live"
    assert agent_config.observation_t3_concurrency == 4


def test_template_observe_cli_always_routes_to_live_stage_runner(monkeypatch, tmp_path) -> None:
    source = tmp_path / "template.docx"
    Document().save(source)
    captured = {}

    def fake_run_live_template_observation_stage(**kwargs):
        captured.update(kwargs)
        return {"stage": "T3", "llm_mode": "live_api"}

    monkeypatch.setattr(
        cli_main,
        "run_live_template_observation_stage",
        fake_run_live_template_observation_stage,
    )

    result = CliRunner().invoke(
        cli_main.app,
        [
            "eval",
            "template-observe",
            "--stage",
            "t3",
            "--template",
            str(source),
            "--out",
            str(tmp_path / "observe"),
        ],
    )

    assert result.exit_code == 0, result.output
    assert captured["stage"] == "t3"
    assert captured["t2_observation_path"] is None
    assert "llm_mode = live_api" in result.output


def test_canonical_template_generate_defaults_to_minimax_live(monkeypatch, tmp_path) -> None:
    source = tmp_path / "template.docx"
    Document().save(source)
    captured = {}

    def fake_run(
        _root,
        _template,
        _out,
        *,
        agent_config=None,
        run_context=None,
    ):
        captured["agent_config"] = agent_config
        captured["run_context"] = run_context
        return StageResult("template_generate", Status.PASS)

    monkeypatch.setattr(cli_main, "run_template_generate_eval", fake_run)
    monkeypatch.setenv("DOCFIT_TEMPLATE_AGENT_ENABLED", "1")
    monkeypatch.setenv("DOCFIT_TEMPLATE_AGENT_OBSERVATION_MODE", "live")

    result = CliRunner().invoke(
        cli_main.app,
        [
            "template",
            "generate",
            "--template",
            str(source),
            "--out",
            str(tmp_path / "out"),
        ],
    )

    assert result.exit_code == 0, result.output
    assert captured["agent_config"].enabled is True
    assert captured["agent_config"].observation_mode == "live"
    assert captured["agent_config"].text_provider == "minimax"
    assert captured["run_context"]["entrypoint"] == "cli.template.generate"


def test_canonical_template_generate_unifies_replay_mode(monkeypatch, tmp_path) -> None:
    source = tmp_path / "template.docx"
    replay = tmp_path / "replay.json"
    Document().save(source)
    replay.write_text("{}", encoding="utf-8")
    captured = {}

    def fake_run(
        _root,
        _template,
        _out,
        *,
        agent_config=None,
        run_context=None,
    ):
        captured["agent_config"] = agent_config
        return StageResult("template_generate", Status.PASS)

    monkeypatch.setattr(cli_main, "run_template_generate_eval", fake_run)
    result = CliRunner().invoke(
        cli_main.app,
        [
            "template",
            "generate",
            "--template",
            str(source),
            "--out",
            str(tmp_path / "out"),
            "--ai",
            "replay",
            "--replay",
            str(replay),
        ],
    )

    assert result.exit_code == 0, result.output
    config = captured["agent_config"]
    assert config.enabled is True
    assert config.observation_mode == "replay"
    assert config.observation_transcript_path == replay


def test_canonical_t3_stage_requires_pinned_upstream_by_default(
    monkeypatch,
    tmp_path,
) -> None:
    source = tmp_path / "template.docx"
    Document().save(source)
    captured = {}

    def fake_stage(**kwargs):
        captured.update(kwargs)
        return {"stage": "T3", "ai_mode": "live"}

    monkeypatch.setattr(cli_main, "run_template_observation_stage", fake_stage)
    result = CliRunner().invoke(
        cli_main.app,
        [
            "template",
            "stage",
            "t3",
            "--template",
            str(source),
            "--out",
            str(tmp_path / "stage"),
            "--with-upstream",
        ],
    )

    assert result.exit_code == 0, result.output
    assert captured["ai_mode"] == "live"
    assert captured["with_upstream"] is True
    assert captured["entrypoint"] == "cli.template.stage.t3"


def test_canonical_t3_stage_forwards_signed_gold_standard(
    monkeypatch,
    tmp_path,
) -> None:
    source = tmp_path / "template.docx"
    Document().save(source)
    standard = tmp_path / "t2.standard.yaml"
    standard.write_text("stage_id: T2\nstandard_state: signed_active\n", encoding="utf-8")
    t3_standard = tmp_path / "t3.standard.yaml"
    t3_standard.write_text(
        "stage_id: T3\nstandard_state: signed_active\n",
        encoding="utf-8",
    )
    captured = {}

    def fake_stage(**kwargs):
        captured.update(kwargs)
        return {"stage": "T3", "ai_mode": "live"}

    monkeypatch.setattr(cli_main, "run_template_observation_stage", fake_stage)
    result = CliRunner().invoke(
        cli_main.app,
        [
            "template",
            "stage",
            "t3",
            "--template",
            str(source),
            "--out",
            str(tmp_path / "stage"),
            "--t2-gold-standard",
            str(standard),
            "--t3-gold-standard",
            str(t3_standard),
        ],
    )

    assert result.exit_code == 0, result.output
    assert captured["t2_gold_standard_path"] == standard
    assert captured["t3_gold_standard_path"] == t3_standard
    assert captured["with_upstream"] is False


def test_template_inspect_is_read_only(monkeypatch, tmp_path) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "run_manifest.json").write_text(
        '{"status":"PASS","ai_mode":"off","api_call_count":0}',
        encoding="utf-8",
    )
    before = sorted(path.name for path in run_dir.iterdir())

    result = CliRunner().invoke(
        cli_main.app,
        ["template", "inspect", "--run", str(run_dir)],
    )

    assert result.exit_code == 0, result.output
    assert "ai_mode = off" in result.output
    assert sorted(path.name for path in run_dir.iterdir()) == before
