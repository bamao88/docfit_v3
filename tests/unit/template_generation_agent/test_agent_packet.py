from __future__ import annotations

import shutil

import pytest
from docx import Document

from docfit.template_generation.agent.packet import (
    _find_entry_page_match,
    _infer_ambiguous_source_bindings,
    _infer_stable_anchor_bindings,
    _page_text_index,
    build_template_agent_render_packet,
)
from docfit.template_generation.source_tree import inspect_document_facts_docx

from .helpers import packet


def test_packet_exposes_projection_fallback_status() -> None:
    render_packet = packet()

    assert render_packet["artifact_type"] == "template_agent_render_packet"
    assert render_packet["status_authority"] == "verify_template_parse_build"
    assert render_packet["advisory_only"] is True
    assert render_packet["render_status"] == "projection_fallback"
    assert render_packet["render_artifacts"]["render_status"] == "projection_fallback"
    assert render_packet["render_artifacts"]["clean_page_images"] == []
    assert render_packet["page_text_index"][0]["source_seq"] == 1
    assert render_packet["page_text_index"][0]["page_no"] == 1
    assert "bbox" in render_packet["page_text_index"][0]
    assert render_packet["page_layout_index"][0]["render_target_id"] == "source_seq:1"
    assert render_packet["source_render_hash"].startswith("sha256:")


def test_packet_can_build_real_render_images_and_page_bindings(tmp_path) -> None:
    if not all(shutil.which(name) for name in ("soffice", "pdftoppm", "pdftotext")):
        pytest.skip("real render packet requires soffice, pdftoppm, and pdftotext")
    source = tmp_path / "two-page-template.docx"
    doc = Document()
    doc.add_paragraph("第一页标题")
    doc.add_paragraph("第一页正文")
    doc.add_page_break()
    doc.add_paragraph("第二页标题")
    doc.add_paragraph("第二页正文")
    doc.save(source)
    facts = inspect_document_facts_docx(source)

    render_packet = build_template_agent_render_packet(
        document_facts=facts,
        source_template_docx=source,
        render_artifacts_dir=tmp_path / "agent_render",
    )

    assert render_packet["render_status"] == "real_render"
    assert render_packet["render_artifacts"]["clean_page_images"]
    assert render_packet["render_artifacts"]["annotated_page_images"]
    pages = {item["page_no"] for item in render_packet["page_text_index"]}
    assert 2 in pages
    assert any(item.get("bbox") for item in render_packet["page_text_index"])
    assert any(
        item.get("page_top_ratio") is not None
        for item in render_packet["page_layout_index"]
    )


def test_packet_infers_ambiguous_merged_cell_from_same_table_row() -> None:
    entries = [
        {
            "source_seq": 10,
            "table_id": "tbl_0001",
            "cell_id": "tbl_0001.r_003.c_001",
        },
        {
            "source_seq": 11,
            "table_id": "tbl_0001",
            "cell_id": "tbl_0001.r_003.c_002",
        },
    ]
    bindings = {
        10: {"binding_status": "exact", "page_no": 7},
        11: {"binding_status": "ambiguous"},
    }

    inferred = _infer_ambiguous_source_bindings(entries, bindings)

    assert inferred[11] == {
        "binding_status": "inferred_table_row",
        "page_no": 7,
    }


def test_packet_does_not_infer_across_different_neighbor_pages() -> None:
    entries = [
        {"source_seq": 1},
        {"source_seq": 2},
        {"source_seq": 3},
    ]
    bindings = {
        1: {"binding_status": "exact", "page_no": 1},
        2: {"binding_status": "ambiguous"},
        3: {"binding_status": "exact", "page_no": 2},
    }

    inferred = _infer_ambiguous_source_bindings(entries, bindings)

    assert inferred[2] == {"binding_status": "ambiguous"}


def test_packet_infers_ambiguous_entry_before_document_page_boundary() -> None:
    entries = [
        {"source_seq": 1, "source_ref": "word/document.xml:p[40]"},
        {"source_seq": 2, "source_ref": "word/document.xml:p[41]"},
        {"source_seq": 3, "source_ref": "word/document.xml:p[46]"},
    ]
    bindings = {
        1: {"binding_status": "exact", "page_no": 2},
        2: {"binding_status": "ambiguous"},
        3: {"binding_status": "exact", "page_no": 3},
    }

    inferred = _infer_ambiguous_source_bindings(
        entries,
        bindings,
        boundary_paragraph_indexes={45},
    )

    assert inferred[2] == {
        "binding_status": "inferred_document_segment",
        "page_no": 2,
    }


def test_packet_infers_repeated_heading_from_unique_document_segment_page() -> None:
    entries = [
        {"source_seq": 31, "source_ref": "word/document.xml:p[104]"},
        {"source_seq": 32, "source_ref": "word/document.xml:p[105]"},
    ]
    bindings = {
        31: {"binding_status": "ambiguous"},
        32: {"binding_status": "inferred_stable_anchors", "page_no": 8},
    }

    inferred = _infer_ambiguous_source_bindings(
        entries,
        bindings,
        boundary_paragraph_indexes={103, 120},
    )

    assert inferred[31] == {
        "binding_status": "inferred_document_segment",
        "page_no": 8,
    }


def test_packet_infers_keep_next_heading_from_following_paragraph() -> None:
    entries = [
        {
            "source_seq": 99,
            "source_ref": "word/document.xml:p[184]",
            "style_details": {"paragraph": {"keep_next": True}},
        },
        {"source_seq": 100, "source_ref": "word/document.xml:p[185]"},
    ]
    bindings = {
        99: {"binding_status": "ambiguous"},
        100: {"binding_status": "inferred_stable_anchors", "page_no": 15},
    }

    inferred = _infer_ambiguous_source_bindings(
        entries,
        bindings,
        boundary_paragraph_indexes={174, 312},
    )

    assert inferred[99] == {
        "binding_status": "inferred_keep_next",
        "page_no": 15,
    }


def test_packet_text_binding_respects_minimum_page() -> None:
    pdf_layout = {
        "pages": [
            {"page_no": 1, "text": "标题", "words": []},
            {
                "page_no": 2,
                "text": "标题",
                "words": [
                    {
                        "page_start": 0,
                        "page_end": 2,
                        "bbox": {
                            "x_min": 0,
                            "y_min": 0,
                            "x_max": 10,
                            "y_max": 10,
                            "page_width": 100,
                            "page_height": 100,
                        },
                    }
                ],
            },
        ],
    }

    match = _find_entry_page_match(pdf_layout, "标题", minimum_page=2)

    assert match is not None
    assert match["page_no"] == 2


def test_packet_text_binding_abstains_when_exact_text_occurs_on_multiple_pages() -> None:
    pdf_layout = {
        "pages": [
            {
                "page_no": page_no,
                "text": "图目录",
                "words": [
                    {
                        "page_start": 0,
                        "page_end": 3,
                        "bbox": {
                            "x_min": 0,
                            "y_min": 0,
                            "x_max": 10,
                            "y_max": 10,
                            "page_width": 100,
                            "page_height": 100,
                        },
                    }
                ],
            }
            for page_no in (6, 8)
        ],
    }

    match = _find_entry_page_match(
        pdf_layout,
        "图目录",
        minimum_page=1,
        abstain_on_multiple_matches=True,
    )

    assert match == {"binding_status": "ambiguous"}


def test_packet_text_binding_uses_stable_anchors_for_updated_word_fields() -> None:
    rendered = "图2.9手动添加跨页的续表"
    pdf_layout = {
        "pages": [
            {
                "page_no": 16,
                "text": rendered,
                "words": [
                    {
                        "page_start": 0,
                        "page_end": len(rendered),
                        "bbox": {
                            "x_min": 0,
                            "y_min": 0,
                            "x_max": 10,
                            "y_max": 10,
                            "page_width": 100,
                            "page_height": 100,
                        },
                    }
                ],
            }
        ],
    }

    match = _find_entry_page_match(
        pdf_layout,
        "图2.4手动添加跨页的续表",
        minimum_page=15,
        allow_stable_anchors=True,
    )

    assert match is not None
    assert match["binding_status"] == "inferred_stable_anchors"
    assert match["page_no"] == 16


def test_stable_anchor_binding_cannot_jump_past_next_exact_page() -> None:
    entries = [
        {"source_seq": 1, "text": "第一页"},
        {"source_seq": 2, "text": "年级专业及班级"},
        {"source_seq": 3, "text": "第二页"},
    ]
    bindings = {
        1: {"binding_status": "exact", "page_no": 1},
        2: {"binding_status": "ambiguous"},
        3: {"binding_status": "exact", "page_no": 2},
    }
    pdf_layout = {
        "pages": [
            {"page_no": 1, "text": "第一页", "words": []},
            {"page_no": 2, "text": "第二页", "words": []},
            {
                "page_no": 6,
                "text": "年级专业及班级",
                "words": [
                    {
                        "page_start": 0,
                        "page_end": 7,
                        "bbox": {
                            "x_min": 0,
                            "y_min": 0,
                            "x_max": 10,
                            "y_max": 10,
                            "page_width": 100,
                            "page_height": 100,
                        },
                    }
                ],
            },
        ]
    }

    inferred = _infer_stable_anchor_bindings(entries, bindings, pdf_layout)

    assert inferred[2] == {"binding_status": "ambiguous"}


def test_packet_page_index_excludes_repeating_header_footer_parts() -> None:
    facts = {
        "body_flow": [
            {
                "source_seq": 1,
                "source_ref": "word/document.xml:p[1]",
                "structure_layer": "body_flow",
                "text": "正文",
            },
            {
                "source_seq": 2,
                "source_ref": "word/header1.xml",
                "structure_layer": "header_footer",
                "text": "动态页眉",
            },
        ]
    }

    rows = _page_text_index(
        facts,
        render={
            "source_bindings": {
                1: {"binding_status": "exact", "page_no": 1},
                2: {"binding_status": "ambiguous"},
            }
        },
    )

    assert [row["source_seq"] for row in rows] == [1]
