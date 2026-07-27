from __future__ import annotations

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from docfit.core.io import sha256_file
from docfit.core.status import Status
from docfit.template_generation.artifacts import (
    build_element_spec,
    build_template_spec,
)
from docfit.template_generation.docx_effects import inspect_docx_layout_effects
from docfit.template_generation.executor import execute_template_generation_plan
from docfit.template_generation.final_results import (
    NOT_AVAILABLE,
    publish_final_stage_result,
)
from docfit.template_generation.input_contract import build_l1_input_contract
from docfit.template_generation.manifest import build_template_generation_manifest
from docfit.template_generation.plan import build_template_generation_plan
from docfit.template_generation.verifier import (
    _verify_t4_global_spec,
    _verify_t5_template_spec,
    _verify_t6_build,
)
from docfit.template_generation.refs import _paragraph_map_by_ooxml_index
from docfit.template_generation.word_ops import _insert_page_break_before


def _template_final(template_spec: dict):
    return publish_final_stage_result(
        {**template_spec, "artifact_type": "template_spec"},
        stage_id="T5",
        artifact_type="template_spec",
        artifact_name="05_template_spec.yaml",
        producer_mode="fixture",
    )


def _build_plan(request: dict, template_spec: dict) -> dict:
    return build_template_generation_plan(
        request,
        template_final=_template_final(template_spec),
    )


def _build_template_spec(
    l1_input_contract: dict,
    unit_map: dict,
    element_spec: dict,
    global_spec: dict,
) -> dict:
    l1_final = publish_final_stage_result(
        l1_input_contract,
        stage_id="L1",
        artifact_type="template_generation_l1_input_contract",
        artifact_name="01.5_l1_input_contract.json",
        producer_mode="fixture",
    )
    input_refs = {"l1": l1_final.input_ref()}
    results = (
        publish_final_stage_result(
            unit_map,
            stage_id="T2",
            artifact_type="unit_map",
            artifact_name="02_unit_map.yaml",
            producer_mode="fixture",
            input_refs=input_refs,
        ),
        publish_final_stage_result(
            element_spec,
            stage_id="T3",
            artifact_type="element_spec",
            artifact_name="03_element_spec.yaml",
            producer_mode="fixture",
            input_refs=input_refs,
        ),
        publish_final_stage_result(
            global_spec,
            stage_id="T4",
            artifact_type="global_spec",
            artifact_name="04_global_spec.yaml",
            producer_mode="fixture",
            input_refs=input_refs,
        ),
    )
    return build_template_spec(l1_final, *results).payload


def _ai_global_spec_with_two_sections() -> dict:
    page_field = {
        "source_ref": "word/footer1.xml:p[1]/field[1]",
        "field_type": "PAGE",
    }
    return {
        "artifact_type": "global_spec",
        "artifact_version": "2.0",
        "section_profiles": [
            {
                "section_profile_id": "section_001",
                "source_ref": "word/document.xml:p[2]/pPr/sectPr",
                "boundary": {
                    "status": "detected",
                    "start_source_seq": 1,
                    "end_source_seq": 2,
                },
                "source_seq_refs": [1, 2],
                "page_setup": {},
                "header_footer": {
                    "references": [],
                    "effective_references": [],
                    "parts": [],
                },
                "page_numbering": {
                    "declared": {"status": "detected", "format": "upperRoman"},
                    "fields": [page_field],
                    "display": {
                        "status": "detected",
                        "has_page_field": True,
                        "checked_scopes": {
                            "body_source_seq_range": {"start": 1, "end": 2},
                        },
                    },
                    "flags": [],
                },
                "flags": [],
                "origin": "ai_layout_observation",
            },
            {
                "section_profile_id": "section_002",
                "source_ref": "word/document.xml:body/sectPr",
                "boundary": {
                    "status": "detected",
                    "start_source_seq": 3,
                    "end_source_seq": 4,
                },
                "source_seq_refs": [3, 4],
                "page_setup": {},
                "header_footer": {
                    "references": [],
                    "effective_references": [],
                    "parts": [],
                },
                "page_numbering": {
                    "declared": {"status": "none"},
                    "fields": [],
                    "display": {
                        "status": "no_page_field",
                        "has_page_field": False,
                        "checked_scopes": {
                            "body_source_seq_range": {"start": 3, "end": 4},
                        },
                    },
                    "flags": [],
                },
                "flags": [],
                "origin": "ai_layout_observation",
            },
        ],
        "default_font": None,
        "page_numbering": {"status": "mixed"},
        "header_footer": [],
        "numbering_rules": {"definitions": [], "refs": []},
        "flags": [],
    }


def _document_facts_with_two_sections() -> dict:
    return {
        "artifact_type": "document_facts",
        "body_flow": [
            {
                "source_seq": 1,
                "source_ref": "word/document.xml:p[1]",
                "part_name": "word/document.xml",
                "paragraph_id": "p_0001",
            },
            {
                "source_seq": 2,
                "source_ref": "word/document.xml:p[2]",
                "part_name": "word/document.xml",
                "paragraph_id": "p_0002",
            },
            {
                "source_seq": 3,
                "source_ref": "word/document.xml:p[3]",
                "part_name": "word/document.xml",
                "paragraph_id": "p_0003",
            },
            {
                "source_seq": 4,
                "source_ref": "word/document.xml:p[4]",
                "part_name": "word/document.xml",
                "paragraph_id": "p_0004",
            },
        ],
        "runs": [],
        "unknown_objects": [],
        "data": {
            "sections": [
                {
                    "index": 1,
                    "paragraph_index": 2,
                    "source_ref": "word/document.xml:p[2]/sectPr",
                    "references": [
                        {
                            "kind": "footer",
                            "type": "default",
                            "part_name": "word/footer1.xml",
                            "source_ref": "word/document.xml:p[2]/sectPr/footerReference[1]",
                        }
                    ],
                    "effective_references": [
                        {
                            "kind": "footer",
                            "type": "default",
                            "part_name": "word/footer1.xml",
                            "source_ref": "word/document.xml:p[2]/sectPr/footerReference[1]",
                        }
                    ],
                    "page_numbering": {"format": "upperRoman", "start": 1},
                    "page_size": {},
                    "page_margins": {},
                },
                {
                    "index": 2,
                    "paragraph_index": None,
                    "source_ref": "word/document.xml:body/sectPr",
                    "references": [],
                    "effective_references": [],
                    "page_numbering": {},
                    "page_size": {},
                    "page_margins": {},
                },
            ],
            "headers_footers": [
                {
                    "kind": "footer",
                    "part_name": "word/footer1.xml",
                    "text": "I",
                    "source_ref": "word/footer1.xml",
                }
            ],
            "fields": [
                {
                    "kind": "fldSimple",
                    "field_type": "PAGE",
                    "instruction": "PAGE",
                    "part_name": "word/footer1.xml",
                    "paragraph_index": 1,
                    "end_paragraph_index": 1,
                    "source_ref": "word/footer1.xml:p[1]/field[1]",
                }
            ],
            "numbering_definitions": [],
            "numbering_refs": [],
        },
    }


def test_global_spec_builds_section_boundaries_and_page_evidence() -> None:
    global_spec = _ai_global_spec_with_two_sections()

    first, second = global_spec["section_profiles"]
    assert global_spec["artifact_version"] == "2.0"
    assert first["boundary"]["start_source_seq"] == 1
    assert first["boundary"]["end_source_seq"] == 2
    assert first["page_numbering"]["declared"]["format"] == "upperRoman"
    assert first["page_numbering"]["display"]["status"] == "detected"
    assert first["page_numbering"]["fields"][0]["source_ref"] == (
        "word/footer1.xml:p[1]/field[1]"
    )
    assert second["boundary"]["start_source_seq"] == 3
    assert second["boundary"]["end_source_seq"] == 4
    assert second["page_numbering"]["display"]["status"] == "no_page_field"
    assert global_spec["page_numbering"]["status"] == "mixed"

    findings = _verify_t4_global_spec(global_spec, start_index=1)
    assert findings == []


def test_global_spec_accepts_lowercase_flag_status() -> None:
    findings = _verify_t4_global_spec(
        {
            "section_profiles": [],
            "header_footer": [],
            "flags": [
                {
                    "status": "unknown",
                    "type": "layout_observation_unavailable",
                    "reason": "replay intentionally omitted T4 decisions",
                }
            ],
        },
        start_index=1,
    )

    flag_finding = next(
        finding
        for finding in findings
        if finding.type == "t4_layout_observation_unavailable"
    )
    assert flag_finding.status == Status.UNKNOWN


def test_template_spec_binds_units_to_overlapping_section_profiles() -> None:
    document_facts = _document_facts_with_two_sections()
    global_spec = _ai_global_spec_with_two_sections()
    unit_map = {
        "artifact_type": "unit_map",
        "flags": [],
        "units": [
            {
                "unit_id": "toc",
                "source_seq_refs": [3, 4],
                "source_seq_range": {"start": 3, "end": 4},
                "section_profile": "section_001",
                "flags": [],
            }
        ],
    }

    template_spec = _build_template_spec(
        build_l1_input_contract(document_facts=document_facts),
        unit_map,
        {"artifact_type": "element_spec", "elements": [], "flags": []},
        global_spec,
    )

    unit = template_spec["units"][0]
    assert unit["section_profile"] == "section_002"
    assert [ref["section_profile_id"] for ref in unit["section_profile_refs"]] == [
        "section_002"
    ]
    assert unit["section_profile_refs"][0]["overlap_source_seq_range"] == {
        "start": 3,
        "end": 4,
    }
    assert _verify_t5_template_spec(template_spec, start_index=1) == []


def test_template_spec_records_cross_section_unit_flag() -> None:
    document_facts = _document_facts_with_two_sections()
    global_spec = _ai_global_spec_with_two_sections()
    unit_map = {
        "artifact_type": "unit_map",
        "flags": [],
        "units": [
            {
                "unit_id": "front_matter",
                "source_seq_refs": [2, 3],
                "source_seq_range": {"start": 2, "end": 3},
                "section_profile": "section_001",
                "flags": [],
            }
        ],
    }

    template_spec = _build_template_spec(
        build_l1_input_contract(document_facts=document_facts),
        unit_map,
        {"artifact_type": "element_spec", "elements": [], "flags": []},
        global_spec,
    )

    unit = template_spec["units"][0]
    assert [ref["section_profile_id"] for ref in unit["section_profile_refs"]] == [
        "section_001",
        "section_002",
    ]
    assert any(
        flag["type"] == "unit_crosses_section_profiles"
        for flag in template_spec["review_flags"]
    )


def test_t4_verifier_rejects_detected_page_numbering_without_page_field() -> None:
    global_spec = _ai_global_spec_with_two_sections()
    first = global_spec["section_profiles"][0]
    first["page_numbering"]["fields"] = []
    first["page_numbering"]["display"]["status"] = "detected"
    first["page_numbering"]["display"]["has_page_field"] = True

    findings = _verify_t4_global_spec(global_spec, start_index=1)

    assert any(
        finding.type == "global_spec_page_numbering_detected_without_page_field"
        and finding.status == Status.FAIL
        for finding in findings
    )


def test_t5_verifier_rejects_missing_section_profile_ref() -> None:
    global_spec = _ai_global_spec_with_two_sections()
    template_spec = {
        "artifact_type": "template_spec",
        "global": global_spec,
        "units": [
            {
                "unit_id": "toc",
                "source_seq_refs": [1],
                "section_profile_refs": [
                    {
                        "section_profile_id": "section_missing",
                        "overlap_source_seq_range": {"start": 1, "end": 1},
                    }
                ],
                "elements": [],
            }
        ],
        "review_flags": [],
    }

    findings = _verify_t5_template_spec(template_spec, start_index=1)

    assert any(
        finding.type == "template_spec_unit_section_profile_ref_missing"
        and finding.status == Status.FAIL
        for finding in findings
    )


def test_t5_verifier_reviews_duplicate_other_units_without_fail() -> None:
    global_spec = _ai_global_spec_with_two_sections()
    other_unit = {
        "unit_id": "other",
        "source_seq_refs": [1],
        "section_profile_refs": [
            {
                "section_profile_id": "section_001",
                "overlap_source_seq_range": {"start": 1, "end": 1},
            }
        ],
        "elements": [],
    }
    template_spec = {
        "artifact_type": "template_spec",
        "global": global_spec,
        "units": [
            other_unit,
            {
                **other_unit,
                "source_seq_refs": [2],
                "section_profile_refs": [
                    {
                        "section_profile_id": "section_001",
                        "overlap_source_seq_range": {"start": 2, "end": 2},
                    }
                ],
            },
        ],
        "review_flags": [],
    }

    findings = _verify_t5_template_spec(template_spec, start_index=1)

    assert any(
        finding.type == "template_spec_duplicate_other_unit_id"
        and finding.status == Status.UNKNOWN
        for finding in findings
    )
    assert not any(
        finding.type == "template_spec_duplicate_unit_id"
        and finding.status == Status.FAIL
        for finding in findings
    )


def test_t6_rejects_manifest_page_action_without_final_docx_effect(tmp_path) -> None:
    fillable = tmp_path / "fillable.docx"
    document = Document()
    document.add_paragraph("正文")
    document.save(fillable)
    manifest = {
        "artifact_type": "build_manifest",
        "output": {"fillable_template_docx_hash": sha256_file(fillable)},
        "actions_executed": [
            {
                "action_id": "a_page_001",
                "action_type": "insert_page_break_before_unit",
                "status": "executed",
                "output_ref": "word/document.xml:p[1]/pageBreakBefore",
            }
        ],
        "actions_requiring_review": [],
        "slots": [],
        "page_policy_results": [],
    }

    findings = _verify_t6_build(
        {"artifact_type": "template_spec", "units": []},
        manifest,
        fillable,
        start_index=1,
    )

    assert any(
        finding.type == "fillable_template_page_action_effect_missing"
        and finding.status == Status.FAIL
        for finding in findings
    )


def test_injected_t2_page_policy_drives_t5_t6_manifest_and_docx_effects(tmp_path) -> None:
    source_docx = tmp_path / "source.docx"
    source = Document()
    source.add_paragraph("封面")
    source.add_paragraph("目录")
    source.add_paragraph("中文摘要")
    source.add_paragraph("正文")
    source.save(source_docx)
    source_hash = sha256_file(source_docx)

    document_facts = {
        "artifact_type": "document_facts",
        "metadata": {"source_template_hash": source_hash},
        "body_flow": [
            _body_flow_item(1, "封面"),
            _body_flow_item(2, "目录"),
            _body_flow_item(3, "中文摘要"),
            _body_flow_item(4, "正文"),
        ],
        "runs": [],
        "unknown_objects": [],
        "data": {
            "sections": [
                {
                    "index": 1,
                    "paragraph_index": None,
                    "source_ref": "word/document.xml:body/sectPr",
                    "references": [],
                    "effective_references": [],
                    "page_numbering": {},
                    "page_size": {},
                    "page_margins": {},
                }
            ],
            "headers_footers": [],
            "fields": [],
            "numbering_definitions": [],
            "numbering_refs": [],
        },
    }
    l1_input_contract = build_l1_input_contract(document_facts=document_facts)
    unit_map = {
        "artifact_type": "unit_map",
        "artifact_version": "1.1",
        "flags": [],
        "units": [
            {
                "unit_id": "cover",
                "order": 1,
                "source_refs": ["word/document.xml:p[1]"],
                "source_seq_refs": [1],
                "source_seq_range": {"start": 1, "end": 1},
                "page_policy": {"start": "document_start", "scope": "page_range_exclusive"},
                "flags": [],
            },
            {
                "unit_id": "toc",
                "order": 2,
                "source_refs": ["word/document.xml:p[2]"],
                "source_seq_refs": [2],
                "source_seq_range": {"start": 2, "end": 2},
                "page_policy": {"start": "new_page", "scope": "page_range_exclusive"},
                "flags": [],
            },
            {
                "unit_id": "abstract_cn",
                "order": 3,
                "source_refs": ["word/document.xml:p[3]"],
                "source_seq_refs": [3],
                "source_seq_range": {"start": 3, "end": 3},
                "page_policy": {"start": "new_page", "scope": "page_range_exclusive"},
                "flags": [],
            },
            {
                "unit_id": "body_main",
                "order": 4,
                "source_refs": ["word/document.xml:p[4]"],
                "source_seq_refs": [4],
                "source_seq_range": {"start": 4, "end": 4},
                "page_policy": {"start": "new_page", "scope": "page_range_exclusive"},
                "flags": [],
            },
        ],
    }
    element_spec = {"artifact_type": "element_spec", "elements": [], "flags": []}
    global_spec = _ai_global_spec_with_two_sections()

    template_spec = _build_template_spec(
        l1_input_contract,
        unit_map,
        element_spec,
        global_spec,
    )
    plan = _build_plan(
        {"source_template_docx": str(source_docx), "source_template_hash": source_hash},
        template_spec,
    )
    generated_docx = tmp_path / "fillable.docx"
    execution = execute_template_generation_plan(
        source_docx,
        generated_docx,
        plan,
        l1_input_contract=l1_input_contract,
    )
    manifest = build_template_generation_manifest(
        request={"source_template_docx": str(source_docx), "source_template_hash": source_hash},
        l1_input_contract=l1_input_contract,
        template_final=_template_final(template_spec),
        plan=plan,
        fillable_template_docx=generated_docx,
        execution=execution,
    )

    action_types = {
        action["action_type"]
        for action in plan["actions"]
        if action["action_type"] != "copy_source_docx"
    }
    assert {
        "insert_page_break_before_unit",
        "ensure_body_slot",
    } <= action_types
    assert not execution["actions_requiring_review"]
    assert manifest["page_breaks"][0]["unit_id"] == "toc"
    assert manifest["page_breaks"][0]["page_policy_unit_id"] == "toc"
    assert manifest["page_breaks"][1]["unit_id"] == "abstract_cn"
    assert manifest["page_breaks"][1]["page_policy_unit_id"] == "abstract_cn"
    assert manifest["page_breaks"][2]["unit_id"] == "body_main"
    assert manifest["page_breaks"][2]["page_policy_unit_id"] == "body_main"
    assert manifest["section_breaks"] == []
    assert manifest["keep_together"] == []

    statuses = {
        result["unit_id"]: result["status"]
        for result in manifest["page_policy_results"]
    }
    assert statuses == {
        "cover": "no_action_required",
        "toc": "executed",
        "abstract_cn": "executed",
        "body_main": "executed",
    }
    assert {
        output_ref
        for result in manifest["page_policy_results"]
        for output_ref in result.get("output_refs", [])
    } >= {
        "word/document.xml:p[2]/pageBreakBefore",
        "word/document.xml:p[3]/pageBreakBefore",
        "word/document.xml:p[4]/pageBreakBefore",
    }

    effects = inspect_docx_layout_effects(generated_docx)
    assert effects["observable_page_break_count"] >= 1
    assert effects["paragraph_section_break_count"] == 0
    assert effects["keep_together_count"] == 0
    assert manifest["observed_layout_effects"] == effects
    assert _verify_t6_build(
        template_spec,
        manifest,
        generated_docx,
        start_index=1,
    ) == []


def test_t5_projects_t2_unit_page_policy_to_executable_page_actions() -> None:
    document_facts = {
        "artifact_type": "document_facts",
        "metadata": {},
        "body_flow": [
            _body_flow_item(1, "封面"),
            _body_flow_item(2, "目录"),
            _body_flow_item(3, "正文"),
        ],
        "runs": [],
        "unknown_objects": [],
        "data": {"sections": [], "headers_footers": []},
    }
    l1_input_contract = build_l1_input_contract(document_facts=document_facts)
    unit_map = {
        "artifact_type": "unit_map",
        "units": [
            {
                "unit_id": "cover",
                "order": 1,
                "source_refs": ["word/document.xml:p[1]"],
                "source_seq_refs": [1],
                "page_policy": {
                    "start": "document_start",
                    "scope": "page_range_exclusive",
                },
            },
            {
                "unit_id": "toc",
                "order": 2,
                "source_refs": ["word/document.xml:p[2]"],
                "source_seq_refs": [2],
                "page_policy": {
                    "start": "new_page",
                    "scope": "page_range_exclusive",
                },
            },
            {
                "unit_id": "body_main",
                "order": 3,
                "source_refs": ["word/document.xml:p[3]"],
                "source_seq_refs": [3],
                "page_policy": {
                    "start": "new_page",
                    "scope": "page_range_exclusive",
                },
            },
        ],
        "flags": [],
    }

    assert unit_map["units"][0]["page_policy"] == {
        "start": "document_start",
        "scope": "page_range_exclusive",
    }
    assert unit_map["units"][1]["page_policy"]["start"] == "new_page"
    assert unit_map["units"][2]["page_policy"]["scope"] == "page_range_exclusive"
    assert all("page" not in unit for unit in unit_map["units"])

    legacy_unit_map_without_page = {
        **unit_map,
        "units": [
            {key: value for key, value in unit.items() if key != "page"}
            for unit in unit_map["units"]
        ],
    }
    template_spec = _build_template_spec(
        l1_input_contract,
        legacy_unit_map_without_page,
        {"artifact_type": "element_spec", "elements": [], "flags": []},
        {"artifact_type": "global_spec", "section_profiles": []},
    )
    plan = _build_plan({}, template_spec)

    action_types = [
        action["action_type"]
        for action in plan["actions"]
        if action["action_type"] != "copy_source_docx"
    ]
    assert action_types.count("insert_page_break_before_unit") == 2
    assert action_types.count("set_keep_together_unit") == 0
    assert {
        result["unit_id"]: result["status"]
        for result in plan["page_policy_results"]
    } == {
        "cover": "no_action_required",
        "toc": "action_planned",
        "body_main": "action_planned",
    }


def test_t5_uses_unit_index_not_order_to_mark_document_start() -> None:
    l1_input_contract = build_l1_input_contract(
        document_facts={
            "artifact_type": "document_facts",
            "metadata": {},
            "body_flow": [_body_flow_item(1, "封面"), _body_flow_item(2, "目录")],
            "runs": [],
            "unknown_objects": [],
            "data": {"sections": [], "headers_footers": []},
        }
    )
    unit_map = {
        "artifact_type": "unit_map",
        "units": [
            {
                "unit_id": "cover",
                "order": 1,
                "source_refs": ["word/document.xml:p[1]"],
                "source_seq_refs": [1],
                "page": _injected_page_policy(
                    page_break="document_start",
                    page_isolation=False,
                    allow_multi_page=False,
                    keep_together=True,
                    proposal_id="fixture.cover.page",
                ),
            },
            {
                "unit_id": "toc",
                "order": 10,
                "source_refs": ["word/document.xml:p[2]"],
                "source_seq_refs": [2],
                "page": {
                    "page_break": "unknown",
                    "page_isolation": "unknown",
                    "allow_multi_page": "unknown",
                    "keep_together": "unknown",
                    "decision": {
                        "origin": "unknown",
                        "confidence": "low",
                        "evidence_refs": [],
                        "conflict_status": "none",
                        "proposal_ids": [],
                    },
                },
            },
        ],
        "flags": [],
    }

    template_spec = _build_template_spec(
        l1_input_contract,
        unit_map,
        {"artifact_type": "element_spec", "elements": [], "flags": []},
        {"artifact_type": "global_spec", "section_profiles": []},
    )

    assert template_spec["units"][0]["page_policy"]["start"] == "document_start"
    assert template_spec["units"][1]["page_policy"]["start"] == "new_page"
    assert all("page" not in unit for unit in template_spec["units"])


def test_t6_dedupes_exclusive_scope_when_next_unit_requires_new_page() -> None:
    template_spec = {
        "units": [
            {
                "unit_id": "cover",
                "source_refs": ["word/document.xml:p[1]"],
                "source_seq_refs": [1],
                "page_policy": {"start": "document_start", "scope": "page_range_exclusive"},
            },
            {
                "unit_id": "abstract_cn",
                "source_refs": ["word/document.xml:p[2]"],
                "source_seq_refs": [2],
                "page_policy": {"start": "new_page", "scope": "page_range_exclusive"},
            },
        ],
    }

    plan = _build_plan({}, template_spec)

    boundary_actions = [
        action
        for action in plan["actions"]
        if action["action_type"]
        in {"insert_page_break_before_unit", "insert_section_break_before_unit"}
    ]
    assert [action["action_type"] for action in boundary_actions] == [
        "insert_page_break_before_unit"
    ]
    planned_ids = {
        result["unit_id"]: result["planned_action_ids"]
        for result in plan["page_policy_results"]
    }
    assert planned_ids["cover"] == []
    assert planned_ids["abstract_cn"] == [boundary_actions[0]["action_id"]]


def test_t6_preserves_t2_observed_page_boundary_without_duplicate_break() -> None:
    template_spec = {
        "units": [
            {
                "unit_id": "cover",
                "source_refs": ["word/document.xml:p[1]"],
                "source_seq_refs": [1],
                "boundary": {"start_page": 1, "end_page": 1},
                "page_policy": {
                    "start": "document_start",
                    "scope": "page_range_exclusive",
                },
            },
            {
                "unit_id": "toc",
                "source_refs": [],
                "source_seq_refs": [],
                "boundary": {"start_page": 2, "end_page": 3},
                "page_policy": {
                    "start": "new_page",
                    "scope": "page_range_exclusive",
                },
            },
        ],
    }

    plan = _build_plan({}, template_spec)

    assert not any(
        action["action_type"] == "insert_page_break_before_unit"
        for action in plan["actions"]
    )
    assert plan["page_policy_results"][1]["status"] == "already_satisfied"
    assert plan["page_policy_results"][1]["observed_boundary"] == {
        "start_page": 2,
        "end_page": 3,
    }


def test_page_break_action_preserves_existing_section_boundary(tmp_path) -> None:
    source = tmp_path / "section-boundary.docx"
    doc = Document()
    doc.add_paragraph("封面")
    doc.add_section(WD_SECTION.NEW_PAGE)
    target = doc.add_paragraph("目录")
    doc.save(source)

    loaded = Document(source)
    paragraph_map = _paragraph_map_by_ooxml_index(loaded)
    target_index = next(
        index
        for index, paragraph in paragraph_map.items()
        if paragraph.text == target.text
    )
    output_ref = _insert_page_break_before(
        loaded,
        paragraph_map,
        f"word/document.xml:p[{target_index}]",
    )

    assert output_ref is not None
    assert output_ref.endswith("/sectPr")
    assert paragraph_map[target_index].paragraph_format.page_break_before is not True


def test_t6_replaces_cover_value_with_one_inline_slot_and_preserves_layout(tmp_path) -> None:
    source_docx = tmp_path / "cover-source.docx"
    source = Document()
    paragraph = source.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.add_run("学生姓名：")
    value = paragraph.add_run("张三")
    value.underline = True
    source.add_paragraph("下一行")
    source.save(source_docx)
    source_hash = sha256_file(source_docx)
    document_facts = {
        "artifact_type": "document_facts",
        "metadata": {"source_template_hash": source_hash},
        "body_flow": [
            {
                **_body_flow_item(1, "学生姓名：张三"),
                "raw_run_ids": ["p_0001.r_001", "p_0001.r_002"],
                "style_details": {
                    "runs": [
                        {"source_ref": "word/document.xml:p[1]", "text": "学生姓名："},
                        {"source_ref": "word/document.xml:p[1]", "text": "张三"},
                    ]
                },
            },
            _body_flow_item(2, "下一行"),
        ],
        "runs": [],
        "unknown_objects": [],
        "data": {"sections": [], "headers_footers": []},
    }
    l1_input_contract = build_l1_input_contract(document_facts=document_facts)
    template_spec = {
        "units": [
            {
                "unit_id": "cover",
                "elements": [
                    {
                        "stable_id": "student_name",
                        "policy": "fill",
                        "source_ref": "word/document.xml:p[1]",
                        "source_seq_refs": [1],
                        "raw_run_ids": ["p_0001.r_002"],
                        "spans": [
                            {
                                "span_id": "value",
                                "span_type": "sample_value",
                                "raw_run_ids": ["p_0001.r_002"],
                                "char_ranges": [
                                    {
                                        "raw_run_id": "p_0001.r_002",
                                        "start": 0,
                                        "end": 2,
                                        "replacement": "",
                                    }
                                ],
                            }
                        ],
                    },
                ],
            }
        ],
    }
    plan = _build_plan(
        {"source_template_docx": str(source_docx), "source_template_hash": source_hash},
        template_spec,
    )

    slot_actions = [
        action
        for action in plan["actions"]
        if action.get("unit_id") == "cover"
        and action.get("element_id") == "student_name"
        and action.get("action_type") in {"create_fillable_slot", "replace_span_with_slot"}
    ]
    assert [action["action_type"] for action in slot_actions] == ["replace_span_with_slot"]

    generated_docx = tmp_path / "cover-fillable.docx"
    execution = execute_template_generation_plan(
        source_docx,
        generated_docx,
        plan,
        l1_input_contract=l1_input_contract,
    )

    assert not execution["actions_requiring_review"]
    output = Document(generated_docx)
    assert len(output.paragraphs) == 2
    assert output.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert output.paragraphs[0].paragraph_format.space_after == Pt(9)
    inline_slots = output.paragraphs[0]._p.findall(qn("w:sdt"))
    assert len(inline_slots) == 1
    slot_run = inline_slots[0].find(f"{qn('w:sdtContent')}/{qn('w:r')}")
    assert slot_run is not None
    assert slot_run.find(f"{qn('w:rPr')}/{qn('w:u')}") is not None


def test_t6_uses_t5_element_id_instead_of_prefixed_stable_id() -> None:
    plan = _build_plan(
        {},
        {
            "units": [
                {
                    "unit_id": "body_main",
                    "source_refs": ["word/document.xml:p[1]"],
                    "source_seq_refs": [1],
                    "page_policy": {
                        "start": "document_start",
                        "scope": "page_range_exclusive",
                    },
                    "elements": [
                        {
                            "element_id": "e_001",
                            "stable_id": "body_main.e_001",
                            "policy": "fill",
                            "source_refs": ["word/document.xml:p[1]"],
                            "source_seq_refs": [1],
                        }
                    ],
                }
            ]
        },
    )

    action = next(
        item
        for item in plan["actions"]
        if item["action_type"] == "create_fillable_slot"
    )
    assert action["element_id"] == "e_001"
    assert action["target_ref"] == "sdt:body_main.e_001"


def test_t6_synthesizes_source_less_fixed_element_from_t5() -> None:
    plan = _build_plan(
        {},
        {
            "units": [
                {
                    "unit_id": "toc",
                    "source_refs": ["word/document.xml:p[2]"],
                    "source_seq_refs": [2],
                    "page_policy": {
                        "start": "document_start",
                        "scope": "page_range_exclusive",
                    },
                    "elements": [
                        {
                            "element_id": "e_001",
                            "order": 1,
                            "policy": "fixed",
                            "content": "目录",
                            "source_refs": [],
                            "source_seq_refs": [],
                        }
                    ],
                }
            ]
        },
    )

    assert any(
        item["action_type"] == "insert_fixed_text"
        and item["source_ref"] == "word/document.xml:p[2]"
        and item["target_ref"] == "目录"
        for item in plan["actions"]
    )


def test_t6_recognizes_open_t2_id_for_table_of_contents_title() -> None:
    plan = _build_plan(
        {},
        {
            "units": [
                {
                    "unit_id": "school_contents_pages",
                    "unit_name": "目录",
                    "source_refs": [],
                    "source_seq_refs": [1],
                    "page_policy": {
                        "start": "document_start",
                        "scope": "page_range_exclusive",
                    },
                    "elements": [
                        {
                            "element_id": "e_001",
                            "order": 1,
                            "policy": "generated",
                            "content": "目录",
                            "source_refs": [],
                            "source_seq_refs": [1],
                        }
                    ],
                },
                {
                    "unit_id": "school_body_pages",
                    "unit_name": "正文",
                    "source_refs": ["word/document.xml:p[2]"],
                    "source_seq_refs": [2],
                    "page_policy": {
                        "start": "new_page",
                        "scope": "page_range_exclusive",
                    },
                    "elements": [],
                },
            ]
        },
    )

    assert any(
        item["action_type"] == "insert_synthetic_unit_title_before"
        and item["unit_id"] == "school_contents_pages"
        and item["target_ref"] == "目录"
        for item in plan["actions"]
    )


def test_t3_generated_field_uses_unit_name_instead_of_fixed_t2_id() -> None:
    element_spec = build_element_spec(
        {
            "artifact_type": "template_generation_model",
            "units": [
                {
                    "unit_id": "school_contents_pages",
                    "unit_name": "目录",
                    "elements": [
                        {
                            "element_id": "e_001",
                            "policy": "generated",
                            "content": "",
                            "name": "",
                            "source_refs": [],
                            "source_seq_refs": [],
                        }
                    ],
                }
            ],
        }
    )

    assert element_spec["elements"][0]["generated"]["field_type"] == "TOC"


def test_t6_gates_actions_by_their_direct_final_upstream() -> None:
    template_spec = {
        "artifact_type": "template_spec",
        "units": [
            {
                "unit_id": "cover",
                "source_refs": ["word/document.xml:p[1]"],
                "source_seq_refs": [1],
                "page_policy": {
                    "start": "document_start",
                    "scope": "page_range_exclusive",
                },
                "elements": [],
            },
            {
                "unit_id": "body_main",
                "source_refs": ["word/document.xml:p[2]"],
                "source_seq_refs": [2],
                "page_policy": {
                    "start": "new_page",
                    "scope": "page_range_exclusive",
                },
                "elements": [
                    {
                        "element_id": "e_001",
                        "policy": "fill",
                        "source_refs": ["word/document.xml:p[2]"],
                        "source_seq_refs": [2],
                    }
                ],
            }
        ],
    }
    t5_final = publish_final_stage_result(
        template_spec,
        stage_id="T5",
        artifact_type="template_spec",
        artifact_name="05_template_spec.yaml",
        availability=NOT_AVAILABLE,
        reason="T3: fixture unavailable",
        producer_mode="fixture",
        input_refs={
            "t2_final": {"availability": "AVAILABLE"},
            "t3_final": {"availability": "NOT_AVAILABLE"},
        },
    )

    plan = build_template_generation_plan({}, template_final=t5_final)
    action_types = {item["action_type"] for item in plan["actions"]}

    assert "insert_page_break_before_unit" in action_types
    assert "create_fillable_slot" not in action_types
    assert "ensure_body_slot" in action_types
    assert plan["page_policy_results"][1]["status"] == "action_planned"

    t2_unavailable = publish_final_stage_result(
        template_spec,
        stage_id="T5",
        artifact_type="template_spec",
        artifact_name="05_template_spec.yaml",
        availability=NOT_AVAILABLE,
        reason="T2: fixture unavailable",
        producer_mode="fixture",
        input_refs={
            "t2_final": {"availability": "NOT_AVAILABLE"},
            "t3_final": {"availability": "AVAILABLE"},
        },
    )
    blocked_plan = build_template_generation_plan(
        {},
        template_final=t2_unavailable,
    )

    assert not any(
        item["action_type"] == "insert_page_break_before_unit"
        for item in blocked_plan["actions"]
    )
    assert all(
        item["status"] == "manual_review"
        for item in blocked_plan["page_policy_results"]
    )


def test_t6_verifier_rejects_page_break_effect_on_wrong_paragraph(tmp_path) -> None:
    fillable = tmp_path / "wrong-ref.docx"
    doc = Document()
    first = doc.add_paragraph("封面")
    first.paragraph_format.page_break_before = True
    doc.add_paragraph("目录")
    doc.save(fillable)
    effects = inspect_docx_layout_effects(fillable)
    template_spec = {
        "units": [
            {
                "unit_id": "cover",
                "page": _injected_page_policy(
                    page_break="document_start",
                    page_isolation=False,
                    allow_multi_page=True,
                    keep_together=False,
                    proposal_id="fixture.cover.page",
                ),
            },
            {
                "unit_id": "toc",
                "page": _injected_page_policy(
                    page_break=True,
                    page_isolation=False,
                    allow_multi_page=True,
                    keep_together=False,
                    proposal_id="fixture.toc.page",
                ),
            },
        ]
    }
    manifest = {
        "output": {"fillable_template_docx_hash": sha256_file(fillable)},
        "slots": [],
        "actions_requiring_review": [],
        "actions_executed": [
            {
                "action_id": "a_002",
                "action_type": "insert_page_break_before_unit",
                "unit_id": "toc",
                "output_ref": "word/document.xml:p[2]/pageBreakBefore",
            }
        ],
        "page_policy_results": [
            {"unit_id": "cover", "status": "no_action_required", "executed_action_ids": []},
            {"unit_id": "toc", "status": "executed", "executed_action_ids": ["a_002"]},
        ],
        "observed_layout_effects": effects,
    }

    findings = _verify_t6_build(
        template_spec,
        manifest,
        fillable,
        start_index=1,
    )

    assert any(
        finding.type == "fillable_template_page_action_effect_missing"
        for finding in findings
    )


def _body_flow_item(source_seq: int, text: str) -> dict:
    return {
        "source_seq": source_seq,
        "source_ref": f"word/document.xml:p[{source_seq}]",
        "part_name": "word/document.xml",
        "paragraph_id": f"p_{source_seq:04d}",
        "kind": "paragraph",
        "flow_item_type": "paragraph",
        "structure_layer": "body_flow",
        "order": source_seq,
        "text": text,
    }


def _injected_page_policy(
    *,
    page_break,
    page_isolation: bool,
    allow_multi_page: bool,
    keep_together,
    proposal_id: str,
    enforcement_hint: str | None = None,
) -> dict:
    page = {
        "page_break": page_break,
        "page_isolation": page_isolation,
        "allow_multi_page": allow_multi_page,
        "keep_together": keep_together,
        "decision": {
            "origin": "fixture_realistic_t2_page_policy",
            "confidence": "high",
            "evidence_refs": [f"manual-fixture:{proposal_id}"],
            "conflict_status": "none",
            "proposal_ids": [proposal_id],
        },
    }
    if enforcement_hint is not None:
        page["page_policy"] = {
            "generation_policy": {
                "requires_new_page": page_break is True,
                "source": "fixture_realistic_t2_page_policy",
                "confidence": "high",
                "enforcement_hint": enforcement_hint,
                "evidence_refs": [f"manual-fixture:{proposal_id}"],
            }
        }
    return page
