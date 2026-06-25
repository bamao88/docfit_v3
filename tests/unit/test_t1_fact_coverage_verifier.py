from __future__ import annotations

from docx import Document

from docfit.template_generation.verifier import _verify_t1_document_facts


def _finding_types(document_facts: dict) -> set[str]:
    return {
        finding.type
        for finding in _verify_t1_document_facts(document_facts, start_index=1)
    }


def test_t1_verifier_flags_paragraph_trace_mismatch_and_semantic_fields() -> None:
    document_facts = {
        "artifact_type": "document_facts",
        "body_flow": [
            {
                "node_id": "body_0001",
                "kind": "paragraph",
                "visible": True,
                "text": "摘  要\tⅠ",
                "source_ref": "word/document.xml:p[1]",
                "paragraph_id": "p_0001",
                "raw_run_ids": ["p_0002.r_001"],
                "structural_signals": {"is_toc_entry": True},
            }
        ],
        "runs": [],
        "unknown_objects": [],
    }

    finding_types = _finding_types(document_facts)

    assert "document_facts_paragraph_trace_mismatch" in finding_types
    assert "document_facts_semantic_field_in_t1" in finding_types


def test_t1_verifier_flags_visible_items_without_trace() -> None:
    document_facts = {
        "artifact_type": "document_facts",
        "body_flow": [
            {
                "node_id": "body_0001",
                "kind": "paragraph",
                "visible": True,
                "text": "正文",
                "source_ref": "word/document.xml:p[1]",
                "raw_run_ids": [],
            },
            {
                "node_id": "body_0002",
                "kind": "table_cell",
                "visible": True,
                "text": "封面",
                "source_ref": "word/document.xml:tbl[1]/tr[1]/tc[1]",
                "raw_run_ids": [],
            },
            {
                "node_id": "body_0003",
                "kind": "footer",
                "visible": True,
                "text": "第 1 页",
                "source_ref": "word/footer1.xml",
                "raw_run_ids": [],
            },
        ],
        "runs": [],
        "unknown_objects": [],
    }

    finding_types = _finding_types(document_facts)

    assert "document_facts_visible_paragraph_trace_missing" in finding_types
    assert "document_facts_visible_table_cell_trace_missing" in finding_types
    assert "document_facts_header_footer_trace_missing" in finding_types


def test_t1_verifier_flags_source_ref_text_mismatch(tmp_path) -> None:
    source = tmp_path / "source_template.docx"
    document = Document()
    document.add_paragraph("正确段落")
    document.add_paragraph("错误引用目标")
    document.save(source)
    document_facts = {
        "artifact_type": "document_facts",
        "metadata": {"source_template_docx": str(source)},
        "body_flow": [
            {
                "node_id": "body_0001",
                "kind": "paragraph",
                "visible": True,
                "text": "错误引用目标",
                "source_ref": "word/document.xml:p[1]",
                "paragraph_id": "p_0001",
                "raw_run_ids": ["p_0001.r_001"],
            }
        ],
        "runs": [],
        "unknown_objects": [],
    }

    finding_types = _finding_types(document_facts)

    assert "document_facts_source_ref_text_mismatch" in finding_types
