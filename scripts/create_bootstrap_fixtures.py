from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from pathlib import Path
from zipfile import ZipFile

import yaml
from docx import Document

from docfit.harness.profiles import (
    BOOTSTRAP_PASS_STUDENT_DOCX,
    BOOTSTRAP_PROFILE,
    BOOTSTRAP_SILENT_DROP_DOCX,
    BOOTSTRAP_TEMPLATE_DOCX,
    BOOTSTRAP_UNSUPPORTED_TEXTBOX_DOCX,
)


ROOT = Path(__file__).resolve().parents[1]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return "sha256:" + digest.hexdigest()


def sha256_text(text: str) -> str:
    return "sha256:" + hashlib.sha256(text.encode("utf-8")).hexdigest()


def sha256_json(data) -> str:
    payload = json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return sha256_text(payload)


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def create_template(path: Path) -> None:
    ensure_dir(path.parent)
    doc = Document()
    doc.add_heading("Demo School Thesis Template", level=1)
    doc.add_paragraph("This synthetic template is owned by the bootstrap standard.")
    doc.add_paragraph("[[DOCFIT_SLOT:body]]")
    doc.save(path)


def create_demo_student(path: Path) -> list[str]:
    ensure_dir(path.parent)
    doc = Document()
    doc.add_heading("Chapter 1 Introduction", level=1)
    doc.add_paragraph(
        "DocFit bootstrap proves that visible paragraphs can be extracted and placed."
    )
    doc.add_heading("Methods", level=2)
    doc.add_paragraph(
        "The harness keeps a ledger for every visible block before rendering."
    )
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    rows = [
        ["Metric", "Value"],
        ["Visible paragraphs", "4"],
        ["Simple tables", "1"],
    ]
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    doc.save(path)
    return [
        sha256_text("Chapter 1 Introduction"),
        sha256_text("DocFit bootstrap proves that visible paragraphs can be extracted and placed."),
        sha256_text("Methods"),
        sha256_text("The harness keeps a ledger for every visible block before rendering."),
        sha256_json(rows),
    ]


def create_silent_drop_student(source: Path, target: Path) -> None:
    ensure_dir(target.parent)
    shutil.copyfile(source, target)


def create_textbox_student(path: Path) -> None:
    ensure_dir(path.parent)
    doc = Document()
    doc.add_heading("Unsupported Object Case", level=1)
    doc.add_paragraph("This paragraph is extractable, but a visible textbox follows.")
    doc.save(path)
    inject_textbox(path)


def inject_textbox(path: Path) -> None:
    tmp = path.with_suffix(".tmp.docx")
    textbox_xml = """
    <w:p>
      <w:r>
        <w:pict>
          <v:shape id="TextBox1" type="#_x0000_t202" style="width:180pt;height:45pt">
            <v:textbox>
              <w:txbxContent>
                <w:p>
                  <w:r><w:t>Visible textbox content unsupported by bootstrap.</w:t></w:r>
                </w:p>
              </w:txbxContent>
            </v:textbox>
          </v:shape>
        </w:pict>
      </w:r>
    </w:p>
    """.strip()
    with ZipFile(path, "r") as source, ZipFile(tmp, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                text = data.decode("utf-8")
                text = text.replace("<w:sectPr", textbox_xml + "<w:sectPr", 1)
                data = text.encode("utf-8")
            target.writestr(item, data)
    tmp.replace(path)


def write_json(path: Path, data) -> None:
    ensure_dir(path.parent)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def create_contracts(standard_dir: Path) -> None:
    ensure_dir(standard_dir)
    common = {
        "contract_version": "1.0",
        "owner": "docfit-core",
        "required_invariants": [],
        "required_capabilities": [],
        "unsupported_policy": {"blocking_unknown": True},
        "coverage_requirements": {"profile": "bootstrap-core"},
        "verifier_refs": [],
    }
    contracts = {
        "template_contract.json": {
            **common,
            "contract_type": "template",
            "required_invariants": ["required slot exists", "required region exists"],
            "required_capabilities": BOOTSTRAP_PROFILE.capabilities_for_stage("template"),
            "required_slots": ["slot_body_start"],
            "required_regions": ["body"],
            "verifier_refs": ["docfit.stages.template_parse.verify_template_artifact"],
        },
        "student_content_contract.json": {
            **common,
            "contract_type": "student_content",
            "required_invariants": ["visible content ledger complete"],
            "required_capabilities": BOOTSTRAP_PROFILE.capabilities_for_stage("content"),
            "verifier_refs": ["docfit.stages.content_extract.verify_student_content_artifact"],
        },
        "placement_contract.json": {
            **common,
            "contract_type": "placement",
            "required_invariants": ["no silent drop", "slot compatibility"],
            "required_capabilities": BOOTSTRAP_PROFILE.capabilities_for_stage("placement"),
            "verifier_refs": ["docfit.stages.placement.verify_placement_plan"],
        },
        "render_contract.json": {
            **common,
            "contract_type": "render",
            "required_invariants": ["valid docx", "plan coverage", "signed golden exists"],
            "required_capabilities": BOOTSTRAP_PROFILE.capabilities_for_stage("render"),
            "verifier_refs": ["docfit.stages.render.verify_render_outputs"],
        },
    }
    for filename, payload in contracts.items():
        write_json(standard_dir / filename, payload)
    (standard_dir / "exceptions.yaml").write_text("exceptions: []\n", encoding="utf-8")


def create_standard(output_root: Path, template_path: Path, expected_hashes: list[str]) -> None:
    standard_dir = output_root / "standards/schools/demo-school/v1"
    create_contracts(standard_dir)
    golden_dir = standard_dir / "golden"
    write_json(
        golden_dir / "feature_snapshot.json",
        {
            "artifact_type": "feature_snapshot_golden",
            "artifact_version": "1.0",
            "required_content_hashes": expected_hashes,
            "change_control": {
                "auto_update_allowed": False,
                "requires_review": True,
            },
        },
    )
    signed_standard = {
        "standard_id": "demo-school-v1",
        "school_id": "demo-school",
        "template_version": "v1",
        "status": "signed",
        "owner": "docfit-core",
        "approved_at": "2026-06-14T00:00:00+00:00",
        "source": {
            "template_docx": str(BOOTSTRAP_TEMPLATE_DOCX),
            "template_docx_sha256": sha256_file(template_path),
        },
        "contracts": {
            "template_contract": "template_contract.json",
            "student_content_contract": "student_content_contract.json",
            "placement_contract": "placement_contract.json",
            "render_contract": "render_contract.json",
        },
        "goldens": {
            "feature_snapshot": "golden/feature_snapshot.json",
            "expected_docx": "golden/expected.docx",
        },
        "coverage_requirements": {
            "profile": BOOTSTRAP_PROFILE.profile_id,
            "required_capabilities": BOOTSTRAP_PROFILE.all_required_capabilities(),
        },
        "change_control": {
            "auto_update_allowed": False,
            "requires_review": True,
            "change_reason": "initial bootstrap standard with explicit capability profile",
        },
    }
    (standard_dir / "signed_standard.yaml").write_text(
        yaml.safe_dump(signed_standard, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )


def create_expected_files(output_root: Path, expected_hashes: list[str]) -> None:
    write_json(
        output_root / BOOTSTRAP_PROFILE.expected_feature_snapshot,
        {"required_content_hashes": expected_hashes},
    )
    write_json(
        output_root / BOOTSTRAP_PROFILE.expected_placement_plan,
        {
            "expected_action_count": len(expected_hashes),
            "required_disposition": "place",
        },
    )


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--root",
        type=Path,
        default=ROOT / "test_outputs/workbench/bootstrap-fixtures",
        help="Output root. Defaults to ignored generated output.",
    )
    parser.add_argument(
        "--write-reviewed-assets",
        action="store_true",
        help="Allow writing signed standards, goldens, and input assets in the repo root.",
    )
    return parser.parse_args(argv)


def resolve_output_root(root: Path, *, write_reviewed_assets: bool) -> Path:
    output_root = root if root.is_absolute() else ROOT / root
    if output_root.resolve() == ROOT.resolve() and not write_reviewed_assets:
        raise SystemExit(
            "Refusing to rewrite reviewed repo assets without --write-reviewed-assets."
        )
    return output_root


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    output_root = resolve_output_root(
        args.root,
        write_reviewed_assets=args.write_reviewed_assets,
    )
    template_path = output_root / BOOTSTRAP_TEMPLATE_DOCX
    student_path = output_root / BOOTSTRAP_PASS_STUDENT_DOCX
    textbox_path = output_root / BOOTSTRAP_UNSUPPORTED_TEXTBOX_DOCX
    silent_drop_path = output_root / BOOTSTRAP_SILENT_DROP_DOCX
    create_template(template_path)
    expected_hashes = create_demo_student(student_path)
    create_silent_drop_student(student_path, silent_drop_path)
    create_textbox_student(textbox_path)
    create_standard(output_root, template_path, expected_hashes)
    create_expected_files(output_root, expected_hashes)


if __name__ == "__main__":
    main()
